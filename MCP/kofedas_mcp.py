from __future__ import annotations

import json
import os
import re
import base64
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable
from xml.etree import ElementTree

SERVER_VERSION = "0.1.0"
PUBLIC_CONTRACT_VERSION = "1.0"
DEFAULT_EMPRESA = 1
DEFAULT_CENTRO = 0
MAX_ROWS_DEFAULT = 100
MAX_ROWS_LIMIT = 1000
ACCESS_LEVELS = {"read": 0, "write": 1, "critical": 2}


class KofedasError(Exception):
    pass


def _env_int(name: str, default: int) -> int:
    raw = os.getenv(name)
    if raw in (None, ""):
        return default
    try:
        return int(raw)
    except ValueError as exc:
        raise KofedasError(f"{name} debe ser entero") from exc


def _normalize(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value


def _row_to_dict(cursor: Any, row: Any) -> dict[str, Any]:
    columns = [column[0].lower() for column in cursor.description]
    return {column: _normalize(value) for column, value in zip(columns, row)}


def _like(value: str) -> str:
    return f"%{value.strip().upper()}%"


def _positive_limit(value: Any, default: int = MAX_ROWS_DEFAULT) -> int:
    if value in (None, ""):
        return default
    try:
        limit = int(value)
    except (TypeError, ValueError) as exc:
        raise KofedasError("limite debe ser entero") from exc
    if limit <= 0:
        raise KofedasError("limite debe ser mayor que cero")
    return min(limit, MAX_ROWS_LIMIT)


class KofedasDatabase:
    def __init__(self) -> None:
        self.dsn = os.getenv("KOFEDAS_ODBC_DSN") or os.getenv("KRONOS_ODBC_DSN") or "Kronos"
        self.user = os.getenv("KOFEDAS_DB_USER") or os.getenv("KRONOS_DB_USER") or "SYSDBA"
        self.password = os.getenv("KOFEDAS_DB_PASSWORD") or os.getenv("KRONOS_DB_PASSWORD") or "masterkey"
        self.timeout = _env_int("KOFEDAS_DB_TIMEOUT", 30)

    @contextmanager
    def connect(self):
        try:
            import pyodbc
        except ImportError as exc:
            raise KofedasError("Falta pyodbc. Instala requirements.txt en el entorno del MCP.") from exc

        connection = pyodbc.connect(
            f"DSN={self.dsn};UID={self.user};PWD={self.password}",
            autocommit=True,
            timeout=self.timeout,
        )
        try:
            yield connection
        finally:
            connection.close()

    def query(self, sql: str, params: tuple[Any, ...] = (), limit: int | None = None) -> list[dict[str, Any]]:
        with self.connect() as connection:
            cursor = connection.cursor()
            cursor.execute(sql, params)
            rows = cursor.fetchmany(limit or MAX_ROWS_LIMIT)
            return [_row_to_dict(cursor, row) for row in rows]

    def one(self, sql: str, params: tuple[Any, ...] = ()) -> dict[str, Any] | None:
        with self.connect() as connection:
            cursor = connection.cursor()
            cursor.execute(sql, params)
            row = cursor.fetchone()
            return _row_to_dict(cursor, row) if row else None

    def columns(self, table: str) -> list[str]:
        with self.connect() as connection:
            cursor = connection.cursor()
            sql = """
            SELECT TRIM(rf.rdb$field_name) AS field_name
            FROM rdb$relation_fields rf
            WHERE rf.rdb$relation_name = ?
            ORDER BY rf.rdb$field_position
            """
            cursor.execute(sql, (table.upper(),))
            return [str(row[0]).strip().upper() for row in cursor.fetchall()]

    def primary_key(self, table: str) -> list[str]:
        with self.connect() as connection:
            cursor = connection.cursor()
            sql = """
            SELECT TRIM(s.rdb$field_name) AS field_name
            FROM rdb$relation_constraints rc
            JOIN rdb$index_segments s ON s.rdb$index_name = rc.rdb$index_name
            WHERE rc.rdb$relation_name = ? AND rc.rdb$constraint_type = 'PRIMARY KEY'
            ORDER BY s.rdb$field_position
            """
            cursor.execute(sql, (table.upper(),))
            return [str(row[0]).strip().upper() for row in cursor.fetchall()]

    def text_columns(self, table: str) -> list[str]:
        with self.connect() as connection:
            cursor = connection.cursor()
            sql = """
            SELECT TRIM(rf.rdb$field_name) AS field_name
            FROM rdb$relation_fields rf
            JOIN rdb$fields f ON f.rdb$field_name = rf.rdb$field_source
            WHERE rf.rdb$relation_name = ? AND f.rdb$field_type IN (14, 37)
            ORDER BY rf.rdb$field_position
            """
            cursor.execute(sql, (table.upper(),))
            return [str(row[0]).strip().upper() for row in cursor.fetchall()]

    def execute(self, sql: str, params: tuple[Any, ...] = ()) -> int:
        try:
            import pyodbc
        except ImportError as exc:
            raise KofedasError("Falta pyodbc. Instala requirements.txt en el entorno del MCP.") from exc

        connection = pyodbc.connect(
            f"DSN={self.dsn};UID={self.user};PWD={self.password}",
            autocommit=False,
            timeout=self.timeout,
        )
        try:
            cursor = connection.cursor()
            cursor.execute(sql, params)
            affected = cursor.rowcount
            connection.commit()
            return affected
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def execute_transaction(self, statements: list[tuple[str, tuple[Any, ...]]]) -> list[int]:
        try:
            import pyodbc
        except ImportError as exc:
            raise KofedasError("Falta pyodbc. Instala requirements.txt en el entorno del MCP.") from exc

        connection = pyodbc.connect(
            f"DSN={self.dsn};UID={self.user};PWD={self.password}",
            autocommit=False,
            timeout=self.timeout,
        )
        try:
            cursor = connection.cursor()
            counts: list[int] = []
            for sql, params in statements:
                cursor.execute(sql, params)
                counts.append(cursor.rowcount)
            connection.commit()
            return counts
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()


def _string_schema(description: str = "") -> dict[str, Any]:
    schema: dict[str, Any] = {"type": "string"}
    if description:
        schema["description"] = description
    return schema


def _int_schema(description: str = "") -> dict[str, Any]:
    schema: dict[str, Any] = {"type": "integer"}
    if description:
        schema["description"] = description
    return schema


def _tool(name: str, description: str, properties: dict[str, Any], required: list[str] | None = None) -> dict[str, Any]:
    return {
        "name": name,
        "description": description,
        "inputSchema": {
            "type": "object",
            "properties": properties,
            "required": required or [],
            "additionalProperties": False,
        },
    }


def _data_schema(description: str) -> dict[str, Any]:
    return {"type": "object", "description": description, "additionalProperties": True}


EMPRESA_FIELDS = [
    "EMP_NUMEMP", "EMP_CODSOC", "EMP_NOMEMP", "EMP_NOMFIS", "EMP_DOMFIS1",
    "EMP_DOMFIS2", "EMP_CODPOS", "EMP_POBLAC", "EMP_CIF", "EMP_REGIVA",
    "EMP_EJEEUR", "EMP_EMAIL",
]
CENTRO_FIELDS = [
    "CEN_NUMEMP", "CEN_CODCEN", "CEN_TIPCEN", "CEN_NOMCEN", "CEN_DOMIC1",
    "CEN_DOMIC2", "CEN_CODPOS", "CEN_POBLAC", "CEN_TELEF", "CEN_FAX",
    "CEN_EMAIL", "CEN_CLIINI", "CEN_CLIFIN", "CEN_CLICINI", "CEN_CLICFIN",
]
USUARIO_FIELDS = [
    "USU_NUMEMP", "USU_CODCEN", "USU_NOMUSU", "USU_PASSWORD", "USU_NOMBRE",
    "USU_IDGRUPO", "USU_CONDGRU", "USU_NIVEL", "USU_NIVPRI", "USU_MENUI",
    "USU_CAMFEC", "USU_FECMIN", "USU_FECMAX", "USU_AUTVEN", "USU_AUTDTO",
]
GRUPO_USUARIO_FIELDS = [
    "GRU_NUMEMP", "GRU_IDGRUPO", "GRU_NIVEL", "GRU_NIVPRI", "GRU_MENUI",
    "GRU_CAMFEC", "GRU_FECMIN", "GRU_FECMAX", "GRU_AUTVEN", "GRU_AUTDTO",
]
PARAMETRO_FIELDS = ["PAR_NUMEMP", "PAR_CODIGO", "PAR_VALOR", "PAR_DESCRI"]

AUXILIARY_TABLES: dict[str, str] = {
    "AGRUP1": "Agrupacion nivel 1",
    "AGRUP2": "Agrupacion nivel 2",
    "AGRUP3": "Agrupacion nivel 3",
    "BANCOS": "Bancos",
    "CAJAS": "Cajas",
    "CATALO": "Catalogos",
    "COMISI": "Comisiones",
    "COMISI2": "Comisiones detalle",
    "CUEREM": "Cuentas de remesa",
    "DETCAT": "Detalle de catalogos",
    "ENTIDA": "Entidades bancarias",
    "ENVIOS": "Envios",
    "FAMILI": "Familias",
    "FORENV": "Formas de envio",
    "FORPAG": "Formas de pago",
    "FORPAGA": "Aplazamientos de forma de pago",
    "IMPRESORAS": "Impresoras",
    "MEDIDAS": "Unidades de medida",
    "MENU1": "Menu nivel 1",
    "MENU2": "Menu nivel 2",
    "NUMERA": "Numeradores",
    "PGMCAT": "Programas de catalogo",
    "PROVIN": "Provincias",
    "PUNPEN": "Puntos pendientes",
    "REPRESE": "Representantes",
    "SSUBFAM": "Subfamilias nivel 3",
    "SUBFAM": "Subfamilias",
    "TABPREC": "Tablas de precio",
    "TARJET": "Tarjetas",
    "TIPIVA": "Tipos de IVA",
    "TIPVEN": "Tipos de venta",
    "ZONAS": "Zonas",
}

CLIENT_TABLES: dict[str, str] = {
    "CLIEN": "Ficha principal de cliente",
    "CLIENI": "Informacion adicional de cliente",
    "CLIFAM": "Descuentos por familia",
    "CLIART": "Precios especiales por articulo",
    "CLIACT": "Actividad por seccion",
    "CLIAGR": "Visibilidad por agrupacion",
    "CLITAR": "Tarjetas/fidelizacion del cliente",
    "CLIDIR3": "Direcciones adicionales/roles Dir3",
}

ARTICLE_TABLES: dict[str, str] = {
    "ARTICUL": "Ficha principal de articulo",
    "ARTICULC": "Codigos de barras",
    "ARTICULE": "Existencias por centro",
    "ARTICULP": "Ficha de compra por proveedor",
    "ARTICULI": "Informacion adicional",
    "ARTICULA": "Articulos alternativos",
    "ARTICULB": "Componentes/blister",
    "ELIMA": "Registro de anulaciones/bajas de articulo",
}

ARTICLE_WRITABLE_TABLES = {"ARTICUL", "ARTICULC", "ARTICULE", "ARTICULP", "ARTICULI", "ARTICULA", "ARTICULB"}

ARTICLE_ADDITIONAL_CODES = {
    "IMAGE": "Imagen",
    "INFO": "Informacion tecnica",
    "PWEB": "Publicar web",
    "NODTO": "No aplicar descuentos",
    "UBICA": "Ubicacion",
    "UBIC1": "Ubicacion secundaria 1",
    "UBIC2": "Ubicacion secundaria 2",
    "SSUBF": "Tercer nivel de familia",
    "CANON": "Canon",
    "REGAL": "Articulo regalo",
}

CLIENT_ADDITIONAL_CODES = {
    "ALB": "Copias de albaran",
    "IMPMI": "Importe minimo",
    "SERIE": "Serie por defecto",
    "CANPM": "Cantidad/precio minimo",
    "ACECR": "Acepta creditos",
    "ACEOF": "Acepta ofertas",
    "TIVEN": "Formato albaran",
    "TIVFA": "Formato factura",
    "TIVFC": "Formato factura contado",
    "COBAL": "Cobrar albaranes",
    "PUBLI": "Enviar publicidad",
    "FEMAI": "Enviar factura por email",
    "IBAN": "IBAN",
    "DOMIC": "Domicilio adicional",
    "DOMEN": "Domicilio envio adicional",
    "WEB": "Web",
    "PAIS": "Pais",
    "ADV": "Advertencia",
    "OBS": "Observaciones",
    "VISIT": "Ultima visita",
    "PERIO": "Periodicidad visitas",
    "EMAIP": "Email publicidad/pedido",
    "EMAIA": "Email albaran",
    "EMAIF": "Email factura",
}

PROVIDER_TABLES: dict[str, str] = {
    "PROVEE": "Ficha principal de proveedor",
    "PROVEEI": "Informacion adicional de proveedor",
    "ARTICULP": "Articulos/referencias de proveedor",
    "CABDOCM": "Documentos de compra del proveedor",
    "CABORC": "Pedidos/presupuestos de compra del proveedor",
    "ELIMP": "Registro de anulaciones/bajas de proveedor",
}

PROVIDER_WRITABLE_TABLES = {"PROVEE", "PROVEEI", "ARTICULP"}

PROVIDER_ADDITIONAL_CODES = {
    "EMAIL": "Email adicional",
    "WEB": "Web",
    "FILE": "Fichero",
    "TEXDE": "Texto destino",
    "TELEF": "Telefono adicional",
    "ACEDE": "Acepta devoluciones",
    "CONTA": "Contacto",
    "OBS": "Observaciones",
    "DOMIC": "Domicilio adicional",
    "SERIE": "Serie por defecto",
    "DESAC": "Descuento activo",
    "CCC": "Cuenta bancaria",
    "CENCO": "Centro contable",
    "CLASI": "Clasificacion",
    "CLNEX": "Clave Nex",
    "DIAFI": "Dia fijo",
    "DTO": "Descuento",
    "DTOPP": "Descuento pronto pago",
    "ENVP": "Envio proveedor",
    "NCLI": "Numero cliente en proveedor",
    "NONEX": "No enlazar",
    "PEDMI": "Pedido minimo",
    "PLAZO": "Plazo",
    "REIVA": "Recargo IVA",
    "SEROC": "Serie orden compra",
}

OFFER_TABLES: dict[str, str] = {
    "OFERTAS": "Cabecera de oferta",
    "DETOFER": "Articulos de la oferta",
}

PURCHASE_ORDER_TABLES: dict[str, str] = {
    "CABORC": "Cabecera de orden de compra",
    "DETORC": "Lineas de orden de compra",
    "ELIMOC": "Registro de anulaciones/bajas de orden de compra",
}

PURCHASE_ORDER_STATUS: dict[str, str] = {
    "A": "A",
    "ABIERTO": "A",
    "ABIERTA": "A",
    "ABIERTOS": "A",
    "ABIERTAS": "A",
    "P": "P",
    "PENDIENTE": "P",
    "PENDIENTES": "P",
    "C": "C",
    "CERRADO": "C",
    "CERRADA": "C",
    "CERRADOS": "C",
    "CERRADAS": "C",
}

PURCHASE_ORDER_STATUS_LABELS: dict[str, str] = {
    "A": "Abierto",
    "P": "Pendiente recibir",
    "C": "Cerrado",
}

WAREHOUSE_ENTRY_TABLES: dict[str, str] = {
    "CABDOCM": "Cabecera de entrada de almacen / documento de compra",
    "DETMOVM": "Lineas de entrada de almacen",
    "DETMOVMC": "Control auxiliar de lineas de entrada verificadas",
    "ARTICULE": "Existencias de articulo por centro",
    "STOCKS": "Historico diario de stock",
}

WAREHOUSE_ENTRY_STATUS: dict[str, str] = {
    "P": "P",
    "PENDIENTE": "P",
    "PENDIENTE_FACTURAR": "P",
    "PENDIENTES_FACTURAR": "P",
    "F": "F",
    "FACTURADA": "F",
    "FACTURADO": "F",
    "PENDIENTE_CONTABILIZAR": "F",
    "PENDIENTES_CONTABILIZAR": "F",
    "C": "C",
    "CONTABILIZADA": "C",
    "CONTABILIZADO": "C",
    "H": "H",
    "HISTORICO": "H",
}

WAREHOUSE_ENTRY_STATUS_LABELS: dict[str, str] = {
    "P": "Pendiente de facturar",
    "F": "Pendiente de contabilizar",
    "C": "Contabilizada",
    "H": "Historico",
}

REGULARIZATION_TABLES: dict[str, str] = {
    "CABDOCR": "Cabecera de regularizaciones y trasvases",
    "DETMOVR": "Lineas de regularizacion/trasvase",
    "RECUENTO": "Recuentos de inventario pendientes de aplicar",
    "ARTICULE": "Existencias actuales por centro",
    "STOCKS": "Stock registrado por fecha",
}

REGULARIZATION_TYPE_LABELS: dict[str, str] = {
    "R": "Regularizacion",
    "S": "Trasvase salida",
    "E": "Trasvase entrada",
}

SALES_TABLES: dict[str, str] = {
    "CABDOCV": "Cabecera de documentos de venta",
    "DETMOV": "Lineas de documentos de venta",
    "NUMERA": "Numeradores de documentos",
    "CLIEN": "Ficha de cliente usada para cabecera y tarifas",
    "CLIART": "Precios/codigos especiales de cliente",
    "CLIFAM": "Descuentos por familia de cliente",
    "CLIACT": "Tarifa por actividad/seccion",
    "OFERTAS": "Cabeceras de ofertas de venta",
    "DETOFER": "Lineas de ofertas de venta",
}

SALES_DOCUMENT_TYPES: dict[str, str] = {
    "P": "Pedido",
    "R": "Presupuesto",
    "S": "Pedido/presupuesto cerrado",
    "A": "Albaran",
    "F": "Factura",
    "T": "Ticket",
    "C": "Credito/abono",
}

SALES_STATUS_LABELS: dict[str, str] = {
    "P": "Pendiente",
    "S": "Servido",
    "C": "Cerrado",
    "F": "Facturado",
    "H": "Historico",
}

CARTERA_TABLES: dict[str, str] = {
    "CABDOCVE": "Efectos/vencimientos de documentos de venta",
    "CABDOCV": "Cabecera del documento origen",
    "CLIEN": "Ficha de cliente",
    "FORPAG": "Formas de pago y codigo de aceptacion",
    "TIPVEN": "Tipos de documento/venta",
}

CARTERA_EFFECT_TYPE_LABELS: dict[str, str] = {
    "C": "Cobro contado",
    "D": "Domiciliacion/recibo",
    "G": "Giro",
    "N": "Giro no negociado",
    "O": "Reembolso",
    "P": "Pagare",
    "R": "Recibo",
    "T": "Transferencia",
}

SALES_COLLECTION_LABELS: dict[str, str] = {
    "E": "Efectivo",
    "T": "Tarjeta",
    "C": "Cheque",
    "F": "Transferencia",
    "W": "Web",
    "O": "Otros",
}

DASHBOARD_SALE_DOCUMENTS = ("F", "T", "A", "C")


PUBLIC_TOOL_DEFINITIONS: dict[str, dict[str, Any]] = {
    "sistema_estado": _tool(
        "sistema_estado",
        "Devuelve version, DSN y contexto por defecto del servidor Kofedas MCP.",
        {},
    ),
    "empresa_listar": _tool(
        "empresa_listar",
        "Configuracion. Lista empresas desde EMPRES.",
        {"limite": _int_schema("Maximo de filas.")},
    ),
    "empresa_obtener": _tool(
        "empresa_obtener",
        "Configuracion. Obtiene una empresa desde EMPRES.",
        {"empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1.")},
    ),
    "empresa_guardar": _tool(
        "empresa_guardar",
        "Configuracion. ESCRITURA. Crea o actualiza una fila de EMPRES. Las claves no se modifican.",
        {
            "empresa": _int_schema("EMP_NUMEMP. Por defecto KOFEDAS_EMPRESA o 1."),
            "datos": _data_schema("Campos EMP_* a guardar, con nombres de base de datos o sin prefijo."),
        },
        ["datos"],
    ),
    "centro_listar": _tool(
        "centro_listar",
        "Lista centros de la empresa desde CENTROS, segun CENTROS_UDM.pas.",
        {"empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1.")},
    ),
    "centro_obtener": _tool(
        "centro_obtener",
        "Configuracion. Obtiene un centro desde CENTROS.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CEN_CODCEN. Por defecto KOFEDAS_CENTRO o 0."),
        },
    ),
    "centro_guardar": _tool(
        "centro_guardar",
        "Configuracion. ESCRITURA. Crea o actualiza una fila de CENTROS. Las claves no se modifican.",
        {
            "empresa": _int_schema("CEN_NUMEMP. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CEN_CODCEN. Por defecto KOFEDAS_CENTRO o 0."),
            "datos": _data_schema("Campos CEN_* a guardar, con nombres de base de datos o sin prefijo."),
        },
        ["datos"],
    ),
    "usuario_listar": _tool(
        "usuario_listar",
        "Configuracion. Lista usuarios desde USUAR sin devolver contraseñas.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro opcional por centro."),
            "texto": _string_schema("Filtro opcional por usuario, nombre o grupo."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "usuario_obtener": _tool(
        "usuario_obtener",
        "Configuracion. Obtiene un usuario desde USUAR sin devolver contraseña.",
        {
            "usuario": _string_schema("USU_NOMUSU."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("USU_CODCEN. Por defecto KOFEDAS_CENTRO o 0."),
        },
        ["usuario"],
    ),
    "usuario_guardar": _tool(
        "usuario_guardar",
        "Configuracion. ESCRITURA. Crea o actualiza una fila de USUAR. Admite USU_PASSWORD para cambiar contraseña, pero nunca la devuelve.",
        {
            "usuario": _string_schema("USU_NOMUSU."),
            "empresa": _int_schema("USU_NUMEMP. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("USU_CODCEN. Por defecto KOFEDAS_CENTRO o 0."),
            "datos": _data_schema("Campos USU_* a guardar, con nombres de base de datos o sin prefijo."),
        },
        ["usuario", "datos"],
    ),
    "grupo_usuario_listar": _tool(
        "grupo_usuario_listar",
        "Configuracion. Lista grupos de usuario desde GRUPUSU.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "texto": _string_schema("Filtro opcional por id de grupo, menu o nivel."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "grupo_usuario_obtener": _tool(
        "grupo_usuario_obtener",
        "Configuracion. Obtiene un grupo de usuario desde GRUPUSU.",
        {
            "grupo": _string_schema("GRU_IDGRUPO."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
        },
        ["grupo"],
    ),
    "grupo_usuario_guardar": _tool(
        "grupo_usuario_guardar",
        "Configuracion. ESCRITURA. Crea o actualiza una fila de GRUPUSU. Las claves no se modifican.",
        {
            "grupo": _string_schema("GRU_IDGRUPO."),
            "empresa": _int_schema("GRU_NUMEMP. Por defecto KOFEDAS_EMPRESA o 1."),
            "datos": _data_schema("Campos GRU_* a guardar, con nombres de base de datos o sin prefijo."),
        },
        ["grupo", "datos"],
    ),
    "parametro_listar": _tool(
        "parametro_listar",
        "Configuracion. Lista parametros desde PARAMETROS.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "texto": _string_schema("Filtro opcional por codigo, valor o descripcion."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "parametro_obtener": _tool(
        "parametro_obtener",
        "Configuracion. Obtiene un parametro desde PARAMETROS.",
        {
            "codigo": _string_schema("PAR_CODIGO."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
        },
        ["codigo"],
    ),
    "parametro_guardar": _tool(
        "parametro_guardar",
        "Configuracion. ESCRITURA. Crea o actualiza una fila de PARAMETROS.",
        {
            "codigo": _string_schema("PAR_CODIGO."),
            "empresa": _int_schema("PAR_NUMEMP. Por defecto KOFEDAS_EMPRESA o 1."),
            "valor": _string_schema("PAR_VALOR. Alternativa comoda a datos.PAR_VALOR."),
            "descripcion": _string_schema("PAR_DESCRI. Alternativa comoda a datos.PAR_DESCRI."),
            "datos": _data_schema("Campos PAR_* a guardar."),
        },
        ["codigo"],
    ),
    "auxiliar_tablas": _tool(
        "auxiliar_tablas",
        "Auxiliares. Lista las tablas auxiliares disponibles, sus columnas y claves primarias.",
        {},
    ),
    "auxiliar_listar": _tool(
        "auxiliar_listar",
        "Auxiliares. Lista filas de una tabla auxiliar permitida con filtros opcionales.",
        {
            "tabla": _string_schema("Nombre de tabla auxiliar permitida."),
            "empresa": _int_schema("Filtro por empresa si la tabla tiene columna *_NUMEMP."),
            "filtros": _data_schema("Filtros exactos por columna."),
            "texto": _string_schema("Busqueda textual opcional en columnas de texto."),
            "limite": _int_schema("Maximo de filas."),
        },
        ["tabla"],
    ),
    "auxiliar_obtener": _tool(
        "auxiliar_obtener",
        "Auxiliares. Obtiene una fila de una tabla auxiliar por sus claves primarias.",
        {
            "tabla": _string_schema("Nombre de tabla auxiliar permitida."),
            "claves": _data_schema("Valores de clave primaria por columna."),
        },
        ["tabla", "claves"],
    ),
    "auxiliar_guardar": _tool(
        "auxiliar_guardar",
        "Auxiliares. ESCRITURA. Crea o actualiza una fila de una tabla auxiliar permitida usando su clave primaria real.",
        {
            "tabla": _string_schema("Nombre de tabla auxiliar permitida."),
            "claves": _data_schema("Valores de clave primaria por columna."),
            "datos": _data_schema("Campos a guardar. Puede incluir tambien las claves."),
        },
        ["tabla", "datos"],
    ),
    "familia_listar": _tool(
        "familia_listar",
        "Lista familias, subfamilias o tercer nivel desde FAMILI/SUBFAM/SSUBFAM.",
        {
            "tipo": {
                "type": "string",
                "enum": ["familia", "subfamilia", "ssubfamilia"],
                "description": "Nivel a listar.",
            },
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "familia": _int_schema("Filtro de familia para subfamilias/ssubfamilias."),
            "subfamilia": _int_schema("Filtro de subfamilia para ssubfamilias."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "articulo_buscar": _tool(
        "articulo_buscar",
        "Busca articulos por codigo, descripcion, EAN o referencia de proveedor.",
        {
            "texto": _string_schema("Texto a buscar."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "incluir_baja": {"type": "boolean", "description": "Incluye articulos con ART_FEBAJA informada."},
            "limite": _int_schema("Maximo de filas."),
        },
        ["texto"],
    ),
    "articulo_obtener": _tool(
        "articulo_obtener",
        "Obtiene ficha basica de ARTICUL con codigos de barras, stock por centro y proveedores.",
        {
            "articulo": _string_schema("Codigo del articulo."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
        },
        ["articulo"],
    ),
    "stock_consultar": _tool(
        "stock_consultar",
        "Consulta existencias actuales desde ARTICULE y, opcionalmente, historico STOCKS.",
        {
            "articulo": _string_schema("Codigo o fragmento de articulo."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "incluir_historico": {"type": "boolean", "description": "Incluye ultimas filas de STOCKS."},
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "articulo_tablas": _tool(
        "articulo_tablas",
        "Articulos. Lista tablas relacionadas, claves reales y columnas disponibles.",
        {},
    ),
    "articulo_relacion_listar": _tool(
        "articulo_relacion_listar",
        "Articulos. Lista registros de una tabla relacionada filtrando por articulo, empresa, texto o columnas.",
        {
            "tabla": _string_schema("Tabla relacionada: ARTICUL, ARTICULC, ARTICULE, ARTICULP, ARTICULI, ARTICULA, ARTICULB o ELIMA."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "articulo": _string_schema("Codigo de articulo para aplicar al campo *_CODART."),
            "texto": _string_schema("Filtro de texto sobre columnas alfanumericas."),
            "filtros": _data_schema("Filtros exactos por columna real."),
            "limite": _int_schema("Maximo de filas."),
        },
        ["tabla"],
    ),
    "articulo_relacion_obtener": _tool(
        "articulo_relacion_obtener",
        "Articulos. Obtiene un registro relacionado usando la clave primaria real de la tabla.",
        {
            "tabla": _string_schema("Tabla relacionada."),
            "claves": _data_schema("Valores de clave primaria por columna real."),
        },
        ["tabla", "claves"],
    ),
    "articulo_relacion_guardar": _tool(
        "articulo_relacion_guardar",
        "Articulos. ESCRITURA. Crea o actualiza una tabla relacionada escribible usando su clave primaria real.",
        {
            "tabla": _string_schema("Tabla relacionada escribible."),
            "claves": _data_schema("Claves primarias por columna real."),
            "datos": _data_schema("Columnas a insertar/actualizar."),
        },
        ["tabla", "datos"],
    ),
    "articulo_completo": _tool(
        "articulo_completo",
        "Articulos. Devuelve ficha ARTICUL y tablas relacionadas principales.",
        {
            "articulo": _string_schema("Codigo del articulo."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "limite_detalle": _int_schema("Maximo de filas por tabla relacionada."),
        },
        ["articulo"],
    ),
    "articulo_alta_preparar": _tool(
        "articulo_alta_preparar",
        "Articulos. Calcula defaults para dar de alta un articulo sin escribir.",
        {
            "articulo": _string_schema("Codigo de articulo."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro para usuario de modificacion. Por defecto KOFEDAS_CENTRO o 0."),
            "datos": _data_schema("Campos ART_* o alias de articulo para superponer a los defaults."),
        },
        ["articulo"],
    ),
    "articulo_alta": _tool(
        "articulo_alta",
        "Articulos. ESCRITURA. Da de alta un articulo con ficha de compra, codigos de barras, stock inicial e informacion adicional opcional.",
        {
            "articulo": _string_schema("Codigo de articulo."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro para stock inicial y usuario de modificacion."),
            "datos": _data_schema("Campos del articulo. Requiere descripcion o ART_DESCRI."),
            "compra": _data_schema("Datos ARTP_* o alias de ficha de compra."),
            "codigos_barras": {"type": "array", "description": "Codigos de barras o objetos {codigo,cantidad}.", "items": {"type": ["string", "object"]}},
            "stock": _data_schema("Stock inicial opcional: existencias, minimo, maximo, centro."),
            "informacion_adicional": _data_schema("Mapa CODINF -> valor para crear ARTICULI."),
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["articulo", "datos"],
    ),
    "articulo_tarifa_excel_previsualizar": _tool(
        "articulo_tarifa_excel_previsualizar",
        "Articulos. Lee una tarifa .xlsx y devuelve el plan de altas/actualizaciones sin escribir.",
        {
            "archivo": _string_schema("Ruta local del .xlsx."),
            "hoja": _string_schema("Nombre de hoja opcional; por defecto la activa."),
            "fila_cabecera": _int_schema("Fila de cabecera. Por defecto 1."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "proveedor": _int_schema("Proveedor por defecto si no viene en la hoja."),
            "seccion": _string_schema("Seccion por defecto para altas/generacion de codigo."),
            "mapeo": _data_schema("Mapa canonico -> cabecera Excel. Ej: articulo->Codigo, descripcion->Descripcion."),
            "generar_codigos": {"type": "boolean", "description": "Genera ART_CODART si no viene en Excel usando seccion/proveedor/contador."},
            "numero_inicial": _int_schema("Contador inicial para generar codigos. Por defecto 1."),
            "digitos": _int_schema("Longitud total del codigo generado. Por defecto 13."),
            "limite": _int_schema("Maximo de filas a analizar."),
        },
        ["archivo"],
    ),
    "articulo_tarifa_excel_importar": _tool(
        "articulo_tarifa_excel_importar",
        "Articulos. ESCRITURA. Importa tarifa .xlsx creando/actualizando ARTICUL, ARTICULP, ARTICULC y ARTICULE.",
        {
            "archivo": _string_schema("Ruta local del .xlsx."),
            "hoja": _string_schema("Nombre de hoja opcional; por defecto la activa."),
            "fila_cabecera": _int_schema("Fila de cabecera. Por defecto 1."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "proveedor": _int_schema("Proveedor por defecto si no viene en la hoja."),
            "seccion": _string_schema("Seccion por defecto para altas/generacion de codigo."),
            "mapeo": _data_schema("Mapa canonico -> cabecera Excel."),
            "generar_codigos": {"type": "boolean", "description": "Genera ART_CODART si no viene en Excel usando seccion/proveedor/contador."},
            "numero_inicial": _int_schema("Contador inicial para generar codigos. Por defecto 1."),
            "digitos": _int_schema("Longitud total del codigo generado. Por defecto 13."),
            "crear_articulos": {"type": "boolean", "description": "Permite altas de ARTICUL. Por defecto true."},
            "actualizar_articulos": {"type": "boolean", "description": "Permite actualizar ARTICUL existente. Por defecto true."},
            "crear_barras": {"type": "boolean", "description": "Permite crear codigos de barras. Por defecto true."},
            "actualizar_barras": {"type": "boolean", "description": "Permite reasignar codigos de barras existentes. Por defecto false."},
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
            "limite": _int_schema("Maximo de filas a procesar."),
        },
        ["archivo"],
    ),
    "cliente_buscar": _tool(
        "cliente_buscar",
        "Busca clientes por codigo, nombre, razon social, CIF, telefono o email.",
        {
            "texto": _string_schema("Texto a buscar."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "incluir_baja": {"type": "boolean", "description": "Incluye clientes con CLI_FEBAJA informada."},
            "limite": _int_schema("Maximo de filas."),
        },
        ["texto"],
    ),
    "cliente_obtener": _tool(
        "cliente_obtener",
        "Obtiene ficha de CLIEN con informacion adicional CLIENI.",
        {
            "cliente": _int_schema("Codigo de cliente."),
            "subcliente": _int_schema("Subcliente. Por defecto 0."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
        },
        ["cliente"],
    ),
    "cliente_tablas": _tool(
        "cliente_tablas",
        "Clientes. Lista las tablas relacionadas con clientes, columnas y claves primarias.",
        {},
    ),
    "cliente_relacion_listar": _tool(
        "cliente_relacion_listar",
        "Clientes. Lista filas de una tabla relacionada con clientes.",
        {
            "tabla": _string_schema("Tabla relacionada permitida: CLIEN, CLIENI, CLIFAM, CLIART, CLIACT, CLIAGR, CLITAR o CLIDIR3."),
            "cliente": _int_schema("Codigo de cliente."),
            "subcliente": _int_schema("Subcliente. Por defecto 0 cuando la tabla lo tiene."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "filtros": _data_schema("Filtros exactos adicionales por columna."),
            "texto": _string_schema("Busqueda textual opcional en columnas de texto."),
            "limite": _int_schema("Maximo de filas."),
        },
        ["tabla"],
    ),
    "cliente_relacion_obtener": _tool(
        "cliente_relacion_obtener",
        "Clientes. Obtiene una fila relacionada por clave primaria real.",
        {
            "tabla": _string_schema("Tabla relacionada permitida."),
            "claves": _data_schema("Valores de clave primaria por columna."),
        },
        ["tabla", "claves"],
    ),
    "cliente_relacion_guardar": _tool(
        "cliente_relacion_guardar",
        "Clientes. ESCRITURA. Crea o actualiza una fila relacionada usando su clave primaria real.",
        {
            "tabla": _string_schema("Tabla relacionada permitida."),
            "claves": _data_schema("Valores de clave primaria por columna."),
            "datos": _data_schema("Campos a guardar. Puede incluir tambien claves."),
        },
        ["tabla", "datos"],
    ),
    "cliente_completo": _tool(
        "cliente_completo",
        "Clientes. Obtiene CLIEN y todas sus tablas relacionadas en una sola respuesta.",
        {
            "cliente": _int_schema("Codigo de cliente."),
            "subcliente": _int_schema("Subcliente. Por defecto 0."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "limite_detalle": _int_schema("Maximo de filas por tabla relacionada."),
        },
        ["cliente"],
    ),
    "cliente_alta_preparar": _tool(
        "cliente_alta_preparar",
        "Clientes. Calcula codigo y defaults para dar de alta un cliente sin escribir.",
        {
            "cliente": _int_schema("Codigo deseado. 0/omitido calcula el siguiente."),
            "subcliente": _int_schema("Subcliente deseado. Para cliente 99999, 0 calcula siguiente subcliente."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro para rangos de numeracion. Por defecto KOFEDAS_CENTRO o 0."),
            "datos": _data_schema("Campos CLI_* o alias de cliente para superponer a los defaults."),
        },
    ),
    "cliente_alta": _tool(
        "cliente_alta",
        "Clientes. ESCRITURA. Da de alta un cliente replicando numeracion y defaults principales del mantenimiento Delphi.",
        {
            "cliente": _int_schema("Codigo deseado. 0/omitido calcula el siguiente."),
            "subcliente": _int_schema("Subcliente deseado. Para cliente 99999, 0 calcula siguiente subcliente."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro para rangos de numeracion. Por defecto KOFEDAS_CENTRO o 0."),
            "datos": _data_schema("Campos del cliente. Requiere nombre o CLI_NOMCLI; razon_social usa nombre si se omite."),
            "informacion_adicional": _data_schema("Mapa CODINF -> valor para crear CLIENI."),
            "heredar_actividades": {"type": "boolean", "description": "Si es subcliente, copia CLIACT del subcliente 0."},
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["datos"],
    ),
    "proveedor_buscar": _tool(
        "proveedor_buscar",
        "Busca proveedores por codigo, nombre fiscal, nombre comercial, abreviado, CIF o email.",
        {
            "texto": _string_schema("Texto a buscar."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "incluir_baja": {"type": "boolean", "description": "Incluye proveedores con PRO_FEBAJA informada."},
            "limite": _int_schema("Maximo de filas."),
        },
        ["texto"],
    ),
    "proveedor_obtener": _tool(
        "proveedor_obtener",
        "Obtiene ficha de PROVEE con informacion adicional PROVEEI.",
        {
            "proveedor": _int_schema("Codigo de proveedor."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
        },
        ["proveedor"],
    ),
    "proveedor_articulos_listar": _tool(
        "proveedor_articulos_listar",
        "Lista referencias ARTICULP de un proveedor con ficha basica del articulo.",
        {
            "proveedor": _int_schema("Codigo de proveedor."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "texto": _string_schema("Filtro opcional por articulo, referencia o descripcion."),
            "limite": _int_schema("Maximo de filas."),
        },
        ["proveedor"],
    ),
    "proveedor_tablas": _tool(
        "proveedor_tablas",
        "Proveedores. Lista tablas relacionadas, claves reales y columnas disponibles.",
        {},
    ),
    "proveedor_relacion_listar": _tool(
        "proveedor_relacion_listar",
        "Proveedores. Lista registros de una tabla relacionada filtrando por proveedor, empresa, texto o columnas.",
        {
            "tabla": _string_schema("Tabla relacionada: PROVEE, PROVEEI, ARTICULP, CABDOCM, CABORC o ELIMP."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "proveedor": _int_schema("Codigo de proveedor para aplicar al campo *_CODPRO."),
            "texto": _string_schema("Filtro de texto sobre columnas alfanumericas."),
            "filtros": _data_schema("Filtros exactos por columna real."),
            "limite": _int_schema("Maximo de filas."),
        },
        ["tabla"],
    ),
    "proveedor_relacion_obtener": _tool(
        "proveedor_relacion_obtener",
        "Proveedores. Obtiene un registro relacionado usando la clave primaria real de la tabla.",
        {
            "tabla": _string_schema("Tabla relacionada."),
            "claves": _data_schema("Valores de clave primaria por columna real."),
        },
        ["tabla", "claves"],
    ),
    "proveedor_relacion_guardar": _tool(
        "proveedor_relacion_guardar",
        "Proveedores. ESCRITURA. Crea o actualiza PROVEE, PROVEEI o ARTICULP usando su clave primaria real.",
        {
            "tabla": _string_schema("Tabla relacionada escribible: PROVEE, PROVEEI o ARTICULP."),
            "claves": _data_schema("Claves primarias por columna real."),
            "datos": _data_schema("Columnas a insertar/actualizar."),
        },
        ["tabla", "datos"],
    ),
    "proveedor_completo": _tool(
        "proveedor_completo",
        "Proveedores. Devuelve ficha PROVEE y tablas relacionadas principales.",
        {
            "proveedor": _int_schema("Codigo de proveedor."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "limite_detalle": _int_schema("Maximo de filas por tabla relacionada."),
        },
        ["proveedor"],
    ),
    "proveedor_alta_preparar": _tool(
        "proveedor_alta_preparar",
        "Proveedores. Calcula codigo y defaults para dar de alta un proveedor sin escribir.",
        {
            "proveedor": _int_schema("Codigo deseado. 0/omitido calcula el siguiente."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro para usuario de modificacion. Por defecto KOFEDAS_CENTRO o 0."),
            "datos": _data_schema("Campos PRO_* o alias de proveedor para superponer a los defaults."),
        },
    ),
    "proveedor_alta": _tool(
        "proveedor_alta",
        "Proveedores. ESCRITURA. Da de alta un proveedor replicando numeracion y defaults principales del mantenimiento Delphi.",
        {
            "proveedor": _int_schema("Codigo deseado. 0/omitido calcula el siguiente."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro para usuario de modificacion. Por defecto KOFEDAS_CENTRO o 0."),
            "datos": _data_schema("Campos del proveedor. Requiere nombre, PRO_NOMCOR o PRO_NOMFIS."),
            "informacion_adicional": _data_schema("Mapa CODINF -> valor para crear PROVEEI."),
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["datos"],
    ),
    "oferta_tablas": _tool(
        "oferta_tablas",
        "Ofertas. Lista tablas relacionadas, claves reales y columnas disponibles.",
        {},
    ),
    "oferta_listar": _tool(
        "oferta_listar",
        "Ofertas. Lista cabeceras de OFERTAS por fechas, proveedor, texto o articulo incluido.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "proveedor": _int_schema("Filtro por proveedor."),
            "articulo": _string_schema("Filtro por articulo incluido en DETOFER."),
            "desde": _string_schema("Fecha minima de inicio/fecha oferta en formato YYYY-MM-DD."),
            "hasta": _string_schema("Fecha maxima de fin/fecha oferta en formato YYYY-MM-DD."),
            "solo_vigentes": {"type": "boolean", "description": "Si true, solo ofertas vigentes a fecha_consulta o hoy."},
            "fecha_consulta": _string_schema("Fecha para vigencia. Por defecto hoy."),
            "texto": _string_schema("Filtro por descripcion."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "oferta_obtener": _tool(
        "oferta_obtener",
        "Ofertas. Obtiene cabecera OFERTAS y lineas DETOFER.",
        {
            "ejercicio": _int_schema("OFE_EJERCI."),
            "oferta": _int_schema("OFE_NUMOFE."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
        },
        ["ejercicio", "oferta"],
    ),
    "oferta_articulos_listar": _tool(
        "oferta_articulos_listar",
        "Ofertas. Lista articulos en oferta, opcionalmente solo vigentes, con datos de articulo y proveedor.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "oferta": _int_schema("Filtro por numero de oferta."),
            "proveedor": _int_schema("Filtro por proveedor."),
            "articulo": _string_schema("Filtro por codigo de articulo."),
            "texto": _string_schema("Filtro por articulo o descripcion."),
            "solo_vigentes": {"type": "boolean", "description": "Si true, filtra DOF_FECINI/DOF_FECFIN por fecha_consulta u hoy."},
            "fecha_consulta": _string_schema("Fecha para vigencia. Por defecto hoy."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "oferta_alta_preparar": _tool(
        "oferta_alta_preparar",
        "Ofertas. Calcula numero, cabecera y lineas para dar de alta una oferta sin escribir.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro para usuario de modificacion. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año actual."),
            "oferta": _int_schema("Numero de oferta. 0/omitido calcula el siguiente."),
            "proveedor": _int_schema("Proveedor de la oferta."),
            "descripcion": _string_schema("Descripcion de la oferta."),
            "fecha": _string_schema("Fecha de oferta. Por defecto hoy."),
            "fecha_inicio": _string_schema("Inicio de vigencia. Por defecto fecha."),
            "fecha_fin": _string_schema("Fin de vigencia. Por defecto fecha_inicio."),
            "gastos": {"type": "number", "description": "Gastos de la oferta. Por defecto 0."},
            "moneda": _string_schema("Moneda. Por defecto E."),
            "tipo": _string_schema("Tipo de oferta. Por defecto T."),
            "articulos": {"type": "array", "description": "Lineas: string codigo o objeto {articulo,codigo_barras,referencia,pvp,precio,precos,dto1,dto2}.", "items": {"type": ["string", "object"]}},
        },
        ["proveedor", "descripcion", "articulos"],
    ),
    "oferta_alta": _tool(
        "oferta_alta",
        "Ofertas. ESCRITURA. Da de alta OFERTAS y DETOFER a partir de una lista de articulos.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro para usuario de modificacion. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año actual."),
            "oferta": _int_schema("Numero de oferta. 0/omitido calcula el siguiente."),
            "proveedor": _int_schema("Proveedor de la oferta."),
            "descripcion": _string_schema("Descripcion de la oferta."),
            "fecha": _string_schema("Fecha de oferta. Por defecto hoy."),
            "fecha_inicio": _string_schema("Inicio de vigencia. Por defecto fecha."),
            "fecha_fin": _string_schema("Fin de vigencia. Por defecto fecha_inicio."),
            "gastos": {"type": "number", "description": "Gastos de la oferta. Por defecto 0."},
            "moneda": _string_schema("Moneda. Por defecto E."),
            "tipo": _string_schema("Tipo de oferta. Por defecto T."),
            "articulos": {"type": "array", "description": "Lineas: string codigo o objeto {articulo,codigo_barras,referencia,pvp,precio,precos,dto1,dto2}.", "items": {"type": ["string", "object"]}},
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["proveedor", "descripcion", "articulos"],
    ),
    "orden_compra_tablas": _tool(
        "orden_compra_tablas",
        "Ordenes de Compra. Describe las tablas CABORC, DETORC y ELIMOC.",
        {},
    ),
    "orden_compra_listar": _tool(
        "orden_compra_listar",
        "Ordenes de Compra. Lista pedidos por estado: abierto, pendiente o cerrado, con filtros de proveedor, fechas, serie, numero o texto.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro. Por defecto sin filtrar."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "serie": _string_schema("Filtro por serie."),
            "numero": _int_schema("Filtro por COC_NUMDOC."),
            "estado": _string_schema("abierto/pendiente/cerrado o A/P/C. Omitido devuelve todos."),
            "proveedor": _int_schema("Filtro por proveedor."),
            "desde": _string_schema("Fecha minima de pedido en formato YYYY-MM-DD."),
            "hasta": _string_schema("Fecha maxima de pedido en formato YYYY-MM-DD."),
            "texto": _string_schema("Filtro por nombre proveedor, observaciones o articulo incluido."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "orden_compra_obtener": _tool(
        "orden_compra_obtener",
        "Ordenes de Compra. Obtiene cabecera CABORC y lineas DETORC de un pedido.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("COC_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("COC_EJERCI."),
            "serie": _string_schema("COC_SERIE."),
            "numero": _int_schema("COC_NUMDOC."),
        },
        ["ejercicio", "serie", "numero"],
    ),
    "orden_compra_lineas_listar": _tool(
        "orden_compra_lineas_listar",
        "Ordenes de Compra. Lista lineas DETORC con filtros por pedido, estado, proveedor, articulo o texto.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "serie": _string_schema("Filtro por serie."),
            "numero": _int_schema("Filtro por COC/DOC_NUMDOC."),
            "estado": _string_schema("Estado de linea: abierto/pendiente/cerrado o A/P/C."),
            "proveedor": _int_schema("Filtro por proveedor de cabecera."),
            "articulo": _string_schema("Filtro por codigo de articulo."),
            "texto": _string_schema("Filtro por articulo, referencia o descripcion."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "orden_compra_alta_preparar": _tool(
        "orden_compra_alta_preparar",
        "Ordenes de Compra. Calcula numero, cabecera y lineas para dar de alta un pedido sin escribir.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año actual."),
            "serie": _string_schema("Serie. Por defecto vacia."),
            "numero": _int_schema("Numero de pedido. 0/omitido calcula el siguiente."),
            "proveedor": _int_schema("Proveedor del pedido."),
            "fecha": _string_schema("Fecha del pedido. Por defecto hoy."),
            "fecha_envio": _string_schema("Fecha de envio al proveedor."),
            "fecha_entrega": _string_schema("Fecha prevista de entrega."),
            "forma_pago": _int_schema("Forma de pago. Por defecto la del proveedor."),
            "representante": _int_schema("Representante. Por defecto 0."),
            "moneda": _string_schema("Moneda. Por defecto proveedor o E."),
            "observaciones": _string_schema("Observaciones de cabecera."),
            "articulos": {"type": "array", "description": "Lineas: string codigo o objeto {articulo,codigo_barras,referencia,cantidad,precio,dto1..dto6,unidad,descripcion,observaciones}.", "items": {"type": ["string", "object"]}},
        },
        ["proveedor", "articulos"],
    ),
    "orden_compra_alta": _tool(
        "orden_compra_alta",
        "Ordenes de Compra. ESCRITURA. Da de alta CABORC y DETORC a partir de cabecera y lista de articulos.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año actual."),
            "serie": _string_schema("Serie. Por defecto vacia."),
            "numero": _int_schema("Numero de pedido. 0/omitido calcula el siguiente."),
            "proveedor": _int_schema("Proveedor del pedido."),
            "fecha": _string_schema("Fecha del pedido. Por defecto hoy."),
            "fecha_envio": _string_schema("Fecha de envio al proveedor."),
            "fecha_entrega": _string_schema("Fecha prevista de entrega."),
            "forma_pago": _int_schema("Forma de pago. Por defecto la del proveedor."),
            "representante": _int_schema("Representante. Por defecto 0."),
            "moneda": _string_schema("Moneda. Por defecto proveedor o E."),
            "observaciones": _string_schema("Observaciones de cabecera."),
            "articulos": {"type": "array", "description": "Lineas: string codigo o objeto {articulo,codigo_barras,referencia,cantidad,precio,dto1..dto6,unidad,descripcion,observaciones}.", "items": {"type": ["string", "object"]}},
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["proveedor", "articulos"],
    ),
    "orden_compra_cerrar": _tool(
        "orden_compra_cerrar",
        "Ordenes de Compra. ESCRITURA. Cierra un pedido de compra, marcando lineas cerradas y pendiente a cero.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("COC_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("COC_EJERCI."),
            "serie": _string_schema("COC_SERIE."),
            "numero": _int_schema("COC_NUMDOC."),
            "usuario": _string_schema("Usuario de modificacion. Por defecto '<centro> MCP'."),
            "simular": {"type": "boolean", "description": "Si true, devuelve las sentencias previstas sin escribir."},
        },
        ["ejercicio", "serie", "numero"],
    ),
    "entrada_almacen_tablas": _tool(
        "entrada_almacen_tablas",
        "Entradas de Almacen. Describe CABDOCM, DETMOVM y tablas auxiliares de stock relacionadas.",
        {},
    ),
    "entrada_almacen_listar": _tool(
        "entrada_almacen_listar",
        "Entradas de Almacen. Lista CABDOCM con filtros por situacion, proveedor, fechas, serie, numero, albaran o factura.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "serie": _string_schema("Filtro por serie."),
            "numero": _int_schema("Filtro por CBM_NUMDOC."),
            "situacion": _string_schema("P/F/C/H o pendiente_facturar/pendiente_contabilizar/contabilizada."),
            "proveedor": _int_schema("Filtro por proveedor."),
            "desde": _string_schema("Fecha minima de entrada en formato YYYY-MM-DD."),
            "hasta": _string_schema("Fecha maxima de entrada en formato YYYY-MM-DD."),
            "albaran": _string_schema("Filtro por albaran de proveedor."),
            "factura": _string_schema("Filtro por factura de proveedor."),
            "texto": _string_schema("Filtro por proveedor, observaciones o articulo incluido."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "entrada_almacen_obtener": _tool(
        "entrada_almacen_obtener",
        "Entradas de Almacen. Obtiene cabecera CABDOCM y lineas DETMOVM.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBM_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("CBM_EJERCI."),
            "serie": _string_schema("CBM_SERIE."),
            "numero": _int_schema("CBM_NUMDOC."),
        },
        ["ejercicio", "serie", "numero"],
    ),
    "entrada_almacen_lineas_listar": _tool(
        "entrada_almacen_lineas_listar",
        "Entradas de Almacen. Lista lineas DETMOVM con filtros por documento, proveedor, articulo o texto.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "serie": _string_schema("Filtro por serie."),
            "numero": _int_schema("Filtro por documento."),
            "proveedor": _int_schema("Filtro por proveedor de cabecera."),
            "articulo": _string_schema("Filtro por articulo."),
            "texto": _string_schema("Filtro por articulo, referencia o descripcion."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "entrada_almacen_pendientes_facturar": _tool(
        "entrada_almacen_pendientes_facturar",
        "Entradas de Almacen. Lista y resume documentos pendientes de facturar: CBM_SITUAC='P'.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "proveedor": _int_schema("Filtro por proveedor."),
            "desde": _string_schema("Fecha minima de entrada."),
            "hasta": _string_schema("Fecha maxima de entrada."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "entrada_almacen_pendientes_contabilizar": _tool(
        "entrada_almacen_pendientes_contabilizar",
        "Entradas de Almacen. Lista y resume documentos facturados pendientes de contabilizar: CBM_SITUAC='F'.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "proveedor": _int_schema("Filtro por proveedor."),
            "desde": _string_schema("Fecha minima de entrada."),
            "hasta": _string_schema("Fecha maxima de entrada."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "entrada_almacen_alta_preparar": _tool(
        "entrada_almacen_alta_preparar",
        "Entradas de Almacen. Calcula cabecera, lineas, IVA y totales para alta de CABDOCM/DETMOVM sin escribir.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año de fecha."),
            "serie": _string_schema("Serie. Por defecto parametro E o E<centro>."),
            "numero": _int_schema("Numero. 0/omitido calcula el siguiente."),
            "proveedor": _int_schema("Proveedor."),
            "cif": _string_schema("CIF para resolver proveedor si no se informa proveedor."),
            "fecha": _string_schema("Fecha entrada. Por defecto hoy."),
            "fecha_recepcion": _string_schema("Fecha recepcion. Por defecto fecha."),
            "albaran": _string_schema("Albaran proveedor."),
            "factura": _string_schema("Factura proveedor. Si se informa, situacion F."),
            "fecha_factura": _string_schema("Fecha factura."),
            "moneda": _string_schema("Moneda. Por defecto proveedor o E."),
            "descuento": {"type": "number", "description": "Descuento de cabecera."},
            "portes": {"type": "number", "description": "Portes/importe adicional."},
            "observaciones": _string_schema("Observaciones."),
            "lineas": {"type": "array", "description": "Lineas {articulo|referencia|codigo_barras,cantidad,precio,dto1..dto6,iva,recargo,descripcion,unidad,pedido_*}.", "items": {"type": "object"}},
        },
        ["lineas"],
    ),
    "entrada_almacen_alta": _tool(
        "entrada_almacen_alta",
        "Entradas de Almacen. ESCRITURA. Crea CABDOCM y DETMOVM a partir de cabecera y lineas; actualiza stock existente.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año de fecha."),
            "serie": _string_schema("Serie. Por defecto parametro E o E<centro>."),
            "numero": _int_schema("Numero. 0/omitido calcula el siguiente."),
            "proveedor": _int_schema("Proveedor."),
            "cif": _string_schema("CIF para resolver proveedor si no se informa proveedor."),
            "fecha": _string_schema("Fecha entrada. Por defecto hoy."),
            "fecha_recepcion": _string_schema("Fecha recepcion. Por defecto fecha."),
            "albaran": _string_schema("Albaran proveedor."),
            "factura": _string_schema("Factura proveedor. Si se informa, situacion F."),
            "fecha_factura": _string_schema("Fecha factura."),
            "moneda": _string_schema("Moneda. Por defecto proveedor o E."),
            "descuento": {"type": "number", "description": "Descuento de cabecera."},
            "portes": {"type": "number", "description": "Portes/importe adicional."},
            "observaciones": _string_schema("Observaciones."),
            "lineas": {"type": "array", "description": "Lineas {articulo|referencia|codigo_barras,cantidad,precio,dto1..dto6,iva,recargo,descripcion,unidad,pedido_*}.", "items": {"type": "object"}},
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["lineas"],
    ),
    "entrada_almacen_pdf_previsualizar": _tool(
        "entrada_almacen_pdf_previsualizar",
        "Entradas de Almacen. Extrae texto de un PDF de proveedor y propone cabecera/lineas para revisar antes del alta.",
        {
            "ruta_pdf": _string_schema("Ruta local del PDF."),
            "content_base64": _string_schema("Contenido PDF en Base64 si no se usa ruta_pdf."),
            "nombre_fichero": _string_schema("Nombre del PDF cuando se usa content_base64."),
            "proveedor": _int_schema("Proveedor esperado para resolver referencias."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "fecha": _string_schema("Fecha entrada por defecto."),
            "limite_lineas": _int_schema("Maximo de lineas candidatas."),
        },
    ),
    "entrada_almacen_desde_pdf": _tool(
        "entrada_almacen_desde_pdf",
        "Entradas de Almacen. ESCRITURA. Crea una entrada a partir de un PDF y ajustes manuales sobre la propuesta extraida.",
        {
            "ruta_pdf": _string_schema("Ruta local del PDF."),
            "content_base64": _string_schema("Contenido PDF en Base64 si no se usa ruta_pdf."),
            "nombre_fichero": _string_schema("Nombre del PDF cuando se usa content_base64."),
            "proveedor": _int_schema("Proveedor esperado para resolver referencias."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "fecha": _string_schema("Fecha entrada por defecto."),
            "cabecera": _data_schema("Campos de cabecera que sustituyen lo extraido del PDF."),
            "lineas": {"type": "array", "description": "Lineas revisadas; si se omite se usan candidatas resueltas del PDF.", "items": {"type": "object"}},
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
    ),
    "venta_tablas": _tool(
        "venta_tablas",
        "Ventas. Describe CABDOCV, DETMOV y tablas comerciales relacionadas usadas por pedidos/precios.",
        {},
    ),
    "venta_listar": _tool(
        "venta_listar",
        "Ventas. Lista cabeceras CABDOCV con filtros por documento, cliente, fechas, estado, serie, numero o articulo incluido.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "tipo_documento": _string_schema("CBV_TIPDOC: P pedido, R presupuesto, A albaran, F factura, T ticket, C credito, S cerrado."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto sin filtrar."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "serie": _string_schema("Filtro por serie."),
            "numero": _int_schema("Filtro por numero de documento."),
            "cliente": _int_schema("Filtro por cliente."),
            "subcliente": _int_schema("Filtro por subcliente."),
            "estado": _string_schema("Filtro CBV_SITUAC."),
            "desde": _string_schema("Fecha minima en formato YYYY-MM-DD."),
            "hasta": _string_schema("Fecha maxima en formato YYYY-MM-DD."),
            "articulo": _string_schema("Filtro por articulo incluido en DETMOV."),
            "texto": _string_schema("Filtro por nombre cliente, referencia, observaciones o descripcion de linea."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "venta_obtener": _tool(
        "venta_obtener",
        "Ventas. Obtiene cabecera CABDOCV y lineas DETMOV de un documento.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC."),
        },
        ["tipo_documento", "ejercicio", "serie", "numero"],
    ),
    "venta_lineas_listar": _tool(
        "venta_lineas_listar",
        "Ventas. Lista lineas DETMOV con filtros por documento, cliente, articulo, fechas o texto.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "tipo_documento": _string_schema("Filtro DMV_TIPDOC."),
            "tipo_accion": _string_schema("Filtro DMV_TIPAC."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "serie": _string_schema("Filtro por serie."),
            "numero": _int_schema("Filtro por numero."),
            "cliente": _int_schema("Filtro por cliente de cabecera."),
            "subcliente": _int_schema("Filtro por subcliente de cabecera."),
            "articulo": _string_schema("Filtro por articulo."),
            "desde": _string_schema("Fecha minima de movimiento."),
            "hasta": _string_schema("Fecha maxima de movimiento."),
            "texto": _string_schema("Filtro por articulo, descripcion o cliente."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "venta_precio_articulo": _tool(
        "venta_precio_articulo",
        "Ventas. Calcula precio de venta de un articulo para cliente, fecha y cantidad usando CLIART, ofertas, CLIFAM, tarifa de cliente e IVA.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "cliente": _int_schema("Codigo de cliente."),
            "subcliente": _int_schema("Subcliente. Por defecto 0."),
            "articulo": _string_schema("Codigo de articulo."),
            "codigo_barras": _string_schema("EAN/codigo de barras para resolver articulo."),
            "codigo_cliente": _string_schema("Codigo propio del cliente en CLIART.CLIA_CODARTC."),
            "fecha": _string_schema("Fecha de venta YYYY-MM-DD. Por defecto hoy."),
            "cantidad": {"type": "number", "description": "Cantidad solicitada. Por defecto 1."},
            "tipo_documento": _string_schema("Tipo de documento para reglas de precio. Por defecto P."),
        },
        ["cliente"],
    ),
    "rentabilidad_articulo_ventas": _tool(
        "rentabilidad_articulo_ventas",
        "Ventas/Rentabilidad. Muestra ventas de un articulo en un periodo y calcula margen segun ANAVEN/RENTABILIDAD_LINEA.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "articulo": _string_schema("Codigo de articulo."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD."),
            "centro": _int_schema("Filtro por centro."),
            "cliente": _int_schema("Filtro por cliente."),
            "subcliente": _int_schema("Filtro por subcliente."),
            "serie": _string_schema("Filtro por serie."),
            "numero_desde": _int_schema("Numero documento inicial."),
            "numero_hasta": _int_schema("Numero documento final."),
            "tipos_documento": {"type": "array", "description": "Tipos de venta a incluir. Por defecto F,A,C,T.", "items": {"type": "string"}},
            "incluir_pedidos": {"type": "boolean", "description": "Si true no fuerza DMV_SIGNO='1'. Por defecto false."},
            "modo_coste": _string_schema("Opcional. Fuerza PBASE, PMEDIO, ULTIMO o PREBAS."),
            "limite": _int_schema("Maximo de lineas."),
        },
        ["articulo"],
    ),
    "rentabilidad_articulos_resumen": _tool(
        "rentabilidad_articulos_resumen",
        "Ventas/Rentabilidad. Resume unidades, ventas, coste, margen y porcentaje de rentabilidad por articulo en un periodo.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD."),
            "centro": _int_schema("Filtro por centro."),
            "articulo": _string_schema("Filtro por codigo exacto de articulo."),
            "texto": _string_schema("Filtro por codigo o descripcion."),
            "proveedor": _int_schema("Filtro por proveedor principal ART_CODPRO."),
            "familia": _string_schema("Filtro por ART_CODFAM."),
            "subfamilia": _string_schema("Filtro por ART_SUBFAM."),
            "cliente": _int_schema("Filtro por cliente."),
            "subcliente": _int_schema("Filtro por subcliente."),
            "serie": _string_schema("Filtro por serie."),
            "representante": _int_schema("Filtro por CBV_CODREP."),
            "en_oferta": {"type": "boolean", "description": "Si true, solo lineas con DMV_EJEOFE > 0."},
            "tipos_documento": {"type": "array", "description": "Tipos de venta a incluir. Por defecto F,A,C,T.", "items": {"type": "string"}},
            "incluir_pedidos": {"type": "boolean", "description": "Si true no fuerza DMV_SIGNO='1'. Por defecto false."},
            "modo_coste": _string_schema("Opcional. Fuerza PBASE, PMEDIO, ULTIMO o PREBAS."),
            "orden": _string_schema("Orden: margen, ventas, rentabilidad, unidades. Por defecto margen."),
            "limite": _int_schema("Maximo de articulos."),
            "limite_lineas": _int_schema("Maximo de lineas DETMOV a analizar antes de agregar. Por defecto 1000."),
        },
    ),
    "ventas_documentos_detalle": _tool(
        "ventas_documentos_detalle",
        "Ventas/ANADOC. Lista documentos de venta de un periodo con importes, cobros, forma de pago, agente, poblacion y vencimiento.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD."),
            "centro": _int_schema("Filtro por centro."),
            "tipos_documento": {"type": "array", "description": "Tipos CBV_TIPDOC a incluir. Por defecto T,A,F,C.", "items": {"type": "string"}},
            "tipo_accion": _string_schema("Filtro CBV_TIPAC. Usa 9 o vacio para todos."),
            "situacion": _string_schema("Filtro CBV_SITUAC."),
            "cliente_desde": _int_schema("Cliente inicial."),
            "cliente_hasta": _int_schema("Cliente final."),
            "subcliente_desde": _int_schema("Subcliente inicial."),
            "subcliente_hasta": _int_schema("Subcliente final."),
            "representante": _int_schema("Filtro CBV_CODREP."),
            "serie_desde": _string_schema("Serie inicial."),
            "serie_hasta": _string_schema("Serie final."),
            "numero_desde": _int_schema("Numero documento inicial."),
            "numero_hasta": _int_schema("Numero documento final."),
            "tarjeta": _string_schema("Filtro CBV_CODTAR."),
            "importe_pendiente_min": {"type": "number", "description": "Importe pendiente minimo."},
            "importe_pendiente_max": {"type": "number", "description": "Importe pendiente maximo."},
            "tipo_factura": _string_schema("Para CBV_TIPDOC='F': contado, tickets, albaran o todas. Por defecto todas."),
            "limite": _int_schema("Maximo de documentos."),
        },
    ),
    "ventas_documentos_resumen": _tool(
        "ventas_documentos_resumen",
        "Ventas/ANADOC. Agrupa documentos de venta por tipo, cliente, forma de pago, forma de cobro, agente, poblacion, etc.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD."),
            "agrupar_por": _string_schema("tipo_documento, cliente, forma_pago, forma_cobro, agente, poblacion, centro, serie, mes, dia_semana, hora, situacion, tarjeta."),
            "centro": _int_schema("Filtro por centro."),
            "tipos_documento": {"type": "array", "description": "Tipos CBV_TIPDOC a incluir. Por defecto T,A,F,C.", "items": {"type": "string"}},
            "tipo_accion": _string_schema("Filtro CBV_TIPAC. Usa 9 o vacio para todos."),
            "situacion": _string_schema("Filtro CBV_SITUAC."),
            "cliente_desde": _int_schema("Cliente inicial."),
            "cliente_hasta": _int_schema("Cliente final."),
            "subcliente_desde": _int_schema("Subcliente inicial."),
            "subcliente_hasta": _int_schema("Subcliente final."),
            "representante": _int_schema("Filtro CBV_CODREP."),
            "serie_desde": _string_schema("Serie inicial."),
            "serie_hasta": _string_schema("Serie final."),
            "numero_desde": _int_schema("Numero documento inicial."),
            "numero_hasta": _int_schema("Numero documento final."),
            "tarjeta": _string_schema("Filtro CBV_CODTAR."),
            "importe_pendiente_min": {"type": "number", "description": "Importe pendiente minimo."},
            "importe_pendiente_max": {"type": "number", "description": "Importe pendiente maximo."},
            "tipo_factura": _string_schema("Para CBV_TIPDOC='F': contado, tickets, albaran o todas. Por defecto todas."),
            "orden": _string_schema("Orden: total, base, pendiente, documentos. Por defecto total."),
            "limite": _int_schema("Maximo de grupos."),
        },
    ),
    "venta_documento_alta_preparar": _tool(
        "venta_documento_alta_preparar",
        "Ventas. Calcula cabecera, lineas, IVA y totales para alta de CABDOCV/DETMOV sin escribir.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("Tipo: P pedido, R presupuesto, A albaran, F factura. Por defecto P."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año de fecha."),
            "serie": _string_schema("Serie. Por defecto CLIENI SERIE o parametro <tipo><centro>/<tipo>, o PM para pedidos."),
            "numero": _int_schema("Numero. 0/omitido calcula desde NUMERA."),
            "cliente": _int_schema("Cliente."),
            "subcliente": _int_schema("Subcliente. Por defecto 0."),
            "fecha": _string_schema("Fecha documento. Por defecto hoy."),
            "fecha_entrega": _string_schema("Fecha entrega/emision. Por defecto fecha."),
            "observaciones": _string_schema("Observaciones."),
            "referencia_cliente": _string_schema("Referencia de cliente."),
            "retira": _string_schema("Persona o texto de retirada."),
            "descuento": {"type": "number", "description": "Descuento de cabecera. Por defecto CLI_DTOESP."},
            "portes": {"type": "number", "description": "Portes/importe adicional."},
            "lineas": {"type": "array", "description": "Lineas {articulo|codigo_barras|codigo_cliente,cantidad,precio,dto1,dto2,descripcion,unidad,tipo_linea}. Si no hay precio se calcula.", "items": {"type": "object"}},
        },
        ["cliente", "lineas"],
    ),
    "venta_documento_alta": _tool(
        "venta_documento_alta",
        "Ventas. ESCRITURA. Crea CABDOCV y DETMOV a partir de datos de cabecera y lineas; para pedidos no toca stock.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("Tipo: P pedido, R presupuesto, A albaran, F factura. Por defecto P."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año de fecha."),
            "serie": _string_schema("Serie. Por defecto CLIENI SERIE o parametro <tipo><centro>/<tipo>, o PM para pedidos."),
            "numero": _int_schema("Numero. 0/omitido calcula desde NUMERA."),
            "cliente": _int_schema("Cliente."),
            "subcliente": _int_schema("Subcliente. Por defecto 0."),
            "fecha": _string_schema("Fecha documento. Por defecto hoy."),
            "fecha_entrega": _string_schema("Fecha entrega/emision. Por defecto fecha."),
            "observaciones": _string_schema("Observaciones."),
            "referencia_cliente": _string_schema("Referencia de cliente."),
            "retira": _string_schema("Persona o texto de retirada."),
            "descuento": {"type": "number", "description": "Descuento de cabecera. Por defecto CLI_DTOESP."},
            "portes": {"type": "number", "description": "Portes/importe adicional."},
            "lineas": {"type": "array", "description": "Lineas {articulo|codigo_barras|codigo_cliente,cantidad,precio,dto1,dto2,descripcion,unidad,tipo_linea}. Si no hay precio se calcula.", "items": {"type": "object"}},
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["cliente", "lineas"],
    ),
    "venta_pedido_alta": _tool(
        "venta_pedido_alta",
        "Ventas. ESCRITURA. Atajo para crear un pedido de cliente (CBV_TIPDOC='P') con cabecera y lineas.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año de fecha."),
            "serie": _string_schema("Serie. Por defecto CLIENI SERIE, parametro P<centro>/P o PM."),
            "numero": _int_schema("Numero. 0/omitido calcula desde NUMERA."),
            "cliente": _int_schema("Cliente."),
            "subcliente": _int_schema("Subcliente. Por defecto 0."),
            "fecha": _string_schema("Fecha documento. Por defecto hoy."),
            "fecha_entrega": _string_schema("Fecha entrega. Por defecto fecha."),
            "observaciones": _string_schema("Observaciones."),
            "referencia_cliente": _string_schema("Referencia de cliente."),
            "retira": _string_schema("Persona o texto de retirada."),
            "lineas": {"type": "array", "description": "Lineas {articulo|codigo_barras|codigo_cliente,cantidad,precio,dto1,dto2,descripcion,unidad}. Si no hay precio se calcula.", "items": {"type": "object"}},
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["cliente", "lineas"],
    ),
    "pedido_crear": _tool(
        "pedido_crear",
        "Pedidos. ESCRITURA. Crea un pedido de cliente con sus lineas y valoracion comercial; alias compatible con Kronos sobre CABDOCV/DETMOV.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio. Por defecto año de fecha."),
            "serie": _string_schema("Serie. Por defecto CLIENI SERIE, parametro P<centro>/P o PM."),
            "numero": _int_schema("Numero. 0/omitido calcula desde NUMERA."),
            "cliente": _int_schema("Cliente. Tambien acepta codcli."),
            "codcli": _int_schema("Alias Kronos de cliente."),
            "subcliente": _int_schema("Subcliente. Tambien acepta subcli. Por defecto 0."),
            "subcli": _int_schema("Alias Kronos de subcliente."),
            "fecha": _string_schema("Fecha documento. Por defecto hoy."),
            "fecha_entrega": _string_schema("Fecha entrega. Por defecto fecha."),
            "observaciones": _string_schema("Observaciones de cabecera."),
            "referencia_cliente": _string_schema("Referencia del cliente."),
            "retira": _string_schema("Persona o texto de retirada."),
            "lineas": {"type": "array", "description": "Lineas estructuradas. Si se omite y se informa texto, intenta parsear lineas JSON.", "items": {"type": "object"}},
            "texto": _string_schema("Compatibilidad Kronos: JSON con lineas o texto informativo."),
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        [],
    ),
    "pedido_listar": _tool(
        "pedido_listar",
        "Pedidos. Lista pedidos y presupuestos de cliente desde CABDOCV, filtrando por cliente, estado, fecha, articulo o documento.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro. Por defecto KOFEDAS_CENTRO si no se indica cliente."),
            "tipo_documento": _string_schema("P pedido, R presupuesto, S cerrado/historico. Por defecto P,R."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto todos."),
            "ejercicio": _int_schema("Filtro por ejercicio."),
            "serie": _string_schema("Filtro por serie."),
            "numero": _int_schema("Filtro por numero."),
            "cliente": _int_schema("Filtro por cliente. Tambien acepta codcli."),
            "codcli": _int_schema("Alias Kronos de cliente."),
            "subcliente": _int_schema("Filtro por subcliente. Tambien acepta subcli."),
            "subcli": _int_schema("Alias Kronos de subcliente."),
            "estado": _string_schema("Filtro CBV_SITUAC."),
            "desde": _string_schema("Fecha desde."),
            "hasta": _string_schema("Fecha hasta."),
            "articulo": _string_schema("Filtro por articulo incluido."),
            "texto": _string_schema("Filtro por cliente, referencia u observaciones."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "pedido_detalle": _tool(
        "pedido_detalle",
        "Pedidos. Devuelve cabecera y lineas de un pedido; modo preparacion incluye cantidades pendientes/preparadas cuando existen.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "modo": _string_schema("normal o preparacion."),
        },
        ["serie"],
    ),
    "pedido_cerrar": _tool(
        "pedido_cerrar",
        "Pedidos. ESCRITURA. Cierra un pedido/presupuesto de cliente marcandolo como historico S y cerrando sus lineas.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("P pedido o R presupuesto. Tambien acepta tipdoc. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "usuario": _string_schema("Usuario de modificacion."),
            "simular": {"type": "boolean", "description": "Si true, no escribe y devuelve el plan."},
        },
        ["serie"],
    ),
    "pedido_situacion_actualizar": _tool(
        "pedido_situacion_actualizar",
        "Pedidos. ESCRITURA. Actualiza CBV_SITUAC de un pedido.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC. Tambien acepta tipdoc. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "situacion": _string_schema("Nuevo valor CBV_SITUAC. Tambien acepta situac."),
            "situac": _string_schema("Alias Kronos de situacion."),
            "usuario": _string_schema("Usuario de modificacion."),
            "simular": {"type": "boolean", "description": "Si true, no escribe y devuelve el plan."},
        },
        ["serie"],
    ),
    "pedido_retirada_actualizar": _tool(
        "pedido_retirada_actualizar",
        "Pedidos. ESCRITURA. Actualiza datos de retirada y referencia de cliente de un pedido.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC. Tambien acepta tipdoc. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "retirado": _string_schema("Texto de retirada para CBV_RETIRA."),
            "retira": _string_schema("Alias de retirado."),
            "referencia": _string_schema("Referencia cliente para CBV_REFCLI."),
            "referencia_cliente": _string_schema("Alias de referencia."),
            "usuario": _string_schema("Usuario de modificacion."),
            "simular": {"type": "boolean", "description": "Si true, no escribe y devuelve el plan."},
        },
        ["serie"],
    ),
    "pedido_albaranar": _tool(
        "pedido_albaranar",
        "Pedidos. ESCRITURA. Crea un albaran desde un pedido usando sus lineas pendientes o las cantidades indicadas.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro del pedido. Por defecto KOFEDAS_CENTRO o 0."),
            "ejercicio": _int_schema("Ejercicio pedido. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("Serie pedido."),
            "numero": _int_schema("Numero pedido. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "tipo_accion": _string_schema("CBV_TIPAC pedido. Por defecto 0."),
            "serie_albaran": _string_schema("Serie del albaran. Por defecto parametros A<centro>/A."),
            "fecha": _string_schema("Fecha del albaran. Por defecto hoy."),
            "lineas": {"type": "array", "description": "Opcional: {linea,cantidad}. Si se omite usa cantidades pendientes/cantidad del pedido.", "items": {"type": "object"}},
            "cerrar_pedido": {"type": "boolean", "description": "Si true, cierra el pedido tras crear el albaran."},
            "usuario": _string_schema("Usuario de modificacion."),
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["serie"],
    ),
    "pedido_marcar_preparado": _tool(
        "pedido_marcar_preparado",
        "Pedidos/almacen. ESCRITURA. Marca cantidades preparadas en lineas de pedido cuando existen campos de preparacion; si no, actualiza la situacion de cabecera.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC. Tambien acepta tipdoc. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "lineas": {"type": "array", "description": "Opcional: {linea,cantidad_preparada}.", "items": {"type": "object"}},
            "situacion": _string_schema("Situacion de cabecera a dejar. Por defecto P."),
            "usuario": _string_schema("Usuario de modificacion."),
            "simular": {"type": "boolean", "description": "Si true, no escribe y devuelve el plan."},
        },
        ["serie"],
    ),
    "pedido_finalizar": _tool(
        "pedido_finalizar",
        "Pedidos. ESCRITURA. Finaliza/revisa la preparacion de un pedido y opcionalmente cambia su situacion.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC. Tambien acepta tipdoc. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "situacion": _string_schema("Situacion final opcional. Por defecto no cambia."),
            "usuario": _string_schema("Usuario de modificacion."),
            "simular": {"type": "boolean", "description": "Si true, no escribe y devuelve diagnostico."},
        },
        ["serie"],
    ),
    "pedido_linea_mover": _tool(
        "pedido_linea_mover",
        "Pedidos/almacen. ESCRITURA. Mueve una linea de pedido entre zonas de preparacion si DETMOV dispone de campos de zona.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC. Tambien acepta tipdoc. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "linea": _int_schema("DMV_NUMLIN."),
            "zona_destino": _int_schema("Zona destino."),
            "zona_origen": _int_schema("Zona origen opcional."),
            "usuario": _string_schema("Usuario de modificacion."),
            "simular": {"type": "boolean", "description": "Si true, no escribe y devuelve el plan."},
        },
        ["serie", "linea", "zona_destino"],
    ),
    "pedido_pdf_gestion": _tool(
        "pedido_pdf_gestion",
        "Pedidos. Genera o recupera una representacion HTML/Base64 del pedido para integraciones sin depender del motor de informes Delphi.",
        {
            "accion": _string_schema("generar u obtener. Ambas devuelven contenido HTML en Base64."),
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC. Tambien acepta tipdoc. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
        },
        ["accion", "serie"],
    ),
    "pedido_enviar": _tool(
        "pedido_enviar",
        "Pedidos. Prepara el envio por email de un pedido devolviendo asunto, destinatario sugerido y HTML/Base64; no envia correo directamente.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("CBV_CENTRO. Por defecto KOFEDAS_CENTRO o 0."),
            "tipo_documento": _string_schema("CBV_TIPDOC. Tambien acepta tipdoc. Por defecto P."),
            "tipdoc": _string_schema("Alias Kronos de tipo_documento."),
            "tipo_accion": _string_schema("CBV_TIPAC. Por defecto 0."),
            "ejercicio": _int_schema("CBV_EJERCI. Tambien acepta ejerci."),
            "ejerci": _int_schema("Alias Kronos de ejercicio."),
            "serie": _string_schema("CBV_SERIE."),
            "numero": _int_schema("CBV_NUMDOC. Tambien acepta numdoc."),
            "numdoc": _int_schema("Alias Kronos de numero."),
            "email": _string_schema("Destinatario. Si se omite, intenta usar el email del cliente."),
            "asunto": _string_schema("Asunto del mensaje."),
        },
        ["serie"],
    ),
    "cartera_tablas": _tool(
        "cartera_tablas",
        "Cartera. Describe CABDOCVE y tablas relacionadas para deuda, vencimientos y remesas.",
        {},
    ),
    "cartera_efectos_detalle": _tool(
        "cartera_efectos_detalle",
        "Cartera. Lista efectos/vencimientos CABDOCVE con filtros por cliente, documento, tipo, situacion, remesa y fechas.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "tipo_documento": _string_schema("Filtro CBVE_TIPDOC: F factura, T ticket, C credito, etc."),
            "tipo_accion": _string_schema("Filtro CBVE_TIPAC."),
            "cliente": _int_schema("Filtro por cliente."),
            "subcliente": _int_schema("Filtro por subcliente."),
            "cliente_desde": _int_schema("Cliente inicial."),
            "cliente_hasta": _int_schema("Cliente final."),
            "ejercicio": _int_schema("Filtro por ejercicio documento."),
            "serie": _string_schema("Filtro por serie."),
            "numero": _int_schema("Filtro por documento."),
            "orden": _int_schema("Filtro por numero de efecto."),
            "tipo_efecto": _string_schema("Filtro CBVE_TIPOEF."),
            "situacion": _string_schema("pendiente, vencido, cobrado/ cancelado, impagado o todos."),
            "remesado": _string_schema("S/N para efectos remesados o pendientes de remesar."),
            "remesa_ejercicio": _int_schema("Filtro CBVE_EJEREM."),
            "remesa_codigo": _int_schema("Filtro CBVE_CODREM."),
            "fecha_desde": _string_schema("Fecha emision desde."),
            "fecha_hasta": _string_schema("Fecha emision hasta."),
            "vencimiento_desde": _string_schema("Vencimiento desde."),
            "vencimiento_hasta": _string_schema("Vencimiento hasta."),
            "fecha_referencia": _string_schema("Fecha para decidir vencidos. Por defecto hoy."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "cartera_deuda_cliente": _tool(
        "cartera_deuda_cliente",
        "Cartera. Consulta deuda de un cliente, con totales y detalle de efectos pendientes/vencidos/remesados.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "cliente": _int_schema("Cliente."),
            "subcliente": _int_schema("Subcliente. Por defecto todos los subclientes."),
            "tipo_documento": _string_schema("Filtro CBVE_TIPDOC."),
            "tipo_efecto": _string_schema("Filtro CBVE_TIPOEF."),
            "situacion": _string_schema("pendiente, vencido, cobrado, impagado o todos. Por defecto pendiente."),
            "remesado": _string_schema("S/N."),
            "fecha_referencia": _string_schema("Fecha para decidir vencidos."),
            "limite": _int_schema("Maximo de efectos devueltos."),
        },
        ["cliente"],
    ),
    "cartera_deuda_por_cliente": _tool(
        "cartera_deuda_por_cliente",
        "Cartera. Resume deuda por cliente con pendiente, vencido, remesado y no remesado.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "cliente_desde": _int_schema("Cliente inicial."),
            "cliente_hasta": _int_schema("Cliente final."),
            "tipo_documento": _string_schema("Filtro CBVE_TIPDOC."),
            "tipo_efecto": _string_schema("Filtro CBVE_TIPOEF."),
            "situacion": _string_schema("pendiente, vencido, cobrado, impagado o todos. Por defecto pendiente."),
            "remesado": _string_schema("S/N."),
            "fecha_referencia": _string_schema("Fecha para decidir vencidos."),
            "limite": _int_schema("Maximo de efectos leidos."),
            "limite_clientes": _int_schema("Maximo de clientes devueltos."),
        },
    ),
    "cartera_pendiente_remesar": _tool(
        "cartera_pendiente_remesar",
        "Cartera. Lista efectos pendientes no remesados: CBVE_FECCAN IS NULL y CBVE_EJEREM/CBVE_CODREM a cero.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "cliente": _int_schema("Filtro por cliente."),
            "tipo_documento": _string_schema("Filtro CBVE_TIPDOC."),
            "tipo_efecto": _string_schema("Filtro CBVE_TIPOEF."),
            "vencimiento_hasta": _string_schema("Vencimiento maximo."),
            "fecha_referencia": _string_schema("Fecha para decidir vencidos."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "cartera_deuda_por_tipo": _tool(
        "cartera_deuda_por_tipo",
        "Cartera. Resume deuda por tipo de documento, tipo de efecto, situacion y remesado.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "cliente": _int_schema("Filtro por cliente."),
            "cliente_desde": _int_schema("Cliente inicial."),
            "cliente_hasta": _int_schema("Cliente final."),
            "tipo_documento": _string_schema("Filtro CBVE_TIPDOC."),
            "tipo_efecto": _string_schema("Filtro CBVE_TIPOEF."),
            "situacion": _string_schema("pendiente, vencido, cobrado, impagado o todos. Por defecto pendiente."),
            "remesado": _string_schema("S/N."),
            "fecha_referencia": _string_schema("Fecha para decidir vencidos."),
            "limite": _int_schema("Maximo de efectos leidos."),
        },
    ),
    "dashboard_resumen": _tool(
        "dashboard_resumen",
        "Dashboard ERP. Resume ventas, compras, cartera, pedidos de compra y entradas pendientes para un periodo.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD. Por defecto inicio del año actual."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD. Por defecto hoy."),
            "tipos_documento": {"type": "array", "description": "Tipos de venta a incluir. Por defecto F,T,A,C.", "items": {"type": "string"}},
        },
    ),
    "ventas_resumen": _tool(
        "ventas_resumen",
        "Dashboard ERP. Resume ventas por periodo con agrupacion opcional por mes, año, cliente, articulo, familia, centro o tipo_documento.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD."),
            "agrupar_por": _string_schema("mes, anio, cliente, articulo, familia, centro o tipo_documento. Por defecto mes."),
            "tipos_documento": {"type": "array", "description": "Tipos de venta a incluir. Por defecto F,T,A,C.", "items": {"type": "string"}},
            "limite": _int_schema("Maximo de grupos."),
        },
    ),
    "compras_resumen": _tool(
        "compras_resumen",
        "Dashboard ERP. Resume compras/entradas por periodo agrupando por mes, año, proveedor, articulo, familia, centro o situacion.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD."),
            "agrupar_por": _string_schema("mes, anio, proveedor, articulo, familia, centro o situacion. Por defecto mes."),
            "limite": _int_schema("Maximo de grupos."),
        },
    ),
    "dashboard_evolucion_anual": _tool(
        "dashboard_evolucion_anual",
        "Dashboard ERP. Evolucion por años de ventas, compras, documentos y margen aproximado.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "anio_desde": _int_schema("Año inicial. Por defecto año actual - 4."),
            "anio_hasta": _int_schema("Año final. Por defecto año actual."),
            "tipos_documento": {"type": "array", "description": "Tipos de venta a incluir. Por defecto F,T,A,C.", "items": {"type": "string"}},
        },
    ),
    "dashboard_series_temporales": _tool(
        "dashboard_series_temporales",
        "Dashboard ERP. Series temporales de ventas y compras agrupadas por mes o año.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD."),
            "agrupar_por": _string_schema("mes o anio. Por defecto mes."),
            "tipos_documento": {"type": "array", "description": "Tipos de venta a incluir. Por defecto F,T,A,C.", "items": {"type": "string"}},
        },
    ),
    "dashboard_rankings": _tool(
        "dashboard_rankings",
        "Dashboard ERP. Rankings de clientes, proveedores y articulos de venta/compra del periodo.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "desde": _string_schema("Fecha inicial YYYY-MM-DD."),
            "hasta": _string_schema("Fecha final YYYY-MM-DD."),
            "tipos_documento": {"type": "array", "description": "Tipos de venta a incluir. Por defecto F,T,A,C.", "items": {"type": "string"}},
            "limite": _int_schema("Maximo de filas por ranking."),
        },
    ),
    "regularizacion_tablas": _tool(
        "regularizacion_tablas",
        "Regularizaciones. Describe CABDOCR, DETMOVR, RECUENTO, ARTICULE y STOCKS.",
        {},
    ),
    "regularizacion_listar": _tool(
        "regularizacion_listar",
        "Regularizaciones. Lista documentos CABDOCR y lineas DETMOVR por tipo, centro, articulo o fecha.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Filtro por centro."),
            "tipo": _string_schema("Tipo CBR_TIPO: R regularizacion, S salida trasvase, E entrada trasvase."),
            "desde": _string_schema("Fecha minima."),
            "hasta": _string_schema("Fecha maxima."),
            "articulo": _string_schema("Filtro por articulo incluido."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "stock_por_almacen": _tool(
        "stock_por_almacen",
        "Stock. Muestra stock actual de un articulo por almacen/centro desde ARTICULE.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "articulo": _string_schema("Codigo de articulo."),
            "centro": _int_schema("Filtro opcional por centro."),
        },
        ["articulo"],
    ),
    "stock_a_fecha": _tool(
        "stock_a_fecha",
        "Stock. Calcula stock de un articulo a fecha usando ARTICULE y movimientos posteriores, segun UTL_STOCK.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "articulo": _string_schema("Codigo de articulo."),
            "centro": _int_schema("Centro; usa -1 para todos. Por defecto KOFEDAS_CENTRO."),
            "fecha": _string_schema("Fecha YYYY-MM-DD. Omitida devuelve stock actual."),
        },
        ["articulo"],
    ),
    "inventario_valorar_articulos": _tool(
        "inventario_valorar_articulos",
        "Stock. Valora existencias a coste respetando PARAMETROS.RENTAB: PBASE, PMEDIO o ULTIMO.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro; usa -1 para todos. Por defecto KOFEDAS_CENTRO."),
            "fecha": _string_schema("Fecha YYYY-MM-DD para coste medio/ultimo y stock historico. Por defecto hoy."),
            "articulo": _string_schema("Codigo exacto de articulo."),
            "texto": _string_schema("Filtro por codigo o descripcion."),
            "familia": _string_schema("Filtro por ART_CODFAM."),
            "subfamilia": _string_schema("Filtro por ART_SUBFAM."),
            "proveedor": _int_schema("Filtro por ART_CODPRO."),
            "modo_coste": _string_schema("Opcional. Fuerza PBASE, PMEDIO, ULTIMO, ART_PRECOS, COSTE_MEDIO o ULTIMO_PRECIO."),
            "solo_con_stock": {"type": "boolean", "description": "Si true, excluye articulos sin stock. Por defecto true."},
            "incluir_no_inventariables": {"type": "boolean", "description": "Si true, incluye ART_INDINV='N'. Por defecto false."},
            "limite": _int_schema("Maximo de articulos devueltos."),
        },
    ),
    "articulo_regularizar": _tool(
        "articulo_regularizar",
        "Regularizaciones. ESCRITURA. Ajusta el stock actual de un articulo a una cantidad objetivo y graba DETMOVR.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "articulo": _string_schema("Codigo de articulo."),
            "cantidad": {"type": "number", "description": "Stock objetivo final."},
            "fecha": _string_schema("Fecha del documento. Por defecto ayer, como Kronos."),
            "serie": _string_schema("Serie regularizacion. Por defecto parametro R<centro> o R."),
            "observaciones": _string_schema("Observaciones de cabecera."),
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["articulo", "cantidad"],
    ),
    "trasvase_generar": _tool(
        "trasvase_generar",
        "Regularizaciones. ESCRITURA. Genera trasvase entre tiendas: salida en origen y entrada espejo en destino.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro_origen": _int_schema("Centro origen."),
            "centro_destino": _int_schema("Centro destino."),
            "fecha": _string_schema("Fecha del trasvase. Por defecto hoy."),
            "serie": _string_schema("Serie. Por defecto parametro R<origen> o R."),
            "lineas": {"type": "array", "description": "Lineas {articulo,cantidad,descripcion,unidad}.", "items": {"type": "object"}},
            "observaciones": _string_schema("Observaciones."),
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["centro_origen", "centro_destino", "lineas"],
    ),
    "recuento_listar": _tool(
        "recuento_listar",
        "Recuentos. Lista recuentos pendientes en RECUENTO por centro/articulo.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "articulo": _string_schema("Filtro por articulo."),
            "limite": _int_schema("Maximo de filas."),
        },
    ),
    "recuento_grabar": _tool(
        "recuento_grabar",
        "Recuentos. ESCRITURA. Crea o actualiza un recuento en RECUENTO; no regulariza stock.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "articulo": _string_schema("Codigo de articulo."),
            "cantidad": {"type": "number", "description": "Cantidad contada."},
            "aumentar": {"type": "boolean", "description": "Si true suma a un recuento existente; si false sustituye."},
            "fecha": _string_schema("Fecha de recuento. Por defecto hoy."),
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["articulo", "cantidad"],
    ),
    "recuento_borrar": _tool(
        "recuento_borrar",
        "Recuentos. ESCRITURA. Borra un recuento pendiente de RECUENTO.",
        {
            "empresa": _int_schema("Empresa. Por defecto KOFEDAS_EMPRESA o 1."),
            "centro": _int_schema("Centro. Por defecto KOFEDAS_CENTRO o 0."),
            "articulo": _string_schema("Codigo de articulo."),
            "simular": {"type": "boolean", "description": "Si true, devuelve el plan sin escribir."},
        },
        ["articulo"],
    ),
}

TOOL_PROFILES: dict[str, set[str]] = {
    "core": set(PUBLIC_TOOL_DEFINITIONS),
    "all": set(PUBLIC_TOOL_DEFINITIONS),
}


def tool_definitions(profile: str | None = None) -> list[dict[str, Any]]:
    names = TOOL_PROFILES.get((profile or "core").lower(), TOOL_PROFILES["core"])
    return [json.loads(json.dumps(PUBLIC_TOOL_DEFINITIONS[name], ensure_ascii=False)) for name in sorted(names)]


class KofedasToolRuntime:
    def __init__(self, db: KofedasDatabase | None = None) -> None:
        self.db = db or KofedasDatabase()
        self.empresa = _env_int("KOFEDAS_EMPRESA", DEFAULT_EMPRESA)
        self.centro = _env_int("KOFEDAS_CENTRO", DEFAULT_CENTRO)
        self.tool_profile = os.getenv("KOFEDAS_MCP_TOOL_PROFILE", "core").strip().lower() or "core"
        self.access_level = os.getenv("KOFEDAS_MCP_ACCESS_LEVEL", "critical").strip().lower() or "critical"
        self._handlers: dict[str, Callable[[dict[str, Any]], Any]] = {
            "sistema_estado": self.sistema_estado,
            "empresa_listar": self.empresa_listar,
            "empresa_obtener": self.empresa_obtener,
            "empresa_guardar": self.empresa_guardar,
            "centro_listar": self.centro_listar,
            "centro_obtener": self.centro_obtener,
            "centro_guardar": self.centro_guardar,
            "usuario_listar": self.usuario_listar,
            "usuario_obtener": self.usuario_obtener,
            "usuario_guardar": self.usuario_guardar,
            "grupo_usuario_listar": self.grupo_usuario_listar,
            "grupo_usuario_obtener": self.grupo_usuario_obtener,
            "grupo_usuario_guardar": self.grupo_usuario_guardar,
            "parametro_listar": self.parametro_listar,
            "parametro_obtener": self.parametro_obtener,
            "parametro_guardar": self.parametro_guardar,
            "auxiliar_tablas": self.auxiliar_tablas,
            "auxiliar_listar": self.auxiliar_listar,
            "auxiliar_obtener": self.auxiliar_obtener,
            "auxiliar_guardar": self.auxiliar_guardar,
            "familia_listar": self.familia_listar,
            "articulo_buscar": self.articulo_buscar,
            "articulo_obtener": self.articulo_obtener,
            "stock_consultar": self.stock_consultar,
            "articulo_tablas": self.articulo_tablas,
            "articulo_relacion_listar": self.articulo_relacion_listar,
            "articulo_relacion_obtener": self.articulo_relacion_obtener,
            "articulo_relacion_guardar": self.articulo_relacion_guardar,
            "articulo_completo": self.articulo_completo,
            "articulo_alta_preparar": self.articulo_alta_preparar,
            "articulo_alta": self.articulo_alta,
            "articulo_tarifa_excel_previsualizar": self.articulo_tarifa_excel_previsualizar,
            "articulo_tarifa_excel_importar": self.articulo_tarifa_excel_importar,
            "cliente_buscar": self.cliente_buscar,
            "cliente_obtener": self.cliente_obtener,
            "cliente_tablas": self.cliente_tablas,
            "cliente_relacion_listar": self.cliente_relacion_listar,
            "cliente_relacion_obtener": self.cliente_relacion_obtener,
            "cliente_relacion_guardar": self.cliente_relacion_guardar,
            "cliente_completo": self.cliente_completo,
            "cliente_alta_preparar": self.cliente_alta_preparar,
            "cliente_alta": self.cliente_alta,
            "proveedor_buscar": self.proveedor_buscar,
            "proveedor_obtener": self.proveedor_obtener,
            "proveedor_articulos_listar": self.proveedor_articulos_listar,
            "proveedor_tablas": self.proveedor_tablas,
            "proveedor_relacion_listar": self.proveedor_relacion_listar,
            "proveedor_relacion_obtener": self.proveedor_relacion_obtener,
            "proveedor_relacion_guardar": self.proveedor_relacion_guardar,
            "proveedor_completo": self.proveedor_completo,
            "proveedor_alta_preparar": self.proveedor_alta_preparar,
            "proveedor_alta": self.proveedor_alta,
            "oferta_tablas": self.oferta_tablas,
            "oferta_listar": self.oferta_listar,
            "oferta_obtener": self.oferta_obtener,
            "oferta_articulos_listar": self.oferta_articulos_listar,
            "oferta_alta_preparar": self.oferta_alta_preparar,
            "oferta_alta": self.oferta_alta,
            "orden_compra_tablas": self.orden_compra_tablas,
            "orden_compra_listar": self.orden_compra_listar,
            "orden_compra_obtener": self.orden_compra_obtener,
            "orden_compra_lineas_listar": self.orden_compra_lineas_listar,
            "orden_compra_alta_preparar": self.orden_compra_alta_preparar,
            "orden_compra_alta": self.orden_compra_alta,
            "orden_compra_cerrar": self.orden_compra_cerrar,
            "entrada_almacen_tablas": self.entrada_almacen_tablas,
            "entrada_almacen_listar": self.entrada_almacen_listar,
            "entrada_almacen_obtener": self.entrada_almacen_obtener,
            "entrada_almacen_lineas_listar": self.entrada_almacen_lineas_listar,
            "entrada_almacen_pendientes_facturar": self.entrada_almacen_pendientes_facturar,
            "entrada_almacen_pendientes_contabilizar": self.entrada_almacen_pendientes_contabilizar,
            "entrada_almacen_alta_preparar": self.entrada_almacen_alta_preparar,
            "entrada_almacen_alta": self.entrada_almacen_alta,
            "entrada_almacen_pdf_previsualizar": self.entrada_almacen_pdf_previsualizar,
            "entrada_almacen_desde_pdf": self.entrada_almacen_desde_pdf,
            "venta_tablas": self.venta_tablas,
            "venta_listar": self.venta_listar,
            "venta_obtener": self.venta_obtener,
            "venta_lineas_listar": self.venta_lineas_listar,
            "venta_precio_articulo": self.venta_precio_articulo,
            "rentabilidad_articulo_ventas": self.rentabilidad_articulo_ventas,
            "rentabilidad_articulos_resumen": self.rentabilidad_articulos_resumen,
            "ventas_documentos_detalle": self.ventas_documentos_detalle,
            "ventas_documentos_resumen": self.ventas_documentos_resumen,
            "venta_documento_alta_preparar": self.venta_documento_alta_preparar,
            "venta_documento_alta": self.venta_documento_alta,
            "venta_pedido_alta": self.venta_pedido_alta,
            "pedido_crear": self.pedido_crear,
            "pedido_listar": self.pedido_listar,
            "pedido_detalle": self.pedido_detalle,
            "pedido_cerrar": self.pedido_cerrar,
            "pedido_situacion_actualizar": self.pedido_situacion_actualizar,
            "pedido_retirada_actualizar": self.pedido_retirada_actualizar,
            "pedido_albaranar": self.pedido_albaranar,
            "pedido_marcar_preparado": self.pedido_marcar_preparado,
            "pedido_finalizar": self.pedido_finalizar,
            "pedido_linea_mover": self.pedido_linea_mover,
            "pedido_pdf_gestion": self.pedido_pdf_gestion,
            "pedido_enviar": self.pedido_enviar,
            "cartera_tablas": self.cartera_tablas,
            "cartera_efectos_detalle": self.cartera_efectos_detalle,
            "cartera_deuda_cliente": self.cartera_deuda_cliente,
            "cartera_deuda_por_cliente": self.cartera_deuda_por_cliente,
            "cartera_pendiente_remesar": self.cartera_pendiente_remesar,
            "cartera_deuda_por_tipo": self.cartera_deuda_por_tipo,
            "dashboard_resumen": self.dashboard_resumen,
            "ventas_resumen": self.ventas_resumen,
            "compras_resumen": self.compras_resumen,
            "dashboard_evolucion_anual": self.dashboard_evolucion_anual,
            "dashboard_series_temporales": self.dashboard_series_temporales,
            "dashboard_rankings": self.dashboard_rankings,
            "regularizacion_tablas": self.regularizacion_tablas,
            "regularizacion_listar": self.regularizacion_listar,
            "stock_por_almacen": self.stock_por_almacen,
            "stock_a_fecha": self.stock_a_fecha,
            "inventario_valorar_articulos": self.inventario_valorar_articulos,
            "articulo_regularizar": self.articulo_regularizar,
            "trasvase_generar": self.trasvase_generar,
            "recuento_listar": self.recuento_listar,
            "recuento_grabar": self.recuento_grabar,
            "recuento_borrar": self.recuento_borrar,
        }

    def invoke_tool(self, name: str, arguments: dict[str, Any], request_id: str | None = None) -> tuple[dict[str, Any], bool]:
        del request_id
        if name not in TOOL_PROFILES.get(self.tool_profile, TOOL_PROFILES["core"]):
            return self._error(name, KofedasError(f"Herramienta no disponible en perfil {self.tool_profile}")), True
        handler = self._handlers.get(name)
        if handler is None:
            return self._error(name, KofedasError(f"Herramienta desconocida: {name}")), True
        try:
            data = handler(arguments or {})
            return {
                "ok": True,
                "data": data,
                "warnings": [],
                "meta": self._meta(name),
            }, False
        except Exception as exc:
            return self._error(name, exc), True

    def _meta(self, name: str) -> dict[str, Any]:
        return {
            "tool": name,
            "profile": self.tool_profile,
            "contract_version": PUBLIC_CONTRACT_VERSION,
            "empresa": self.empresa,
            "centro": self.centro,
        }

    def _error(self, name: str, exc: Exception) -> dict[str, Any]:
        return {
            "ok": False,
            "error": {
                "code": "KOFEDAS_ERROR" if isinstance(exc, KofedasError) else "INTERNAL_ERROR",
                "message": str(exc),
            },
            "warnings": [],
            "meta": self._meta(name),
        }

    def _empresa(self, args: dict[str, Any]) -> int:
        return int(args.get("empresa") or self.empresa)

    def _centro(self, args: dict[str, Any]) -> int:
        return int(args.get("centro") if args.get("centro") is not None else self.centro)

    def _require_write(self) -> None:
        if ACCESS_LEVELS.get(self.access_level, 0) < ACCESS_LEVELS["write"]:
            raise KofedasError("La herramienta requiere KOFEDAS_MCP_ACCESS_LEVEL=write o critical")

    def _clean_data(self, data: Any, fields: list[str], prefix: str, forced: dict[str, Any]) -> dict[str, Any]:
        if data is None:
            data = {}
        if not isinstance(data, dict):
            raise KofedasError("datos debe ser un objeto")
        allowed = {field.upper(): field.upper() for field in fields}
        allowed.update({field.removeprefix(prefix).upper(): field.upper() for field in fields})
        cleaned: dict[str, Any] = {}
        unknown: list[str] = []
        for raw_key, value in data.items():
            key = str(raw_key).strip().upper()
            column = allowed.get(key)
            if not column:
                unknown.append(str(raw_key))
                continue
            cleaned[column] = value
        if unknown:
            raise KofedasError("Campos no permitidos: " + ", ".join(sorted(unknown)))
        cleaned.update(forced)
        return cleaned

    def _upsert(
        self,
        table: str,
        fields: list[str],
        key_fields: list[str],
        data: dict[str, Any],
    ) -> dict[str, Any]:
        where = " AND ".join(f"{field}=?" for field in key_fields)
        key_params = tuple(data[field] for field in key_fields)
        exists = self.db.one(f"SELECT FIRST 1 1 AS EXISTE FROM {table} WHERE {where}", key_params) is not None
        if exists:
            update_fields = [field for field in fields if field not in key_fields and field in data]
            if not update_fields:
                return {"accion": "sin_cambios", "tabla": table, "claves": {field.lower(): data[field] for field in key_fields}}
            sql = f"UPDATE {table} SET " + ", ".join(f"{field}=?" for field in update_fields) + f" WHERE {where}"
            params = tuple(data[field] for field in update_fields) + key_params
            affected = self.db.execute(sql, params)
            return {"accion": "actualizado", "tabla": table, "filas_afectadas": affected}
        insert_fields = [field for field in fields if field in data]
        missing_keys = [field for field in key_fields if field not in data]
        if missing_keys:
            raise KofedasError("Faltan claves: " + ", ".join(missing_keys))
        sql = f"INSERT INTO {table} (" + ", ".join(insert_fields) + ") VALUES (" + ", ".join("?" for _ in insert_fields) + ")"
        affected = self.db.execute(sql, tuple(data[field] for field in insert_fields))
        return {"accion": "insertado", "tabla": table, "filas_afectadas": affected}

    def _redact_usuario(self, row: dict[str, Any] | None) -> dict[str, Any] | None:
        if row is None:
            return None
        redacted = dict(row)
        if "usu_password" in redacted:
            redacted["usu_password"] = "***" if redacted["usu_password"] else ""
        return redacted

    def _aux_table(self, table: Any) -> str:
        name = str(table or "").strip().upper()
        if name not in AUXILIARY_TABLES:
            raise KofedasError("Tabla auxiliar no permitida: " + name)
        return name

    def _client_table(self, table: Any) -> str:
        name = str(table or "").strip().upper()
        if name not in CLIENT_TABLES:
            raise KofedasError("Tabla de cliente no permitida: " + name)
        return name

    def _provider_table(self, table: Any) -> str:
        name = str(table or "").strip().upper()
        if name not in PROVIDER_TABLES:
            raise KofedasError("Tabla de proveedor no permitida: " + name)
        return name

    def _article_table(self, table: Any) -> str:
        name = str(table or "").strip().upper()
        if name not in ARTICLE_TABLES:
            raise KofedasError("Tabla de articulo no permitida: " + name)
        return name

    def _offer_table(self, table: Any) -> str:
        name = str(table or "").strip().upper()
        if name not in OFFER_TABLES:
            raise KofedasError("Tabla de oferta no permitida: " + name)
        return name

    def _purchase_order_table(self, table: Any) -> str:
        name = str(table or "").strip().upper()
        if name not in PURCHASE_ORDER_TABLES:
            raise KofedasError("Tabla de orden de compra no permitida: " + name)
        return name

    def _table_columns(self, table: str) -> list[str]:
        columns = self.db.columns(table)
        if not columns:
            raise KofedasError("No se encontraron columnas para " + table)
        return columns

    def _table_pk(self, table: str) -> list[str]:
        keys = self.db.primary_key(table)
        if not keys:
            raise KofedasError("No se encontro clave primaria para " + table)
        return keys

    def _normalize_column_map(self, values: Any, columns: list[str], required: bool = False) -> dict[str, Any]:
        if values is None:
            if required:
                raise KofedasError("Se requiere un objeto de campos")
            return {}
        if not isinstance(values, dict):
            raise KofedasError("El valor debe ser un objeto")
        column_set = set(columns)
        result: dict[str, Any] = {}
        unknown: list[str] = []
        for raw_key, value in values.items():
            column = str(raw_key).strip().upper()
            if column not in column_set:
                unknown.append(str(raw_key))
                continue
            result[column] = value
        if unknown:
            raise KofedasError("Columnas no permitidas: " + ", ".join(sorted(unknown)))
        return result

    def _empresa_column(self, columns: list[str]) -> str | None:
        for column in columns:
            if column.endswith("_NUMEMP"):
                return column
        return None

    def _parameter_value(self, code: str, default: Any = "", empresa: int | None = None) -> Any:
        row = self.db.one(
            "SELECT PAR_VALOR FROM PARAMETROS WHERE PAR_NUMEMP = ? AND PAR_CODIGO = ?",
            (empresa or self.empresa, code),
        )
        return row["par_valor"] if row and row.get("par_valor") not in (None, "") else default

    def _cliente_field_data(self, values: Any) -> dict[str, Any]:
        columns = self._table_columns("CLIEN")
        aliases = {
            "CLIENTE": "CLI_CODCLI",
            "SUBCLIENTE": "CLI_SUBCLI",
            "NOMBRE": "CLI_NOMCLI",
            "RAZON_SOCIAL": "CLI_RAZSOC",
            "RAZON": "CLI_RAZSOC",
            "CIF": "CLI_CIF",
            "NIF": "CLI_CIF",
            "DOMICILIO": "CLI_DOMICI",
            "CODIGO_POSTAL": "CLI_CODPOS",
            "POBLACION": "CLI_POBLAC",
            "TELEFONO": "CLI_TELEFO",
            "EMAIL": "CLI_EMAIL",
            "FORMA_PAGO": "CLI_FORPAG",
            "TARIFA": "CLI_TIPPRE",
            "RIESGO": "CLI_RIESGO",
            "OBSERVACIONES": "CLI_OBSFOR",
        }
        if values is None:
            return {}
        if not isinstance(values, dict):
            raise KofedasError("datos debe ser un objeto")
        result: dict[str, Any] = {}
        unknown: list[str] = []
        for raw_key, value in values.items():
            key = str(raw_key).strip().upper()
            column = aliases.get(key, key)
            if column not in columns:
                unknown.append(str(raw_key))
                continue
            result[column] = value
        if unknown:
            raise KofedasError("Campos de cliente no permitidos: " + ", ".join(sorted(unknown)))
        return result

    def _cliente_defaults(self, empresa: int, cliente: int, subcliente: int, centro: int, data: dict[str, Any]) -> dict[str, Any]:
        today = date.today().isoformat()
        now = datetime.now().replace(microsecond=0).isoformat(sep=" ")
        riesgo = data.get("CLI_RIESGO", self._parameter_value("RIESGO", 0, empresa))
        tarifa = data.get("CLI_TIPPRE", self._parameter_value("TARIFA", "5", empresa))
        cuenta = data.get("CLI_CUECON", "" if cliente == 99999 else "430" + str(cliente).zfill(7))
        defaults = {
            "CLI_NUMEMP": empresa,
            "CLI_CODCLI": cliente,
            "CLI_SUBCLI": subcliente,
            "CLI_NOMCLI": "",
            "CLI_RAZSOC": "",
            "CLI_FEALTA": today,
            "CLI_FEBAJA": None,
            "CLI_FECCOM": today,
            "CLI_DOMICI": "",
            "CLI_CODPOS": 0,
            "CLI_POBLAC": "",
            "CLI_DOMENV": "",
            "CLI_CODPOSE": 0,
            "CLI_POBLACE": "",
            "CLI_CIF": "",
            "CLI_REGIVA": "N",
            "CLI_TELEFO": "",
            "CLI_FAX": "",
            "CLI_EMAIL": "",
            "CLI_CODBAN": 0,
            "CLI_CODENT": 0,
            "CLI_DIGITO": 0,
            "CLI_CUENTA": "",
            "CLI_CONTAC": "",
            "CLI_CUECON": cuenta,
            "CLI_ZONA": 0,
            "CLI_CODREP": 0,
            "CLI_FORENV": 0,
            "CLI_FORPAG": 0 if cliente == 99999 else data.get("CLI_FORPAG", 0),
            "CLI_MESNOV1": 0,
            "CLI_MESNOV2": 0,
            "CLI_DIAPAG1": 0,
            "CLI_DIAPAG2": 0,
            "CLI_DIAPAG3": 0,
            "CLI_AGRUPAC": "S",
            "CLI_AGRUPAA": "S",
            "CLI_TIPFAC": "",
            "CLI_RIESGO": riesgo,
            "CLI_RIESGOA": 0,
            "CLI_CODMON": "E",
            "CLI_DTOESP": 0,
            "CLI_NUMFAC": 1,
            "CLI_TIPPRE": tarifa,
            "CLI_PORAUM": 0,
            "CLI_OBSINT": "",
            "CLI_OBSFOR": "",
            "CLI_FECMOD": now,
            "CLI_USUMOD": f"{centro} MCP",
        }
        defaults.update(data)
        if not defaults.get("CLI_RAZSOC"):
            defaults["CLI_RAZSOC"] = defaults.get("CLI_NOMCLI", "")
        return defaults

    def _next_cliente_code(self, empresa: int, centro: int, cliente: int | None, subcliente: int | None) -> tuple[int, int]:
        center = self.db.one(
            "SELECT CEN_CLIFIN, CEN_CLICFIN FROM CENTROS WHERE CEN_NUMEMP = ? AND CEN_CODCEN = ?",
            (empresa, centro),
        ) or {}
        client_limit = int(center.get("cen_clifin") or 99999)
        subclient_limit = int(center.get("cen_clicfin") or 999999)
        code = int(cliente or 0)
        sub = int(subcliente or 0)
        if code == 0:
            row = self.db.one(
                "SELECT MAX(CLI_CODCLI) AS MAXIMO FROM CLIEN WHERE CLI_NUMEMP = ? AND CLI_CODCLI < ? AND CLI_CODCLI < 99999",
                (empresa, client_limit),
            )
            code = int((row or {}).get("maximo") or 0) + 1
        if code == 99999 and sub == 0:
            row = self.db.one(
                "SELECT MAX(CLI_SUBCLI) AS MAXIMO FROM CLIEN WHERE CLI_NUMEMP = ? AND CLI_CODCLI = 99999 AND CLI_SUBCLI < ?",
                (empresa, subclient_limit),
            )
            sub = int((row or {}).get("maximo") or 0) + 1
        if sub != 0 and code != 99999:
            principal = self.db.one(
                "SELECT FIRST 1 1 AS EXISTE FROM CLIEN WHERE CLI_NUMEMP = ? AND CLI_CODCLI = ? AND CLI_SUBCLI = 0",
                (empresa, code),
            )
            if principal is None:
                raise KofedasError("No existe cliente principal con subcliente 0")
        return code, sub

    def _cliente_exists(self, empresa: int, cliente: int, subcliente: int) -> bool:
        return self.db.one(
            "SELECT FIRST 1 1 AS EXISTE FROM CLIEN WHERE CLI_NUMEMP = ? AND CLI_CODCLI = ? AND CLI_SUBCLI = ?",
            (empresa, cliente, subcliente),
        ) is not None

    def _proveedor_field_data(self, values: Any) -> dict[str, Any]:
        columns = self._table_columns("PROVEE")
        aliases = {
            "PROVEEDOR": "PRO_CODPRO",
            "NOMBRE": "PRO_NOMCOR",
            "NOMBRE_COMERCIAL": "PRO_NOMCOR",
            "NOMBRE_FISCAL": "PRO_NOMFIS",
            "RAZON_SOCIAL": "PRO_NOMFIS",
            "ABREVIADO": "PRO_NOMABR",
            "CIF": "PRO_CIF",
            "NIF": "PRO_CIF",
            "DOMICILIO": "PRO_DOMICI",
            "CODIGO_POSTAL": "PRO_CODPOS",
            "POBLACION": "PRO_POBLAC",
            "TELEFONO": "PRO_TELEFO",
            "FAX": "PRO_FAX",
            "EMAIL": "PRO_EMAIL",
            "CONTACTO": "PRO_CONTAC",
            "REPRESENTANTE": "PRO_CODREP",
            "MONEDA": "PRO_CODMON",
            "PORTES": "PRO_PORTES",
            "FORMA_PAGO": "PRO_CODPAG",
            "CUENTA_CONTABLE": "PRO_CUECON",
            "OBSERVACIONES_INTERNAS": "PRO_OBSINT",
            "OBSERVACIONES": "PRO_OBSFOR",
        }
        if values is None:
            return {}
        if not isinstance(values, dict):
            raise KofedasError("datos debe ser un objeto")
        result: dict[str, Any] = {}
        unknown: list[str] = []
        for raw_key, value in values.items():
            key = str(raw_key).strip().upper()
            column = aliases.get(key, key)
            if column not in columns:
                unknown.append(str(raw_key))
                continue
            result[column] = value
        if unknown:
            raise KofedasError("Campos de proveedor no permitidos: " + ", ".join(sorted(unknown)))
        return result

    def _proveedor_defaults(self, empresa: int, proveedor: int, centro: int, data: dict[str, Any]) -> dict[str, Any]:
        today = date.today().isoformat()
        now = datetime.now().replace(microsecond=0).isoformat(sep=" ")
        defaults = {
            "PRO_NUMEMP": empresa,
            "PRO_CODPRO": proveedor,
            "PRO_NOMFIS": "",
            "PRO_NOMCOR": "",
            "PRO_NOMABR": "",
            "PRO_DOMICI": "",
            "PRO_CODPOS": 0,
            "PRO_POBLAC": "",
            "PRO_CIF": "",
            "PRO_TELEFO": "",
            "PRO_FAX": "",
            "PRO_EMAIL": "",
            "PRO_CONTAC": "",
            "PRO_CODREP": 0,
            "PRO_NOMDEL": "",
            "PRO_DOMDEL": "",
            "PRO_CODPOSD": 0,
            "PRO_POBDEL": "",
            "PRO_TELDEL": "",
            "PRO_FAXDEL": "",
            "PRO_CONDEL": "",
            "PRO_COMMIN": "",
            "PRO_CODMON": "E",
            "PRO_PORTES": "N",
            "PRO_CODPAG": 0,
            "PRO_CUECON": "400" + str(proveedor).zfill(7),
            "PRO_OBSINT": "",
            "PRO_OBSFOR": "",
            "PRO_FECMOD": now,
            "PRO_USUMOD": f"{centro} MCP",
            "PRO_FEALTA": today,
            "PRO_FEBAJA": None,
        }
        defaults.update(data)
        if not defaults.get("PRO_NOMFIS"):
            defaults["PRO_NOMFIS"] = defaults.get("PRO_NOMCOR", "")
        if not defaults.get("PRO_NOMCOR"):
            defaults["PRO_NOMCOR"] = defaults.get("PRO_NOMFIS", "")
        if not defaults.get("PRO_NOMABR"):
            defaults["PRO_NOMABR"] = str(defaults.get("PRO_NOMCOR") or defaults.get("PRO_NOMFIS") or "")[:15]
        return defaults

    def _next_proveedor_code(self, empresa: int, proveedor: int | None) -> int:
        code = int(proveedor or 0)
        if code != 0:
            return code
        row = self.db.one(
            "SELECT MAX(PRO_CODPRO) AS MAXIMO FROM PROVEE WHERE PRO_NUMEMP = ?",
            (empresa,),
        )
        return int((row or {}).get("maximo") or 0) + 1

    def _proveedor_exists(self, empresa: int, proveedor: int) -> bool:
        return self.db.one(
            "SELECT FIRST 1 1 AS EXISTE FROM PROVEE WHERE PRO_NUMEMP = ? AND PRO_CODPRO = ?",
            (empresa, proveedor),
        ) is not None

    def _article_field_data(self, values: Any) -> dict[str, Any]:
        columns = self._table_columns("ARTICUL")
        aliases = {
            "ARTICULO": "ART_CODART",
            "CODIGO": "ART_CODART",
            "CODART": "ART_CODART",
            "DESCRIPCION": "ART_DESCRI",
            "DESCRI": "ART_DESCRI",
            "SECCION": "ART_SECCIO",
            "PROPIO": "ART_INDPROP",
            "BLISTER": "ART_INDBLI",
            "TIPO_PRECIO": "ART_TIPPRE",
            "INVENTARIABLE": "ART_INDINV",
            "OBSOLETO": "ART_OBSOL",
            "FAMILIA": "ART_CODFAM",
            "SUBFAMILIA": "ART_SUBFAM",
            "IVA": "ART_TIPIVA",
            "PRECIO_TARIFA": "ART_PRETAR",
            "PRECIO_BASE": "ART_PREBAS",
            "PRECIO_COSTE": "ART_PRECOS",
            "TABLA_PRECIO": "ART_TABPREC",
            "PRECIO_VENTA1": "ART_PREVEN1",
            "PRECIO_VENTA2": "ART_PREVEN2",
            "PRECIO_VENTA3": "ART_PREVEN3",
            "PRECIO_VENTA4": "ART_PREVEN4",
            "PVP": "ART_PVP",
            "MONEDA": "ART_CODMON",
            "UNIDAD": "ART_UNIMED",
            "CANTIDAD_PRECIO": "ART_CANPRE",
            "CANTIDAD_MINIMA": "ART_CANPMI",
            "PROVEEDOR": "ART_CODPRO",
            "OBSERVACIONES_INTERNAS": "ART_OBSINT",
            "OBSERVACIONES": "ART_OBSFOR",
            "AGRUPACION1": "ART_AGRUP1",
            "AGRUPACION2": "ART_AGRUP2",
            "AGRUPACION3": "ART_AGRUP3",
            "NORMA": "ART_NORMA",
        }
        if values is None:
            return {}
        if not isinstance(values, dict):
            raise KofedasError("datos debe ser un objeto")
        result: dict[str, Any] = {}
        unknown: list[str] = []
        for raw_key, value in values.items():
            key = str(raw_key).strip().upper()
            column = aliases.get(key, key)
            if column not in columns:
                unknown.append(str(raw_key))
                continue
            result[column] = value
        if unknown:
            raise KofedasError("Campos de articulo no permitidos: " + ", ".join(sorted(unknown)))
        return result

    def _article_purchase_data(self, values: Any) -> dict[str, Any]:
        columns = self._table_columns("ARTICULP")
        aliases = {
            "ARTICULO": "ARTP_CODART",
            "CODIGO": "ARTP_CODART",
            "CODART": "ARTP_CODART",
            "PROVEEDOR": "ARTP_CODPRO",
            "REFERENCIA": "ARTP_REFPRO",
            "REFERENCIA_PROVEEDOR": "ARTP_REFPRO",
            "REFPRO": "ARTP_REFPRO",
            "DESCRIPCION": "ARTP_DESCRI",
            "DESCRIPCION_PROVEEDOR": "ARTP_DESCRI",
            "UNIDAD": "ARTP_UNIMED",
            "UNIDAD_COMPRA": "ARTP_UNIMED",
            "CANTIDAD_COMPRA": "ARTP_CANCON",
            "CANTIDAD_VENTA": "ARTP_CANVEN",
            "UNIDADES_PAQUETE": "ARTP_UNIPAQ",
            "PRECIO": "ARTP_PREBAS",
            "PRECIO_COMPRA": "ARTP_PREBAS",
            "PRECIO_BASE": "ARTP_PREBAS",
            "DTO1": "ARTP_DTOAUM1",
            "DTO2": "ARTP_DTOAUM2",
            "DTO3": "ARTP_DTOAUM3",
            "DTO4": "ARTP_DTOAUM4",
            "DTO5": "ARTP_DTOAUM5",
            "DTO6": "ARTP_DTOAUM6",
            "MONEDA": "ARTP_CODMON",
            "CANTIDAD_PRECIO": "ARTP_CANPRE",
            "UBICACION": "ARTP_UBICA",
            "AMPUNIV": "ARTP_AMPUNIV",
            "AJUSTE": "ARTP_AJUSTE",
        }
        if values is None:
            return {}
        if not isinstance(values, dict):
            raise KofedasError("compra debe ser un objeto")
        result: dict[str, Any] = {}
        unknown: list[str] = []
        for raw_key, value in values.items():
            key = str(raw_key).strip().upper()
            column = aliases.get(key, key)
            if column not in columns:
                unknown.append(str(raw_key))
                continue
            result[column] = value
        if unknown:
            raise KofedasError("Campos de compra no permitidos: " + ", ".join(sorted(unknown)))
        return result

    def _article_defaults(self, empresa: int, articulo: str, centro: int, data: dict[str, Any]) -> dict[str, Any]:
        today = date.today().isoformat()
        now = datetime.now().replace(microsecond=0).isoformat(sep=" ")
        defaults = {
            "ART_NUMEMP": empresa,
            "ART_CODART": articulo,
            "ART_DESCRI": "",
            "ART_SECCIO": "",
            "ART_INDPROP": "N",
            "ART_INDBLI": "N",
            "ART_TIPPRE": "V",
            "ART_INDINV": "S",
            "ART_OBSOL": "N",
            "ART_FEALTA": today,
            "ART_FEBAJA": None,
            "ART_FECMOV": None,
            "ART_FECVAR": today,
            "ART_CODFAM": 0,
            "ART_SUBFAM": 0,
            "ART_TIPIVA": 1,
            "ART_PRETAR": 0,
            "ART_FECTAR": None,
            "ART_PREBAS": 0,
            "ART_DTOAUM1": 0,
            "ART_DTOAUM2": 0,
            "ART_DTOAUM3": 0,
            "ART_DTOAUM4": 0,
            "ART_DTOAUM5": 0,
            "ART_DTOAUM6": 0,
            "ART_DTOAUM7": 0,
            "ART_IMPFIJ": 0,
            "ART_PRECOS": 0,
            "ART_TABPREC": 0,
            "ART_PREVEN1": 0,
            "ART_PREVEN2": 0,
            "ART_PREVEN3": 0,
            "ART_PREVEN4": 0,
            "ART_PVP": 0,
            "ART_CODMON": "E",
            "ART_UNIMED": "UNID",
            "ART_CANPRE": 1,
            "ART_CANPMI": 0,
            "ART_CODPRO": 0,
            "ART_OBSINT": "",
            "ART_OBSFOR": "",
            "ART_FECMOD": now,
            "ART_USUMOD": f"{centro} MCP",
            "ART_AGRUP1": 0,
            "ART_AGRUP2": 0,
            "ART_AGRUP3": 0,
            "ART_NORMA": "",
        }
        defaults.update(data)
        defaults["ART_CODART"] = str(defaults["ART_CODART"]).strip()
        defaults["ART_DESCRI"] = str(defaults.get("ART_DESCRI") or "")[:50]
        return self._calculate_article_price(defaults, empresa)

    def _purchase_defaults(self, empresa: int, articulo: str, proveedor: int, data: dict[str, Any], article: dict[str, Any]) -> dict[str, Any]:
        defaults = {
            "ARTP_NUMEMP": empresa,
            "ARTP_CODART": articulo,
            "ARTP_CODPRO": proveedor,
            "ARTP_REFPRO": "",
            "ARTP_DESCRI": article.get("ART_DESCRI", ""),
            "ARTP_UNIMED": article.get("ART_UNIMED", "UNID"),
            "ARTP_CANCON": 1,
            "ARTP_CANVEN": 1,
            "ARTP_UNIPAQ": 1,
            "ARTP_PREBAS": article.get("ART_PREBAS", 0),
            "ARTP_DTOAUM1": 0,
            "ARTP_DTOAUM2": 0,
            "ARTP_DTOAUM3": 0,
            "ARTP_DTOAUM4": 0,
            "ARTP_DTOAUM5": 0,
            "ARTP_DTOAUM6": 0,
            "ARTP_CODMON": article.get("ART_CODMON", "E"),
            "ARTP_CANPRE": article.get("ART_CANPRE", 1),
            "ARTP_UBICA": "",
            "ARTP_AMPUNIV": "",
            "ARTP_AJUSTE": "",
        }
        defaults.update(data)
        defaults["ARTP_CODART"] = articulo
        defaults["ARTP_CODPRO"] = proveedor
        defaults["ARTP_DESCRI"] = str(defaults.get("ARTP_DESCRI") or "")[:50]
        return defaults

    def _calculate_article_price(self, article: dict[str, Any], empresa: int) -> dict[str, Any]:
        prebas = self._to_float(article.get("ART_PREBAS"), 0)
        cost = prebas
        for column in ("ART_DTOAUM1", "ART_DTOAUM2", "ART_DTOAUM3", "ART_DTOAUM4", "ART_DTOAUM5", "ART_DTOAUM6", "ART_DTOAUM7"):
            cost *= 1 + self._to_float(article.get(column), 0) / 100
        if article.get("ART_TIPPRE") == "V":
            cost += self._to_float(article.get("ART_IMPFIJ"), 0)
        article["ART_PRECOS"] = round(cost, int(self._parameter_value("NUMDEC", 2, empresa) or 2))
        return article

    def _article_exists(self, empresa: int, articulo: str) -> bool:
        return self.db.one(
            "SELECT FIRST 1 1 AS EXISTE FROM ARTICUL WHERE ART_NUMEMP = ? AND ART_CODART = ?",
            (empresa, articulo),
        ) is not None

    def _barcode_owner(self, empresa: int, barcode: str) -> str | None:
        row = self.db.one(
            "SELECT ARTC_CODART FROM ARTICULC WHERE ARTC_NUMEMP = ? AND ARTC_CODIGO = ?",
            (empresa, barcode),
        )
        return str(row["artc_codart"]) if row else None

    def _next_generated_article_code(self, empresa: int, seccion: str, proveedor: int, counter: int, digits: int, reserved: set[str]) -> tuple[str, int]:
        prefix = str(seccion or "").strip() + str(proveedor).zfill(4)
        suffix_len = max(1, int(digits or 13) - len(prefix))
        while True:
            code = prefix + str(counter).zfill(suffix_len)
            if code not in reserved and not self._article_exists(empresa, code):
                reserved.add(code)
                return code, counter + 1
            counter += 1

    def _to_float(self, value: Any, default: float = 0) -> float:
        if value in (None, ""):
            return default
        if isinstance(value, str):
            value = value.strip().replace(".", "").replace(",", ".") if "," in value else value.strip()
        try:
            return float(value)
        except (TypeError, ValueError):
            return default

    def _to_int(self, value: Any, default: int = 0) -> int:
        if value in (None, ""):
            return default
        try:
            return int(float(str(value).strip().replace(",", ".")))
        except (TypeError, ValueError):
            return default

    def _xlsx_rows(self, path: str, sheet: str | None, header_row: int, limit: int) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            return self._xlsx_rows_stdlib(path, sheet, header_row, limit)
        file_path = Path(path).expanduser()
        if not file_path.exists():
            raise KofedasError("No existe el archivo Excel: " + str(file_path))
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet] if sheet else workbook.active
            rows = worksheet.iter_rows(values_only=True)
            headers: list[str] = []
            result: list[dict[str, Any]] = []
            for row_number, row in enumerate(rows, start=1):
                if row_number < header_row:
                    continue
                if row_number == header_row:
                    headers = [self._header_key(value) for value in row]
                    continue
                if not any(value not in (None, "") for value in row):
                    continue
                item = {headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
                result.append(item)
                if len(result) >= limit:
                    break
            return result
        finally:
            workbook.close()

    def _xlsx_rows_stdlib(self, path: str, sheet: str | None, header_row: int, limit: int) -> list[dict[str, Any]]:
        file_path = Path(path).expanduser()
        if not file_path.exists():
            raise KofedasError("No existe el archivo Excel: " + str(file_path))
        ns = {
            "a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
        }
        with zipfile.ZipFile(file_path) as archive:
            shared: list[str] = []
            if "xl/sharedStrings.xml" in archive.namelist():
                root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
                for item in root.findall("a:si", ns):
                    shared.append("".join(node.text or "" for node in item.findall(".//a:t", ns)))
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            rel_targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels.findall("rel:Relationship", ns)}
            selected_rel = None
            for sheet_node in workbook.findall("a:sheets/a:sheet", ns):
                if sheet is None or sheet_node.attrib.get("name") == sheet:
                    selected_rel = sheet_node.attrib.get(f"{{{ns['r']}}}id")
                    break
            if selected_rel is None:
                raise KofedasError("Hoja Excel no encontrada: " + str(sheet))
            target = rel_targets[selected_rel].lstrip("/")
            sheet_path = target if target.startswith("xl/") else "xl/" + target
            root = ElementTree.fromstring(archive.read(sheet_path))
            raw_rows: dict[int, dict[int, Any]] = {}
            for row_node in root.findall(".//a:sheetData/a:row", ns):
                row_index = int(row_node.attrib["r"])
                values: dict[int, Any] = {}
                for cell in row_node.findall("a:c", ns):
                    ref = cell.attrib.get("r", "")
                    column_letters = re.sub(r"[^A-Z]", "", ref.upper())
                    column_index = 0
                    for char in column_letters:
                        column_index = column_index * 26 + ord(char) - ord("A") + 1
                    value_node = cell.find("a:v", ns)
                    inline_node = cell.find("a:is/a:t", ns)
                    value: Any = ""
                    if inline_node is not None:
                        value = inline_node.text or ""
                    elif value_node is not None:
                        value = value_node.text or ""
                        if cell.attrib.get("t") == "s":
                            value = shared[int(value)] if str(value).isdigit() and int(value) < len(shared) else ""
                    values[column_index] = value
                raw_rows[row_index] = values
            header_values = raw_rows.get(header_row, {})
            headers = {column: self._header_key(value) for column, value in header_values.items()}
            result: list[dict[str, Any]] = []
            for row_index in sorted(index for index in raw_rows if index > header_row):
                values = raw_rows[row_index]
                if not any(value not in (None, "") for value in values.values()):
                    continue
                item = {headers[column]: value for column, value in values.items() if column in headers and headers[column]}
                result.append(item)
                if len(result) >= limit:
                    break
            return result

    def _header_key(self, value: Any) -> str:
        text = str(value or "").strip().lower()
        replacements = str.maketrans("áéíóúüñ", "aeiouun")
        text = text.translate(replacements)
        return re.sub(r"[^a-z0-9]+", "_", text).strip("_")

    def _date_arg(self, value: Any, default: str | None = None) -> str:
        if value in (None, ""):
            return default or date.today().isoformat()
        if isinstance(value, (datetime, date)):
            return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
        text = str(value).strip()
        try:
            return date.fromisoformat(text[:10]).isoformat()
        except ValueError as exc:
            raise KofedasError("Fecha invalida, usa YYYY-MM-DD: " + text) from exc

    def _next_offer_number(self, empresa: int, ejercicio: int, requested: int | None = None) -> int:
        number = int(requested or 0)
        if number:
            return number
        row = self.db.one(
            "SELECT MAX(OFE_NUMOFE) AS MAXIMO FROM OFERTAS WHERE OFE_NUMEMP = ? AND OFE_EJERCI = ?",
            (empresa, ejercicio),
        )
        return int((row or {}).get("maximo") or 0) + 1

    def _offer_exists(self, empresa: int, ejercicio: int, oferta: int) -> bool:
        return self.db.one(
            "SELECT FIRST 1 1 AS EXISTE FROM OFERTAS WHERE OFE_NUMEMP = ? AND OFE_EJERCI = ? AND OFE_NUMOFE = ?",
            (empresa, ejercicio, oferta),
        ) is not None

    def _offer_article_code(self, empresa: int, proveedor: int, item: Any) -> str:
        if isinstance(item, str):
            code = item.strip()
            barcode = code if len(code) == 13 else ""
            reference = ""
        elif isinstance(item, dict):
            code = str(item.get("articulo") or item.get("codigo") or "").strip()
            barcode = str(item.get("codigo_barras") or item.get("ean") or "").strip()
            reference = str(item.get("referencia") or item.get("refpro") or item.get("referencia_proveedor") or "").strip()
        else:
            raise KofedasError("Cada articulo debe ser texto u objeto")
        if code and self._article_exists(empresa, code):
            return code
        if barcode:
            owner = self._barcode_owner(empresa, barcode)
            if owner:
                return owner
        if reference and proveedor:
            row = self.db.one(
                "SELECT FIRST 1 ARTP_CODART FROM ARTICULP WHERE ARTP_NUMEMP = ? AND ARTP_CODPRO = ? AND ARTP_REFPRO = ?",
                (empresa, proveedor, reference),
            )
            if row:
                return str(row["artp_codart"])
        if code and proveedor:
            row = self.db.one(
                "SELECT FIRST 1 ARTP_CODART FROM ARTICULP WHERE ARTP_NUMEMP = ? AND ARTP_CODPRO = ? AND ARTP_REFPRO = ?",
                (empresa, proveedor, code),
            )
            if row:
                return str(row["artp_codart"])
        raise KofedasError("Articulo no encontrado en oferta: " + (code or barcode or reference))

    def _offer_line_defaults(self, empresa: int, ejercicio: int, oferta: int, proveedor: int, fecini: str, fecfin: str, moneda: str, item: Any) -> dict[str, Any]:
        code = self._offer_article_code(empresa, proveedor, item)
        article = self.db.one(
            "SELECT FIRST 1 ART_CODART, ART_DESCRI, ART_PRECOS, ART_PVP, ART_CODMON, ART_CANPRE FROM ARTICUL WHERE ART_NUMEMP = ? AND ART_CODART = ?",
            (empresa, code),
        )
        if article is None:
            raise KofedasError("Articulo no encontrado: " + code)
        data = item if isinstance(item, dict) else {}
        return {
            "DOF_NUMEMP": empresa,
            "DOF_EJERCI": ejercicio,
            "DOF_NUMOFE": oferta,
            "DOF_CODART": code,
            "DOF_FECINI": self._date_arg(data.get("fecha_inicio"), fecini),
            "DOF_FECFIN": self._date_arg(data.get("fecha_fin"), fecfin),
            "DOF_CODPRO": self._to_int(data.get("proveedor"), proveedor),
            "DOF_PRECOS": self._to_float(data.get("precos", data.get("precio_coste")), self._to_float(article.get("art_precos"), 0)),
            "DOF_PRECIO": self._to_float(data.get("precio"), 0),
            "DOF_PVP": self._to_float(data.get("pvp"), self._to_float(article.get("art_pvp"), 0)),
            "DOF_CODMON": str(data.get("moneda") or moneda or article.get("art_codmon") or "E").strip()[:1],
            "DOF_CANPRE": self._to_float(data.get("canpre", data.get("cantidad_precio")), self._to_float(article.get("art_canpre"), 1)),
            "DOF_DTO1": self._to_float(data.get("dto1"), 0),
            "DOF_DTO2": self._to_float(data.get("dto2"), 0),
            "art_descri": article.get("art_descri"),
        }

    def _purchase_order_status(self, value: Any) -> str | None:
        if value in (None, ""):
            return None
        key = str(value).strip().upper()
        result = PURCHASE_ORDER_STATUS.get(key)
        if not result:
            raise KofedasError("Estado de orden de compra no valido: " + str(value))
        return result

    def _next_purchase_order_number(self, empresa: int, centro: int, ejercicio: int, serie: str, requested: int | None = None) -> int:
        number = int(requested or 0)
        if number:
            return number
        row = self.db.one(
            """
            SELECT MAX(COC_NUMDOC) AS MAXIMO
            FROM CABORC
            WHERE COC_NUMEMP = ? AND COC_CENTRO = ? AND COC_EJERCI = ? AND COC_SERIE = ?
            """,
            (empresa, centro, ejercicio, serie),
        )
        return int((row or {}).get("maximo") or 0) + 1

    def _purchase_order_exists(self, empresa: int, centro: int, ejercicio: int, serie: str, numero: int) -> bool:
        return self.db.one(
            """
            SELECT FIRST 1 1 AS EXISTE
            FROM CABORC
            WHERE COC_NUMEMP = ? AND COC_CENTRO = ? AND COC_EJERCI = ? AND COC_SERIE = ? AND COC_NUMDOC = ?
            """,
            (empresa, centro, ejercicio, serie, numero),
        ) is not None

    def _purchase_order_article_code(self, empresa: int, proveedor: int, item: Any) -> str:
        if isinstance(item, str):
            code = item.strip()
            barcode = code if len(code) == 13 else ""
            reference = ""
        elif isinstance(item, dict):
            code = str(item.get("articulo") or item.get("codigo") or "").strip()
            barcode = str(item.get("codigo_barras") or item.get("ean") or "").strip()
            reference = str(item.get("referencia") or item.get("refpro") or item.get("referencia_proveedor") or "").strip()
        else:
            raise KofedasError("Cada articulo debe ser texto u objeto")
        if code and self._article_exists(empresa, code):
            return code
        if barcode:
            owner = self._barcode_owner(empresa, barcode)
            if owner:
                return owner
        if reference and proveedor:
            row = self.db.one(
                "SELECT FIRST 1 ARTP_CODART FROM ARTICULP WHERE ARTP_NUMEMP = ? AND ARTP_CODPRO = ? AND ARTP_REFPRO = ?",
                (empresa, proveedor, reference),
            )
            if row:
                return str(row["artp_codart"])
        if code and proveedor:
            row = self.db.one(
                "SELECT FIRST 1 ARTP_CODART FROM ARTICULP WHERE ARTP_NUMEMP = ? AND ARTP_CODPRO = ? AND ARTP_REFPRO = ?",
                (empresa, proveedor, code),
            )
            if row:
                return str(row["artp_codart"])
        raise KofedasError("Articulo no encontrado en orden de compra: " + (code or barcode or reference))

    def _purchase_order_line_defaults(
        self,
        empresa: int,
        centro: int,
        ejercicio: int,
        serie: str,
        numero: int,
        linea: int,
        proveedor: int,
        fecha: str,
        moneda: str,
        item: Any,
    ) -> dict[str, Any]:
        code = self._purchase_order_article_code(empresa, proveedor, item)
        data = item if isinstance(item, dict) else {}
        article = self.db.one(
            """
            SELECT FIRST 1 ART_CODART, ART_DESCRI, ART_PRECOS, ART_CODMON,
                   ART_UNIMED, ART_TIPIVA
            FROM ARTICUL
            WHERE ART_NUMEMP = ? AND ART_CODART = ?
            """,
            (empresa, code),
        )
        if article is None:
            raise KofedasError("Articulo no encontrado: " + code)
        purchase = self.db.one(
            """
            SELECT FIRST 1 ARTP_REFPRO, ARTP_DESCRI, ARTP_UNIMED, ARTP_PREBAS,
                   ARTP_DTOAUM1, ARTP_DTOAUM2, ARTP_DTOAUM3, ARTP_DTOAUM4,
                   ARTP_DTOAUM5, ARTP_DTOAUM6, ARTP_CODMON
            FROM ARTICULP
            WHERE ARTP_NUMEMP = ? AND ARTP_CODART = ? AND ARTP_CODPRO = ?
            ORDER BY ARTP_REFPRO
            """,
            (empresa, code, proveedor),
        ) or {}
        quantity = self._to_float(data.get("cantidad", data.get("cantid")), 1)
        if quantity <= 0:
            raise KofedasError("La cantidad debe ser mayor que cero para " + code)
        base_price = self._to_float(
            data.get("precio", data.get("prebas", data.get("precio_base"))),
            self._to_float(purchase.get("artp_prebas"), self._to_float(article.get("art_precos"), 0)),
        )
        discounts = [
            self._to_float(data.get(f"dto{index}", data.get(f"descuento{index}")), self._to_float(purchase.get(f"artp_dtoaum{index}"), 0))
            for index in range(1, 7)
        ]
        net_price = base_price
        for discount in discounts:
            net_price *= 1 - (discount / 100)
        pending_quantity = self._to_float(data.get("cantidad_pendiente", data.get("canpen")), quantity)
        line_value = net_price * quantity
        pending_value = net_price * pending_quantity
        tax_percent = self._to_float(data.get("poriva", data.get("iva")), 0)
        if not tax_percent and article.get("art_tipiva") not in (None, ""):
            tax = self.db.one(
                "SELECT FIRST 1 TIV_PORIVA FROM TIPIVA WHERE TIV_NUMEMP = ? AND TIV_TIPIVA = ?",
                (empresa, int(article.get("art_tipiva"))),
            )
            if tax:
                tax_percent = self._to_float(tax.get("tiv_poriva"), 0)
        return {
            "DOC_NUMEMP": empresa,
            "DOC_CENTRO": centro,
            "DOC_EJERCI": ejercicio,
            "DOC_SERIE": serie,
            "DOC_NUMDOC": numero,
            "DOC_NUMLIN": linea,
            "DOC_FECMOV": self._date_arg(data.get("fecha"), fecha),
            "DOC_TIPLIN": str(data.get("tipo_linea") or data.get("tiplin") or "D").strip()[:1] or "D",
            "DOC_CODART": code,
            "DOC_CODARTP": str(data.get("referencia") or data.get("refpro") or purchase.get("artp_refpro") or "")[:20],
            "DOC_DESCRI": str(data.get("descripcion") or purchase.get("artp_descri") or article.get("art_descri") or "")[:50],
            "DOC_CANTID": quantity,
            "DOC_UNIMED": str(data.get("unidad") or purchase.get("artp_unimed") or article.get("art_unimed") or "")[:4],
            "DOC_PREBAS": base_price,
            "DOC_CODMON": str(data.get("moneda") or moneda or purchase.get("artp_codmon") or article.get("art_codmon") or "E").strip()[:1] or "E",
            "DOC_DTOAUM1": discounts[0],
            "DOC_DTOAUM2": discounts[1],
            "DOC_DTOAUM3": discounts[2],
            "DOC_DTOAUM4": discounts[3],
            "DOC_DTOAUM5": discounts[4],
            "DOC_DTOAUM6": discounts[5],
            "DOC_PORIVA": tax_percent,
            "DOC_PORREQ": self._to_float(data.get("porreq", data.get("recargo")), 0),
            "DOC_VALLIN": line_value,
            "DOC_INDOFE": str(data.get("indofe") or "N").strip()[:1] or "N",
            "DOC_CANPEN": pending_quantity,
            "DOC_VALPEN": pending_value,
            "DOC_SITUAC": str(data.get("estado") or "A").strip().upper()[:1] or "A",
            "DOC_CIEMAN": str(data.get("cierre_manual") or data.get("cieman") or "N").strip().upper()[:1] or "N",
            "DOC_OBSERV": str(data.get("observaciones") or data.get("observ") or "")[:60],
            "art_descri": article.get("art_descri"),
        }

    def _purchase_order_totals(self, lines: list[dict[str, Any]]) -> tuple[float, float]:
        ordered = sum(self._to_float(line.get("DOC_VALLIN"), 0) for line in lines)
        pending = sum(self._to_float(line.get("DOC_VALPEN"), 0) for line in lines)
        return ordered, pending

    def _warehouse_entry_status(self, value: Any) -> str | None:
        if value in (None, ""):
            return None
        key = str(value).strip().upper().replace(" ", "_").replace("-", "_")
        result = WAREHOUSE_ENTRY_STATUS.get(key)
        if not result:
            raise KofedasError("Situacion de entrada no valida: " + str(value))
        return result

    def _next_warehouse_entry_number(self, empresa: int, centro: int, ejercicio: int, serie: str, requested: int | None = None) -> int:
        number = int(requested or 0)
        if number:
            return number
        row = self.db.one(
            """
            SELECT MAX(CBM_NUMDOC) AS MAXIMO
            FROM CABDOCM
            WHERE CBM_NUMEMP = ? AND CBM_CENTRO = ? AND CBM_EJERCI = ? AND CBM_SERIE = ?
            """,
            (empresa, centro, ejercicio, serie),
        )
        return int((row or {}).get("maximo") or 0) + 1

    def _warehouse_entry_exists(self, empresa: int, centro: int, ejercicio: int, serie: str, numero: int) -> bool:
        return self.db.one(
            """
            SELECT FIRST 1 1 AS EXISTE
            FROM CABDOCM
            WHERE CBM_NUMEMP = ? AND CBM_CENTRO = ? AND CBM_EJERCI = ? AND CBM_SERIE = ? AND CBM_NUMDOC = ?
            """,
            (empresa, centro, ejercicio, serie, numero),
        ) is not None

    def _warehouse_entry_provider(self, empresa: int, args: dict[str, Any]) -> dict[str, Any]:
        provider = args.get("proveedor")
        if provider not in (None, "", 0):
            row = self.db.one(
                """
                SELECT FIRST 1 PRO_CODPRO, PRO_NOMCOR, PRO_NOMFIS, PRO_DOMICI, PRO_CODPOS,
                       PRO_POBLAC, PRO_CIF, PRO_CODPAG, PRO_CODMON
                FROM PROVEE
                WHERE PRO_NUMEMP = ? AND PRO_CODPRO = ?
                """,
                (empresa, int(provider)),
            )
        elif str(args.get("cif") or "").strip():
            row = self.db.one(
                """
                SELECT FIRST 1 PRO_CODPRO, PRO_NOMCOR, PRO_NOMFIS, PRO_DOMICI, PRO_CODPOS,
                       PRO_POBLAC, PRO_CIF, PRO_CODPAG, PRO_CODMON
                FROM PROVEE
                WHERE PRO_NUMEMP = ? AND PRO_CIF = ?
                """,
                (empresa, str(args.get("cif")).strip()),
            )
        else:
            row = None
        if row is None:
            raise KofedasError("Proveedor no encontrado para la entrada de almacen")
        return row

    def _warehouse_entry_article_code(self, empresa: int, proveedor: int, item: dict[str, Any]) -> str:
        code = str(item.get("articulo") or item.get("codigo") or "").strip()
        barcode = str(item.get("codigo_barras") or item.get("ean") or "").strip()
        reference = str(item.get("referencia") or item.get("referencia_proveedor") or item.get("refpro") or "").strip()
        if code and self._article_exists(empresa, code):
            return code
        if barcode:
            owner = self._barcode_owner(empresa, barcode)
            if owner:
                return owner
        for candidate in (reference, code):
            if candidate and proveedor:
                row = self.db.one(
                    """
                    SELECT FIRST 1 ARTP_CODART
                    FROM ARTICULP
                    WHERE ARTP_NUMEMP = ? AND ARTP_CODPRO = ? AND ARTP_REFPRO = ?
                    """,
                    (empresa, proveedor, candidate),
                )
                if row:
                    return str(row["artp_codart"])
        raise KofedasError("Articulo no encontrado en entrada: " + (code or barcode or reference))

    def _warehouse_entry_line_defaults(
        self,
        empresa: int,
        centro: int,
        ejercicio: int,
        serie: str,
        numero: int,
        linea: int,
        proveedor: int,
        fecha: str,
        moneda: str,
        item: dict[str, Any],
    ) -> dict[str, Any]:
        code = self._warehouse_entry_article_code(empresa, proveedor, item)
        article = self.db.one(
            """
            SELECT FIRST 1 ART_CODART, ART_DESCRI, ART_UNIMED, ART_CODMON, ART_TIPIVA, ART_INDINV
            FROM ARTICUL
            WHERE ART_NUMEMP = ? AND ART_CODART = ?
            """,
            (empresa, code),
        )
        if article is None:
            raise KofedasError("Articulo no encontrado: " + code)
        purchase = self.db.one(
            """
            SELECT FIRST 1 ARTP_REFPRO, ARTP_DESCRI, ARTP_UNIMED, ARTP_CANCON, ARTP_CANVEN,
                   ARTP_PREBAS, ARTP_DTOAUM1, ARTP_DTOAUM2, ARTP_DTOAUM3, ARTP_DTOAUM4,
                   ARTP_DTOAUM5, ARTP_DTOAUM6, ARTP_CODMON
            FROM ARTICULP
            WHERE ARTP_NUMEMP = ? AND ARTP_CODART = ? AND ARTP_CODPRO = ?
            ORDER BY ARTP_REFPRO
            """,
            (empresa, code, proveedor),
        ) or {}
        quantity_provider = self._to_float(item.get("cantidad_proveedor", item.get("cantidad", item.get("cantid"))), 0)
        if quantity_provider <= 0:
            raise KofedasError("La cantidad debe ser mayor que cero para " + code)
        cancon = self._to_float(purchase.get("artp_cancon"), 1) or 1
        canven = self._to_float(purchase.get("artp_canven"), 1) or 1
        quantity_stock = self._to_float(item.get("cantidad_stock", item.get("cantidad_interna")), quantity_provider * canven / cancon)
        price = self._to_float(item.get("precio", item.get("prebas", item.get("precio_base"))), self._to_float(purchase.get("artp_prebas"), 0))
        discounts = [
            self._to_float(item.get(f"dto{index}", item.get(f"descuento{index}")), self._to_float(purchase.get(f"artp_dtoaum{index}"), 0))
            for index in range(1, 7)
        ]
        tax_percent = self._to_float(item.get("iva", item.get("poriva")), 0)
        if not tax_percent and article.get("art_tipiva") not in (None, ""):
            tax = self.db.one(
                "SELECT FIRST 1 TIV_PORIVA FROM TIPIVA WHERE TIV_NUMEMP = ? AND TIV_TIPIVA = ?",
                (empresa, int(article.get("art_tipiva"))),
            )
            if tax:
                tax_percent = self._to_float(tax.get("tiv_poriva"), 0)
        net = price
        for discount in discounts:
            net *= 1 - discount / 100
        return {
            "DMM_NUMEMP": empresa,
            "DMM_CENTRO": centro,
            "DMM_EJERCI": ejercicio,
            "DMM_SERIE": serie,
            "DMM_NUMDOC": numero,
            "DMM_NUMLIN": self._to_int(item.get("linea"), linea),
            "DMM_FECMOV": self._date_arg(item.get("fecha"), fecha),
            "DMM_TIPLIN": str(item.get("tipo_linea") or "D").strip()[:1] or "D",
            "DMM_CODART": code,
            "DMM_CODARP": str(item.get("referencia") or item.get("referencia_proveedor") or purchase.get("artp_refpro") or "")[:20],
            "DMM_DESCRI": str(item.get("descripcion") or purchase.get("artp_descri") or article.get("art_descri") or "")[:100],
            "DMM_CANTIDP": quantity_provider,
            "DMM_UNIMED": str(item.get("unidad") or item.get("unidad_medida") or purchase.get("artp_unimed") or article.get("art_unimed") or "UNID")[:4],
            "DMM_CANTID": quantity_stock,
            "DMM_PREBAS": price,
            "DMM_CODMON": str(item.get("moneda") or moneda or purchase.get("artp_codmon") or article.get("art_codmon") or "E").strip()[:1] or "E",
            "DMM_DTOAUM1": discounts[0],
            "DMM_DTOAUM2": discounts[1],
            "DMM_DTOAUM3": discounts[2],
            "DMM_DTOAUM4": discounts[3],
            "DMM_DTOAUM5": discounts[4],
            "DMM_DTOAUM6": discounts[5],
            "DMM_PORIVA": tax_percent,
            "DMM_PORREQ": self._to_float(item.get("recargo", item.get("porreq")), 0),
            "DMM_VALLIN": net * quantity_provider,
            "DMM_IMPDTO": 0,
            "DMM_OBSERV": str(item.get("observaciones") or item.get("observ") or "")[:40],
            "DMM_EJERCIP": self._to_int(item.get("pedido_ejercicio", item.get("ejercicio_pedido")), 0),
            "DMM_SERIEP": str(item.get("pedido_serie", item.get("serie_pedido")) or "")[:2],
            "DMM_NUMDOCP": self._to_int(item.get("pedido_numero", item.get("numero_pedido")), 0),
            "DMM_NUMLINP": self._to_int(item.get("pedido_linea", item.get("linea_pedido")), 0),
            "DMM_EJEOFE": self._to_int(item.get("oferta_ejercicio"), 0),
            "DMM_NUMOFE": self._to_int(item.get("oferta_numero"), 0),
            "art_indinv": str(article.get("art_indinv") or "S").strip().upper(),
        }

    def _warehouse_entry_totals(self, header: dict[str, Any], lines: list[dict[str, Any]]) -> dict[str, Any]:
        buckets: list[dict[str, float]] = []
        for line in lines:
            if str(line.get("DMM_TIPLIN") or "").upper() == "C":
                continue
            iva = self._to_float(line.get("DMM_PORIVA"), 0)
            req = self._to_float(line.get("DMM_PORREQ"), 0)
            bucket = next((item for item in buckets if item["iva"] == iva and item["req"] == req), None)
            if bucket is None and len(buckets) < 4:
                bucket = {"base": 0.0, "iva": iva, "req": req}
                buckets.append(bucket)
            if bucket is not None:
                bucket["base"] += self._to_float(line.get("DMM_VALLIN"), 0)
        while len(buckets) < 4:
            buckets.append({"base": 0.0, "iva": 0.0, "req": 0.0})
        discount = self._to_float(header.get("CBM_PORDTO"), 0)
        shipping = self._to_float(header.get("CBM_IMPPOR"), 0)
        total_base = 0.0
        total_tax = 0.0
        total_surcharge = 0.0
        for index, bucket in enumerate(buckets[:4], start=1):
            header[f"CBM_BASIMP{index}"] = round(bucket["base"], 4)
            header[f"CBM_PORIVA{index}"] = bucket["iva"]
            header[f"CBM_PORREQ{index}"] = bucket["req"]
            taxable_base = bucket["base"] * (1 - discount / 100) + (shipping if index == 1 else 0)
            total_base += taxable_base
            total_tax += taxable_base * bucket["iva"] / 100
            total_surcharge += taxable_base * bucket["req"] / 100
        header["CBM_TOTALS"] = round(total_base, 2)
        header["CBM_TOTALD"] = round(total_base + total_tax + total_surcharge, 2)
        return header

    def _pdf_text(self, args: dict[str, Any]) -> tuple[str, dict[str, Any]]:
        data: bytes
        source = ""
        if args.get("content_base64"):
            try:
                data = base64.b64decode(str(args.get("content_base64")), validate=True)
            except Exception as exc:
                raise KofedasError("content_base64 no es Base64 valido") from exc
            source = str(args.get("nombre_fichero") or "entrada.pdf")
        elif args.get("ruta_pdf"):
            path = Path(str(args.get("ruta_pdf"))).expanduser()
            if not path.exists():
                raise KofedasError("No existe el PDF: " + str(path))
            data = path.read_bytes()
            source = str(path)
        else:
            raise KofedasError("Debe informar ruta_pdf o content_base64")
        if not data.startswith(b"%PDF"):
            raise KofedasError("El contenido no parece un PDF")
        text = ""
        try:
            from pypdf import PdfReader
            import io
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError as exc:
            raise KofedasError("Para extraer PDF instala pypdf en el entorno del MCP") from exc
        except Exception as exc:
            raise KofedasError("No se pudo extraer texto del PDF: " + str(exc)) from exc
        return text, {"origen": source, "bytes": len(data)}

    def _warehouse_entry_pdf_proposal(self, args: dict[str, Any]) -> dict[str, Any]:
        text, meta = self._pdf_text(args)
        empresa = self._empresa(args)
        provider_code = self._to_int(args.get("proveedor"), 0)
        cif_match = re.search(r"\b(?:CIF|NIF|VAT)\s*[:\-]?\s*([A-Z0-9][A-Z0-9 .\-]{6,15})", text, re.IGNORECASE)
        invoice_match = re.search(r"\b(?:FACTURA|FRA\.?|INVOICE)\s*(?:N[ºO]\.?)?\s*[:\-]?\s*([A-Z0-9][A-Z0-9./\-]{2,25})", text, re.IGNORECASE)
        delivery_match = re.search(r"\b(?:ALBAR[AÁ]N|DELIVERY\s*NOTE)\s*(?:N[ºO]\.?)?\s*[:\-]?\s*([A-Z0-9][A-Z0-9./\-]{2,25})", text, re.IGNORECASE)
        date_match = re.search(r"\b(\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b", text)
        parsed_pdf_date = date.today().isoformat()
        if date_match:
            raw_date = date_match.group(1)
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
                try:
                    parsed_pdf_date = datetime.strptime(raw_date, fmt).date().isoformat()
                    break
                except ValueError:
                    pass
        cabecera: dict[str, Any] = {
            "centro": self._centro(args),
            "proveedor": provider_code or None,
            "cif": cif_match.group(1).replace(" ", "").replace("-", "").strip().upper() if cif_match else "",
            "factura": invoice_match.group(1).strip() if invoice_match else "",
            "albaran": delivery_match.group(1).strip() if delivery_match else "",
            "fecha": args.get("fecha") or parsed_pdf_date,
            "observaciones": "Entrada propuesta desde PDF",
        }
        candidates: list[dict[str, Any]] = []
        number = r"[-+]?\d+(?:[.,]\d+)?"
        line_re = re.compile(rf"^\s*([A-Z0-9][A-Z0-9./_-]{{2,20}})\s+(.+?)\s+({number})\s+({number})(?:\s+({number}))?\s*$")
        for raw_line in text.splitlines():
            clean = re.sub(r"\s+", " ", raw_line).strip()
            if len(clean) < 8:
                continue
            match = line_re.match(clean)
            if not match:
                continue
            reference, description, quantity, price, tax = match.groups()
            item: dict[str, Any] = {
                "referencia": reference,
                "descripcion": description[:100],
                "cantidad": self._to_float(quantity, 0),
                "precio": self._to_float(price, 0),
            }
            if tax:
                parsed_tax = self._to_float(tax, 0)
                if parsed_tax in (0, 4, 5, 10, 21):
                    item["iva"] = parsed_tax
            if provider_code:
                try:
                    item["articulo"] = self._warehouse_entry_article_code(empresa, provider_code, item)
                    item["resuelto"] = True
                except Exception as exc:
                    item["resuelto"] = False
                    item["error"] = str(exc)
            else:
                item["resuelto"] = False
            candidates.append(item)
            if len(candidates) >= _positive_limit(args.get("limite_lineas"), 100):
                break
        return {
            "pdf": meta,
            "texto_extraido_caracteres": len(text),
            "cabecera": cabecera,
            "lineas": candidates,
            "advertencias": [
                "La extraccion de PDF es heuristica: revisa proveedor, albaran/factura, cantidades y precios antes de grabar.",
            ],
        }

    def _regularization_series(self, empresa: int, centro: int, value: Any = None) -> str:
        series = str(value or self._parameter_value(f"R{centro}", "", empresa) or self._parameter_value("R", "", empresa) or "").strip()[:2]
        if not series:
            raise KofedasError("No existe parametro de serie para regularizaciones: R<centro>/R; informa serie")
        return series

    def _next_regularization_number(self, empresa: int, centro: int, ejercicio: int, serie: str) -> int:
        row = self.db.one(
            """
            SELECT MAX(CBR_NUMDOC) AS MAXIMO
            FROM CABDOCR
            WHERE CBR_NUMEMP = ? AND CBR_CENTRO = ? AND CBR_EJERCI = ? AND CBR_SERIE = ?
            """,
            (empresa, centro, ejercicio, serie),
        )
        return int((row or {}).get("maximo") or 0) + 1

    def _next_regularization_line(self, empresa: int, centro: int, ejercicio: int, serie: str, numero: int) -> int:
        row = self.db.one(
            """
            SELECT MAX(DMR_NUMLIN) AS MAXIMO
            FROM DETMOVR
            WHERE DMR_NUMEMP = ? AND DMR_CENTRO = ? AND DMR_EJERCI = ? AND DMR_SERIE = ? AND DMR_NUMDOC = ?
            """,
            (empresa, centro, ejercicio, serie, numero),
        )
        return int((row or {}).get("maximo") or 0) + 1

    def _regularization_header(self, empresa: int, centro: int, fecha: str, tipo: str, serie: str, observ: str, cenrel: int = 0, numdoce: int = 0) -> dict[str, Any]:
        ejercicio = int(fecha[:4])
        numero = self._next_regularization_number(empresa, centro, ejercicio, serie)
        return {
            "CBR_NUMEMP": empresa,
            "CBR_CENTRO": centro,
            "CBR_EJERCI": ejercicio,
            "CBR_SERIE": serie,
            "CBR_NUMDOC": numero,
            "CBR_FECHA": fecha,
            "CBR_TIPO": tipo,
            "CBR_CENREL": cenrel,
            "CBR_FECINF": None,
            "CBR_FECSUP": None,
            "CBR_OBSERV": str(observ or "")[:50],
            "CBR_NUMDOCE": numdoce,
            "CBR_FECMOD": datetime.now().replace(microsecond=0).isoformat(sep=" "),
            "CBR_USUMOD": f"{centro} MCP",
        }

    def _article_basic(self, empresa: int, articulo: str) -> dict[str, Any]:
        row = self.db.one(
            """
            SELECT FIRST 1 ART_CODART, ART_DESCRI, ART_UNIMED, ART_INDINV
            FROM ARTICUL
            WHERE ART_NUMEMP = ? AND ART_CODART = ?
            """,
            (empresa, articulo),
        )
        if row is None:
            raise KofedasError("Articulo no encontrado: " + articulo)
        return row

    def _stock_current(self, empresa: int, articulo: str, centro: int) -> float:
        if centro == -1:
            row = self.db.one(
                "SELECT SUM(ARTE_EXIST) AS EXISTENCIAS FROM ARTICULE WHERE ARTE_NUMEMP = ? AND ARTE_CODART = ?",
                (empresa, articulo),
            )
        else:
            row = self.db.one(
                "SELECT SUM(ARTE_EXIST) AS EXISTENCIAS FROM ARTICULE WHERE ARTE_NUMEMP = ? AND ARTE_CENTRO = ? AND ARTE_CODART = ?",
                (empresa, centro, articulo),
            )
        return self._to_float((row or {}).get("existencias"), 0)

    def _inventory_cost_mode(self, empresa: int, requested: Any = None) -> dict[str, Any]:
        parameter = str(self._parameter_value("RENTAB", "PBASE", empresa) or "PBASE").strip().upper()
        raw = str(requested or parameter or "PBASE").strip().upper()
        aliases = {
            "PBASE": "PBASE",
            "PCOSTE": "PBASE",
            "COSTE": "PBASE",
            "PRECOST": "PBASE",
            "PRECIO_COSTE": "PBASE",
            "ART_PRECOS": "PBASE",
            "PMEDIO": "PMEDIO",
            "MEDIO": "PMEDIO",
            "PRECIO_MEDIO": "PMEDIO",
            "COSTE_MEDIO": "PMEDIO",
            "ULTIMO": "ULTIMO",
            "ULTIMO_PRECIO": "ULTIMO",
            "ULTIMO_PRECIO_COMPRA": "ULTIMO",
            "COSTE_ULTIMO": "ULTIMO",
            "PREBAS": "PREBAS",
            "BASE": "PREBAS",
            "PRECIO_BASE": "PREBAS",
            "ART_PREBAS": "PREBAS",
        }
        normalized = aliases.get(raw)
        if normalized is None:
            raise KofedasError("modo_coste no valido. Usa PBASE, PMEDIO, ULTIMO o PREBAS")
        labels = {
            "PBASE": "Coste ficha articulo (ART_PRECOS)",
            "PMEDIO": "Precio medio (ARTICULM)",
            "ULTIMO": "Ultimo precio de compra (DETMOVM)",
            "PREBAS": "Precio base articulo (ART_PREBAS)",
        }
        return {"modo": normalized, "parametro_codigo": "RENTAB", "parametro_valor": parameter, "origen": "argumento" if requested not in (None, "") else "PARAMETROS", "descripcion": labels[normalized]}

    def _article_unit_field_price(self, empresa: int, article: dict[str, Any], field: str) -> float:
        price = self._to_float(article.get(field.lower()), 0)
        canpre = self._to_float(article.get("art_canpre"), 0)
        if canpre > 0:
            return price / canpre
        provider = self._to_int(article.get("art_codpro"), 0)
        if provider:
            purchase = self.db.one(
                """
                SELECT FIRST 1 ARTP_CANCON, ARTP_CANVEN
                FROM ARTICULP
                WHERE ARTP_NUMEMP = ? AND ARTP_CODART = ? AND ARTP_CODPRO = ?
                """,
                (empresa, article.get("art_codart"), provider),
            )
            if purchase:
                cancon = self._to_float(purchase.get("artp_cancon"), 1) or 1
                canven = self._to_float(purchase.get("artp_canven"), 1) or 1
                return price * cancon / canven
        return price

    def _article_last_purchase_cost(self, empresa: int, articulo: str, query_date: str, fallback: float) -> tuple[float, str]:
        row = self.db.one(
            """
            SELECT FIRST 1 DMM_FECMOV, DMM_CANTID, DMM_VALLIN, DMM_IMPDTO
            FROM DETMOVM
            WHERE DMM_NUMEMP = ? AND DMM_CODART = ? AND DMM_FECMOV <= ?
            ORDER BY DMM_FECMOV DESC
            """,
            (empresa, articulo, query_date),
        )
        if not row:
            return fallback, "ART_PRECOS sin entradas anteriores"
        quantity = self._to_float(row.get("dmm_cantid"), 0) or 1
        value = self._to_float(row.get("dmm_vallin"), 0) - self._to_float(row.get("dmm_impdto"), 0)
        return value / quantity, "DETMOVM " + str(row.get("dmm_fecmov") or "")

    def _article_average_cost(self, empresa: int, articulo: str, query_date: str, fallback: float) -> tuple[float, str]:
        if not self.db.columns("ARTICULM"):
            return fallback, "ARTICULM no disponible; fallback ART_PRECOS"
        row = self.db.one(
            """
            SELECT FIRST 1 ARTM_COSMED, ARTM_FECFIN
            FROM ARTICULM
            WHERE ARTM_NUMEMP = ? AND ARTM_CODART = ? AND ARTM_FECFIN <= ?
            ORDER BY ARTM_FECFIN DESC
            """,
            (empresa, articulo, query_date),
        )
        if not row:
            return self._article_last_purchase_cost(empresa, articulo, query_date, fallback)
        return self._to_float(row.get("artm_cosmed"), fallback), "ARTICULM " + str(row.get("artm_fecfin") or "")

    def _article_inventory_cost(self, empresa: int, article: dict[str, Any], query_date: str, mode: str) -> dict[str, Any]:
        fallback = self._article_unit_field_price(empresa, article, "ART_PRECOS")
        if mode == "PBASE":
            cost, source = fallback, "ART_PRECOS"
        elif mode == "PREBAS":
            cost, source = self._article_unit_field_price(empresa, article, "ART_PREBAS"), "ART_PREBAS"
        elif mode == "PMEDIO":
            cost, source = self._article_average_cost(empresa, str(article.get("art_codart") or ""), query_date, fallback)
        else:
            cost, source = self._article_last_purchase_cost(empresa, str(article.get("art_codart") or ""), query_date, fallback)
        return {"coste_unitario": round(cost, 6), "origen_coste": source}

    def _profit_article(self, empresa: int, articulo: str, cache: dict[str, dict[str, Any]]) -> dict[str, Any]:
        key = articulo.strip()
        if key not in cache:
            row = self.db.one(
                """
                SELECT FIRST 1 ART_CODART, ART_DESCRI, ART_INDINV, ART_CODFAM, ART_SUBFAM,
                       ART_CODPRO, ART_PRECOS, ART_PREBAS, ART_CANPRE, ART_CODMON
                FROM ARTICUL
                WHERE ART_NUMEMP = ? AND ART_CODART = ?
                """,
                (empresa, key),
            )
            cache[key] = row or {
                "art_codart": key,
                "art_descri": "",
                "art_codfam": None,
                "art_subfam": None,
                "art_codpro": None,
                "art_precos": 0,
                "art_prebas": 0,
                "art_canpre": 1,
            }
        return cache[key]

    def _profit_line_values(self, empresa: int, row: dict[str, Any], mode: str, article_cache: dict[str, dict[str, Any]]) -> dict[str, Any]:
        tiplin = str(row.get("dmv_tiplin") or "").strip().upper()
        description = str(row.get("dmv_descri") or "")
        if any(marker in description for marker in ("Linea Factura de Tickets", "Tickets Cobro", "IVA:")):
            return {"incluida": False, "motivo": "linea_resumen_ticket"}
        excluded_series = str(self._parameter_value("DESSER", "", empresa) or "").strip()
        if excluded_series and str(row.get("dmv_serie") or "").strip() == excluded_series:
            return {"incluida": False, "motivo": "serie_excluida_DESSER"}
        sale_value = 0.0
        cost_value = 0.0
        unit_cost = 0.0
        cost_source = ""
        quantity = self._to_float(row.get("dmv_cantid"), 0)
        if tiplin == "D":
            sale_value = self._to_float(row.get("dmv_vallins"), 0) - self._to_float(row.get("dmv_impdto"), 0)
            article = self._profit_article(empresa, str(row.get("dmv_codart") or ""), article_cache)
            cost = self._article_inventory_cost(empresa, article, self._date_arg(row.get("dmv_fecmov")), mode)
            unit_cost = self._to_float(cost.get("coste_unitario"), 0)
            cost_value = unit_cost * quantity
            cost_source = str(cost.get("origen_coste") or "")
        elif tiplin == "X":
            sale_value = self._to_float(row.get("dmv_vallins"), 0) - self._to_float(row.get("dmv_impdto"), 0)
            canpre = self._to_float(row.get("dmv_canpre"), 0)
            if canpre not in (0, 1):
                unit_cost = canpre
                cost_value = canpre * quantity
                cost_source = "DMV_CANPRE"
            else:
                target_margin = self._to_float(self._parameter_value("RENTAF", "25", empresa), 25) / 100
                cost_value = sale_value * (1 - target_margin)
                unit_cost = cost_value / quantity if quantity else 0
                cost_source = "PARAMETROS.RENTAF"
        else:
            return {"incluida": False, "motivo": "tipo_linea_no_valorado"}
        margin = sale_value - cost_value
        profitability = (margin * 100 / sale_value) if sale_value > 0 else 0
        return {
            "incluida": True,
            "ventas_raw": sale_value,
            "coste_raw": cost_value,
            "margen_raw": margin,
            "ventas": round(sale_value, 2),
            "coste": round(cost_value, 2),
            "margen": round(margin, 2),
            "rentabilidad": round(profitability, 4),
            "coste_unitario": round(unit_cost, 6),
            "origen_coste": cost_source,
        }

    def _profit_sales_rows(self, args: dict[str, Any], require_article: bool, limit_default: int) -> tuple[int, dict[str, Any], list[dict[str, Any]]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"), limit_default)
        mode_info = self._inventory_cost_mode(empresa, args.get("modo_coste"))
        where = ["D.DMV_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("D.DMV_CENTRO = ?")
            params.append(int(args["centro"]))
        article = str(args.get("articulo") or "").strip()
        if article:
            where.append("D.DMV_CODART = ?")
            params.append(article)
        elif require_article:
            raise KofedasError("articulo es obligatorio")
        if args.get("desde"):
            where.append("D.DMV_FECMOV >= ?")
            params.append(self._date_arg(args["desde"]))
        if args.get("hasta"):
            where.append("D.DMV_FECMOV <= ?")
            params.append(self._date_arg(args["hasta"]))
        if args.get("cliente"):
            where.append("C.CBV_CODCLI = ?")
            params.append(int(args["cliente"]))
        if args.get("subcliente") not in (None, ""):
            where.append("C.CBV_SUBCLI = ?")
            params.append(int(args["subcliente"]))
        if args.get("representante") is not None:
            where.append("C.CBV_CODREP = ?")
            params.append(int(args["representante"]))
        if args.get("serie") not in (None, ""):
            where.append("D.DMV_SERIE = ?")
            params.append(str(args["serie"]).strip())
        if args.get("numero_desde") is not None:
            where.append("D.DMV_NUMDOC >= ?")
            params.append(int(args["numero_desde"]))
        if args.get("numero_hasta") is not None:
            where.append("D.DMV_NUMDOC <= ?")
            params.append(int(args["numero_hasta"]))
        if args.get("proveedor") is not None:
            where.append("A.ART_CODPRO = ?")
            params.append(int(args["proveedor"]))
        if args.get("familia") not in (None, ""):
            where.append("A.ART_CODFAM = ?")
            params.append(str(args["familia"]).strip())
        if args.get("subfamilia") not in (None, ""):
            where.append("A.ART_SUBFAM = ?")
            params.append(str(args["subfamilia"]).strip())
        if str(args.get("texto") or "").strip():
            where.append("(UPPER(D.DMV_CODART) LIKE ? OR UPPER(D.DMV_DESCRI) LIKE ?)")
            like = _like(str(args["texto"]))
            params.extend([like, like])
        if args.get("en_oferta"):
            where.append("D.DMV_EJEOFE > 0")
        raw_types = args.get("tipos_documento") or DASHBOARD_SALE_DOCUMENTS
        if isinstance(raw_types, str):
            raw_types = [item.strip() for item in raw_types.split(",") if item.strip()]
        document_types = [self._sale_document_type(value) for value in raw_types]
        if document_types:
            where.append("D.DMV_TIPDOC IN (" + ", ".join("?" for _ in document_types) + ")")
            params.extend(document_types)
        if not args.get("incluir_pedidos"):
            where.append("D.DMV_SIGNO = '1'")
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   D.DMV_NUMEMP, D.DMV_CENTRO, D.DMV_TIPDOC, D.DMV_TIPAC, D.DMV_EJERCI,
                   D.DMV_SERIE, D.DMV_NUMDOC, D.DMV_NUMLIN, D.DMV_SIGNO, D.DMV_TIPLIN,
                   D.DMV_FECMOV, D.DMV_CODART, D.DMV_DESCRI, D.DMV_CODMON, D.DMV_PREVEN,
                   D.DMV_CANTID, D.DMV_CANPRE, D.DMV_VALLINS, D.DMV_IMPDTO, D.DMV_EJEOFE,
                   C.CBV_FECHA, C.CBV_CODCLI, C.CBV_SUBCLI, C.CBV_NOMCLI, C.CBV_CODREP,
                   C.CBV_REFCLI, A.ART_CODFAM, A.ART_SUBFAM, A.ART_CODPRO
            FROM DETMOV D
            LEFT JOIN CABDOCV C ON C.CBV_NUMEMP = D.DMV_NUMEMP
                               AND C.CBV_CENTRO = D.DMV_CENTRO
                               AND C.CBV_TIPDOC = D.DMV_TIPDOC
                               AND C.CBV_TIPAC = D.DMV_TIPAC
                               AND C.CBV_EJERCI = D.DMV_EJERCI
                               AND C.CBV_SERIE = D.DMV_SERIE
                               AND C.CBV_NUMDOC = D.DMV_NUMDOC
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = D.DMV_NUMEMP AND A.ART_CODART = D.DMV_CODART
            WHERE {' AND '.join(where)}
            ORDER BY D.DMV_FECMOV DESC, D.DMV_EJERCI DESC, D.DMV_SERIE, D.DMV_NUMDOC DESC, D.DMV_NUMLIN
            """,
            tuple(params),
            limit,
        )
        return empresa, mode_info, rows

    def _sales_document_types_arg(self, value: Any, default: tuple[str, ...] = DASHBOARD_SALE_DOCUMENTS) -> list[str]:
        raw = value or default
        if isinstance(raw, str):
            raw = [item.strip() for item in raw.split(",") if item.strip()]
        return [self._sale_document_type(item) for item in raw]

    def _anadoc_base_expr(self, alias: str = "C") -> str:
        return (
            f"(COALESCE({alias}.CBV_BASIMP1,0) * (1 - COALESCE({alias}.CBV_PORDTO,0)/100) + COALESCE({alias}.CBV_IMPPOR,0) + "
            f"COALESCE({alias}.CBV_BASIMP2,0) * (1 - COALESCE({alias}.CBV_PORDTO,0)/100) + "
            f"COALESCE({alias}.CBV_BASIMP3,0) * (1 - COALESCE({alias}.CBV_PORDTO,0)/100) + "
            f"COALESCE({alias}.CBV_BASIMP4,0) * (1 - COALESCE({alias}.CBV_PORDTO,0)/100))"
        )

    def _anadoc_tax_expr(self, alias: str = "C", tax: str = "IVA") -> str:
        column = "PORIVA" if tax.upper() == "IVA" else "PORREQ"
        return (
            f"((COALESCE({alias}.CBV_BASIMP1,0) * (1 - COALESCE({alias}.CBV_PORDTO,0)/100) + COALESCE({alias}.CBV_IMPPOR,0)) * COALESCE({alias}.CBV_{column}1,0)/100 + "
            f"COALESCE({alias}.CBV_BASIMP2,0) * (1 - COALESCE({alias}.CBV_PORDTO,0)/100) * COALESCE({alias}.CBV_{column}2,0)/100 + "
            f"COALESCE({alias}.CBV_BASIMP3,0) * (1 - COALESCE({alias}.CBV_PORDTO,0)/100) * COALESCE({alias}.CBV_{column}3,0)/100 + "
            f"COALESCE({alias}.CBV_BASIMP4,0) * (1 - COALESCE({alias}.CBV_PORDTO,0)/100) * COALESCE({alias}.CBV_{column}4,0)/100)"
        )

    def _anadoc_where(self, args: dict[str, Any], alias: str = "C") -> tuple[list[str], list[Any]]:
        empresa = self._empresa(args)
        where = [f"{alias}.CBV_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append(f"{alias}.CBV_CENTRO = ?")
            params.append(int(args["centro"]))
        types = self._sales_document_types_arg(args.get("tipos_documento"))
        if types:
            where.append(f"{alias}.CBV_TIPDOC IN (" + ", ".join("?" for _ in types) + ")")
            params.extend(types)
        if args.get("desde"):
            where.append(f"{alias}.CBV_FECHA >= ?")
            params.append(self._date_arg(args["desde"]))
        if args.get("hasta"):
            where.append(f"{alias}.CBV_FECHA <= ?")
            params.append(self._date_arg(args["hasta"]))
        if args.get("tipo_accion") not in (None, "", "9"):
            where.append(f"{alias}.CBV_TIPAC = ?")
            params.append(str(args["tipo_accion"]).strip()[:1])
        if args.get("situacion") not in (None, ""):
            where.append(f"{alias}.CBV_SITUAC = ?")
            params.append(str(args["situacion"]).strip().upper()[:1])
        if args.get("cliente_desde") is not None:
            where.append(f"{alias}.CBV_CODCLI >= ?")
            params.append(int(args["cliente_desde"]))
        if args.get("cliente_hasta") is not None:
            where.append(f"{alias}.CBV_CODCLI <= ?")
            params.append(int(args["cliente_hasta"]))
        if args.get("subcliente_desde") is not None:
            where.append(f"{alias}.CBV_SUBCLI >= ?")
            params.append(int(args["subcliente_desde"]))
        if args.get("subcliente_hasta") is not None:
            where.append(f"{alias}.CBV_SUBCLI <= ?")
            params.append(int(args["subcliente_hasta"]))
        if args.get("representante") is not None:
            where.append(f"{alias}.CBV_CODREP = ?")
            params.append(int(args["representante"]))
        if args.get("serie_desde") not in (None, ""):
            where.append(f"{alias}.CBV_SERIE >= ?")
            params.append(str(args["serie_desde"]).strip())
        if args.get("serie_hasta") not in (None, ""):
            where.append(f"{alias}.CBV_SERIE <= ?")
            params.append(str(args["serie_hasta"]).strip())
        if args.get("numero_desde") is not None:
            where.append(f"{alias}.CBV_NUMDOC >= ?")
            params.append(int(args["numero_desde"]))
        if args.get("numero_hasta") is not None:
            where.append(f"{alias}.CBV_NUMDOC <= ?")
            params.append(int(args["numero_hasta"]))
        if args.get("tarjeta") not in (None, ""):
            where.append(f"{alias}.CBV_CODTAR = ?")
            params.append(str(args["tarjeta"]).strip())
        pending_expr = f"(COALESCE({alias}.CBV_TOTALD,0) - COALESCE({alias}.CBV_IMPCOB,0))"
        if args.get("importe_pendiente_min") is not None:
            where.append(f"{pending_expr} >= ?")
            params.append(self._to_float(args["importe_pendiente_min"], 0))
        if args.get("importe_pendiente_max") is not None:
            where.append(f"{pending_expr} <= ?")
            params.append(self._to_float(args["importe_pendiente_max"], 0))
        invoice_kind = str(args.get("tipo_factura") or "todas").strip().lower()
        contado = self._to_int(self._parameter_value("FPGCON", "0", empresa), 0)
        if invoice_kind in {"contado", "factura_contado"}:
            where.append(f"({alias}.CBV_TIPDOC <> 'F' OR {alias}.CBV_CODPAG = ?)")
            params.append(contado)
        elif invoice_kind in {"tickets", "ticket", "factura_tickets"}:
            where.append(f"({alias}.CBV_TIPDOC <> 'F' OR {alias}.CBV_CODPAG = -1)")
        elif invoice_kind in {"albaran", "albaranes", "factura_albaran"}:
            where.append(f"({alias}.CBV_TIPDOC <> 'F' OR ({alias}.CBV_CODPAG > 0 AND {alias}.CBV_CODPAG <> ?))")
            params.append(contado)
        return where, params

    def _anadoc_group_expr(self, group_by: str) -> tuple[str, str, str, str]:
        group = group_by.strip().lower()
        mapping = {
            "tipo_documento": ("C.CBV_TIPDOC", "C.CBV_TIPDOC", "CODIGO", "NOMBRE"),
            "documento": ("C.CBV_TIPDOC", "C.CBV_TIPDOC", "CODIGO", "NOMBRE"),
            "cliente": ("C.CBV_CODCLI || '/' || C.CBV_SUBCLI", "NULLIF(TRIM(COALESCE(CL.CLI_NOMCLI, C.CBV_NOMCLI, '')), '')", "CODIGO", "NOMBRE"),
            "forma_pago": ("C.CBV_CODPAG", "COALESCE(FP.FPG_DESCRI, CAST(C.CBV_CODPAG AS VARCHAR(20)))", "CODIGO", "NOMBRE"),
            "forma_cobro": ("C.CBV_FORCOB", "C.CBV_FORCOB", "CODIGO", "NOMBRE"),
            "agente": ("C.CBV_CODREP", "COALESCE(R.REP_NOMBRE, CAST(C.CBV_CODREP AS VARCHAR(20)))", "CODIGO", "NOMBRE"),
            "representante": ("C.CBV_CODREP", "COALESCE(R.REP_NOMBRE, CAST(C.CBV_CODREP AS VARCHAR(20)))", "CODIGO", "NOMBRE"),
            "poblacion": ("UPPER(TRIM(COALESCE(NULLIF(C.CBV_POBLAC, ''), CL.CLI_POBLAC, '')))", "UPPER(TRIM(COALESCE(NULLIF(C.CBV_POBLAC, ''), CL.CLI_POBLAC, '')))", "CODIGO", "NOMBRE"),
            "centro": ("C.CBV_CENTRO", "COALESCE(CE.CEN_NOMCEN, CAST(C.CBV_CENTRO AS VARCHAR(20)))", "CODIGO", "NOMBRE"),
            "serie": ("C.CBV_SERIE", "C.CBV_SERIE", "CODIGO", "NOMBRE"),
            "mes": ("EXTRACT(MONTH FROM C.CBV_FECHA)", "EXTRACT(MONTH FROM C.CBV_FECHA)", "CODIGO", "NOMBRE"),
            "dia_semana": ("EXTRACT(WEEKDAY FROM C.CBV_FECHA)", "EXTRACT(WEEKDAY FROM C.CBV_FECHA)", "CODIGO", "NOMBRE"),
            "hora": ("EXTRACT(HOUR FROM C.CBV_FECMOD)", "EXTRACT(HOUR FROM C.CBV_FECMOD)", "CODIGO", "NOMBRE"),
            "situacion": ("C.CBV_SITUAC", "C.CBV_SITUAC", "CODIGO", "NOMBRE"),
            "tarjeta": ("C.CBV_CODTAR", "COALESCE(T.TAR_NOMBRE, C.CBV_CODTAR)", "CODIGO", "NOMBRE"),
        }
        if group not in mapping:
            raise KofedasError("agrupar_por no valido para ventas_documentos_resumen")
        return mapping[group]

    def _stock_at_date(self, empresa: int, articulo: str, centro: int, query_date: str) -> float:
        if query_date == date.today().isoformat():
            return self._stock_current(empresa, articulo, centro)
        current = self._stock_current(empresa, articulo, centro)
        if centro == -1:
            entrada_filter = "DMM_NUMEMP = ? AND DMM_CODART = ? AND DMM_FECMOV >= ?"
            regular_filter = "DMR_NUMEMP = ? AND DMR_CODART = ? AND DMR_FECMOV >= ?"
            venta_filter = "DMV_NUMEMP = ? AND DMV_CODART = ? AND DMV_SIGNO = '1' AND DMV_FECMOV >= ?"
            movement_params = (empresa, articulo, query_date)
        else:
            entrada_filter = "DMM_NUMEMP = ? AND DMM_CENTRO = ? AND DMM_CODART = ? AND DMM_FECMOV >= ?"
            regular_filter = "DMR_NUMEMP = ? AND DMR_CENTRO = ? AND DMR_CODART = ? AND DMR_FECMOV >= ?"
            venta_filter = "DMV_NUMEMP = ? AND DMV_CENTRO = ? AND DMV_CODART = ? AND DMV_SIGNO = '1' AND DMV_FECMOV >= ?"
            movement_params = (empresa, centro, articulo, query_date)
        entradas = self._movement_sum(f"SELECT SUM(DMM_CANTID) AS EXISTENCIAS FROM DETMOVM WHERE {entrada_filter}", movement_params)
        regularizaciones = self._movement_sum(f"SELECT SUM(DMR_CANTID) AS EXISTENCIAS FROM DETMOVR WHERE {regular_filter}", movement_params)
        ventas = self._movement_sum(f"SELECT SUM(DMV_CANTID) AS EXISTENCIAS FROM DETMOV WHERE {venta_filter}", movement_params)
        return current - entradas - regularizaciones + ventas

    def _movement_sum(self, sql: str, params: tuple[Any, ...]) -> float:
        row = self.db.one(sql, params)
        return self._to_float((row or {}).get("existencias"), 0)

    def _sales_table(self, table: Any) -> str:
        name = str(table or "").strip().upper()
        if name not in SALES_TABLES:
            raise KofedasError("Tabla de ventas no permitida: " + name)
        return name

    def _sale_document_type(self, value: Any, default: str = "P") -> str:
        tipdoc = str(value or default or "P").strip().upper()[:1]
        if tipdoc not in SALES_DOCUMENT_TYPES:
            raise KofedasError("Tipo de documento de venta no valido: " + str(value))
        return tipdoc

    def _sale_client_extra(self, empresa: int, cliente: int, subcliente: int, code: str, default: Any = "") -> Any:
        row = self.db.one(
            """
            SELECT CLII_DESCRI
            FROM CLIENI
            WHERE CLII_NUMEMP = ? AND CLII_CODCLI = ? AND CLII_SUBCLI = ? AND CLII_CODINF = ?
            """,
            (empresa, cliente, subcliente, code),
        )
        return row["clii_descri"] if row and row.get("clii_descri") not in (None, "") else default

    def _sale_client(self, empresa: int, cliente: int, subcliente: int) -> dict[str, Any]:
        if cliente == 99999 and subcliente == 0:
            return {
                "cli_codcli": 99999,
                "cli_subcli": 0,
                "cli_codrep": 0,
                "cli_nomcli": "Clientes Caja",
                "cli_razsoc": "Clientes Caja",
                "cli_cif": "",
                "cli_domici": "",
                "cli_codpos": 0,
                "cli_poblac": "",
                "cli_forpag": 0,
                "cli_forenv": 0,
                "cli_dtoesp": 0,
                "cli_regiva": "N",
                "cli_tippre": "4",
                "cli_poraum": 0,
                "cli_codmon": "E",
            }
        row = self.db.one(
            "SELECT FIRST 1 * FROM CLIEN WHERE CLI_NUMEMP = ? AND CLI_CODCLI = ? AND CLI_SUBCLI = ?",
            (empresa, cliente, subcliente),
        )
        if row is None:
            raise KofedasError(f"Cliente {cliente}/{subcliente} no encontrado")
        return row

    def _sale_article_code(self, empresa: int, cliente: int, args: dict[str, Any]) -> str:
        code = str(args.get("articulo") or args.get("codigo") or "").strip()
        barcode = str(args.get("codigo_barras") or args.get("ean") or "").strip()
        client_code = str(args.get("codigo_cliente") or args.get("articulo_cliente") or "").strip()
        if code:
            row = self.db.one(
                """
                SELECT FIRST 1 ART_CODART
                FROM ARTICUL
                WHERE ART_NUMEMP = ? AND (ART_CODART = ? OR ART_CODART LIKE ?)
                ORDER BY ART_CODART
                """,
                (empresa, code, "%" + code + "%"),
            )
            if row and row.get("art_codart"):
                return str(row["art_codart"]).strip()
        if barcode:
            owner = self._barcode_owner(empresa, barcode)
            if owner:
                return owner
        for candidate in (client_code, code):
            if candidate:
                row = self.db.one(
                    """
                    SELECT FIRST 1 CLIA_CODART
                    FROM CLIART
                    WHERE CLIA_NUMEMP = ? AND CLIA_CODCLI = ? AND CLIA_CODARTC = ?
                    """,
                    (empresa, cliente, candidate),
                )
                if row and row.get("clia_codart"):
                    return str(row["clia_codart"]).strip()
        raise KofedasError("Articulo no encontrado en venta: " + (code or barcode or client_code))

    def _sale_tax(self, empresa: int, tipiva: Any, recargo: bool) -> tuple[float, float]:
        row = self.db.one(
            "SELECT FIRST 1 TIV_PORIVA, TIV_PORREQ FROM TIPIVA WHERE TIV_NUMEMP = ? AND TIV_TIPIVA = ?",
            (empresa, self._to_int(tipiva, 0)),
        )
        if not row:
            return 0.0, 0.0
        return self._to_float(row.get("tiv_poriva"), 0), self._to_float(row.get("tiv_porreq"), 0) if recargo else 0.0

    def _sale_price_level(self, article: dict[str, Any], level: str, customer: dict[str, Any]) -> tuple[float, float, str]:
        level = str(level or "4").strip().upper()
        increase = 1 + self._to_float(customer.get("cli_poraum"), 0) / 100
        if level == "1":
            return self._to_float(article.get("art_preven1"), 0) * increase, self._to_float(article.get("art_pvp"), 0) * increase, "1"
        if level == "2":
            return self._to_float(article.get("art_preven2"), 0) * increase, self._to_float(article.get("art_pvp"), 0) * increase, "2"
        if level == "3":
            return self._to_float(article.get("art_preven3"), 0) * increase, self._to_float(article.get("art_pvp"), 0) * increase, "3"
        if level == "T" and self._to_float(article.get("art_pretar"), 0) != 0:
            price = self._to_float(article.get("art_pretar"), 0) * increase
            return price, self._to_float(article.get("art_pvp"), 0) * increase, "T"
        if level == "8":
            return self._to_float(article.get("art_prebas"), 0) * increase, self._to_float(article.get("art_pvp"), 0), "8"
        if level == "9":
            return self._to_float(article.get("art_precos"), 0) * increase, self._to_float(article.get("art_pvp"), 0), "9"
        return self._to_float(article.get("art_preven4"), 0) * increase, self._to_float(article.get("art_pvp"), 0) * increase, "4"

    def _sale_special_price(self, empresa: int, cliente: int, articulo: str, cantidad: float) -> dict[str, Any] | None:
        row = self.db.one(
            """
            SELECT FIRST 1 *
            FROM CLIART
            WHERE CLIA_NUMEMP = ? AND (CLIA_CODART = ? OR CLIA_CODART LIKE ?) AND CLIA_CODCLI = ?
            """,
            (empresa, articulo, "%" + articulo + "%", cliente),
        )
        if not row:
            return None
        canpre = self._to_float(row.get("clia_canpre"), 1) or 1
        if canpre > 1 and cantidad < canpre:
            return None
        price = self._to_float(row.get("clia_precio"), 0)
        discount = self._to_float(row.get("clia_descue"), 0)
        if price == 0 and discount == 0:
            return None
        return {
            "origen": "CLIART",
            "precio": price,
            "dto1": discount,
            "canpre": canpre,
            "moneda": str(row.get("clia_codmon") or "").strip(),
            "descripcion": str(row.get("clia_descri") or "").strip(),
        }

    def _sale_active_offer(self, empresa: int, articulo: str, fecha: str, cantidad: float, cliente: dict[str, Any]) -> dict[str, Any] | None:
        if str(self._sale_client_extra(empresa, int(cliente.get("cli_codcli") or 0), int(cliente.get("cli_subcli") or 0), "ACEOF", "S")).upper() == "N":
            return None
        row = self.db.one(
            """
            SELECT FIRST 1 D.*, O.OFE_TIPOFE
            FROM DETOFER D
            LEFT JOIN OFERTAS O ON O.OFE_NUMEMP = D.DOF_NUMEMP
                              AND O.OFE_EJERCI = D.DOF_EJERCI
                              AND O.OFE_NUMOFE = D.DOF_NUMOFE
            WHERE D.DOF_NUMEMP = ? AND (D.DOF_CODART = ? OR D.DOF_CODART LIKE ?)
              AND D.DOF_FECINI <= ? AND D.DOF_FECFIN >= ?
            ORDER BY D.DOF_PVP, D.DOF_PRECIO
            """,
            (empresa, articulo, "%" + articulo + "%", fecha, fecha),
        )
        if not row:
            return None
        tipo = str(row.get("ofe_tipofe") or "").strip().upper()
        if tipo in {"P", "S"}:
            return None
        canpre = self._to_float(row.get("dof_canpre"), 1) or 1
        if canpre > 1 and cantidad < canpre:
            return None
        price = self._to_float(row.get("dof_precio"), 0)
        pvp = self._to_float(row.get("dof_pvp"), 0)
        if price == 0 and pvp == 0:
            return None
        return {
            "origen": "DETOFER",
            "precio": price,
            "pvp": pvp,
            "dto1": self._to_float(row.get("dof_dto1"), 0),
            "dto2": self._to_float(row.get("dof_dto2"), 0),
            "canpre": canpre,
            "moneda": str(row.get("dof_codmon") or "").strip(),
            "ejercicio_oferta": self._to_int(row.get("dof_ejerci"), 0),
            "numero_oferta": self._to_int(row.get("dof_numofe"), 0),
        }

    def _sale_family_discount(self, empresa: int, cliente: int, codfam: Any, subfam: Any) -> float | None:
        candidates = [
            (self._to_int(codfam, 0), self._to_int(subfam, 0)),
            (self._to_int(codfam, 0), 0),
            (0, 0),
        ]
        seen: set[tuple[int, int]] = set()
        for family, subfamily in candidates:
            if (family, subfamily) in seen:
                continue
            seen.add((family, subfamily))
            row = self.db.one(
                """
                SELECT FIRST 1 CLIF_DTO
                FROM CLIFAM
                WHERE CLIF_NUMEMP = ? AND CLIF_CODCLI = ? AND CLIF_CODFAM = ? AND CLIF_SUBFAM = ?
                """,
                (empresa, cliente, family, subfamily),
            )
            if row:
                return self._to_float(row.get("clif_dto"), 0)
        return None

    def _sale_price(self, empresa: int, cliente: int, subcliente: int, item: dict[str, Any], fecha: str, tipo_documento: str) -> dict[str, Any]:
        cantidad = self._to_float(item.get("cantidad", item.get("cantid")), 1)
        if cantidad <= 0:
            raise KofedasError("La cantidad debe ser mayor que cero")
        customer = self._sale_client(empresa, cliente, subcliente)
        articulo = self._sale_article_code(empresa, cliente, item)
        article = self.db.one(
            """
            SELECT FIRST 1 *
            FROM ARTICUL
            WHERE ART_NUMEMP = ? AND (ART_CODART = ? OR ART_CODART LIKE ?)
            ORDER BY ART_CODART
            """,
            (empresa, articulo, "%" + articulo + "%"),
        )
        if not article:
            raise KofedasError("Articulo no encontrado: " + articulo)
        recargo = str(customer.get("cli_regiva") or "").strip().upper() == "R" and not (cliente == 99999 and subcliente == 0)
        poriva, porreq = self._sale_tax(empresa, article.get("art_tipiva"), recargo)
        moneda = str(article.get("art_codmon") or customer.get("cli_codmon") or "E").strip()[:1] or "E"
        pvp = self._to_float(article.get("art_pvp"), 0)
        canpre = self._to_float(article.get("art_canpre"), 1) or 1
        dto1 = self._to_float(item.get("dto1", item.get("descuento1")), 0)
        dto2 = self._to_float(item.get("dto2", item.get("descuento2")), 0)
        tippre = str(customer.get("cli_tippre") or "4").strip().upper() or "4"
        source = "ARTICUL/CLIEN"
        offer_number = 0
        offer_year = 0
        preiva = "N"

        special = None if cliente == 99999 else self._sale_special_price(empresa, cliente, articulo, cantidad)
        if special:
            if special["precio"] != 0:
                price = special["precio"]
            else:
                price, pvp, tippre = self._sale_price_level(article, "T", customer)
            dto1 = self._to_float(item.get("dto1", item.get("descuento1")), special["dto1"])
            canpre = special["canpre"] or canpre
            moneda = (special["moneda"] or moneda)[:1]
            source = "CLIART"
        else:
            offer = self._sale_active_offer(empresa, articulo, fecha, cantidad, customer)
            if offer:
                moneda = (offer["moneda"] or moneda)[:1]
                canpre = offer["canpre"] or canpre
                dto1 = self._to_float(item.get("dto1", item.get("descuento1")), offer["dto1"])
                dto2 = self._to_float(item.get("dto2", item.get("descuento2")), offer["dto2"])
                offer_number = offer["numero_oferta"]
                offer_year = offer["ejercicio_oferta"]
                source = "DETOFER"
                if offer["pvp"] != 0:
                    pvp = offer["pvp"]
                    price = pvp / (1 + poriva / 100) if poriva != -100 else 0
                    preiva = "S" if cliente == 99999 else "N"
                else:
                    price = offer["precio"]
            else:
                family_discount = None if cliente == 99999 else self._sale_family_discount(empresa, cliente, article.get("art_codfam"), article.get("art_subfam"))
                if family_discount is not None and dto1 == 0:
                    dto1 = family_discount
                    source = "CLIFAM"
                price, pvp, tippre = self._sale_price_level(article, tippre, customer)

        explicit_price = item.get("precio", item.get("preven"))
        if explicit_price not in (None, ""):
            price = self._to_float(explicit_price, price)
            source = "manual"
        explicit_pvp = item.get("pvp")
        if explicit_pvp not in (None, ""):
            pvp = self._to_float(explicit_pvp, pvp)
        base = price * cantidad / (canpre or 1)
        base_after_discounts = base * (1 - dto1 / 100) * (1 - dto2 / 100)
        total = base_after_discounts * (1 + poriva / 100 + porreq / 100)
        return {
            "encontrado": True,
            "articulo": articulo,
            "descripcion": str(item.get("descripcion") or item.get("descri") or (special or {}).get("descripcion") or article.get("art_descri") or "")[:100],
            "cantidad": round(cantidad, 4),
            "moneda": moneda,
            "tipo_precio": tippre,
            "precio": round(price, 6),
            "pvp": round(pvp, 6),
            "canpre": round(canpre or 1, 4),
            "unidad": str(item.get("unidad") or article.get("art_unimed") or "")[:4],
            "dto1": round(dto1, 4),
            "dto2": round(dto2, 4),
            "poriva": round(poriva, 4),
            "porreq": round(porreq, 4),
            "base": round(base, 4),
            "base_neta": round(base_after_discounts, 4),
            "total": round(total, 2),
            "preiva": preiva,
            "origen_precio": source,
            "ejercicio_oferta": offer_year,
            "numero_oferta": offer_number,
            "tipo_documento": tipo_documento,
        }

    def _sale_series(self, empresa: int, centro: int, cliente: int, subcliente: int, tipdoc: str, value: Any = None) -> str:
        if value not in (None, ""):
            return str(value).strip()[:2]
        client_series = str(self._sale_client_extra(empresa, cliente, subcliente, "SERIE", "") or "").strip()
        series = client_series or str(self._parameter_value(f"{tipdoc}{centro}", "", empresa) or "").strip()
        series = series or str(self._parameter_value(tipdoc, "", empresa) or "").strip()
        if not series and tipdoc == "P":
            series = "PM"
        return series[:2]

    def _next_sale_number(self, empresa: int, centro: int, tipdoc: str, tipac: str, ejercicio: int, serie: str, requested: int | None = None) -> int:
        number = int(requested or 0)
        if number:
            return number
        row = self.db.one(
            """
            SELECT MAX(NUM_NUMERO) AS MAXIMO
            FROM NUMERA
            WHERE NUM_NUMEMP = ? AND NUM_CENTRO = ? AND NUM_TIPAC = ? AND NUM_TIPDOC = ?
              AND NUM_EJERCI = ? AND NUM_SERIE = ?
            """,
            (empresa, centro, tipac, tipdoc, ejercicio, serie),
        )
        if row and row.get("maximo") is not None:
            return int(row["maximo"]) + 1
        row = self.db.one(
            """
            SELECT MAX(CBV_NUMDOC) AS MAXIMO
            FROM CABDOCV
            WHERE CBV_NUMEMP = ? AND CBV_CENTRO = ? AND CBV_TIPDOC = ? AND CBV_TIPAC = ?
              AND CBV_EJERCI = ? AND CBV_SERIE = ?
            """,
            (empresa, centro, tipdoc, tipac, ejercicio, serie),
        )
        return int((row or {}).get("maximo") or 0) + 1

    def _sale_document_exists(self, empresa: int, centro: int, tipdoc: str, tipac: str, ejercicio: int, serie: str, numero: int) -> bool:
        return self.db.one(
            """
            SELECT FIRST 1 1 AS EXISTE
            FROM CABDOCV
            WHERE CBV_NUMEMP = ? AND CBV_CENTRO = ? AND CBV_TIPDOC = ? AND CBV_TIPAC = ?
              AND CBV_EJERCI = ? AND CBV_SERIE = ? AND CBV_NUMDOC = ?
            """,
            (empresa, centro, tipdoc, tipac, ejercicio, serie, numero),
        ) is not None

    def _sale_tipven(self, empresa: int, tipdoc: str, tipac: str) -> int:
        row = self.db.one(
            "SELECT FIRST 1 TIV_CODIGO FROM TIPVEN WHERE TIV_NUMEMP = ? AND TIV_TIPDOC = ? AND TIV_TIPAC = ? ORDER BY TIV_CODIGO",
            (empresa, tipdoc, tipac),
        )
        return self._to_int((row or {}).get("tiv_codigo"), -1)

    def _sale_line_defaults(self, header: dict[str, Any], line_number: int, item: dict[str, Any], customer: dict[str, Any]) -> dict[str, Any]:
        tiplin = str(item.get("tipo_linea") or item.get("tiplin") or "D").strip().upper()[:1] or "D"
        if tiplin == "C":
            return {
                "DMV_NUMEMP": header["CBV_NUMEMP"], "DMV_CENTRO": header["CBV_CENTRO"], "DMV_TIPDOC": header["CBV_TIPDOC"],
                "DMV_TIPAC": header["CBV_TIPAC"], "DMV_EJERCI": header["CBV_EJERCI"], "DMV_SERIE": header["CBV_SERIE"],
                "DMV_NUMDOC": header["CBV_NUMDOC"], "DMV_NUMLIN": line_number, "DMV_SIGNO": "0",
                "DMV_CAJA": header["CBV_CAJA"], "DMV_USUAR": header["CBV_USUMOD"], "DMV_TIPLIN": "C",
                "DMV_FECMOV": header["CBV_FECHA"], "DMV_CODART": "", "DMV_DESCRI": str(item.get("texto") or item.get("descripcion") or "")[:100],
                "DMV_CODMON": header["CBV_CODMON"], "DMV_TIPPRE": "", "DMV_PREVEN": 0, "DMV_PORIVA": 0, "DMV_PORREQ": 0,
                "DMV_PVP": 0, "DMV_CANTID": 0, "DMV_CANPRE": 1, "DMV_UNIMED": "", "DMV_DTO1": 0, "DMV_DTO2": 0,
                "DMV_VALLIN": 0, "DMV_VALLINS": 0, "DMV_IMPDTO": 0, "DMV_EJEOFE": 0, "DMV_NUMOFE": 0,
                "DMV_EJERCIO": 0, "DMV_TIPDOCO": "", "DMV_SERIEO": "", "DMV_NUMDOCO": 0, "DMV_NUMLINO": 0, "DMV_PREIVA": "",
            }
        price = self._sale_price(
            int(header["CBV_NUMEMP"]),
            int(header["CBV_CODCLI"]),
            int(header["CBV_SUBCLI"]),
            item,
            str(header["CBV_FECHA"]),
            str(header["CBV_TIPDOC"]),
        )
        sign = "0" if str(header["CBV_TIPDOC"]) in {"P", "R", "S"} else ("-1" if str(header["CBV_TIPDOC"]) == "C" else "1")
        return {
            "DMV_NUMEMP": header["CBV_NUMEMP"], "DMV_CENTRO": header["CBV_CENTRO"], "DMV_TIPDOC": header["CBV_TIPDOC"],
            "DMV_TIPAC": header["CBV_TIPAC"], "DMV_EJERCI": header["CBV_EJERCI"], "DMV_SERIE": header["CBV_SERIE"],
            "DMV_NUMDOC": header["CBV_NUMDOC"], "DMV_NUMLIN": self._to_int(item.get("linea"), line_number), "DMV_SIGNO": sign,
            "DMV_CAJA": header["CBV_CAJA"], "DMV_USUAR": header["CBV_USUMOD"], "DMV_TIPLIN": "D",
            "DMV_FECMOV": header["CBV_FECHA"], "DMV_CODART": price["articulo"], "DMV_DESCRI": price["descripcion"][:100],
            "DMV_CODMON": price["moneda"], "DMV_TIPPRE": price["tipo_precio"], "DMV_PREVEN": price["precio"],
            "DMV_PORIVA": price["poriva"], "DMV_PORREQ": price["porreq"], "DMV_PVP": price["pvp"],
            "DMV_CANTID": price["cantidad"], "DMV_CANPRE": price["canpre"], "DMV_UNIMED": price["unidad"],
            "DMV_DTO1": price["dto1"], "DMV_DTO2": price["dto2"], "DMV_VALLIN": price["total"],
            "DMV_VALLINS": price["base_neta"], "DMV_IMPDTO": 0, "DMV_EJEOFE": price["ejercicio_oferta"],
            "DMV_NUMOFE": price["numero_oferta"], "DMV_EJERCIO": self._to_int(item.get("origen_ejercicio"), 0),
            "DMV_TIPDOCO": str(item.get("origen_tipo_documento") or "")[:1], "DMV_SERIEO": str(item.get("origen_serie") or "")[:2],
            "DMV_NUMDOCO": self._to_int(item.get("origen_numero"), 0), "DMV_NUMLINO": self._to_int(item.get("origen_linea"), 0),
            "DMV_PREIVA": price["preiva"],
            "precio_calculado": price,
        }

    def _sale_totals(self, header: dict[str, Any], lines: list[dict[str, Any]]) -> dict[str, Any]:
        buckets: list[dict[str, float]] = []
        for line in lines:
            if str(line.get("DMV_TIPLIN") or "").upper() not in {"D", "X"}:
                continue
            iva = self._to_float(line.get("DMV_PORIVA"), 0)
            req = self._to_float(line.get("DMV_PORREQ"), 0)
            bucket = next((item for item in buckets if item["iva"] == iva and item["req"] == req), None)
            if bucket is None and len(buckets) < 4:
                bucket = {"base": 0.0, "iva": iva, "req": req}
                buckets.append(bucket)
            if bucket is not None:
                bucket["base"] += self._to_float(line.get("DMV_VALLINS"), 0)
        while len(buckets) < 4:
            buckets.append({"base": 0.0, "iva": 0.0, "req": 0.0})
        discount = self._to_float(header.get("CBV_PORDTO"), 0)
        shipping = self._to_float(header.get("CBV_IMPPOR"), 0)
        total_base = 0.0
        total_doc = 0.0
        for index, bucket in enumerate(buckets[:4], start=1):
            header[f"CBV_BASIMP{index}"] = round(bucket["base"], 4)
            header[f"CBV_PORIVA{index}"] = bucket["iva"]
            header[f"CBV_PORREQ{index}"] = bucket["req"]
            taxable = bucket["base"] * (1 - discount / 100) + (shipping if index == 1 else 0)
            total_base += taxable
            total_doc += taxable * (1 + bucket["iva"] / 100 + bucket["req"] / 100)
        header["CBV_TOTALS"] = round(total_base, 2)
        header["CBV_TOTALD"] = round(total_doc, 2)
        return header

    def _sale_numera_statement(self, header: dict[str, Any]) -> tuple[str, tuple[Any, ...]]:
        key = (
            header["CBV_NUMEMP"], header["CBV_CENTRO"], header["CBV_TIPAC"], header["CBV_TIPDOC"],
            header["CBV_EJERCI"], header["CBV_SERIE"],
        )
        exists = self.db.one(
            """
            SELECT FIRST 1 1 AS EXISTE
            FROM NUMERA
            WHERE NUM_NUMEMP = ? AND NUM_CENTRO = ? AND NUM_TIPAC = ? AND NUM_TIPDOC = ?
              AND NUM_EJERCI = ? AND NUM_SERIE = ?
            """,
            key,
        )
        if exists:
            return (
                """
                UPDATE NUMERA SET NUM_NUMERO = ?
                WHERE NUM_NUMEMP = ? AND NUM_CENTRO = ? AND NUM_TIPAC = ? AND NUM_TIPDOC = ?
                  AND NUM_EJERCI = ? AND NUM_SERIE = ?
                """,
                (header["CBV_NUMDOC"], *key),
            )
        return (
            "INSERT INTO NUMERA (NUM_NUMEMP, NUM_CENTRO, NUM_TIPAC, NUM_TIPDOC, NUM_EJERCI, NUM_SERIE, NUM_NUMERO) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (*key, header["CBV_NUMDOC"]),
        )

    def _pedido_key(self, args: dict[str, Any], default_tipdoc: str = "P") -> dict[str, Any]:
        tipdoc = self._sale_document_type(args.get("tipo_documento", args.get("tipdoc")), default_tipdoc)
        return {
            "empresa": self._empresa(args),
            "centro": self._centro(args),
            "tipdoc": tipdoc,
            "tipac": str(args.get("tipo_accion") or args.get("tipac") or "0").strip()[:1] or "0",
            "ejercicio": self._to_int(args.get("ejercicio", args.get("ejerci")), 0),
            "serie": str(args.get("serie") or "").strip()[:2],
            "numero": self._to_int(args.get("numero", args.get("numdoc")), 0),
        }

    def _pedido_where_from_key(self, key: dict[str, Any], alias: str = "C") -> tuple[str, tuple[Any, ...]]:
        missing = [name for name in ("ejercicio", "serie", "numero") if key.get(name) in (None, "", 0)]
        if missing:
            raise KofedasError("Faltan claves de pedido: " + ", ".join(missing))
        prefix = alias + "." if alias else ""
        return (
            f"{prefix}CBV_NUMEMP = ? AND {prefix}CBV_CENTRO = ? AND {prefix}CBV_TIPDOC = ? "
            f"AND {prefix}CBV_TIPAC = ? AND {prefix}CBV_EJERCI = ? AND {prefix}CBV_SERIE = ? AND {prefix}CBV_NUMDOC = ?",
            (key["empresa"], key["centro"], key["tipdoc"], key["tipac"], key["ejercicio"], key["serie"], key["numero"]),
        )

    def _pedido_header(self, key: dict[str, Any]) -> dict[str, Any]:
        where, params = self._pedido_where_from_key(key)
        row = self.db.one(
            f"""
            SELECT FIRST 1 C.*, CL.CLI_EMAIL, CL.CLI_NOMCLI, CL.CLI_RAZSOC
            FROM CABDOCV C
            LEFT JOIN CLIEN CL ON CL.CLI_NUMEMP = C.CBV_NUMEMP
                              AND CL.CLI_CODCLI = C.CBV_CODCLI
                              AND CL.CLI_SUBCLI = C.CBV_SUBCLI
            WHERE {where}
            """,
            params,
        )
        if row is None:
            raise KofedasError("Pedido no encontrado")
        return row

    def _pedido_lines(self, key: dict[str, Any]) -> list[dict[str, Any]]:
        return self.db.query(
            """
            SELECT D.*
            FROM DETMOV D
            WHERE D.DMV_NUMEMP = ? AND D.DMV_CENTRO = ? AND D.DMV_TIPDOC = ?
              AND D.DMV_TIPAC = ? AND D.DMV_EJERCI = ? AND D.DMV_SERIE = ? AND D.DMV_NUMDOC = ?
            ORDER BY D.DMV_NUMLIN
            """,
            (key["empresa"], key["centro"], key["tipdoc"], key["tipac"], key["ejercicio"], key["serie"], key["numero"]),
            MAX_ROWS_LIMIT,
        )

    def _pedido_document_result(self, key: dict[str, Any], modo: str = "normal") -> dict[str, Any]:
        header = self._pedido_header(key)
        lines = self._pedido_lines(key)
        items: list[dict[str, Any]] = []
        total_quantity = 0.0
        total_prepared = 0.0
        line_columns = set(lines[0].keys()) if lines else set()
        prepared_field = next((field for field in ("dmv_canprea", "dmv_canpre", "dmv_canser", "dmv_canpreparada") if field in line_columns), None)
        pending_field = next((field for field in ("dmv_canpen", "dmv_canpte", "dmv_pendie") if field in line_columns), None)
        zone_field = next((field for field in ("dmv_zona", "dmv_zonpre", "dmv_zonprep") if field in line_columns), None)
        for row in lines:
            qty = self._to_float(row.get("dmv_cantid"), 0)
            prepared = self._to_float(row.get(prepared_field), 0) if prepared_field else 0.0
            pending = self._to_float(row.get(pending_field), max(qty - prepared, 0)) if pending_field else max(qty - prepared, 0)
            total_quantity += qty
            total_prepared += prepared
            item = {
                "linea": row.get("dmv_numlin"),
                "tipo_linea": row.get("dmv_tiplin"),
                "articulo": row.get("dmv_codart"),
                "descripcion": row.get("dmv_descri"),
                "cantidad": qty,
                "cantidad_preparada": prepared,
                "cantidad_pendiente": pending,
                "unidad": row.get("dmv_unimed"),
                "precio": row.get("dmv_preven"),
                "dto1": row.get("dmv_dto1"),
                "dto2": row.get("dmv_dto2"),
                "base": row.get("dmv_vallins"),
                "total": row.get("dmv_vallin"),
            }
            if zone_field:
                item["zona"] = row.get(zone_field)
            if modo == "normal":
                item["raw"] = row
            items.append(item)
        return {
            "modo": modo,
            "cabecera": header,
            "lineas": items,
            "totales_preparacion": {
                "lineas": len(items),
                "cantidad": round(total_quantity, 4),
                "cantidad_preparada": round(total_prepared, 4),
                "cantidad_pendiente": round(max(total_quantity - total_prepared, 0), 4),
            },
        }

    def _pedido_update_statement(
        self,
        table: str,
        updates: dict[str, Any],
        where: str,
        params: tuple[Any, ...],
    ) -> tuple[str, tuple[Any, ...]] | None:
        columns = set(self._table_columns(table))
        applied = {column: value for column, value in updates.items() if column in columns}
        if not applied:
            return None
        sql = f"UPDATE {table} SET " + ", ".join(f"{column}=?" for column in applied) + f" WHERE {where}"
        return sql, tuple(applied.values()) + params

    def _pedido_html(self, key: dict[str, Any]) -> str:
        detail = self._pedido_document_result(key, "preparacion")
        header = detail["cabecera"]
        lines_html = []
        for line in detail["lineas"]:
            lines_html.append(
                "<tr>"
                f"<td>{line.get('linea') or ''}</td>"
                f"<td>{line.get('articulo') or ''}</td>"
                f"<td>{line.get('descripcion') or ''}</td>"
                f"<td style=\"text-align:right\">{line.get('cantidad') or 0}</td>"
                f"<td style=\"text-align:right\">{line.get('precio') or 0}</td>"
                f"<td style=\"text-align:right\">{line.get('base') or 0}</td>"
                "</tr>"
            )
        title = f"Pedido {key['ejercicio']}/{key['serie']}/{key['numero']}"
        return (
            "<!doctype html><html><head><meta charset=\"utf-8\">"
            f"<title>{title}</title>"
            "<style>body{font-family:Arial,sans-serif;margin:32px;color:#1f2937}"
            "table{width:100%;border-collapse:collapse;margin-top:24px}"
            "th,td{border-bottom:1px solid #d1d5db;padding:8px;text-align:left}"
            "th{background:#f3f4f6}</style></head><body>"
            f"<h1>{title}</h1>"
            f"<p><strong>Cliente:</strong> {header.get('cbv_codcli')}/{header.get('cbv_subcli')} "
            f"{header.get('cbv_nomcli') or header.get('cli_razsoc') or header.get('cli_nomcli') or ''}</p>"
            f"<p><strong>Fecha:</strong> {header.get('cbv_fecha') or ''} "
            f"<strong>Total:</strong> {header.get('cbv_totald') or 0}</p>"
            "<table><thead><tr><th>Linea</th><th>Articulo</th><th>Descripcion</th>"
            "<th>Cantidad</th><th>Precio</th><th>Base</th></tr></thead><tbody>"
            + "".join(lines_html)
            + "</tbody></table></body></html>"
        )

    def _cartera_effect_type_label(self, value: Any) -> str:
        key = str(value or "").strip().upper()
        return CARTERA_EFFECT_TYPE_LABELS.get(key, key or "Sin tipo")

    def _cartera_reference_date(self, args: dict[str, Any]) -> date:
        text = self._date_arg(args.get("fecha_referencia")) if args.get("fecha_referencia") else date.today().isoformat()
        return date.fromisoformat(text)

    def _cartera_status(self, row: dict[str, Any], reference: date) -> str:
        if row.get("cbve_feccan"):
            return "cobrado"
        if row.get("cbve_fecimp"):
            return "impagado"
        due = row.get("cbve_fecvto")
        if due:
            try:
                due_date = date.fromisoformat(str(due)[:10])
                if due_date <= reference:
                    return "vencido"
            except ValueError:
                pass
        return "pendiente"

    def _cartera_is_remitted(self, row: dict[str, Any]) -> bool:
        return self._to_int(row.get("cbve_ejerem"), 0) != 0 or self._to_int(row.get("cbve_codrem"), 0) != 0

    def _cartera_pending_amount(self, row: dict[str, Any]) -> float:
        if row.get("cbve_feccan"):
            return 0.0
        return round(
            self._to_float(row.get("cbve_import"), 0)
            + self._to_float(row.get("cbve_impgas"), 0)
            - self._to_float(row.get("cbve_impcob"), 0),
            2,
        )

    def _cartera_where(self, args: dict[str, Any]) -> tuple[list[str], list[Any]]:
        empresa = self._empresa(args)
        where = ["E.CBVE_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("E.CBVE_CENTRO = ?")
            params.append(int(args["centro"]))
        if str(args.get("tipo_documento") or "").strip():
            where.append("E.CBVE_TIPDOC = ?")
            params.append(str(args["tipo_documento"]).strip().upper()[:1])
        if str(args.get("tipo_accion") or "").strip():
            where.append("E.CBVE_TIPAC = ?")
            params.append(str(args["tipo_accion"]).strip()[:1])
        if args.get("cliente") not in (None, "", 0):
            where.append("E.CBVE_CODCLI = ?")
            params.append(int(args["cliente"]))
        if args.get("subcliente") not in (None, ""):
            where.append("E.CBVE_SUBCLI = ?")
            params.append(int(args["subcliente"]))
        if args.get("cliente_desde") not in (None, "", 0):
            where.append("E.CBVE_CODCLI >= ?")
            params.append(int(args["cliente_desde"]))
        if args.get("cliente_hasta") not in (None, "", 0):
            where.append("E.CBVE_CODCLI <= ?")
            params.append(int(args["cliente_hasta"]))
        if args.get("ejercicio") not in (None, "", 0):
            where.append("E.CBVE_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        if str(args.get("serie") or "").strip():
            where.append("E.CBVE_SERIE = ?")
            params.append(str(args["serie"]).strip())
        if args.get("numero") not in (None, "", 0):
            where.append("E.CBVE_NUMDOC = ?")
            params.append(int(args["numero"]))
        if args.get("orden") not in (None, "", 0):
            where.append("E.CBVE_NUMORD = ?")
            params.append(int(args["orden"]))
        if str(args.get("tipo_efecto") or "").strip():
            where.append("E.CBVE_TIPOEF = ?")
            params.append(str(args["tipo_efecto"]).strip().upper()[:1])
        if args.get("remesa_ejercicio") not in (None, "", 0):
            where.append("E.CBVE_EJEREM = ?")
            params.append(int(args["remesa_ejercicio"]))
        if args.get("remesa_codigo") not in (None, "", 0):
            where.append("E.CBVE_CODREM = ?")
            params.append(int(args["remesa_codigo"]))
        for arg_name, field, op in (
            ("fecha_desde", "E.CBVE_FECHA", ">="),
            ("fecha_hasta", "E.CBVE_FECHA", "<="),
            ("vencimiento_desde", "E.CBVE_FECVTO", ">="),
            ("vencimiento_hasta", "E.CBVE_FECVTO", "<="),
        ):
            if args.get(arg_name):
                where.append(f"{field} {op} ?")
                params.append(self._date_arg(args[arg_name]))
        situacion = str(args.get("situacion") or "").strip().lower()
        if situacion in {"pendiente", "pendientes"}:
            where.append("E.CBVE_FECCAN IS NULL")
        elif situacion in {"cobrado", "cobrados", "cancelado", "cancelados"}:
            where.append("E.CBVE_FECCAN IS NOT NULL")
        elif situacion in {"impagado", "impagados"}:
            where.append("E.CBVE_FECCAN IS NULL")
            where.append("E.CBVE_FECIMP IS NOT NULL")
        elif situacion in {"vencido", "vencidos"}:
            where.append("E.CBVE_FECCAN IS NULL")
            where.append("E.CBVE_FECVTO <= ?")
            params.append(self._date_arg(args.get("fecha_referencia"), date.today().isoformat()))
        remitted = str(args.get("remesado") or "").strip().upper()
        if remitted == "S":
            where.append("(E.CBVE_EJEREM <> 0 OR E.CBVE_CODREM <> 0)")
        elif remitted == "N":
            where.append("E.CBVE_EJEREM = 0")
            where.append("E.CBVE_CODREM = 0")
        return where, params

    def _cartera_effect_rows(self, args: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        limit = _positive_limit(args.get("limite"), 500)
        reference = self._cartera_reference_date(args)
        where, params = self._cartera_where(args)
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   E.CBVE_NUMEMP, E.CBVE_CENTRO, E.CBVE_TIPDOC, E.CBVE_TIPAC,
                   E.CBVE_EJERCI, E.CBVE_SERIE, E.CBVE_NUMDOC, E.CBVE_NUMORD,
                   E.CBVE_TIPGES, E.CBVE_TIPOEF, E.CBVE_CODACEP, E.CBVE_FECHA,
                   E.CBVE_FECVTO, E.CBVE_IMPORT, E.CBVE_CODMON, E.CBVE_CODCLI,
                   E.CBVE_SUBCLI, E.CBVE_FECCAN, E.CBVE_IMPCOB, E.CBVE_IMPGAS,
                   E.CBVE_FECIMP, E.CBVE_INDEDI, E.CBVE_OBSERV, E.CBVE_EJEREM,
                   E.CBVE_CODREM, C.CLI_NOMCLI, C.CLI_RAZSOC, C.CLI_CIF,
                   V.CBV_NOMCLI, V.CBV_REFCLI, V.CBV_TOTALD
            FROM CABDOCVE E
            LEFT JOIN CLIEN C ON C.CLI_NUMEMP = E.CBVE_NUMEMP
                              AND C.CLI_CODCLI = E.CBVE_CODCLI
                              AND C.CLI_SUBCLI = E.CBVE_SUBCLI
            LEFT JOIN CABDOCV V ON V.CBV_NUMEMP = E.CBVE_NUMEMP
                               AND V.CBV_CENTRO = E.CBVE_CENTRO
                               AND V.CBV_TIPDOC = E.CBVE_TIPDOC
                               AND V.CBV_TIPAC = E.CBVE_TIPAC
                               AND V.CBV_EJERCI = E.CBVE_EJERCI
                               AND V.CBV_SERIE = E.CBVE_SERIE
                               AND V.CBV_NUMDOC = E.CBVE_NUMDOC
            WHERE {' AND '.join(where)}
            ORDER BY E.CBVE_FECVTO, E.CBVE_CODCLI, E.CBVE_SUBCLI, E.CBVE_EJERCI, E.CBVE_SERIE, E.CBVE_NUMDOC, E.CBVE_NUMORD
            """,
            tuple(params),
            limit,
        )
        effects: list[dict[str, Any]] = []
        totals = {"nominal": 0.0, "gastos": 0.0, "cobrado": 0.0, "pendiente": 0.0, "vencido": 0.0, "remesado": 0.0, "no_remesado": 0.0}
        for row in rows:
            status = self._cartera_status(row, reference)
            remitted = self._cartera_is_remitted(row)
            nominal = self._to_float(row.get("cbve_import"), 0)
            gastos = self._to_float(row.get("cbve_impgas"), 0)
            cobrado = self._to_float(row.get("cbve_impcob"), 0)
            pending = self._cartera_pending_amount(row)
            totals["nominal"] += nominal
            totals["gastos"] += gastos
            totals["cobrado"] += cobrado
            totals["pendiente"] += pending
            if status == "vencido":
                totals["vencido"] += pending
            if remitted:
                totals["remesado"] += pending
            else:
                totals["no_remesado"] += pending
            effects.append({
                "documento": {
                    "centro": self._to_int(row.get("cbve_centro"), 0),
                    "tipo": str(row.get("cbve_tipdoc") or ""),
                    "tipo_accion": str(row.get("cbve_tipac") or ""),
                    "ejercicio": self._to_int(row.get("cbve_ejerci"), 0),
                    "serie": str(row.get("cbve_serie") or ""),
                    "numero": self._to_int(row.get("cbve_numdoc"), 0),
                    "orden": self._to_int(row.get("cbve_numord"), 0),
                    "referencia_cliente": row.get("cbv_refcli"),
                    "total_documento": row.get("cbv_totald"),
                },
                "cliente": {
                    "codigo": self._to_int(row.get("cbve_codcli"), 0),
                    "subcliente": self._to_int(row.get("cbve_subcli"), 0),
                    "nombre": row.get("cli_razsoc") or row.get("cli_nomcli") or row.get("cbv_nomcli"),
                    "cif": row.get("cli_cif"),
                },
                "tipo_efecto": str(row.get("cbve_tipoef") or ""),
                "tipo_efecto_nombre": self._cartera_effect_type_label(row.get("cbve_tipoef")),
                "tipo_gestion": str(row.get("cbve_tipges") or ""),
                "aceptacion": str(row.get("cbve_codacep") or ""),
                "fecha": row.get("cbve_fecha"),
                "vencimiento": row.get("cbve_fecvto"),
                "fecha_impagado": row.get("cbve_fecimp"),
                "fecha_cancelacion": row.get("cbve_feccan"),
                "situacion": status,
                "remesado": remitted,
                "remesa": {"ejercicio": self._to_int(row.get("cbve_ejerem"), 0), "codigo": self._to_int(row.get("cbve_codrem"), 0)},
                "moneda": str(row.get("cbve_codmon") or ""),
                "nominal": round(nominal, 2),
                "gastos": round(gastos, 2),
                "cobrado": round(cobrado, 2),
                "pendiente": pending,
                "editable": str(row.get("cbve_indedi") or ""),
                "observaciones": row.get("cbve_observ"),
            })
        totals = {key: round(value, 2) for key, value in totals.items()}
        totals["efectos"] = len(effects)
        return effects, {"referencia": reference.isoformat(), "totales": totals}

    def _dashboard_period(self, args: dict[str, Any]) -> tuple[str, str]:
        today = date.today()
        start = self._date_arg(args.get("desde") or args.get("fecha_desde"), date(today.year, 1, 1).isoformat())
        end = self._date_arg(args.get("hasta") or args.get("fecha_hasta"), today.isoformat())
        if start > end:
            raise KofedasError("La fecha inicial no puede ser posterior a la fecha final")
        return start, end

    def _dashboard_sale_types(self, args: dict[str, Any]) -> list[str]:
        raw = args.get("tipos_documento")
        if raw in (None, ""):
            return list(DASHBOARD_SALE_DOCUMENTS)
        if isinstance(raw, str):
            values = [part.strip().upper()[:1] for part in raw.split(",") if part.strip()]
        elif isinstance(raw, list):
            values = [str(part).strip().upper()[:1] for part in raw if str(part).strip()]
        else:
            raise KofedasError("tipos_documento debe ser lista o texto separado por comas")
        values = [value for value in values if value]
        if not values:
            raise KofedasError("tipos_documento no puede estar vacio")
        return values

    def _dashboard_center_filter(self, args: dict[str, Any], alias: str, params: list[Any]) -> list[str]:
        if args.get("centro") is None:
            return []
        params.append(int(args["centro"]))
        return [f"{alias}_CENTRO = ?"]

    def _dashboard_sales_where(self, args: dict[str, Any], start: str, end: str, alias: str = "C") -> tuple[list[str], list[Any], list[str]]:
        empresa = self._empresa(args)
        types = self._dashboard_sale_types(args)
        where = [f"{alias}.CBV_NUMEMP = ?", f"{alias}.CBV_FECHA >= ?", f"{alias}.CBV_FECHA <= ?"]
        params: list[Any] = [empresa, start, end]
        if args.get("centro") is not None:
            where.append(f"{alias}.CBV_CENTRO = ?")
            params.append(int(args["centro"]))
        placeholders = ", ".join("?" for _ in types)
        where.append(f"{alias}.CBV_TIPDOC IN ({placeholders})")
        params.extend(types)
        return where, params, types

    def _dashboard_purchase_where(self, args: dict[str, Any], start: str, end: str, alias: str = "C") -> tuple[list[str], list[Any]]:
        empresa = self._empresa(args)
        where = [f"{alias}.CBM_NUMEMP = ?", f"{alias}.CBM_FECHA >= ?", f"{alias}.CBM_FECHA <= ?"]
        params: list[Any] = [empresa, start, end]
        if args.get("centro") is not None:
            where.append(f"{alias}.CBM_CENTRO = ?")
            params.append(int(args["centro"]))
        return where, params

    def _dashboard_sales_totals(self, args: dict[str, Any], start: str, end: str) -> dict[str, Any]:
        where, params, types = self._dashboard_sales_where(args, start, end, "C")
        row = self.db.one(
            f"""
            SELECT COUNT(*) AS DOCUMENTOS,
                   SUM(CASE WHEN C.CBV_TIPDOC = 'C' THEN -C.CBV_TOTALS ELSE C.CBV_TOTALS END) AS BASE,
                   SUM(CASE WHEN C.CBV_TIPDOC = 'C' THEN -C.CBV_TOTALD ELSE C.CBV_TOTALD END) AS TOTAL,
                   SUM(CASE WHEN C.CBV_TIPDOC = 'C' THEN -C.CBV_IMPCOB ELSE C.CBV_IMPCOB END) AS COBRADO
            FROM CABDOCV C
            WHERE {' AND '.join(where)}
            """,
            tuple(params),
        ) or {}
        return {
            "documentos": self._to_int(row.get("documentos"), 0),
            "base": round(self._to_float(row.get("base"), 0), 2),
            "total": round(self._to_float(row.get("total"), 0), 2),
            "cobrado": round(self._to_float(row.get("cobrado"), 0), 2),
            "pendiente": round(self._to_float(row.get("total"), 0) - self._to_float(row.get("cobrado"), 0), 2),
            "tipos_documento": types,
        }

    def _dashboard_purchase_totals(self, args: dict[str, Any], start: str, end: str) -> dict[str, Any]:
        where, params = self._dashboard_purchase_where(args, start, end, "C")
        row = self.db.one(
            f"""
            SELECT COUNT(*) AS DOCUMENTOS, SUM(C.CBM_TOTALS) AS BASE, SUM(C.CBM_TOTALD) AS TOTAL
            FROM CABDOCM C
            WHERE {' AND '.join(where)}
            """,
            tuple(params),
        ) or {}
        return {
            "documentos": self._to_int(row.get("documentos"), 0),
            "base": round(self._to_float(row.get("base"), 0), 2),
            "total": round(self._to_float(row.get("total"), 0), 2),
        }

    def _dashboard_group_expr(self, area: str, group_by: str) -> tuple[str, str, str, str]:
        group = group_by.strip().lower()
        if area == "ventas":
            mapping = {
                "mes": ("EXTRACT(YEAR FROM C.CBV_FECHA) * 100 + EXTRACT(MONTH FROM C.CBV_FECHA)", "EXTRACT(YEAR FROM C.CBV_FECHA) * 100 + EXTRACT(MONTH FROM C.CBV_FECHA)", "PERIODO", "PERIODO"),
                "anio": ("EXTRACT(YEAR FROM C.CBV_FECHA)", "EXTRACT(YEAR FROM C.CBV_FECHA)", "ANIO", "ANIO"),
                "centro": ("C.CBV_CENTRO", "C.CBV_CENTRO", "CENTRO", "CENTRO"),
                "tipo_documento": ("C.CBV_TIPDOC", "C.CBV_TIPDOC", "TIPO_DOCUMENTO", "TIPO_DOCUMENTO"),
                "cliente": ("C.CBV_CODCLI", "C.CBV_NOMCLI", "CODIGO", "NOMBRE"),
                "articulo": ("D.DMV_CODART", "D.DMV_DESCRI", "CODIGO", "NOMBRE"),
                "familia": ("A.ART_CODFAM", "F.FAM_DESCRI", "CODIGO", "NOMBRE"),
            }
        else:
            mapping = {
                "mes": ("EXTRACT(YEAR FROM C.CBM_FECHA) * 100 + EXTRACT(MONTH FROM C.CBM_FECHA)", "EXTRACT(YEAR FROM C.CBM_FECHA) * 100 + EXTRACT(MONTH FROM C.CBM_FECHA)", "PERIODO", "PERIODO"),
                "anio": ("EXTRACT(YEAR FROM C.CBM_FECHA)", "EXTRACT(YEAR FROM C.CBM_FECHA)", "ANIO", "ANIO"),
                "centro": ("C.CBM_CENTRO", "C.CBM_CENTRO", "CENTRO", "CENTRO"),
                "situacion": ("C.CBM_SITUAC", "C.CBM_SITUAC", "SITUACION", "SITUACION"),
                "proveedor": ("C.CBM_CODPRO", "C.CBM_NOMPRO", "CODIGO", "NOMBRE"),
                "articulo": ("D.DMM_CODART", "D.DMM_DESCRI", "CODIGO", "NOMBRE"),
                "familia": ("A.ART_CODFAM", "F.FAM_DESCRI", "CODIGO", "NOMBRE"),
            }
        if group not in mapping:
            raise KofedasError(f"agrupar_por no soportado para {area}: {group_by}")
        return mapping[group]

    def _regularization_stock_statements(self, empresa: int, centro: int, articulo: str, cantidad: float, fecha: str) -> list[tuple[str, tuple[Any, ...]]]:
        now = datetime.now().replace(microsecond=0).isoformat(sep=" ")
        exists = self.db.one(
            "SELECT FIRST 1 1 AS EXISTE FROM ARTICULE WHERE ARTE_NUMEMP = ? AND ARTE_CODART = ? AND ARTE_CENTRO = ?",
            (empresa, articulo, centro),
        )
        if exists:
            return [(
                """
                UPDATE ARTICULE
                SET ARTE_EXIST = ARTE_EXIST + ?, ARTE_FECMOV = ?
                WHERE ARTE_NUMEMP = ? AND ARTE_CODART = ? AND ARTE_CENTRO = ?
                """,
                (cantidad, now, empresa, articulo, centro),
            )]
        return [(
            """
            INSERT INTO ARTICULE (ARTE_NUMEMP, ARTE_CODART, ARTE_CENTRO, ARTE_EXIST, ARTE_MINIMO, ARTE_MAXIMO, ARTE_FECCOM, ARTE_FECVEN, ARTE_FECMOV)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (empresa, articulo, centro, cantidad, 0, 0, None, None, now),
        ), (
            "INSERT INTO STOCKS (STO_NUMEMP, STO_CODART, STO_CENTRO, STO_FECHA, STO_EXIST) VALUES (?, ?, ?, ?, ?)",
            (empresa, articulo, centro, fecha, 0),
        )]

    def _detmovr_row(self, empresa: int, centro: int, ejercicio: int, serie: str, numero: int, linea: int, fecha: str, articulo: str, descripcion: str, unidad: str, cantidad: float, origen: dict[str, Any] | None = None) -> dict[str, Any]:
        origen = origen or {}
        return {
            "DMR_NUMEMP": empresa,
            "DMR_CENTRO": centro,
            "DMR_EJERCI": ejercicio,
            "DMR_SERIE": serie,
            "DMR_NUMDOC": numero,
            "DMR_NUMLIN": linea,
            "DMR_FECMOV": fecha,
            "DMR_CODART": articulo,
            "DMR_DESCRI": descripcion[:100],
            "DMR_UNIMED": unidad[:4],
            "DMR_CANTID": cantidad,
            "DMR_EJERCIO": self._to_int(origen.get("ejercicio"), 0),
            "DMR_SERIEO": str(origen.get("serie") or "")[:2],
            "DMR_NUMDOCO": self._to_int(origen.get("numero"), 0),
            "DMR_NUMLINO": self._to_int(origen.get("linea"), 0),
        }

    def sistema_estado(self, args: dict[str, Any]) -> dict[str, Any]:
        del args
        return {
            "servidor": "kofedas-mcp",
            "version": SERVER_VERSION,
            "contrato": PUBLIC_CONTRACT_VERSION,
            "dsn": self.db.dsn,
            "empresa": self.empresa,
            "centro": self.centro,
            "perfil": self.tool_profile,
            "nivel_acceso": self.access_level,
        }

    def empresa_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        limit = _positive_limit(args.get("limite"))
        return self.db.query(
            """
            SELECT EMP_NUMEMP, EMP_CODSOC, EMP_NOMEMP, EMP_NOMFIS, EMP_DOMFIS1,
                   EMP_DOMFIS2, EMP_CODPOS, EMP_POBLAC, EMP_CIF, EMP_REGIVA,
                   EMP_EJEEUR, EMP_EMAIL
            FROM EMPRES
            ORDER BY EMP_NUMEMP
            """,
            (),
            limit,
        )

    def empresa_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        row = self.db.one(
            """
            SELECT EMP_NUMEMP, EMP_CODSOC, EMP_NOMEMP, EMP_NOMFIS, EMP_DOMFIS1,
                   EMP_DOMFIS2, EMP_CODPOS, EMP_POBLAC, EMP_CIF, EMP_REGIVA,
                   EMP_EJEEUR, EMP_EMAIL
            FROM EMPRES
            WHERE EMP_NUMEMP = ?
            """,
            (empresa,),
        )
        if row is None:
            raise KofedasError("Empresa no encontrada")
        return row

    def empresa_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        data = self._clean_data(args.get("datos"), EMPRESA_FIELDS, "EMP_", {"EMP_NUMEMP": empresa})
        return self._upsert("EMPRES", EMPRESA_FIELDS, ["EMP_NUMEMP"], data)

    def centro_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        return self.db.query(
            """
            SELECT CEN_NUMEMP, CEN_CODCEN, CEN_TIPCEN, CEN_NOMCEN, CEN_DOMIC1,
                   CEN_DOMIC2, CEN_CODPOS, CEN_POBLAC, CEN_TELEF, CEN_EMAIL,
                   CEN_CLIINI, CEN_CLIFIN, CEN_CLICINI, CEN_CLICFIN
            FROM CENTROS
            WHERE CEN_NUMEMP = ?
            ORDER BY CEN_CODCEN
            """,
            (empresa,),
        )

    def centro_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        row = self.db.one(
            """
            SELECT CEN_NUMEMP, CEN_CODCEN, CEN_TIPCEN, CEN_NOMCEN, CEN_DOMIC1,
                   CEN_DOMIC2, CEN_CODPOS, CEN_POBLAC, CEN_TELEF, CEN_FAX,
                   CEN_EMAIL, CEN_CLIINI, CEN_CLIFIN, CEN_CLICINI, CEN_CLICFIN
            FROM CENTROS
            WHERE CEN_NUMEMP = ? AND CEN_CODCEN = ?
            """,
            (empresa, centro),
        )
        if row is None:
            raise KofedasError("Centro no encontrado")
        return row

    def centro_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        centro = self._centro(args)
        data = self._clean_data(args.get("datos"), CENTRO_FIELDS, "CEN_", {"CEN_NUMEMP": empresa, "CEN_CODCEN": centro})
        return self._upsert("CENTROS", CENTRO_FIELDS, ["CEN_NUMEMP", "CEN_CODCEN"], data)

    def usuario_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["USU_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("USU_CODCEN = ?")
            params.append(int(args["centro"]))
        text = str(args.get("texto") or "").strip()
        if text:
            where.append("(UPPER(USU_NOMUSU) LIKE ? OR UPPER(USU_NOMBRE) LIKE ? OR UPPER(USU_IDGRUPO) LIKE ?)")
            params.extend([_like(text), _like(text), _like(text)])
        rows = self.db.query(
            f"""
            SELECT USU_NUMEMP, USU_CODCEN, USU_NOMUSU, USU_PASSWORD, USU_NOMBRE,
                   USU_IDGRUPO, USU_CONDGRU, USU_NIVEL, USU_NIVPRI, USU_MENUI,
                   USU_CAMFEC, USU_FECMIN, USU_FECMAX, USU_AUTVEN, USU_AUTDTO
            FROM USUAR
            WHERE {' AND '.join(where)}
            ORDER BY USU_CODCEN, USU_NOMUSU
            """,
            tuple(params),
            limit,
        )
        return [self._redact_usuario(row) for row in rows if row is not None]

    def usuario_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        usuario = str(args.get("usuario") or "").strip()
        if not usuario:
            raise KofedasError("usuario es obligatorio")
        row = self.db.one(
            """
            SELECT USU_NUMEMP, USU_CODCEN, USU_NOMUSU, USU_PASSWORD, USU_NOMBRE,
                   USU_IDGRUPO, USU_CONDGRU, USU_NIVEL, USU_NIVPRI, USU_MENUI,
                   USU_CAMFEC, USU_FECMIN, USU_FECMAX, USU_AUTVEN, USU_AUTDTO
            FROM USUAR
            WHERE USU_NUMEMP = ? AND USU_CODCEN = ? AND USU_NOMUSU = ?
            """,
            (empresa, centro, usuario),
        )
        if row is None:
            raise KofedasError("Usuario no encontrado")
        return self._redact_usuario(row) or {}

    def usuario_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        centro = self._centro(args)
        usuario = str(args.get("usuario") or "").strip()
        if not usuario:
            raise KofedasError("usuario es obligatorio")
        data = self._clean_data(
            args.get("datos"),
            USUARIO_FIELDS,
            "USU_",
            {"USU_NUMEMP": empresa, "USU_CODCEN": centro, "USU_NOMUSU": usuario},
        )
        return self._upsert("USUAR", USUARIO_FIELDS, ["USU_NUMEMP", "USU_CODCEN", "USU_NOMUSU"], data)

    def grupo_usuario_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["GRU_NUMEMP = ?"]
        params: list[Any] = [empresa]
        text = str(args.get("texto") or "").strip()
        if text:
            where.append("(UPPER(GRU_IDGRUPO) LIKE ? OR UPPER(GRU_NIVEL) LIKE ? OR UPPER(GRU_MENUI) LIKE ?)")
            params.extend([_like(text), _like(text), _like(text)])
        return self.db.query(
            f"""
            SELECT GRU_NUMEMP, GRU_IDGRUPO, GRU_NIVEL, GRU_NIVPRI, GRU_MENUI,
                   GRU_CAMFEC, GRU_FECMIN, GRU_FECMAX, GRU_AUTVEN, GRU_AUTDTO
            FROM GRUPUSU
            WHERE {' AND '.join(where)}
            ORDER BY GRU_IDGRUPO
            """,
            tuple(params),
            limit,
        )

    def grupo_usuario_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        grupo = str(args.get("grupo") or "").strip()
        if not grupo:
            raise KofedasError("grupo es obligatorio")
        row = self.db.one(
            """
            SELECT GRU_NUMEMP, GRU_IDGRUPO, GRU_NIVEL, GRU_NIVPRI, GRU_MENUI,
                   GRU_CAMFEC, GRU_FECMIN, GRU_FECMAX, GRU_AUTVEN, GRU_AUTDTO
            FROM GRUPUSU
            WHERE GRU_NUMEMP = ? AND GRU_IDGRUPO = ?
            """,
            (empresa, grupo),
        )
        if row is None:
            raise KofedasError("Grupo de usuario no encontrado")
        return row

    def grupo_usuario_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        grupo = str(args.get("grupo") or "").strip()
        if not grupo:
            raise KofedasError("grupo es obligatorio")
        data = self._clean_data(args.get("datos"), GRUPO_USUARIO_FIELDS, "GRU_", {"GRU_NUMEMP": empresa, "GRU_IDGRUPO": grupo})
        return self._upsert("GRUPUSU", GRUPO_USUARIO_FIELDS, ["GRU_NUMEMP", "GRU_IDGRUPO"], data)

    def parametro_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["PAR_NUMEMP = ?"]
        params: list[Any] = [empresa]
        text = str(args.get("texto") or "").strip()
        if text:
            where.append("(UPPER(PAR_CODIGO) LIKE ? OR UPPER(PAR_VALOR) LIKE ? OR UPPER(PAR_DESCRI) LIKE ?)")
            params.extend([_like(text), _like(text), _like(text)])
        return self.db.query(
            f"""
            SELECT PAR_NUMEMP, PAR_CODIGO, PAR_VALOR, PAR_DESCRI
            FROM PARAMETROS
            WHERE {' AND '.join(where)}
            ORDER BY PAR_CODIGO
            """,
            tuple(params),
            limit,
        )

    def parametro_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        codigo = str(args.get("codigo") or "").strip()
        if not codigo:
            raise KofedasError("codigo es obligatorio")
        row = self.db.one(
            """
            SELECT PAR_NUMEMP, PAR_CODIGO, PAR_VALOR, PAR_DESCRI
            FROM PARAMETROS
            WHERE PAR_NUMEMP = ? AND PAR_CODIGO = ?
            """,
            (empresa, codigo),
        )
        if row is None:
            raise KofedasError("Parametro no encontrado")
        return row

    def parametro_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        codigo = str(args.get("codigo") or "").strip()
        if not codigo:
            raise KofedasError("codigo es obligatorio")
        raw_data = dict(args.get("datos") or {})
        if "valor" in args:
            raw_data["PAR_VALOR"] = args["valor"]
        if "descripcion" in args:
            raw_data["PAR_DESCRI"] = args["descripcion"]
        data = self._clean_data(raw_data, PARAMETRO_FIELDS, "PAR_", {"PAR_NUMEMP": empresa, "PAR_CODIGO": codigo})
        return self._upsert("PARAMETROS", PARAMETRO_FIELDS, ["PAR_NUMEMP", "PAR_CODIGO"], data)

    def auxiliar_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(AUXILIARY_TABLES.items()):
            columns = self._table_columns(table)
            result.append({
                "tabla": table,
                "descripcion": description,
                "claves": self._table_pk(table),
                "columnas": columns,
            })
        return result

    def auxiliar_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        table = self._aux_table(args.get("tabla"))
        columns = self._table_columns(table)
        limit = _positive_limit(args.get("limite"))
        where: list[str] = []
        params: list[Any] = []
        empresa_column = self._empresa_column(columns)
        if args.get("empresa") is not None and empresa_column:
            where.append(f"{empresa_column} = ?")
            params.append(int(args["empresa"]))
        filters = self._normalize_column_map(args.get("filtros"), columns)
        for column, value in filters.items():
            where.append(f"{column} = ?")
            params.append(value)
        text = str(args.get("texto") or "").strip()
        if text:
            text_columns = self.db.text_columns(table)
            if not text_columns:
                return []
            text_clauses = [f"{column} CONTAINING ?" for column in text_columns]
            where.append("(" + " OR ".join(text_clauses) + ")")
            params.extend([text] * len(text_clauses))
        sql = f"SELECT FIRST {limit} " + ", ".join(columns) + f" FROM {table}"
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY " + ", ".join(self._table_pk(table))
        return self.db.query(sql, tuple(params), limit)

    def auxiliar_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        table = self._aux_table(args.get("tabla"))
        columns = self._table_columns(table)
        pk = self._table_pk(table)
        keys = self._normalize_column_map(args.get("claves"), columns, required=True)
        missing = [column for column in pk if column not in keys]
        if missing:
            raise KofedasError("Faltan claves: " + ", ".join(missing))
        where = " AND ".join(f"{column}=?" for column in pk)
        row = self.db.one(
            f"SELECT " + ", ".join(columns) + f" FROM {table} WHERE {where}",
            tuple(keys[column] for column in pk),
        )
        if row is None:
            raise KofedasError("Registro auxiliar no encontrado")
        return row

    def auxiliar_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        table = self._aux_table(args.get("tabla"))
        columns = self._table_columns(table)
        pk = self._table_pk(table)
        data = self._normalize_column_map(args.get("datos"), columns, required=True)
        keys = self._normalize_column_map(args.get("claves"), columns)
        data.update(keys)
        missing = [column for column in pk if column not in data]
        if missing:
            raise KofedasError("Faltan claves: " + ", ".join(missing))
        return self._upsert(table, columns, pk, data)

    def familia_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        tipo = str(args.get("tipo") or "familia").lower()
        limit = _positive_limit(args.get("limite"))
        if tipo == "familia":
            return self.db.query(
                """
                SELECT FAM_NUMEMP, FAM_CODIGO, FAM_DESCRI, FAM_TABLA
                FROM FAMILI
                WHERE FAM_NUMEMP = ?
                ORDER BY FAM_CODIGO
                """,
                (empresa,),
                limit,
            )
        if tipo == "subfamilia":
            where = ["SUB_NUMEMP = ?"]
            params: list[Any] = [empresa]
            if args.get("familia") is not None:
                where.append("SUB_CODFAM = ?")
                params.append(int(args["familia"]))
            return self.db.query(
                f"""
                SELECT SUB_NUMEMP, SUB_CODFAM, SUB_CODIGO, SUB_DESCRI, SUB_TABLA
                FROM SUBFAM
                WHERE {' AND '.join(where)}
                ORDER BY SUB_CODFAM, SUB_CODIGO
                """,
                tuple(params),
                limit,
            )
        if tipo == "ssubfamilia":
            where = ["SSUB_NUMEMP = ?"]
            params = [empresa]
            if args.get("familia") is not None:
                where.append("SSUB_CODFAM = ?")
                params.append(int(args["familia"]))
            if args.get("subfamilia") is not None:
                where.append("SSUB_CODSUB = ?")
                params.append(int(args["subfamilia"]))
            return self.db.query(
                f"""
                SELECT SSUB_NUMEMP, SSUB_CODFAM, SSUB_CODSUB, SSUB_CODIGO, SSUB_DESCRI, SSUB_TABLA
                FROM SSUBFAM
                WHERE {' AND '.join(where)}
                ORDER BY SSUB_CODFAM, SSUB_CODSUB, SSUB_CODIGO
                """,
                tuple(params),
                limit,
            )
        raise KofedasError("tipo debe ser familia, subfamilia o ssubfamilia")

    def articulo_buscar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        text = str(args.get("texto") or "").strip()
        if not text:
            raise KofedasError("texto es obligatorio")
        limit = _positive_limit(args.get("limite"))
        baja_filter = "" if args.get("incluir_baja") else "AND (A.ART_FEBAJA IS NULL)"
        fields = """
               A.ART_NUMEMP, A.ART_CODART, A.ART_DESCRI, A.ART_SECCIO,
               A.ART_CODFAM, A.ART_SUBFAM, A.ART_TIPIVA, A.ART_PRECOS,
               A.ART_PREVEN1, A.ART_PREVEN2, A.ART_PREVEN3, A.ART_PREVEN4,
               A.ART_PVP, A.ART_UNIMED, A.ART_CODPRO, A.ART_OBSOL, A.ART_FEBAJA
        """
        rows = self.db.query(
            f"""
            SELECT FIRST {limit} {fields}
            FROM ARTICUL A
            WHERE A.ART_NUMEMP = ?
              AND (A.ART_CODART = ? OR UPPER(A.ART_CODART) LIKE ? OR UPPER(A.ART_DESCRI) LIKE ?)
              {baja_filter}
            ORDER BY A.ART_DESCRI, A.ART_CODART
            """,
            (empresa, text, _like(text), _like(text)),
            limit,
        )
        seen = {row["art_codart"] for row in rows}
        remaining = limit - len(rows)
        if remaining > 0:
            for source_sql, params in (
                (
                    f"""
                    SELECT FIRST {remaining} {fields}
                    FROM ARTICULC C
                    JOIN ARTICUL A ON A.ART_NUMEMP = C.ARTC_NUMEMP AND A.ART_CODART = C.ARTC_CODART
                    WHERE C.ARTC_NUMEMP = ?
                      AND (C.ARTC_CODIGO = ? OR UPPER(C.ARTC_CODIGO) LIKE ?)
                      {baja_filter}
                    ORDER BY A.ART_DESCRI, A.ART_CODART
                    """,
                    (empresa, text, _like(text)),
                ),
                (
                    f"""
                    SELECT FIRST {remaining} {fields}
                    FROM ARTICULP P
                    JOIN ARTICUL A ON A.ART_NUMEMP = P.ARTP_NUMEMP AND A.ART_CODART = P.ARTP_CODART
                    WHERE P.ARTP_NUMEMP = ?
                      AND (P.ARTP_REFPRO = ? OR UPPER(P.ARTP_REFPRO) LIKE ?)
                      {baja_filter}
                    ORDER BY A.ART_DESCRI, A.ART_CODART
                    """,
                    (empresa, text, _like(text)),
                ),
            ):
                for row in self.db.query(source_sql, params, remaining):
                    if row["art_codart"] not in seen:
                        rows.append(row)
                        seen.add(row["art_codart"])
                        remaining -= 1
                    if remaining <= 0:
                        break
                if remaining <= 0:
                    break
        return rows

    def articulo_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        articulo = str(args.get("articulo") or "").strip()
        if not articulo:
            raise KofedasError("articulo es obligatorio")
        ficha = self.db.one(
            """
            SELECT FIRST 1 ART_NUMEMP, ART_CODART, ART_DESCRI, ART_SECCIO, ART_INDPROP, ART_INDBLI,
                   ART_TIPPRE, ART_INDINV, ART_OBSOL, ART_FEALTA, ART_FEBAJA, ART_FECMOV,
                   ART_FECVAR, ART_CODFAM, ART_SUBFAM, ART_TIPIVA, ART_PRETAR, ART_FECTAR,
                   ART_PREBAS, ART_PRECOS, ART_TABPREC, ART_PREVEN1, ART_PREVEN2,
                   ART_PREVEN3, ART_PREVEN4, ART_PVP, ART_CODMON, ART_UNIMED, ART_CANPRE,
                   ART_CANPMI, ART_CODPRO, ART_AGRUP1, ART_AGRUP2, ART_AGRUP3, ART_NORMA
            FROM ARTICUL
            WHERE ART_NUMEMP = ? AND (ART_CODART = ? OR UPPER(ART_CODART) LIKE ?)
            ORDER BY ART_CODART
            """,
            (empresa, articulo, _like(articulo)),
        )
        if ficha is None:
            raise KofedasError("Articulo no encontrado")
        articulo_resuelto = ficha["art_codart"]
        return {
            "articulo": ficha,
            "codigos_barras": self.db.query(
                """
                SELECT ARTC_CODIGO, ARTC_CANTID
                FROM ARTICULC
                WHERE ARTC_NUMEMP = ? AND (ARTC_CODART = ? OR UPPER(ARTC_CODART) LIKE ?)
                ORDER BY ARTC_CODIGO
                """,
                (empresa, articulo_resuelto, _like(articulo_resuelto)),
            ),
            "stock_centros": self.db.query(
                """
                SELECT E.ARTE_CENTRO, C.CEN_NOMCEN, E.ARTE_EXIST, E.ARTE_MINIMO, E.ARTE_MAXIMO,
                       E.ARTE_FECCOM, E.ARTE_FECVEN, E.ARTE_FECMOV
                FROM ARTICULE E
                LEFT JOIN CENTROS C ON C.CEN_NUMEMP = E.ARTE_NUMEMP AND C.CEN_CODCEN = E.ARTE_CENTRO
                WHERE E.ARTE_NUMEMP = ? AND (E.ARTE_CODART = ? OR UPPER(E.ARTE_CODART) LIKE ?)
                ORDER BY E.ARTE_CENTRO
                """,
                (empresa, articulo_resuelto, _like(articulo_resuelto)),
            ),
            "proveedores": self.db.query(
                """
                SELECT P.ARTP_CODPRO, V.PRO_NOMCOR, P.ARTP_REFPRO, P.ARTP_DESCRI,
                       P.ARTP_UNIMED, P.ARTP_CANCON, P.ARTP_CANVEN, P.ARTP_UNIPAQ,
                       P.ARTP_PREBAS, P.ARTP_DTOAUM1, P.ARTP_DTOAUM2, P.ARTP_DTOAUM3,
                       P.ARTP_DTOAUM4, P.ARTP_DTOAUM5, P.ARTP_DTOAUM6, P.ARTP_CODMON,
                       P.ARTP_CANPRE, P.ARTP_UBICA, P.ARTP_AMPUNIV, P.ARTP_AJUSTE
                FROM ARTICULP P
                LEFT JOIN PROVEE V ON V.PRO_NUMEMP = P.ARTP_NUMEMP AND V.PRO_CODPRO = P.ARTP_CODPRO
                WHERE P.ARTP_NUMEMP = ? AND (P.ARTP_CODART = ? OR UPPER(P.ARTP_CODART) LIKE ?)
                ORDER BY P.ARTP_CODPRO
                """,
                (empresa, articulo_resuelto, _like(articulo_resuelto)),
            ),
        }

    def stock_consultar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        articulo = str(args.get("articulo") or "").strip()
        limit = _positive_limit(args.get("limite"))
        where = ["E.ARTE_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if centro != 0:
            where.append("E.ARTE_CENTRO = ?")
            params.append(centro)
        if articulo:
            where.append("(E.ARTE_CODART = ? OR UPPER(E.ARTE_CODART) LIKE ? OR UPPER(A.ART_DESCRI) LIKE ?)")
            params.extend([articulo, _like(articulo), _like(articulo)])
        actual = self.db.query(
            f"""
            SELECT E.ARTE_CODART, A.ART_DESCRI, E.ARTE_CENTRO, C.CEN_NOMCEN,
                   E.ARTE_EXIST, E.ARTE_MINIMO, E.ARTE_MAXIMO, E.ARTE_FECCOM,
                   E.ARTE_FECVEN, E.ARTE_FECMOV
            FROM ARTICULE E
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = E.ARTE_NUMEMP AND A.ART_CODART = E.ARTE_CODART
            LEFT JOIN CENTROS C ON C.CEN_NUMEMP = E.ARTE_NUMEMP AND C.CEN_CODCEN = E.ARTE_CENTRO
            WHERE {' AND '.join(where)}
            ORDER BY E.ARTE_CENTRO, E.ARTE_CODART
            """,
            tuple(params),
            limit,
        )
        result: dict[str, Any] = {"actual": actual}
        if args.get("incluir_historico"):
            hist_where = ["S.STO_NUMEMP = ?"]
            hist_params: list[Any] = [empresa]
            if centro != 0:
                hist_where.append("S.STO_CENTRO = ?")
                hist_params.append(centro)
            if articulo:
                hist_where.append("(S.STO_CODART = ? OR UPPER(S.STO_CODART) LIKE ?)")
                hist_params.extend([articulo, _like(articulo)])
            result["historico"] = self.db.query(
                f"""
                SELECT S.STO_CODART, S.STO_CENTRO, S.STO_FECHA, S.STO_EXIST
                FROM STOCKS S
                WHERE {' AND '.join(hist_where)}
                ORDER BY S.STO_FECHA DESC, S.STO_CENTRO, S.STO_CODART
                """,
                tuple(hist_params),
                limit,
            )
        return result

    def articulo_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(ARTICLE_TABLES.items()):
            columns = self._table_columns(table)
            result.append({
                "tabla": table,
                "descripcion": description,
                "claves": self._table_pk(table),
                "columnas": columns,
                "solo_lectura": table not in ARTICLE_WRITABLE_TABLES,
            })
        return result

    def articulo_relacion_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        table = self._article_table(args.get("tabla"))
        columns = self._table_columns(table)
        limit = _positive_limit(args.get("limite"))
        where: list[str] = []
        params: list[Any] = []
        filters = self._normalize_column_map(args.get("filtros"), columns)
        empresa = args.get("empresa")
        articulo = str(args.get("articulo") or "").strip()
        for column in columns:
            if column.endswith("_NUMEMP") and empresa is not None:
                filters.setdefault(column, int(empresa))
            elif column.endswith("_CODART") and articulo:
                filters.setdefault(column, articulo)
        for column, value in filters.items():
            where.append(f"{column} = ?")
            params.append(value)
        text = str(args.get("texto") or "").strip()
        if text:
            text_columns = self.db.text_columns(table)
            if not text_columns:
                return []
            where.append("(" + " OR ".join(f"{column} CONTAINING ?" for column in text_columns) + ")")
            params.extend([text] * len(text_columns))
        sql = f"SELECT FIRST {limit} " + ", ".join(columns) + f" FROM {table}"
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY " + ", ".join(self._table_pk(table))
        return self.db.query(sql, tuple(params), limit)

    def articulo_relacion_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        table = self._article_table(args.get("tabla"))
        columns = self._table_columns(table)
        pk = self._table_pk(table)
        keys = self._normalize_column_map(args.get("claves"), columns, required=True)
        missing = [column for column in pk if column not in keys]
        if missing:
            raise KofedasError("Faltan claves: " + ", ".join(missing))
        where = " AND ".join(f"{column}=?" for column in pk)
        row = self.db.one(
            "SELECT " + ", ".join(columns) + f" FROM {table} WHERE {where}",
            tuple(keys[column] for column in pk),
        )
        if row is None:
            raise KofedasError("Registro de articulo no encontrado")
        return row

    def articulo_relacion_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        table = self._article_table(args.get("tabla"))
        if table not in ARTICLE_WRITABLE_TABLES:
            raise KofedasError("Tabla de articulo de solo lectura: " + table)
        columns = self._table_columns(table)
        pk = self._table_pk(table)
        data = self._normalize_column_map(args.get("datos"), columns, required=True)
        keys = self._normalize_column_map(args.get("claves"), columns)
        data.update(keys)
        missing = [column for column in pk if column not in data]
        if missing:
            raise KofedasError("Faltan claves: " + ", ".join(missing))
        return self._upsert(table, columns, pk, data)

    def articulo_completo(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        articulo = str(args.get("articulo") or "").strip()
        limit = _positive_limit(args.get("limite_detalle"), default=500)
        main = self.articulo_obtener({"empresa": empresa, "articulo": articulo})
        resolved = main["articulo"]["art_codart"]
        details: dict[str, Any] = {}
        for table in ARTICLE_TABLES:
            if table == "ARTICUL":
                continue
            details[table.lower()] = self.articulo_relacion_listar({
                "tabla": table,
                "empresa": empresa,
                "articulo": resolved,
                "limite": limit,
            })
        return {"articulo": main["articulo"], "basico": main, "relaciones": details}

    def articulo_alta_preparar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        articulo = str(args.get("articulo") or "").strip()
        if not articulo:
            raise KofedasError("articulo es obligatorio")
        data = self._article_field_data(args.get("datos"))
        data["ART_CODART"] = articulo
        return {
            "articulo": articulo,
            "datos": self._article_defaults(empresa, articulo, centro, data),
            "adicionales_reconocidos": ARTICLE_ADDITIONAL_CODES,
            "fuente_delphi": {
                "defaults": "ARTICUL_UDM.INICIALIZAR_ARTICUL",
                "alta": "ARTICUL_UDM.GRABAR_ARTICUL('G')",
                "precios": "ARTICUL_UDM.CALCULAR_PRECIO_ARTICUL, aproximado sin reglas de TABPREC avanzadas",
            },
        }

    def articulo_alta(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        plan = self.articulo_alta_preparar(args)
        article = plan["datos"]
        if not str(article.get("ART_DESCRI") or "").strip():
            raise KofedasError("descripcion o ART_DESCRI es obligatorio")
        empresa = int(article["ART_NUMEMP"])
        articulo = str(article["ART_CODART"])
        if self._article_exists(empresa, articulo):
            raise KofedasError("Ya existe articulo indicado")
        statements = self._article_insert_statements(
            article,
            self._article_purchase_data(args.get("compra")),
            args.get("codigos_barras") or [],
            args.get("stock") or {},
            args.get("informacion_adicional") or {},
            self._centro(args),
        )
        if args.get("simular"):
            return {"simulado": True, "articulo": articulo, "sentencias": len(statements), "datos": article}
        counts = self.db.execute_transaction(statements)
        return {"simulado": False, "articulo": articulo, "filas_afectadas": counts}

    def articulo_tarifa_excel_previsualizar(self, args: dict[str, Any]) -> dict[str, Any]:
        preview_args = dict(args)
        preview_args["simular"] = True
        plan = self._article_tariff_plan(preview_args)
        plan.pop("_statements", None)
        return plan

    def articulo_tarifa_excel_importar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        plan = self._article_tariff_plan(args)
        if args.get("simular"):
            plan.pop("_statements", None)
            plan["simulado"] = True
            return plan
        if plan["resumen"]["errores"]:
            plan.pop("_statements", None)
            raise KofedasError("La tarifa contiene errores; usa articulo_tarifa_excel_previsualizar para revisar el detalle")
        counts = self.db.execute_transaction(plan.pop("_statements"))
        plan["simulado"] = False
        plan["filas_afectadas"] = counts
        return plan

    def _article_insert_statements(
        self,
        article: dict[str, Any],
        purchase_data: dict[str, Any],
        barcodes: Any,
        stock: Any,
        additional: Any,
        centro: int,
    ) -> list[tuple[str, tuple[Any, ...]]]:
        columns = self._table_columns("ARTICUL")
        statements: list[tuple[str, tuple[Any, ...]]] = [(
            "INSERT INTO ARTICUL (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
            tuple(article.get(column) for column in columns),
        )]
        empresa = int(article["ART_NUMEMP"])
        articulo = str(article["ART_CODART"])
        proveedor = self._to_int(purchase_data.get("ARTP_CODPRO"), self._to_int(article.get("ART_CODPRO"), 0))
        if proveedor:
            purchase = self._purchase_defaults(empresa, articulo, proveedor, purchase_data, article)
            if not purchase.get("ARTP_REFPRO"):
                purchase["ARTP_REFPRO"] = articulo
            pcols = self._table_columns("ARTICULP")
            statements.append((
                "INSERT INTO ARTICULP (" + ", ".join(pcols) + ") VALUES (" + ", ".join("?" for _ in pcols) + ")",
                tuple(purchase.get(column) for column in pcols),
            ))
        if isinstance(barcodes, (str, int, float)):
            barcodes = [barcodes]
        if not isinstance(barcodes, list):
            raise KofedasError("codigos_barras debe ser una lista")
        for item in barcodes:
            code = str(item.get("codigo") if isinstance(item, dict) else item).strip()
            if not code:
                continue
            quantity = self._to_float(item.get("cantidad") if isinstance(item, dict) else 1, 1)
            statements.append((
                "INSERT INTO ARTICULC (ARTC_NUMEMP, ARTC_CODART, ARTC_CODIGO, ARTC_CANTID) VALUES (?, ?, ?, ?)",
                (empresa, articulo, code, quantity),
            ))
        if isinstance(stock, dict) and any(value not in (None, "") for value in stock.values()):
            stock_centro = self._to_int(stock.get("centro"), centro)
            statements.append((
                "INSERT INTO ARTICULE (ARTE_NUMEMP, ARTE_CODART, ARTE_CENTRO, ARTE_EXIST, ARTE_MINIMO, ARTE_MAXIMO, ARTE_FECCOM, ARTE_FECVEN, ARTE_FECMOV) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    empresa, articulo, stock_centro,
                    self._to_float(stock.get("existencias"), 0),
                    self._to_float(stock.get("minimo"), 0),
                    self._to_float(stock.get("maximo"), 0),
                    None, None, None,
                ),
            ))
        if not isinstance(additional, dict):
            raise KofedasError("informacion_adicional debe ser un objeto CODINF -> valor")
        line = 1
        for code, value in additional.items():
            code_text = str(code).strip().upper()
            if not code_text or value in (None, ""):
                continue
            statements.append((
                "INSERT INTO ARTICULI (ARTI_NUMEMP, ARTI_CODART, ARTI_NUMLIN, ARTI_CODINF, ARTI_DESCRI) VALUES (?, ?, ?, ?, ?)",
                (empresa, articulo, line, code_text, str(value)),
            ))
            line += 1
        return statements

    def _article_tariff_plan(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        limit = _positive_limit(args.get("limite"), default=500)
        rows = self._xlsx_rows(str(args["archivo"]), args.get("hoja"), int(args.get("fila_cabecera") or 1), limit)
        statements: list[tuple[str, tuple[Any, ...]]] = []
        resumen = {"filas_excel": len(rows), "altas_articulo": 0, "actualizaciones_articulo": 0, "altas_compra": 0, "actualizaciones_compra": 0, "altas_barras": 0, "actualizaciones_barras": 0, "altas_stock": 0, "actualizaciones_stock": 0, "errores": 0}
        detalle: list[dict[str, Any]] = []
        reserved: set[str] = set()
        counter = int(args.get("numero_inicial") or 1)
        create_articles = args.get("crear_articulos", True) is not False
        update_articles = args.get("actualizar_articulos", True) is not False
        create_barcodes = args.get("crear_barras", True) is not False
        update_barcodes = bool(args.get("actualizar_barras"))
        for index, row in enumerate(rows, start=int(args.get("fila_cabecera") or 1) + 1):
            try:
                parsed = self._article_tariff_row(row, args)
                articulo = parsed["articulo"]
                proveedor = self._to_int(parsed.get("proveedor"), self._to_int(args.get("proveedor"), 0))
                seccion = str(parsed.get("seccion") or args.get("seccion") or "").strip()
                if not articulo and args.get("generar_codigos"):
                    if not proveedor or not seccion:
                        raise KofedasError("Para generar codigos se requiere proveedor y seccion")
                    articulo, counter = self._next_generated_article_code(empresa, seccion, proveedor, counter, int(args.get("digitos") or 13), reserved)
                if not articulo:
                    raise KofedasError("Fila sin codigo de articulo")
                article_data = dict(parsed["article"])
                article_data["ART_CODART"] = articulo
                if proveedor:
                    article_data.setdefault("ART_CODPRO", proveedor)
                if seccion:
                    article_data.setdefault("ART_SECCIO", seccion)
                exists = self._article_exists(empresa, articulo)
                article = self._article_defaults(empresa, articulo, centro, article_data)
                row_actions: list[str] = []
                columns = self._table_columns("ARTICUL")
                if exists and update_articles:
                    update_fields = [column for column in columns if column not in ("ART_NUMEMP", "ART_CODART")]
                    statements.append((
                        "UPDATE ARTICUL SET " + ", ".join(f"{column}=?" for column in update_fields) + " WHERE ART_NUMEMP=? AND ART_CODART=?",
                        tuple(article.get(column) for column in update_fields) + (empresa, articulo),
                    ))
                    resumen["actualizaciones_articulo"] += 1
                    row_actions.append("actualizar_articulo")
                elif not exists and create_articles:
                    statements.append((
                        "INSERT INTO ARTICUL (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
                        tuple(article.get(column) for column in columns),
                    ))
                    resumen["altas_articulo"] += 1
                    row_actions.append("alta_articulo")
                elif not exists:
                    raise KofedasError("Articulo no existe y crear_articulos=false")
                purchase = self._purchase_defaults(empresa, articulo, proveedor, parsed["purchase"], article) if proveedor else None
                if purchase:
                    if not purchase.get("ARTP_REFPRO"):
                        purchase["ARTP_REFPRO"] = str(parsed.get("referencia") or articulo)
                    pcols = self._table_columns("ARTICULP")
                    pexists = self.db.one(
                        "SELECT FIRST 1 1 AS EXISTE FROM ARTICULP WHERE ARTP_NUMEMP=? AND ARTP_CODART=? AND ARTP_CODPRO=? AND ARTP_REFPRO=?",
                        (empresa, articulo, proveedor, purchase["ARTP_REFPRO"]),
                    ) is not None
                    if pexists:
                        update_fields = [column for column in pcols if column not in ("ARTP_NUMEMP", "ARTP_CODART", "ARTP_CODPRO", "ARTP_REFPRO")]
                        statements.append((
                            "UPDATE ARTICULP SET " + ", ".join(f"{column}=?" for column in update_fields) + " WHERE ARTP_NUMEMP=? AND ARTP_CODART=? AND ARTP_CODPRO=? AND ARTP_REFPRO=?",
                            tuple(purchase.get(column) for column in update_fields) + (empresa, articulo, proveedor, purchase["ARTP_REFPRO"]),
                        ))
                        resumen["actualizaciones_compra"] += 1
                        row_actions.append("actualizar_compra")
                    else:
                        statements.append((
                            "INSERT INTO ARTICULP (" + ", ".join(pcols) + ") VALUES (" + ", ".join("?" for _ in pcols) + ")",
                            tuple(purchase.get(column) for column in pcols),
                        ))
                        resumen["altas_compra"] += 1
                        row_actions.append("alta_compra")
                barcode = str(parsed.get("codigo_barras") or "").strip()
                if barcode and create_barcodes:
                    owner = self._barcode_owner(empresa, barcode)
                    quantity = self._to_float(parsed.get("cantidad_barras"), 1)
                    if owner is None:
                        statements.append(("INSERT INTO ARTICULC (ARTC_NUMEMP, ARTC_CODART, ARTC_CODIGO, ARTC_CANTID) VALUES (?, ?, ?, ?)", (empresa, articulo, barcode, quantity)))
                        resumen["altas_barras"] += 1
                        row_actions.append("alta_barra")
                    elif owner == articulo or update_barcodes:
                        statements.append(("UPDATE ARTICULC SET ARTC_CODART=?, ARTC_CANTID=? WHERE ARTC_NUMEMP=? AND ARTC_CODIGO=?", (articulo, quantity, empresa, barcode)))
                        resumen["actualizaciones_barras"] += 1
                        row_actions.append("actualizar_barra")
                    else:
                        raise KofedasError(f"Codigo de barras {barcode} ya pertenece a {owner}")
                if parsed["stock"]:
                    stock_data = parsed["stock"]
                    stock_centro = self._to_int(stock_data.get("centro"), centro)
                    sexists = self.db.one(
                        "SELECT FIRST 1 1 AS EXISTE FROM ARTICULE WHERE ARTE_NUMEMP=? AND ARTE_CODART=? AND ARTE_CENTRO=?",
                        (empresa, articulo, stock_centro),
                    ) is not None
                    values = (
                        self._to_float(stock_data.get("existencias"), 0),
                        self._to_float(stock_data.get("minimo"), 0),
                        self._to_float(stock_data.get("maximo"), 0),
                    )
                    if sexists:
                        statements.append(("UPDATE ARTICULE SET ARTE_EXIST=?, ARTE_MINIMO=?, ARTE_MAXIMO=? WHERE ARTE_NUMEMP=? AND ARTE_CODART=? AND ARTE_CENTRO=?", values + (empresa, articulo, stock_centro)))
                        resumen["actualizaciones_stock"] += 1
                        row_actions.append("actualizar_stock")
                    else:
                        statements.append(("INSERT INTO ARTICULE (ARTE_NUMEMP, ARTE_CODART, ARTE_CENTRO, ARTE_EXIST, ARTE_MINIMO, ARTE_MAXIMO, ARTE_FECCOM, ARTE_FECVEN, ARTE_FECMOV) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", (empresa, articulo, stock_centro, *values, None, None, None)))
                        resumen["altas_stock"] += 1
                        row_actions.append("alta_stock")
                detalle.append({"fila": index, "articulo": articulo, "acciones": row_actions})
            except Exception as exc:
                resumen["errores"] += 1
                detalle.append({"fila": index, "error": str(exc)})
        return {"simulado": True, "resumen": resumen, "detalle": detalle[:100], "_statements": statements}

    def _article_tariff_row(self, row: dict[str, Any], args: dict[str, Any]) -> dict[str, Any]:
        mapping = {self._header_key(key): self._header_key(value) for key, value in (args.get("mapeo") or {}).items()}
        aliases = {
            "articulo": ["articulo", "codigo", "codart", "cod_art", "refer", "referencia_articulo"],
            "descripcion": ["descripcion", "descri", "nombre", "art_descri"],
            "seccion": ["seccion", "sec", "art_seccio"],
            "proveedor": ["proveedor", "codpro", "codigo_proveedor", "art_codpro", "artp_codpro"],
            "referencia": ["referencia", "referencia_proveedor", "refpro", "ref_pro", "artp_refpro"],
            "precio_compra": ["precio_compra", "precio", "coste", "artp_prebas", "art_prebas"],
            "dto1": ["dto1", "dto", "descuento", "artp_dtoaum1"],
            "dto2": ["dto2", "artp_dtoaum2"],
            "dto3": ["dto3", "artp_dtoaum3"],
            "dto4": ["dto4", "artp_dtoaum4"],
            "dto5": ["dto5", "artp_dtoaum5"],
            "dto6": ["dto6", "artp_dtoaum6"],
            "codigo_barras": ["codigo_barras", "codbar", "ean", "barcode", "artc_codigo"],
            "cantidad_barras": ["cantidad_barras", "cantidad_barra", "artc_cantid"],
            "unidad": ["unidad", "unimed", "unidad_medida", "art_unimed", "artp_unimed"],
            "familia": ["familia", "codfam", "art_codfam"],
            "subfamilia": ["subfamilia", "subfam", "art_subfam"],
            "iva": ["iva", "tipiva", "art_tipiva"],
            "pvp": ["pvp", "art_pvp"],
            "precio_venta1": ["precio_venta1", "preven1", "art_preven1"],
            "precio_venta2": ["precio_venta2", "preven2", "art_preven2"],
            "precio_venta3": ["precio_venta3", "preven3", "art_preven3"],
            "precio_venta4": ["precio_venta4", "preven4", "art_preven4"],
            "stock": ["stock", "existencias", "arte_exist"],
            "minimo": ["minimo", "stock_minimo", "arte_minimo"],
            "maximo": ["maximo", "stock_maximo", "arte_maximo"],
            "ubicacion": ["ubicacion", "ubica", "artp_ubica"],
        }
        def value(name: str) -> Any:
            keys = [mapping[name]] if name in mapping else []
            keys.extend(aliases.get(name, []))
            for key in keys:
                if key in row and row[key] not in (None, ""):
                    return row[key]
            return None
        articulo = str(value("articulo") or "").strip()
        descripcion = str(value("descripcion") or "").strip()
        proveedor = value("proveedor")
        precio = value("precio_compra")
        article: dict[str, Any] = {}
        if descripcion:
            article["ART_DESCRI"] = descripcion
        if value("seccion") is not None:
            article["ART_SECCIO"] = str(value("seccion")).strip()
        if value("familia") is not None:
            article["ART_CODFAM"] = self._to_int(value("familia"))
        if value("subfamilia") is not None:
            article["ART_SUBFAM"] = self._to_int(value("subfamilia"))
        if value("iva") is not None:
            article["ART_TIPIVA"] = self._to_int(value("iva"), 1)
        if value("unidad") is not None:
            article["ART_UNIMED"] = str(value("unidad")).strip()[:4]
        if proveedor is not None:
            article["ART_CODPRO"] = self._to_int(proveedor)
        if precio is not None:
            article["ART_PREBAS"] = self._to_float(precio)
        for target, source in (("ART_PVP", "pvp"), ("ART_PREVEN1", "precio_venta1"), ("ART_PREVEN2", "precio_venta2"), ("ART_PREVEN3", "precio_venta3"), ("ART_PREVEN4", "precio_venta4")):
            if value(source) is not None:
                article[target] = self._to_float(value(source))
        purchase: dict[str, Any] = {}
        if proveedor is not None:
            purchase["ARTP_CODPRO"] = self._to_int(proveedor)
        if value("referencia") is not None:
            purchase["ARTP_REFPRO"] = str(value("referencia")).strip()
        if descripcion:
            purchase["ARTP_DESCRI"] = descripcion
        if value("unidad") is not None:
            purchase["ARTP_UNIMED"] = str(value("unidad")).strip()[:4]
        if precio is not None:
            purchase["ARTP_PREBAS"] = self._to_float(precio)
        for index in range(1, 7):
            if value(f"dto{index}") is not None:
                purchase[f"ARTP_DTOAUM{index}"] = self._to_float(value(f"dto{index}"))
        if value("ubicacion") is not None:
            purchase["ARTP_UBICA"] = str(value("ubicacion")).strip()[:6]
        stock: dict[str, Any] = {}
        if value("stock") is not None:
            stock["existencias"] = self._to_float(value("stock"))
        if value("minimo") is not None:
            stock["minimo"] = self._to_float(value("minimo"))
        if value("maximo") is not None:
            stock["maximo"] = self._to_float(value("maximo"))
        return {
            "articulo": articulo,
            "proveedor": proveedor,
            "seccion": value("seccion"),
            "referencia": value("referencia"),
            "codigo_barras": value("codigo_barras"),
            "cantidad_barras": value("cantidad_barras"),
            "article": article,
            "purchase": purchase,
            "stock": stock,
        }

    def cliente_buscar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        text = str(args.get("texto") or "").strip()
        if not text:
            raise KofedasError("texto es obligatorio")
        limit = _positive_limit(args.get("limite"))
        baja_filter = "" if args.get("incluir_baja") else "AND (CLI_FEBAJA IS NULL)"
        params: list[Any] = [empresa]
        numeric = int(text) if text.isdigit() else None
        params.extend([numeric if numeric is not None else -1, _like(text), _like(text), _like(text), _like(text), _like(text)])
        return self.db.query(
            f"""
            SELECT CLI_NUMEMP, CLI_CODCLI, CLI_SUBCLI, CLI_NOMCLI, CLI_RAZSOC,
                   CLI_DOMICI, CLI_CODPOS, CLI_POBLAC, CLI_CIF, CLI_TELEFO,
                   CLI_EMAIL, CLI_ZONA, CLI_CODREP, CLI_RIESGO, CLI_CODMON,
                   CLI_TIPPRE, CLI_FEALTA, CLI_FEBAJA
            FROM CLIEN
            WHERE CLI_NUMEMP = ?
              AND (
                CLI_CODCLI = ?
                OR UPPER(CLI_NOMCLI) LIKE ?
                OR UPPER(CLI_RAZSOC) LIKE ?
                OR UPPER(CLI_CIF) LIKE ?
                OR UPPER(CLI_TELEFO) LIKE ?
                OR UPPER(CLI_EMAIL) LIKE ?
              )
              {baja_filter}
            ORDER BY CLI_NOMCLI, CLI_CODCLI, CLI_SUBCLI
            """,
            tuple(params),
            limit,
        )

    def cliente_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        cliente = int(args["cliente"])
        subcliente = int(args.get("subcliente") or 0)
        ficha = self.db.one(
            "SELECT * FROM CLIEN WHERE CLI_NUMEMP = ? AND CLI_CODCLI = ? AND CLI_SUBCLI = ?",
            (empresa, cliente, subcliente),
        )
        if ficha is None:
            raise KofedasError("Cliente no encontrado")
        return {
            "cliente": ficha,
            "informacion_adicional": self.db.query(
                """
                SELECT CLII_NUMLIN, CLII_CODINF, CLII_DESCRI
                FROM CLIENI
                WHERE CLII_NUMEMP = ? AND CLII_CODCLI = ? AND CLII_SUBCLI = ?
                ORDER BY CLII_NUMLIN
                """,
                (empresa, cliente, subcliente),
            ),
        }

    def cliente_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(CLIENT_TABLES.items()):
            columns = self._table_columns(table)
            result.append({
                "tabla": table,
                "descripcion": description,
                "claves": self._table_pk(table),
                "columnas": columns,
            })
        result.append({
            "tabla": "CLIPRO",
            "descripcion": "Descuentos cliente/proveedor/familia definidos en fuentes Delphi, no presente en esta base",
            "claves": [],
            "columnas": [],
            "disponible": False,
        })
        return result

    def cliente_relacion_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        table = self._client_table(args.get("tabla"))
        columns = self._table_columns(table)
        limit = _positive_limit(args.get("limite"))
        where: list[str] = []
        params: list[Any] = []
        filters = self._normalize_column_map(args.get("filtros"), columns)
        empresa = args.get("empresa")
        cliente = args.get("cliente")
        subcliente = args.get("subcliente")
        for column in columns:
            if column.endswith("_NUMEMP") and empresa is not None:
                filters.setdefault(column, int(empresa))
            elif column.endswith("_CODCLI") and cliente is not None:
                filters.setdefault(column, int(cliente))
            elif column.endswith("_SUBCLI") and subcliente is not None:
                filters.setdefault(column, int(subcliente))
        for column, value in filters.items():
            where.append(f"{column} = ?")
            params.append(value)
        text = str(args.get("texto") or "").strip()
        if text:
            text_columns = self.db.text_columns(table)
            if not text_columns:
                return []
            where.append("(" + " OR ".join(f"{column} CONTAINING ?" for column in text_columns) + ")")
            params.extend([text] * len(text_columns))
        sql = f"SELECT FIRST {limit} " + ", ".join(columns) + f" FROM {table}"
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY " + ", ".join(self._table_pk(table))
        return self.db.query(sql, tuple(params), limit)

    def cliente_relacion_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        table = self._client_table(args.get("tabla"))
        columns = self._table_columns(table)
        pk = self._table_pk(table)
        keys = self._normalize_column_map(args.get("claves"), columns, required=True)
        missing = [column for column in pk if column not in keys]
        if missing:
            raise KofedasError("Faltan claves: " + ", ".join(missing))
        where = " AND ".join(f"{column}=?" for column in pk)
        row = self.db.one(
            "SELECT " + ", ".join(columns) + f" FROM {table} WHERE {where}",
            tuple(keys[column] for column in pk),
        )
        if row is None:
            raise KofedasError("Registro de cliente no encontrado")
        return row

    def cliente_relacion_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        table = self._client_table(args.get("tabla"))
        columns = self._table_columns(table)
        pk = self._table_pk(table)
        data = self._normalize_column_map(args.get("datos"), columns, required=True)
        keys = self._normalize_column_map(args.get("claves"), columns)
        data.update(keys)
        missing = [column for column in pk if column not in data]
        if missing:
            raise KofedasError("Faltan claves: " + ", ".join(missing))
        return self._upsert(table, columns, pk, data)

    def cliente_completo(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        cliente = int(args["cliente"])
        subcliente = int(args.get("subcliente") or 0)
        limit = _positive_limit(args.get("limite_detalle"), default=500)
        main = self.cliente_obtener({"empresa": empresa, "cliente": cliente, "subcliente": subcliente})
        details: dict[str, Any] = {}
        for table in CLIENT_TABLES:
            if table == "CLIEN":
                continue
            table_args: dict[str, Any] = {"tabla": table, "empresa": empresa, "cliente": cliente, "limite": limit}
            if "SUBCLI" in " ".join(self._table_columns(table)):
                table_args["subcliente"] = subcliente
            details[table.lower()] = self.cliente_relacion_listar(table_args)
        return {"cliente": main["cliente"], "informacion_adicional": main["informacion_adicional"], "relaciones": details}

    def cliente_alta_preparar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        raw_data = self._cliente_field_data(args.get("datos"))
        requested_cliente = args.get("cliente", raw_data.get("CLI_CODCLI"))
        requested_subcliente = args.get("subcliente", raw_data.get("CLI_SUBCLI"))
        cliente, subcliente = self._next_cliente_code(
            empresa,
            centro,
            int(requested_cliente or 0),
            int(requested_subcliente or 0),
        )
        data = self._cliente_defaults(empresa, cliente, subcliente, centro, raw_data)
        return {
            "cliente": cliente,
            "subcliente": subcliente,
            "datos": data,
            "adicionales_reconocidos": CLIENT_ADDITIONAL_CODES,
            "fuente_delphi": {
                "numeracion": "CLIEN_UDM.NUMERAR_CLIEN / NUMERAR_SUBCLIEN",
                "alta": "MNTCLI_U.VALORES_INICIALES + CLIEN_UDM.GRABAR_CLIEN('G')",
            },
        }

    def cliente_alta(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        plan = self.cliente_alta_preparar(args)
        data = plan["datos"]
        if not str(data.get("CLI_NOMCLI") or "").strip():
            raise KofedasError("nombre o CLI_NOMCLI es obligatorio")
        if not str(data.get("CLI_RAZSOC") or "").strip():
            raise KofedasError("razon_social o CLI_RAZSOC es obligatorio")
        empresa = int(data["CLI_NUMEMP"])
        cliente = int(data["CLI_CODCLI"])
        subcliente = int(data["CLI_SUBCLI"])
        explicit_code = bool(args.get("cliente") or self._cliente_field_data(args.get("datos")).get("CLI_CODCLI"))
        explicit_sub = bool(args.get("subcliente") or self._cliente_field_data(args.get("datos")).get("CLI_SUBCLI"))
        if self._cliente_exists(empresa, cliente, subcliente):
            if explicit_code or explicit_sub:
                raise KofedasError("Ya existe cliente/subcliente indicado")
            while self._cliente_exists(empresa, cliente, subcliente):
                if cliente == 99999:
                    subcliente += 1
                else:
                    cliente += 1
                data["CLI_CODCLI"] = cliente
                data["CLI_SUBCLI"] = subcliente
        columns = self._table_columns("CLIEN")
        statements: list[tuple[str, tuple[Any, ...]]] = [
            (
                "INSERT INTO CLIEN (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
                tuple(data.get(column) for column in columns),
            )
        ]
        additional = args.get("informacion_adicional") or {}
        if not isinstance(additional, dict):
            raise KofedasError("informacion_adicional debe ser un objeto CODINF -> valor")
        line = 1
        for code, value in additional.items():
            code_text = str(code).strip().upper()
            if not code_text or value in (None, ""):
                continue
            statements.append((
                "INSERT INTO CLIENI (CLII_NUMEMP, CLII_CODCLI, CLII_SUBCLI, CLII_NUMLIN, CLII_CODINF, CLII_DESCRI) VALUES (?, ?, ?, ?, ?, ?)",
                (empresa, cliente, subcliente, line, code_text, str(value)),
            ))
            line += 1
        inherited = 0
        if args.get("heredar_actividades") and subcliente != 0 and cliente != 99999:
            rows = self.db.query(
                "SELECT CLIC_SECCIO, CLIC_ACTIVI FROM CLIACT WHERE CLIC_NUMEMP = ? AND CLIC_CODCLI = ? AND CLIC_SUBCLI = 0",
                (empresa, cliente),
            )
            for row in rows:
                statements.append((
                    "INSERT INTO CLIACT (CLIC_NUMEMP, CLIC_CODCLI, CLIC_SUBCLI, CLIC_SECCIO, CLIC_ACTIVI) VALUES (?, ?, ?, ?, ?)",
                    (empresa, cliente, subcliente, row["clic_seccio"], row["clic_activi"]),
                ))
                inherited += 1
        if args.get("simular"):
            return {
                "simulado": True,
                "cliente": cliente,
                "subcliente": subcliente,
                "sentencias": len(statements),
                "datos": data,
                "adicionales": line - 1,
                "actividades_heredadas": inherited,
            }
        counts = self.db.execute_transaction(statements)
        return {
            "simulado": False,
            "cliente": cliente,
            "subcliente": subcliente,
            "filas_afectadas": counts,
            "adicionales": line - 1,
            "actividades_heredadas": inherited,
        }

    def proveedor_buscar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        text = str(args.get("texto") or "").strip()
        if not text:
            raise KofedasError("texto es obligatorio")
        limit = _positive_limit(args.get("limite"))
        baja_filter = "" if args.get("incluir_baja") else "AND (PRO_FEBAJA IS NULL)"
        numeric = int(text) if text.isdigit() else None
        params: list[Any] = [empresa, numeric if numeric is not None else -1, _like(text), _like(text), _like(text), _like(text), _like(text)]
        return self.db.query(
            f"""
            SELECT PRO_NUMEMP, PRO_CODPRO, PRO_NOMFIS, PRO_NOMCOR, PRO_NOMABR,
                   PRO_DOMICI, PRO_CODPOS, PRO_POBLAC, PRO_CIF, PRO_TELEFO,
                   PRO_EMAIL, PRO_CODMON, PRO_PORTES, PRO_CODPAG, PRO_FEALTA, PRO_FEBAJA
            FROM PROVEE
            WHERE PRO_NUMEMP = ?
              AND (
                PRO_CODPRO = ?
                OR UPPER(PRO_NOMFIS) LIKE ?
                OR UPPER(PRO_NOMCOR) LIKE ?
                OR UPPER(PRO_NOMABR) LIKE ?
                OR UPPER(PRO_CIF) LIKE ?
                OR UPPER(PRO_EMAIL) LIKE ?
              )
              {baja_filter}
            ORDER BY PRO_NOMCOR, PRO_CODPRO
            """,
            tuple(params),
            limit,
        )

    def proveedor_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        proveedor = int(args["proveedor"])
        ficha = self.db.one(
            "SELECT * FROM PROVEE WHERE PRO_NUMEMP = ? AND PRO_CODPRO = ?",
            (empresa, proveedor),
        )
        if ficha is None:
            raise KofedasError("Proveedor no encontrado")
        return {
            "proveedor": ficha,
            "informacion_adicional": self.db.query(
                """
                SELECT PROI_NUMLIN, PROI_CODINF, PROI_TEXTO
                FROM PROVEEI
                WHERE PROI_NUMEMP = ? AND PROI_CODPRO = ?
                ORDER BY PROI_NUMLIN
                """,
                (empresa, proveedor),
            ),
        }

    def proveedor_articulos_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        proveedor = int(args["proveedor"])
        limit = _positive_limit(args.get("limite"))
        where = ["P.ARTP_NUMEMP = ?", "P.ARTP_CODPRO = ?"]
        params: list[Any] = [empresa, proveedor]
        text = str(args.get("texto") or "").strip()
        if text:
            where.append("(P.ARTP_CODART = ? OR UPPER(P.ARTP_CODART) LIKE ? OR UPPER(P.ARTP_REFPRO) LIKE ? OR UPPER(P.ARTP_DESCRI) LIKE ? OR UPPER(A.ART_DESCRI) LIKE ?)")
            params.extend([text, _like(text), _like(text), _like(text), _like(text)])
        return self.db.query(
            f"""
            SELECT P.ARTP_CODART, A.ART_DESCRI AS ART_DESCRI_MAESTRA, P.ARTP_REFPRO,
                   P.ARTP_DESCRI, P.ARTP_UNIMED, P.ARTP_CANCON, P.ARTP_CANVEN,
                   P.ARTP_UNIPAQ, P.ARTP_PREBAS, P.ARTP_DTOAUM1, P.ARTP_DTOAUM2,
                   P.ARTP_DTOAUM3, P.ARTP_DTOAUM4, P.ARTP_DTOAUM5, P.ARTP_DTOAUM6,
                   P.ARTP_CODMON, P.ARTP_CANPRE, P.ARTP_UBICA, P.ARTP_AMPUNIV,
                   P.ARTP_AJUSTE
            FROM ARTICULP P
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = P.ARTP_NUMEMP AND A.ART_CODART = P.ARTP_CODART
            WHERE {' AND '.join(where)}
            ORDER BY P.ARTP_CODART
            """,
            tuple(params),
            limit,
        )

    def proveedor_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(PROVIDER_TABLES.items()):
            columns = self._table_columns(table)
            result.append({
                "tabla": table,
                "descripcion": description,
                "claves": self._table_pk(table),
                "columnas": columns,
                "solo_lectura": table not in PROVIDER_WRITABLE_TABLES,
            })
        return result

    def proveedor_relacion_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        table = self._provider_table(args.get("tabla"))
        columns = self._table_columns(table)
        limit = _positive_limit(args.get("limite"))
        where: list[str] = []
        params: list[Any] = []
        filters = self._normalize_column_map(args.get("filtros"), columns)
        empresa = args.get("empresa")
        proveedor = args.get("proveedor")
        for column in columns:
            if column.endswith("_NUMEMP") and empresa is not None:
                filters.setdefault(column, int(empresa))
            elif column.endswith("_CODPRO") and proveedor is not None:
                filters.setdefault(column, int(proveedor))
        for column, value in filters.items():
            where.append(f"{column} = ?")
            params.append(value)
        text = str(args.get("texto") or "").strip()
        if text:
            text_columns = self.db.text_columns(table)
            if not text_columns:
                return []
            where.append("(" + " OR ".join(f"{column} CONTAINING ?" for column in text_columns) + ")")
            params.extend([text] * len(text_columns))
        sql = f"SELECT FIRST {limit} " + ", ".join(columns) + f" FROM {table}"
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY " + ", ".join(self._table_pk(table))
        return self.db.query(sql, tuple(params), limit)

    def proveedor_relacion_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        table = self._provider_table(args.get("tabla"))
        columns = self._table_columns(table)
        pk = self._table_pk(table)
        keys = self._normalize_column_map(args.get("claves"), columns, required=True)
        missing = [column for column in pk if column not in keys]
        if missing:
            raise KofedasError("Faltan claves: " + ", ".join(missing))
        where = " AND ".join(f"{column}=?" for column in pk)
        row = self.db.one(
            "SELECT " + ", ".join(columns) + f" FROM {table} WHERE {where}",
            tuple(keys[column] for column in pk),
        )
        if row is None:
            raise KofedasError("Registro de proveedor no encontrado")
        return row

    def proveedor_relacion_guardar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        table = self._provider_table(args.get("tabla"))
        if table not in PROVIDER_WRITABLE_TABLES:
            raise KofedasError("Tabla de proveedor de solo lectura: " + table)
        columns = self._table_columns(table)
        pk = self._table_pk(table)
        data = self._normalize_column_map(args.get("datos"), columns, required=True)
        keys = self._normalize_column_map(args.get("claves"), columns)
        data.update(keys)
        missing = [column for column in pk if column not in data]
        if missing:
            raise KofedasError("Faltan claves: " + ", ".join(missing))
        return self._upsert(table, columns, pk, data)

    def proveedor_completo(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        proveedor = int(args["proveedor"])
        limit = _positive_limit(args.get("limite_detalle"), default=500)
        main = self.proveedor_obtener({"empresa": empresa, "proveedor": proveedor})
        details: dict[str, Any] = {}
        for table in PROVIDER_TABLES:
            if table == "PROVEE":
                continue
            details[table.lower()] = self.proveedor_relacion_listar({
                "tabla": table,
                "empresa": empresa,
                "proveedor": proveedor,
                "limite": limit,
            })
        return {"proveedor": main["proveedor"], "informacion_adicional": main["informacion_adicional"], "relaciones": details}

    def proveedor_alta_preparar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        raw_data = self._proveedor_field_data(args.get("datos"))
        requested_proveedor = args.get("proveedor", raw_data.get("PRO_CODPRO"))
        proveedor = self._next_proveedor_code(empresa, int(requested_proveedor or 0))
        data = self._proveedor_defaults(empresa, proveedor, centro, raw_data)
        return {
            "proveedor": proveedor,
            "datos": data,
            "adicionales_reconocidos": PROVIDER_ADDITIONAL_CODES,
            "fuente_delphi": {
                "numeracion": "PROVEE_UDM.NUMERAR_PROVEE",
                "alta": "MNTPRO_U.VALORES_INICIALES + PROVEE_UDM.GRABAR_PROVEE('G')",
            },
        }

    def proveedor_alta(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        plan = self.proveedor_alta_preparar(args)
        data = plan["datos"]
        if not str(data.get("PRO_NOMCOR") or data.get("PRO_NOMFIS") or "").strip():
            raise KofedasError("nombre, PRO_NOMCOR o PRO_NOMFIS es obligatorio")
        empresa = int(data["PRO_NUMEMP"])
        proveedor = int(data["PRO_CODPRO"])
        explicit_code = bool(args.get("proveedor") or self._proveedor_field_data(args.get("datos")).get("PRO_CODPRO"))
        if self._proveedor_exists(empresa, proveedor):
            if explicit_code:
                raise KofedasError("Ya existe proveedor indicado")
            while self._proveedor_exists(empresa, proveedor):
                proveedor += 1
                data["PRO_CODPRO"] = proveedor
                data["PRO_CUECON"] = "400" + str(proveedor).zfill(7)
        columns = self._table_columns("PROVEE")
        statements: list[tuple[str, tuple[Any, ...]]] = [
            (
                "INSERT INTO PROVEE (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
                tuple(data.get(column) for column in columns),
            )
        ]
        additional = args.get("informacion_adicional") or {}
        if not isinstance(additional, dict):
            raise KofedasError("informacion_adicional debe ser un objeto CODINF -> valor")
        line = 1
        for code, value in additional.items():
            code_text = str(code).strip().upper()
            if not code_text or value in (None, ""):
                continue
            statements.append((
                "INSERT INTO PROVEEI (PROI_NUMEMP, PROI_CODPRO, PROI_NUMLIN, PROI_CODINF, PROI_TEXTO) VALUES (?, ?, ?, ?, ?)",
                (empresa, proveedor, line, code_text, str(value)),
            ))
            line += 1
        if args.get("simular"):
            return {
                "simulado": True,
                "proveedor": proveedor,
                "sentencias": len(statements),
                "datos": data,
                "adicionales": line - 1,
            }
        counts = self.db.execute_transaction(statements)
        return {
            "simulado": False,
            "proveedor": proveedor,
            "filas_afectadas": counts,
            "adicionales": line - 1,
        }

    def oferta_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(OFFER_TABLES.items()):
            columns = self._table_columns(table)
            result.append({
                "tabla": table,
                "descripcion": description,
                "claves": self._table_pk(table),
                "columnas": columns,
            })
        return result

    def oferta_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["O.OFE_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("ejercicio") is not None:
            where.append("O.OFE_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        if args.get("proveedor") is not None:
            where.append("O.OFE_CODPRO = ?")
            params.append(int(args["proveedor"]))
        if args.get("desde"):
            where.append("O.OFE_FECFIN >= ?")
            params.append(self._date_arg(args.get("desde")))
        if args.get("hasta"):
            where.append("O.OFE_FECINI <= ?")
            params.append(self._date_arg(args.get("hasta")))
        if args.get("solo_vigentes"):
            current = self._date_arg(args.get("fecha_consulta"))
            where.append("O.OFE_FECINI <= ? AND O.OFE_FECFIN >= ?")
            params.extend([current, current])
        text = str(args.get("texto") or "").strip()
        if text:
            where.append("O.OFE_DESCRI CONTAINING ?")
            params.append(text)
        articulo = str(args.get("articulo") or "").strip()
        join = ""
        if articulo:
            join = "JOIN DETOFER D ON D.DOF_NUMEMP = O.OFE_NUMEMP AND D.DOF_EJERCI = O.OFE_EJERCI AND D.DOF_NUMOFE = O.OFE_NUMOFE"
            where.append("D.DOF_CODART = ?")
            params.append(articulo)
        return self.db.query(
            f"""
            SELECT FIRST {limit} DISTINCT O.OFE_NUMEMP, O.OFE_EJERCI, O.OFE_NUMOFE,
                   O.OFE_CODPRO, P.PRO_NOMCOR, O.OFE_DESCRI, O.OFE_FECHA,
                   O.OFE_FECINI, O.OFE_FECFIN, O.OFE_GASTOS, O.OFE_CODMON,
                   O.OFE_TIPOFE, O.OFE_FECMOD, O.OFE_USUMOD
            FROM OFERTAS O
            LEFT JOIN PROVEE P ON P.PRO_NUMEMP = O.OFE_NUMEMP AND P.PRO_CODPRO = O.OFE_CODPRO
            {join}
            WHERE {' AND '.join(where)}
            ORDER BY O.OFE_EJERCI DESC, O.OFE_NUMOFE DESC
            """,
            tuple(params),
            limit,
        )

    def oferta_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        ejercicio = int(args["ejercicio"])
        oferta = int(args["oferta"])
        header = self.db.one(
            """
            SELECT O.*, P.PRO_NOMCOR
            FROM OFERTAS O
            LEFT JOIN PROVEE P ON P.PRO_NUMEMP = O.OFE_NUMEMP AND P.PRO_CODPRO = O.OFE_CODPRO
            WHERE O.OFE_NUMEMP = ? AND O.OFE_EJERCI = ? AND O.OFE_NUMOFE = ?
            """,
            (empresa, ejercicio, oferta),
        )
        if header is None:
            raise KofedasError("Oferta no encontrada")
        lines = self.oferta_articulos_listar({
            "empresa": empresa,
            "ejercicio": ejercicio,
            "oferta": oferta,
            "limite": MAX_ROWS_LIMIT,
        })
        return {"oferta": header, "articulos": lines}

    def oferta_articulos_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["D.DOF_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("ejercicio") is not None:
            where.append("D.DOF_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        if args.get("oferta") is not None:
            where.append("D.DOF_NUMOFE = ?")
            params.append(int(args["oferta"]))
        if args.get("proveedor") is not None:
            where.append("D.DOF_CODPRO = ?")
            params.append(int(args["proveedor"]))
        articulo = str(args.get("articulo") or "").strip()
        if articulo:
            where.append("D.DOF_CODART = ?")
            params.append(articulo)
        if args.get("solo_vigentes"):
            current = self._date_arg(args.get("fecha_consulta"))
            where.append("D.DOF_FECINI <= ? AND D.DOF_FECFIN >= ?")
            params.extend([current, current])
        text = str(args.get("texto") or "").strip()
        if text:
            where.append("(D.DOF_CODART CONTAINING ? OR A.ART_DESCRI CONTAINING ? OR O.OFE_DESCRI CONTAINING ?)")
            params.extend([text, text, text])
        return self.db.query(
            f"""
            SELECT FIRST {limit}
                   D.DOF_NUMEMP, D.DOF_EJERCI, D.DOF_NUMOFE, O.OFE_DESCRI,
                   O.OFE_TIPOFE, D.DOF_CODART, A.ART_DESCRI, D.DOF_FECINI,
                   D.DOF_FECFIN, D.DOF_CODPRO, P.PRO_NOMCOR, D.DOF_PRECOS,
                   D.DOF_PRECIO, D.DOF_PVP, D.DOF_DTO1, D.DOF_DTO2,
                   D.DOF_PVP * (1 - D.DOF_DTO1 / 100) * (1 - D.DOF_DTO2 / 100) AS DOF_PVP_NETO,
                   D.DOF_CODMON, D.DOF_CANPRE, A.ART_PVP AS ART_PVP_ACTUAL,
                   A.ART_PRECOS AS ART_PRECOS_ACTUAL
            FROM DETOFER D
            JOIN OFERTAS O ON O.OFE_NUMEMP = D.DOF_NUMEMP AND O.OFE_EJERCI = D.DOF_EJERCI AND O.OFE_NUMOFE = D.DOF_NUMOFE
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = D.DOF_NUMEMP AND A.ART_CODART = D.DOF_CODART
            LEFT JOIN PROVEE P ON P.PRO_NUMEMP = D.DOF_NUMEMP AND P.PRO_CODPRO = D.DOF_CODPRO
            WHERE {' AND '.join(where)}
            ORDER BY D.DOF_FECINI DESC, D.DOF_EJERCI DESC, D.DOF_NUMOFE DESC, D.DOF_CODART
            """,
            tuple(params),
            limit,
        )

    def oferta_alta_preparar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        ejercicio = int(args.get("ejercicio") or date.today().year)
        oferta = self._next_offer_number(empresa, ejercicio, args.get("oferta"))
        proveedor = int(args["proveedor"])
        description = str(args.get("descripcion") or "").strip()
        if not description:
            raise KofedasError("descripcion es obligatoria")
        offer_date = self._date_arg(args.get("fecha"))
        start_date = self._date_arg(args.get("fecha_inicio"), offer_date)
        end_date = self._date_arg(args.get("fecha_fin"), start_date)
        if end_date < start_date:
            raise KofedasError("fecha_fin no puede ser anterior a fecha_inicio")
        articles = args.get("articulos") or []
        if not isinstance(articles, list) or not articles:
            raise KofedasError("articulos debe ser una lista no vacia")
        currency = str(args.get("moneda") or "E").strip()[:1] or "E"
        header = {
            "OFE_NUMEMP": empresa,
            "OFE_EJERCI": ejercicio,
            "OFE_NUMOFE": oferta,
            "OFE_CODPRO": proveedor,
            "OFE_DESCRI": description,
            "OFE_FECHA": offer_date,
            "OFE_FECINI": start_date,
            "OFE_FECFIN": end_date,
            "OFE_GASTOS": self._to_float(args.get("gastos"), 0),
            "OFE_CODMON": currency,
            "OFE_TIPOFE": str(args.get("tipo") or "T").strip()[:1] or "T",
            "OFE_FECMOD": datetime.now().replace(microsecond=0).isoformat(sep=" "),
            "OFE_USUMOD": f"{centro} MCP",
        }
        seen: set[str] = set()
        lines: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        for index, item in enumerate(articles, start=1):
            try:
                line = self._offer_line_defaults(empresa, ejercicio, oferta, proveedor, start_date, end_date, currency, item)
                code = str(line["DOF_CODART"])
                if code in seen:
                    raise KofedasError("Articulo duplicado en la oferta: " + code)
                seen.add(code)
                lines.append(line)
            except Exception as exc:
                errors.append({"linea": index, "error": str(exc)})
        return {
            "oferta": oferta,
            "ejercicio": ejercicio,
            "cabecera": header,
            "lineas": lines,
            "errores": errors,
            "fuente_delphi": {
                "numeracion": "OFERTAS_UDM.NUMERAR_OFERTAS",
                "cabecera": "MNTOFE_U.VALORES_INICIALES + OFERTAS_UDM.GRABAR_OFERTAS('G')",
                "detalle": "MNTOFE_U.VALORES_INICIALES_DETALLE + OFERTAS_UDM.INSERTAR_DETOFER",
            },
        }

    def oferta_alta(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        plan = self.oferta_alta_preparar(args)
        if plan["errores"]:
            if args.get("simular"):
                return {"simulado": True, **plan}
            raise KofedasError("La oferta contiene errores; usa simular=true para revisar el detalle")
        header = plan["cabecera"]
        empresa = int(header["OFE_NUMEMP"])
        ejercicio = int(header["OFE_EJERCI"])
        oferta = int(header["OFE_NUMOFE"])
        explicit_offer = bool(args.get("oferta"))
        if self._offer_exists(empresa, ejercicio, oferta):
            if explicit_offer:
                raise KofedasError("Ya existe la oferta indicada")
            while self._offer_exists(empresa, ejercicio, oferta):
                oferta += 1
            header["OFE_NUMOFE"] = oferta
            for line in plan["lineas"]:
                line["DOF_NUMOFE"] = oferta
            plan["oferta"] = oferta
        columns = self._table_columns("OFERTAS")
        statements: list[tuple[str, tuple[Any, ...]]] = [(
            "INSERT INTO OFERTAS (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
            tuple(header.get(column) for column in columns),
        )]
        dcols = self._table_columns("DETOFER")
        for line in plan["lineas"]:
            data = {column: line.get(column) for column in dcols}
            statements.append((
                "INSERT INTO DETOFER (" + ", ".join(dcols) + ") VALUES (" + ", ".join("?" for _ in dcols) + ")",
                tuple(data.get(column) for column in dcols),
            ))
        if args.get("simular"):
            return {
                "simulado": True,
                "ejercicio": ejercicio,
                "oferta": oferta,
                "lineas": len(plan["lineas"]),
                "sentencias": len(statements),
                "plan": plan,
            }
        counts = self.db.execute_transaction(statements)
        return {
            "simulado": False,
            "ejercicio": ejercicio,
            "oferta": oferta,
            "lineas": len(plan["lineas"]),
            "filas_afectadas": counts,
        }

    def orden_compra_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(PURCHASE_ORDER_TABLES.items()):
            result.append({
                "tabla": table,
                "descripcion": description,
                "claves": self._table_pk(table),
                "columnas": self._table_columns(table),
            })
        return result

    def orden_compra_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["C.COC_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("C.COC_CENTRO = ?")
            params.append(int(args["centro"]))
        if args.get("ejercicio") is not None:
            where.append("C.COC_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        serie = str(args.get("serie") or "").strip()
        if serie or args.get("serie") == "":
            where.append("C.COC_SERIE = ?")
            params.append(serie[:2])
        if args.get("numero") is not None:
            where.append("C.COC_NUMDOC = ?")
            params.append(int(args["numero"]))
        status = self._purchase_order_status(args.get("estado"))
        if status:
            where.append("C.COC_SITUAC = ?")
            params.append(status)
        if args.get("proveedor") is not None:
            where.append("C.COC_CODPRO = ?")
            params.append(int(args["proveedor"]))
        if args.get("desde"):
            where.append("C.COC_FECHA >= ?")
            params.append(self._date_arg(args.get("desde")))
        if args.get("hasta"):
            where.append("C.COC_FECHA <= ?")
            params.append(self._date_arg(args.get("hasta")))
        text = str(args.get("texto") or "").strip()
        if text:
            where.append(
                """
                (C.COC_NOMPRO CONTAINING ? OR C.COC_OBSERV CONTAINING ?
                 OR EXISTS (
                    SELECT 1 FROM DETORC D
                    WHERE D.DOC_NUMEMP = C.COC_NUMEMP
                      AND D.DOC_CENTRO = C.COC_CENTRO
                      AND D.DOC_EJERCI = C.COC_EJERCI
                      AND D.DOC_SERIE = C.COC_SERIE
                      AND D.DOC_NUMDOC = C.COC_NUMDOC
                      AND (D.DOC_CODART CONTAINING ? OR D.DOC_CODARTP CONTAINING ? OR D.DOC_DESCRI CONTAINING ?)
                 ))
                """
            )
            params.extend([text, text, text, text, text])
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   C.COC_NUMEMP, C.COC_CENTRO, C.COC_EJERCI, C.COC_SERIE,
                   C.COC_NUMDOC, C.COC_FECHA, C.COC_CODPRO, C.COC_NOMPRO,
                   P.PRO_NOMCOR, C.COC_FECENV, C.COC_FECENT, C.COC_CODPAG,
                   C.COC_IMPPED, C.COC_IMPPEN, C.COC_CODMON, C.COC_SITUAC,
                   C.COC_INDEDI, C.COC_OBSERV, C.COC_FECMOD, C.COC_USUMOD
            FROM CABORC C
            LEFT JOIN PROVEE P ON P.PRO_NUMEMP = C.COC_NUMEMP AND P.PRO_CODPRO = C.COC_CODPRO
            WHERE {' AND '.join(where)}
            ORDER BY C.COC_EJERCI DESC, C.COC_FECHA DESC, C.COC_CENTRO, C.COC_SERIE, C.COC_NUMDOC DESC
            """,
            tuple(params),
            limit,
        )
        for row in rows:
            status_code = str(row.get("coc_situac") or "").strip().upper()
            row["estado"] = PURCHASE_ORDER_STATUS_LABELS.get(status_code, status_code)
        return rows

    def orden_compra_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        ejercicio = int(args["ejercicio"])
        serie = str(args.get("serie") or "").strip()[:2]
        numero = int(args["numero"])
        header = self.db.one(
            """
            SELECT C.*, P.PRO_NOMCOR
            FROM CABORC C
            LEFT JOIN PROVEE P ON P.PRO_NUMEMP = C.COC_NUMEMP AND P.PRO_CODPRO = C.COC_CODPRO
            WHERE C.COC_NUMEMP = ? AND C.COC_CENTRO = ? AND C.COC_EJERCI = ? AND C.COC_SERIE = ? AND C.COC_NUMDOC = ?
            """,
            (empresa, centro, ejercicio, serie, numero),
        )
        if header is None:
            raise KofedasError("Orden de compra no encontrada")
        status_code = str(header.get("coc_situac") or "").strip().upper()
        header["estado"] = PURCHASE_ORDER_STATUS_LABELS.get(status_code, status_code)
        lines = self.orden_compra_lineas_listar({
            "empresa": empresa,
            "centro": centro,
            "ejercicio": ejercicio,
            "serie": serie,
            "numero": numero,
            "limite": MAX_ROWS_LIMIT,
        })
        return {"orden_compra": header, "lineas": lines}

    def orden_compra_lineas_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["D.DOC_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("D.DOC_CENTRO = ?")
            params.append(int(args["centro"]))
        if args.get("ejercicio") is not None:
            where.append("D.DOC_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        serie = str(args.get("serie") or "").strip()
        if serie or args.get("serie") == "":
            where.append("D.DOC_SERIE = ?")
            params.append(serie[:2])
        if args.get("numero") is not None:
            where.append("D.DOC_NUMDOC = ?")
            params.append(int(args["numero"]))
        status = self._purchase_order_status(args.get("estado"))
        if status:
            where.append("D.DOC_SITUAC = ?")
            params.append(status)
        if args.get("proveedor") is not None:
            where.append("C.COC_CODPRO = ?")
            params.append(int(args["proveedor"]))
        articulo = str(args.get("articulo") or "").strip()
        if articulo:
            where.append("D.DOC_CODART = ?")
            params.append(articulo)
        text = str(args.get("texto") or "").strip()
        if text:
            where.append("(D.DOC_CODART CONTAINING ? OR D.DOC_CODARTP CONTAINING ? OR D.DOC_DESCRI CONTAINING ? OR C.COC_NOMPRO CONTAINING ?)")
            params.extend([text, text, text, text])
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   D.DOC_NUMEMP, D.DOC_CENTRO, D.DOC_EJERCI, D.DOC_SERIE,
                   D.DOC_NUMDOC, D.DOC_NUMLIN, C.COC_FECHA, C.COC_CODPRO,
                   C.COC_NOMPRO, C.COC_SITUAC AS COC_ESTADO, D.DOC_FECMOV,
                   D.DOC_TIPLIN, D.DOC_CODART, A.ART_DESCRI AS ART_DESCRI_MAESTRA,
                   D.DOC_CODARTP, D.DOC_DESCRI, D.DOC_CANTID, D.DOC_UNIMED,
                   D.DOC_PREBAS, D.DOC_CODMON, D.DOC_DTOAUM1, D.DOC_DTOAUM2,
                   D.DOC_DTOAUM3, D.DOC_DTOAUM4, D.DOC_DTOAUM5, D.DOC_DTOAUM6,
                   D.DOC_PORIVA, D.DOC_PORREQ, D.DOC_VALLIN, D.DOC_CANPEN,
                   D.DOC_VALPEN, D.DOC_SITUAC, D.DOC_CIEMAN, D.DOC_OBSERV
            FROM DETORC D
            JOIN CABORC C ON C.COC_NUMEMP = D.DOC_NUMEMP
                         AND C.COC_CENTRO = D.DOC_CENTRO
                         AND C.COC_EJERCI = D.DOC_EJERCI
                         AND C.COC_SERIE = D.DOC_SERIE
                         AND C.COC_NUMDOC = D.DOC_NUMDOC
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = D.DOC_NUMEMP AND A.ART_CODART = D.DOC_CODART
            WHERE {' AND '.join(where)}
            ORDER BY D.DOC_EJERCI DESC, D.DOC_CENTRO, D.DOC_SERIE, D.DOC_NUMDOC DESC, D.DOC_NUMLIN
            """,
            tuple(params),
            limit,
        )
        for row in rows:
            row_status = str(row.get("doc_situac") or "").strip().upper()
            header_status = str(row.get("coc_estado") or "").strip().upper()
            row["estado_linea"] = PURCHASE_ORDER_STATUS_LABELS.get(row_status, row_status)
            row["estado_cabecera"] = PURCHASE_ORDER_STATUS_LABELS.get(header_status, header_status)
        return rows

    def orden_compra_alta_preparar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        ejercicio = int(args.get("ejercicio") or date.today().year)
        serie = str(args.get("serie") or "").strip()[:2]
        numero = self._next_purchase_order_number(empresa, centro, ejercicio, serie, args.get("numero"))
        proveedor = int(args["proveedor"])
        provider = self.db.one(
            """
            SELECT FIRST 1 PRO_CODPRO, PRO_NOMCOR, PRO_NOMFIS, PRO_CODPAG, PRO_CODMON
            FROM PROVEE
            WHERE PRO_NUMEMP = ? AND PRO_CODPRO = ?
            """,
            (empresa, proveedor),
        )
        if provider is None:
            raise KofedasError("Proveedor no encontrado")
        articles = args.get("articulos") or []
        if not isinstance(articles, list) or not articles:
            raise KofedasError("articulos debe ser una lista no vacia")
        order_date = self._date_arg(args.get("fecha"))
        currency = str(args.get("moneda") or provider.get("pro_codmon") or "E").strip()[:1] or "E"
        lines: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        for index, item in enumerate(articles, start=1):
            try:
                lines.append(self._purchase_order_line_defaults(empresa, centro, ejercicio, serie, numero, index, proveedor, order_date, currency, item))
            except Exception as exc:
                errors.append({"linea": index, "error": str(exc)})
        ordered_total, pending_total = self._purchase_order_totals(lines)
        header = {
            "COC_NUMEMP": empresa,
            "COC_CENTRO": centro,
            "COC_EJERCI": ejercicio,
            "COC_SERIE": serie,
            "COC_NUMDOC": numero,
            "COC_FECHA": order_date,
            "COC_CODPRO": proveedor,
            "COC_NOMPRO": str(provider.get("pro_nomcor") or provider.get("pro_nomfis") or "")[:50],
            "COC_CODREP": self._to_int(args.get("representante"), 0),
            "COC_FECENV": self._date_arg(args.get("fecha_envio"), order_date) if args.get("fecha_envio") else None,
            "COC_FECENT": self._date_arg(args.get("fecha_entrega"), order_date) if args.get("fecha_entrega") else None,
            "COC_CODPAG": self._to_int(args.get("forma_pago"), self._to_int(provider.get("pro_codpag"), 0)),
            "COC_IMPPED": ordered_total,
            "COC_IMPPEN": pending_total,
            "COC_CODMON": currency,
            "COC_SITUAC": "P" if args.get("fecha_envio") else "A",
            "COC_INDEDI": "N",
            "COC_OBSERV": str(args.get("observaciones") or "")[:60],
            "COC_FECMOD": datetime.now().replace(microsecond=0).isoformat(sep=" "),
            "COC_USUMOD": f"{centro} MCP",
        }
        return {
            "numero": numero,
            "ejercicio": ejercicio,
            "serie": serie,
            "cabecera": header,
            "lineas": lines,
            "errores": errors,
            "fuente_delphi": {
                "numeracion": "MNTORC_U.NUMERAR_CABORC / NUMERAR_DETORC",
                "cabecera": "MNTORC_U.VALORES_INICIALES + GRABAR_CABORC('G')",
                "detalle": "MNTORC_U.VALORES_INICIALES_DETALLE + GRABAR_DETORC('G')",
                "valoracion": "MNTORC_U.VALORAR_DETORC / VALORAR_CABORC",
            },
        }

    def orden_compra_alta(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        plan = self.orden_compra_alta_preparar(args)
        if plan["errores"]:
            if args.get("simular"):
                return {"simulado": True, **plan}
            raise KofedasError("La orden de compra contiene errores; usa simular=true para revisar el detalle")
        header = plan["cabecera"]
        empresa = int(header["COC_NUMEMP"])
        centro = int(header["COC_CENTRO"])
        ejercicio = int(header["COC_EJERCI"])
        serie = str(header["COC_SERIE"])
        numero = int(header["COC_NUMDOC"])
        explicit_number = bool(args.get("numero"))
        if self._purchase_order_exists(empresa, centro, ejercicio, serie, numero):
            if explicit_number:
                raise KofedasError("Ya existe la orden de compra indicada")
            while self._purchase_order_exists(empresa, centro, ejercicio, serie, numero):
                numero += 1
            header["COC_NUMDOC"] = numero
            for line in plan["lineas"]:
                line["DOC_NUMDOC"] = numero
            plan["numero"] = numero
        columns = self._table_columns("CABORC")
        statements: list[tuple[str, tuple[Any, ...]]] = [(
            "INSERT INTO CABORC (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
            tuple(header.get(column) for column in columns),
        )]
        dcols = self._table_columns("DETORC")
        for line in plan["lineas"]:
            data = {column: line.get(column) for column in dcols}
            statements.append((
                "INSERT INTO DETORC (" + ", ".join(dcols) + ") VALUES (" + ", ".join("?" for _ in dcols) + ")",
                tuple(data.get(column) for column in dcols),
            ))
        if args.get("simular"):
            return {
                "simulado": True,
                "ejercicio": ejercicio,
                "serie": serie,
                "numero": numero,
                "lineas": len(plan["lineas"]),
                "sentencias": len(statements),
                "plan": plan,
            }
        counts = self.db.execute_transaction(statements)
        return {
            "simulado": False,
            "ejercicio": ejercicio,
            "serie": serie,
            "numero": numero,
            "lineas": len(plan["lineas"]),
            "filas_afectadas": counts,
        }

    def orden_compra_cerrar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        centro = self._centro(args)
        ejercicio = int(args["ejercicio"])
        serie = str(args.get("serie") or "").strip()[:2]
        numero = int(args["numero"])
        header = self.db.one(
            """
            SELECT COC_SITUAC
            FROM CABORC
            WHERE COC_NUMEMP = ? AND COC_CENTRO = ? AND COC_EJERCI = ? AND COC_SERIE = ? AND COC_NUMDOC = ?
            """,
            (empresa, centro, ejercicio, serie, numero),
        )
        if header is None:
            raise KofedasError("Orden de compra no encontrada")
        if str(header.get("coc_situac") or "").strip().upper() == "C":
            return {
                "simulado": bool(args.get("simular")),
                "accion": "sin_cambios",
                "motivo": "La orden de compra ya esta cerrada",
                "ejercicio": ejercicio,
                "serie": serie,
                "numero": numero,
            }
        now = datetime.now().replace(microsecond=0).isoformat(sep=" ")
        user = str(args.get("usuario") or f"{centro} MCP")[:10]
        statements: list[tuple[str, tuple[Any, ...]]] = [
            (
                """
                UPDATE DETORC
                SET DOC_CIEMAN = 'S', DOC_CANPEN = 0, DOC_VALPEN = 0, DOC_SITUAC = 'C'
                WHERE DOC_NUMEMP = ? AND DOC_CENTRO = ? AND DOC_EJERCI = ? AND DOC_SERIE = ? AND DOC_NUMDOC = ?
                """,
                (empresa, centro, ejercicio, serie, numero),
            ),
            (
                """
                UPDATE CABORC
                SET COC_SITUAC = 'C', COC_IMPPEN = 0, COC_FECMOD = ?, COC_USUMOD = ?
                WHERE COC_NUMEMP = ? AND COC_CENTRO = ? AND COC_EJERCI = ? AND COC_SERIE = ? AND COC_NUMDOC = ?
                """,
                (now, user, empresa, centro, ejercicio, serie, numero),
            ),
        ]
        if args.get("simular"):
            return {
                "simulado": True,
                "accion": "cerrar",
                "ejercicio": ejercicio,
                "serie": serie,
                "numero": numero,
                "sentencias": len(statements),
            }
        counts = self.db.execute_transaction(statements)
        return {
            "simulado": False,
            "accion": "cerrado",
            "ejercicio": ejercicio,
            "serie": serie,
            "numero": numero,
            "filas_afectadas": counts,
        }

    def entrada_almacen_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(WAREHOUSE_ENTRY_TABLES.items()):
            result.append({
                "tabla": table,
                "descripcion": description,
                "claves": self._table_pk(table),
                "columnas": self._table_columns(table),
            })
        return result

    def entrada_almacen_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["C.CBM_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("C.CBM_CENTRO = ?")
            params.append(int(args["centro"]))
        if args.get("ejercicio") is not None:
            where.append("C.CBM_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        serie = str(args.get("serie") or "").strip()
        if serie or args.get("serie") == "":
            where.append("C.CBM_SERIE = ?")
            params.append(serie[:2])
        if args.get("numero") is not None:
            where.append("C.CBM_NUMDOC = ?")
            params.append(int(args["numero"]))
        status = self._warehouse_entry_status(args.get("situacion"))
        if status:
            where.append("C.CBM_SITUAC = ?")
            params.append(status)
        if args.get("proveedor") is not None:
            where.append("C.CBM_CODPRO = ?")
            params.append(int(args["proveedor"]))
        if args.get("desde"):
            where.append("C.CBM_FECHA >= ?")
            params.append(self._date_arg(args.get("desde")))
        if args.get("hasta"):
            where.append("C.CBM_FECHA <= ?")
            params.append(self._date_arg(args.get("hasta")))
        if str(args.get("albaran") or "").strip():
            where.append("C.CBM_ALBPRO = ?")
            params.append(str(args.get("albaran")).strip())
        if str(args.get("factura") or "").strip():
            where.append("C.CBM_FACPRO = ?")
            params.append(str(args.get("factura")).strip())
        text = str(args.get("texto") or "").strip()
        if text:
            where.append(
                """
                (C.CBM_NOMPRO CONTAINING ? OR C.CBM_OBSERV CONTAINING ?
                 OR EXISTS (
                    SELECT 1 FROM DETMOVM D
                    WHERE D.DMM_NUMEMP = C.CBM_NUMEMP
                      AND D.DMM_CENTRO = C.CBM_CENTRO
                      AND D.DMM_EJERCI = C.CBM_EJERCI
                      AND D.DMM_SERIE = C.CBM_SERIE
                      AND D.DMM_NUMDOC = C.CBM_NUMDOC
                      AND (D.DMM_CODART CONTAINING ? OR D.DMM_CODARP CONTAINING ? OR D.DMM_DESCRI CONTAINING ?)
                 ))
                """
            )
            params.extend([text, text, text, text, text])
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   C.CBM_NUMEMP, C.CBM_CENTRO, C.CBM_EJERCI, C.CBM_SERIE,
                   C.CBM_NUMDOC, C.CBM_FECHA, C.CBM_FECREC, C.CBM_CODPRO,
                   C.CBM_NOMPRO, P.PRO_NOMCOR, C.CBM_CODPAG, C.CBM_ALBPRO,
                   C.CBM_FACPRO, C.CBM_FECFAC, C.CBM_CODMON, C.CBM_TOTALD,
                   C.CBM_TOTALS, C.CBM_SITUAC, C.CBM_OBSERV, C.CBM_FECMOD,
                   C.CBM_USUMOD
            FROM CABDOCM C
            LEFT JOIN PROVEE P ON P.PRO_NUMEMP = C.CBM_NUMEMP AND P.PRO_CODPRO = C.CBM_CODPRO
            WHERE {' AND '.join(where)}
            ORDER BY C.CBM_FECHA DESC, C.CBM_CENTRO, C.CBM_EJERCI DESC, C.CBM_SERIE, C.CBM_NUMDOC DESC
            """,
            tuple(params),
            limit,
        )
        for row in rows:
            status_code = str(row.get("cbm_situac") or "").strip().upper()
            row["situacion_descripcion"] = WAREHOUSE_ENTRY_STATUS_LABELS.get(status_code, status_code)
        return rows

    def entrada_almacen_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        ejercicio = int(args["ejercicio"])
        serie = str(args.get("serie") or "").strip()[:2]
        numero = int(args["numero"])
        header = self.db.one(
            """
            SELECT C.*, P.PRO_NOMCOR
            FROM CABDOCM C
            LEFT JOIN PROVEE P ON P.PRO_NUMEMP = C.CBM_NUMEMP AND P.PRO_CODPRO = C.CBM_CODPRO
            WHERE C.CBM_NUMEMP = ? AND C.CBM_CENTRO = ? AND C.CBM_EJERCI = ? AND C.CBM_SERIE = ? AND C.CBM_NUMDOC = ?
            """,
            (empresa, centro, ejercicio, serie, numero),
        )
        if header is None:
            raise KofedasError("Entrada de almacen no encontrada")
        status_code = str(header.get("cbm_situac") or "").strip().upper()
        header["situacion_descripcion"] = WAREHOUSE_ENTRY_STATUS_LABELS.get(status_code, status_code)
        lines = self.entrada_almacen_lineas_listar({
            "empresa": empresa,
            "centro": centro,
            "ejercicio": ejercicio,
            "serie": serie,
            "numero": numero,
            "limite": MAX_ROWS_LIMIT,
        })
        return {"entrada": header, "lineas": lines}

    def entrada_almacen_lineas_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["D.DMM_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("D.DMM_CENTRO = ?")
            params.append(int(args["centro"]))
        if args.get("ejercicio") is not None:
            where.append("D.DMM_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        serie = str(args.get("serie") or "").strip()
        if serie or args.get("serie") == "":
            where.append("D.DMM_SERIE = ?")
            params.append(serie[:2])
        if args.get("numero") is not None:
            where.append("D.DMM_NUMDOC = ?")
            params.append(int(args["numero"]))
        if args.get("proveedor") is not None:
            where.append("C.CBM_CODPRO = ?")
            params.append(int(args["proveedor"]))
        articulo = str(args.get("articulo") or "").strip()
        if articulo:
            where.append("D.DMM_CODART = ?")
            params.append(articulo)
        text = str(args.get("texto") or "").strip()
        if text:
            where.append("(D.DMM_CODART CONTAINING ? OR D.DMM_CODARP CONTAINING ? OR D.DMM_DESCRI CONTAINING ? OR C.CBM_NOMPRO CONTAINING ?)")
            params.extend([text, text, text, text])
        return self.db.query(
            f"""
            SELECT FIRST {limit}
                   D.DMM_NUMEMP, D.DMM_CENTRO, D.DMM_EJERCI, D.DMM_SERIE, D.DMM_NUMDOC,
                   D.DMM_NUMLIN, C.CBM_FECHA, C.CBM_CODPRO, C.CBM_NOMPRO, C.CBM_SITUAC,
                   D.DMM_FECMOV, D.DMM_TIPLIN, D.DMM_CODART, A.ART_DESCRI AS ART_DESCRI_MAESTRA,
                   D.DMM_CODARP, D.DMM_DESCRI, D.DMM_CANTIDP, D.DMM_UNIMED, D.DMM_CANTID,
                   D.DMM_PREBAS, D.DMM_CODMON, D.DMM_DTOAUM1, D.DMM_DTOAUM2, D.DMM_DTOAUM3,
                   D.DMM_DTOAUM4, D.DMM_DTOAUM5, D.DMM_DTOAUM6, D.DMM_PORIVA, D.DMM_PORREQ,
                   D.DMM_VALLIN, D.DMM_IMPDTO, D.DMM_OBSERV, D.DMM_EJERCIP, D.DMM_SERIEP,
                   D.DMM_NUMDOCP, D.DMM_NUMLINP, D.DMM_EJEOFE, D.DMM_NUMOFE
            FROM DETMOVM D
            JOIN CABDOCM C ON C.CBM_NUMEMP = D.DMM_NUMEMP
                         AND C.CBM_CENTRO = D.DMM_CENTRO
                         AND C.CBM_EJERCI = D.DMM_EJERCI
                         AND C.CBM_SERIE = D.DMM_SERIE
                         AND C.CBM_NUMDOC = D.DMM_NUMDOC
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = D.DMM_NUMEMP AND A.ART_CODART = D.DMM_CODART
            WHERE {' AND '.join(where)}
            ORDER BY D.DMM_EJERCI DESC, D.DMM_CENTRO, D.DMM_SERIE, D.DMM_NUMDOC DESC, D.DMM_NUMLIN
            """,
            tuple(params),
            limit,
        )

    def _entrada_almacen_pendientes(self, args: dict[str, Any], status: str) -> dict[str, Any]:
        query_args = dict(args)
        query_args["situacion"] = status
        items = self.entrada_almacen_listar(query_args)
        total = sum(self._to_float(item.get("cbm_totald"), 0) for item in items)
        base = sum(self._to_float(item.get("cbm_totals"), 0) for item in items)
        providers: dict[int, dict[str, Any]] = {}
        for item in items:
            code = self._to_int(item.get("cbm_codpro"), 0)
            bucket = providers.setdefault(code, {"proveedor": code, "nombre": item.get("cbm_nompro"), "documentos": 0, "base": 0.0, "total": 0.0})
            bucket["documentos"] += 1
            bucket["base"] += self._to_float(item.get("cbm_totals"), 0)
            bucket["total"] += self._to_float(item.get("cbm_totald"), 0)
        return {
            "criterio": f"CBM_SITUAC='{status}'",
            "situacion": WAREHOUSE_ENTRY_STATUS_LABELS.get(status, status),
            "totales": {"documentos": len(items), "base": round(base, 2), "total": round(total, 2)},
            "proveedores": sorted(providers.values(), key=lambda item: item["total"], reverse=True),
            "documentos": items,
        }

    def entrada_almacen_pendientes_facturar(self, args: dict[str, Any]) -> dict[str, Any]:
        return self._entrada_almacen_pendientes(args, "P")

    def entrada_almacen_pendientes_contabilizar(self, args: dict[str, Any]) -> dict[str, Any]:
        return self._entrada_almacen_pendientes(args, "F")

    def entrada_almacen_alta_preparar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        provider = self._warehouse_entry_provider(empresa, args)
        entry_date = self._date_arg(args.get("fecha"))
        ejercicio = int(args.get("ejercicio") or entry_date[:4])
        serie = str(args.get("serie") or self._parameter_value(f"E{centro}", "", empresa) or self._parameter_value("E", "", empresa) or "").strip()[:2]
        if not serie:
            raise KofedasError("No existe parametro de serie para entradas: E/E<centro>; informa serie")
        numero = self._next_warehouse_entry_number(empresa, centro, ejercicio, serie, args.get("numero"))
        lines_arg = args.get("lineas") or []
        if not isinstance(lines_arg, list) or not lines_arg:
            raise KofedasError("lineas debe ser una lista no vacia")
        currency = str(args.get("moneda") or provider.get("pro_codmon") or "E").strip()[:1] or "E"
        lines: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        for index, item in enumerate(lines_arg, start=1):
            if not isinstance(item, dict):
                errors.append({"linea": index, "error": "La linea debe ser un objeto"})
                continue
            try:
                lines.append(self._warehouse_entry_line_defaults(empresa, centro, ejercicio, serie, numero, index * 10, int(provider["pro_codpro"]), entry_date, currency, item))
            except Exception as exc:
                errors.append({"linea": index, "error": str(exc)})
        header = {
            "CBM_NUMEMP": empresa,
            "CBM_CENTRO": centro,
            "CBM_EJERCI": ejercicio,
            "CBM_SERIE": serie,
            "CBM_NUMDOC": numero,
            "CBM_FECHA": entry_date,
            "CBM_FECREC": self._date_arg(args.get("fecha_recepcion"), entry_date),
            "CBM_CODPRO": int(provider["pro_codpro"]),
            "CBM_NOMPRO": str(provider.get("pro_nomfis") or provider.get("pro_nomcor") or "")[:40],
            "CBM_DOMICI": str(provider.get("pro_domici") or "")[:50],
            "CBM_CODPOS": self._to_int(provider.get("pro_codpos"), 0),
            "CBM_POBLAC": str(provider.get("pro_poblac") or "")[:40],
            "CBM_CIF": str(provider.get("pro_cif") or "")[:15],
            "CBM_CODPAG": self._to_int(args.get("forma_pago"), self._to_int(provider.get("pro_codpag"), 0)),
            "CBM_ALBPRO": str(args.get("albaran") or "")[:20],
            "CBM_FACPRO": str(args.get("factura") or "")[:20],
            "CBM_FECFAC": self._date_arg(args.get("fecha_factura"), entry_date) if args.get("fecha_factura") else None,
            "CBM_CODMON": currency,
            "CBM_PORDTO": self._to_float(args.get("descuento"), 0),
            "CBM_IMPPOR": self._to_float(args.get("portes"), 0),
            "CBM_BASIMP1": 0,
            "CBM_PORIVA1": 0,
            "CBM_PORREQ1": 0,
            "CBM_BASIMP2": 0,
            "CBM_PORIVA2": 0,
            "CBM_PORREQ2": 0,
            "CBM_BASIMP3": 0,
            "CBM_PORIVA3": 0,
            "CBM_PORREQ3": 0,
            "CBM_BASIMP4": 0,
            "CBM_PORIVA4": 0,
            "CBM_PORREQ4": 0,
            "CBM_TOTALD": 0,
            "CBM_TOTALS": 0,
            "CBM_INDEDI": "N",
            "CBM_SITUAC": "F" if str(args.get("factura") or "").strip() else "P",
            "CBM_OBSERV": str(args.get("observaciones") or "")[:50],
            "CBM_FECMOD": datetime.now().replace(microsecond=0).isoformat(sep=" "),
            "CBM_USUMOD": f"{centro} MCP",
        }
        header = self._warehouse_entry_totals(header, lines)
        return {
            "numero": numero,
            "ejercicio": ejercicio,
            "serie": serie,
            "cabecera": header,
            "lineas": lines,
            "errores": errors,
            "fuente_delphi": {
                "cabecera": "CABDOCM_UDM.GRABAR_CABDOCM('G') / PEDPRO_U.B_ACEPTARClick",
                "detalle": "CABDOCM_UDM.GRABAR_DETMOVM('G')",
                "valoracion": "CABDOCM_UDM.VALORAR_DETMOVM + FINALIZAR_CABDOCM",
            },
        }

    def entrada_almacen_alta(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        plan = self.entrada_almacen_alta_preparar(args)
        if plan["errores"]:
            if args.get("simular"):
                return {"simulado": True, **plan}
            raise KofedasError("La entrada contiene errores; usa simular=true para revisar el detalle")
        header = plan["cabecera"]
        empresa = int(header["CBM_NUMEMP"])
        centro = int(header["CBM_CENTRO"])
        ejercicio = int(header["CBM_EJERCI"])
        serie = str(header["CBM_SERIE"])
        numero = int(header["CBM_NUMDOC"])
        explicit_number = bool(args.get("numero"))
        if self._warehouse_entry_exists(empresa, centro, ejercicio, serie, numero):
            if explicit_number:
                raise KofedasError("Ya existe la entrada indicada")
            while self._warehouse_entry_exists(empresa, centro, ejercicio, serie, numero):
                numero += 1
            header["CBM_NUMDOC"] = numero
            for line in plan["lineas"]:
                line["DMM_NUMDOC"] = numero
            plan["numero"] = numero
        columns = self._table_columns("CABDOCM")
        statements: list[tuple[str, tuple[Any, ...]]] = [(
            "INSERT INTO CABDOCM (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
            tuple(header.get(column) for column in columns),
        )]
        dcols = self._table_columns("DETMOVM")
        for line in plan["lineas"]:
            data = {column: line.get(column) for column in dcols}
            statements.append((
                "INSERT INTO DETMOVM (" + ", ".join(dcols) + ") VALUES (" + ", ".join("?" for _ in dcols) + ")",
                tuple(data.get(column) for column in dcols),
            ))
            if str(line.get("DMM_TIPLIN") or "").upper() == "D" and str(line.get("art_indinv") or "S").upper() != "N":
                statements.append((
                    """
                    UPDATE ARTICULE
                    SET ARTE_EXIST = ARTE_EXIST + ?, ARTE_FECCOM = ?, ARTE_FECMOV = ?
                    WHERE ARTE_NUMEMP = ? AND ARTE_CODART = ? AND ARTE_CENTRO = ?
                    """,
                    (
                        line.get("DMM_CANTID"),
                        line.get("DMM_FECMOV"),
                        datetime.now().replace(microsecond=0).isoformat(sep=" "),
                        empresa,
                        line.get("DMM_CODART"),
                        centro,
                    ),
                ))
                statements.append((
                    "UPDATE ARTICUL SET ART_OBSOL = 'N' WHERE ART_NUMEMP = ? AND ART_CODART = ?",
                    (empresa, line.get("DMM_CODART")),
                ))
        if args.get("simular"):
            return {
                "simulado": True,
                "ejercicio": ejercicio,
                "serie": serie,
                "numero": numero,
                "lineas": len(plan["lineas"]),
                "sentencias": len(statements),
                "plan": plan,
            }
        counts = self.db.execute_transaction(statements)
        return {
            "simulado": False,
            "ejercicio": ejercicio,
            "serie": serie,
            "numero": numero,
            "lineas": len(plan["lineas"]),
            "filas_afectadas": counts,
            "totales": {"base": header["CBM_TOTALS"], "total": header["CBM_TOTALD"]},
        }

    def entrada_almacen_pdf_previsualizar(self, args: dict[str, Any]) -> dict[str, Any]:
        return self._warehouse_entry_pdf_proposal(args)

    def entrada_almacen_desde_pdf(self, args: dict[str, Any]) -> dict[str, Any]:
        proposal = self._warehouse_entry_pdf_proposal(args)
        header = dict(proposal["cabecera"])
        if isinstance(args.get("cabecera"), dict):
            header.update(args["cabecera"])
        if args.get("proveedor") not in (None, "", 0):
            header["proveedor"] = int(args["proveedor"])
        lines = args.get("lineas") if isinstance(args.get("lineas"), list) and args.get("lineas") else proposal["lineas"]
        unresolved = [line for line in lines if isinstance(line, dict) and line.get("resuelto") is False and not line.get("articulo")]
        if unresolved and not args.get("simular"):
            raise KofedasError("El PDF contiene lineas sin resolver; usa simular=true o informa lineas revisadas")
        payload = {**header, "lineas": lines, "simular": args.get("simular", True)}
        result = self.entrada_almacen_alta(payload)
        if isinstance(result, dict):
            result["pdf"] = proposal["pdf"]
            result["lineas_pdf_detectadas"] = len(proposal["lineas"])
            result["advertencias_pdf"] = proposal["advertencias"]
        return result

    def venta_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(SALES_TABLES.items()):
            try:
                columns = self._table_columns(table)
                keys = self.db.primary_key(table)
                available = True
            except Exception as exc:
                columns = []
                keys = []
                available = False
                description = f"{description} (no disponible: {exc})"
            result.append({"tabla": table, "descripcion": description, "disponible": available, "claves": keys, "columnas": columns})
        return result

    def venta_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["C.CBV_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("C.CBV_CENTRO = ?")
            params.append(int(args["centro"]))
        if args.get("tipo_documento"):
            where.append("C.CBV_TIPDOC = ?")
            params.append(self._sale_document_type(args.get("tipo_documento")))
        if args.get("tipo_accion") not in (None, ""):
            where.append("C.CBV_TIPAC = ?")
            params.append(str(args["tipo_accion"]).strip()[:1])
        if args.get("ejercicio"):
            where.append("C.CBV_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        if args.get("serie") not in (None, ""):
            where.append("C.CBV_SERIE = ?")
            params.append(str(args["serie"]).strip())
        if args.get("numero"):
            where.append("C.CBV_NUMDOC = ?")
            params.append(int(args["numero"]))
        if args.get("cliente"):
            where.append("C.CBV_CODCLI = ?")
            params.append(int(args["cliente"]))
        if args.get("subcliente") not in (None, ""):
            where.append("C.CBV_SUBCLI = ?")
            params.append(int(args["subcliente"]))
        if args.get("estado") not in (None, ""):
            where.append("C.CBV_SITUAC = ?")
            params.append(str(args["estado"]).strip().upper()[:1])
        if args.get("desde"):
            where.append("C.CBV_FECHA >= ?")
            params.append(self._date_arg(args["desde"]))
        if args.get("hasta"):
            where.append("C.CBV_FECHA <= ?")
            params.append(self._date_arg(args["hasta"]))
        if str(args.get("articulo") or "").strip():
            where.append(
                """
                EXISTS (
                    SELECT 1 FROM DETMOV D
                    WHERE D.DMV_NUMEMP = C.CBV_NUMEMP AND D.DMV_CENTRO = C.CBV_CENTRO
                      AND D.DMV_TIPDOC = C.CBV_TIPDOC AND D.DMV_TIPAC = C.CBV_TIPAC
                      AND D.DMV_EJERCI = C.CBV_EJERCI AND D.DMV_SERIE = C.CBV_SERIE
                      AND D.DMV_NUMDOC = C.CBV_NUMDOC AND D.DMV_CODART = ?
                )
                """
            )
            params.append(str(args["articulo"]).strip())
        if str(args.get("texto") or "").strip():
            where.append(
                """
                (
                    UPPER(C.CBV_NOMCLI) LIKE ? OR UPPER(C.CBV_REFCLI) LIKE ? OR UPPER(C.CBV_OBSERV) LIKE ?
                    OR EXISTS (
                        SELECT 1 FROM DETMOV D
                        WHERE D.DMV_NUMEMP = C.CBV_NUMEMP AND D.DMV_CENTRO = C.CBV_CENTRO
                          AND D.DMV_TIPDOC = C.CBV_TIPDOC AND D.DMV_TIPAC = C.CBV_TIPAC
                          AND D.DMV_EJERCI = C.CBV_EJERCI AND D.DMV_SERIE = C.CBV_SERIE
                          AND D.DMV_NUMDOC = C.CBV_NUMDOC
                          AND (UPPER(D.DMV_CODART) LIKE ? OR UPPER(D.DMV_DESCRI) LIKE ?)
                    )
                )
                """
            )
            like = _like(str(args["texto"]))
            params.extend([like, like, like, like, like])
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   C.CBV_NUMEMP, C.CBV_CENTRO, C.CBV_TIPDOC, C.CBV_TIPAC, C.CBV_EJERCI,
                   C.CBV_SERIE, C.CBV_NUMDOC, C.CBV_FECHA, C.CBV_FECHAE, C.CBV_CODCLI,
                   C.CBV_SUBCLI, C.CBV_NOMCLI, C.CBV_CODPAG, C.CBV_FORENV, C.CBV_PORDTO,
                   C.CBV_TOTALS, C.CBV_TOTALD, C.CBV_CODMON, C.CBV_SITUAC, C.CBV_INDEDI,
                   C.CBV_REFCLI, C.CBV_OBSERV, COUNT(D.DMV_NUMLIN) AS LINEAS
            FROM CABDOCV C
            LEFT JOIN DETMOV D ON D.DMV_NUMEMP = C.CBV_NUMEMP
                              AND D.DMV_CENTRO = C.CBV_CENTRO
                              AND D.DMV_TIPDOC = C.CBV_TIPDOC
                              AND D.DMV_TIPAC = C.CBV_TIPAC
                              AND D.DMV_EJERCI = C.CBV_EJERCI
                              AND D.DMV_SERIE = C.CBV_SERIE
                              AND D.DMV_NUMDOC = C.CBV_NUMDOC
            WHERE {' AND '.join(where)}
            GROUP BY C.CBV_NUMEMP, C.CBV_CENTRO, C.CBV_TIPDOC, C.CBV_TIPAC, C.CBV_EJERCI,
                     C.CBV_SERIE, C.CBV_NUMDOC, C.CBV_FECHA, C.CBV_FECHAE, C.CBV_CODCLI,
                     C.CBV_SUBCLI, C.CBV_NOMCLI, C.CBV_CODPAG, C.CBV_FORENV, C.CBV_PORDTO,
                     C.CBV_TOTALS, C.CBV_TOTALD, C.CBV_CODMON, C.CBV_SITUAC, C.CBV_INDEDI,
                     C.CBV_REFCLI, C.CBV_OBSERV
            ORDER BY C.CBV_FECHA DESC, C.CBV_EJERCI DESC, C.CBV_SERIE, C.CBV_NUMDOC DESC
            """,
            tuple(params),
            limit,
        )
        for row in rows:
            tipdoc = str(row.get("cbv_tipdoc") or "").strip().upper()
            situac = str(row.get("cbv_situac") or "").strip().upper()
            row["tipo_documento_descripcion"] = SALES_DOCUMENT_TYPES.get(tipdoc, tipdoc)
            row["situacion_descripcion"] = SALES_STATUS_LABELS.get(situac, situac)
        return rows

    def venta_obtener(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        tipdoc = self._sale_document_type(args.get("tipo_documento"))
        tipac = str(args.get("tipo_accion") or "0").strip()[:1] or "0"
        ejercicio = int(args["ejercicio"])
        serie = str(args["serie"]).strip()
        numero = int(args["numero"])
        key = (empresa, centro, tipdoc, tipac, ejercicio, serie, numero)
        header = self.db.one(
            """
            SELECT *
            FROM CABDOCV
            WHERE CBV_NUMEMP = ? AND CBV_CENTRO = ? AND CBV_TIPDOC = ? AND CBV_TIPAC = ?
              AND CBV_EJERCI = ? AND CBV_SERIE = ? AND CBV_NUMDOC = ?
            """,
            key,
        )
        if header is None:
            raise KofedasError("Documento de venta no encontrado")
        lines = self.db.query(
            """
            SELECT *
            FROM DETMOV
            WHERE DMV_NUMEMP = ? AND DMV_CENTRO = ? AND DMV_TIPDOC = ? AND DMV_TIPAC = ?
              AND DMV_EJERCI = ? AND DMV_SERIE = ? AND DMV_NUMDOC = ?
            ORDER BY DMV_NUMLIN
            """,
            key,
            MAX_ROWS_LIMIT,
        )
        return {"cabecera": header, "lineas": lines, "lineas_total": len(lines)}

    def venta_lineas_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["D.DMV_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("D.DMV_CENTRO = ?")
            params.append(int(args["centro"]))
        if args.get("tipo_documento"):
            where.append("D.DMV_TIPDOC = ?")
            params.append(self._sale_document_type(args.get("tipo_documento")))
        if args.get("tipo_accion") not in (None, ""):
            where.append("D.DMV_TIPAC = ?")
            params.append(str(args["tipo_accion"]).strip()[:1])
        if args.get("ejercicio"):
            where.append("D.DMV_EJERCI = ?")
            params.append(int(args["ejercicio"]))
        if args.get("serie") not in (None, ""):
            where.append("D.DMV_SERIE = ?")
            params.append(str(args["serie"]).strip())
        if args.get("numero"):
            where.append("D.DMV_NUMDOC = ?")
            params.append(int(args["numero"]))
        if args.get("cliente"):
            where.append("C.CBV_CODCLI = ?")
            params.append(int(args["cliente"]))
        if args.get("subcliente") not in (None, ""):
            where.append("C.CBV_SUBCLI = ?")
            params.append(int(args["subcliente"]))
        if str(args.get("articulo") or "").strip():
            where.append("D.DMV_CODART = ?")
            params.append(str(args["articulo"]).strip())
        if args.get("desde"):
            where.append("D.DMV_FECMOV >= ?")
            params.append(self._date_arg(args["desde"]))
        if args.get("hasta"):
            where.append("D.DMV_FECMOV <= ?")
            params.append(self._date_arg(args["hasta"]))
        if str(args.get("texto") or "").strip():
            where.append("(UPPER(D.DMV_CODART) LIKE ? OR UPPER(D.DMV_DESCRI) LIKE ? OR UPPER(C.CBV_NOMCLI) LIKE ?)")
            like = _like(str(args["texto"]))
            params.extend([like, like, like])
        return self.db.query(
            f"""
            SELECT FIRST {limit}
                   D.*, C.CBV_FECHA, C.CBV_CODCLI, C.CBV_SUBCLI, C.CBV_NOMCLI, C.CBV_REFCLI, C.CBV_SITUAC
            FROM DETMOV D
            LEFT JOIN CABDOCV C ON C.CBV_NUMEMP = D.DMV_NUMEMP
                               AND C.CBV_CENTRO = D.DMV_CENTRO
                               AND C.CBV_TIPDOC = D.DMV_TIPDOC
                               AND C.CBV_TIPAC = D.DMV_TIPAC
                               AND C.CBV_EJERCI = D.DMV_EJERCI
                               AND C.CBV_SERIE = D.DMV_SERIE
                               AND C.CBV_NUMDOC = D.DMV_NUMDOC
            WHERE {' AND '.join(where)}
            ORDER BY D.DMV_FECMOV DESC, D.DMV_EJERCI DESC, D.DMV_SERIE, D.DMV_NUMDOC DESC, D.DMV_NUMLIN
            """,
            tuple(params),
            limit,
        )

    def venta_precio_articulo(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        cliente = int(args["cliente"])
        subcliente = self._to_int(args.get("subcliente"), 0)
        fecha = self._date_arg(args.get("fecha"))
        tipdoc = self._sale_document_type(args.get("tipo_documento"), "P")
        item = dict(args)
        item["cantidad"] = self._to_float(args.get("cantidad"), 1)
        return self._sale_price(empresa, cliente, subcliente, item, fecha, tipdoc)

    def rentabilidad_articulo_ventas(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa, mode_info, rows = self._profit_sales_rows(args, True, 250)
        article_cache: dict[str, dict[str, Any]] = {}
        lines: list[dict[str, Any]] = []
        totals = {"unidades": 0.0, "ventas": 0.0, "coste": 0.0, "margen": 0.0}
        skipped = 0
        for row in rows:
            values = self._profit_line_values(empresa, row, str(mode_info["modo"]), article_cache)
            if not values.get("incluida"):
                skipped += 1
                continue
            quantity = self._to_float(row.get("dmv_cantid"), 0)
            totals["unidades"] += quantity
            totals["ventas"] += self._to_float(values.get("ventas_raw"), 0)
            totals["coste"] += self._to_float(values.get("coste_raw"), 0)
            totals["margen"] += self._to_float(values.get("margen_raw"), 0)
            lines.append({
                "fecha": row.get("dmv_fecmov"),
                "centro": row.get("dmv_centro"),
                "tipo_documento": row.get("dmv_tipdoc"),
                "tipo_accion": row.get("dmv_tipac"),
                "ejercicio": row.get("dmv_ejerci"),
                "serie": row.get("dmv_serie"),
                "numero": row.get("dmv_numdoc"),
                "linea": row.get("dmv_numlin"),
                "cliente": row.get("cbv_codcli"),
                "subcliente": row.get("cbv_subcli"),
                "nombre_cliente": row.get("cbv_nomcli"),
                "articulo": row.get("dmv_codart"),
                "descripcion": row.get("dmv_descri"),
                "tipo_linea": row.get("dmv_tiplin"),
                "cantidad": round(quantity, 4),
                "precio_venta": row.get("dmv_preven"),
                "venta_neta": values["ventas"],
                "coste_unitario": values["coste_unitario"],
                "coste_total": values["coste"],
                "margen": values["margen"],
                "rentabilidad": values["rentabilidad"],
                "origen_coste": values["origen_coste"],
                "en_oferta": self._to_int(row.get("dmv_ejeofe"), 0) > 0,
            })
        profitability = (totals["margen"] * 100 / totals["ventas"]) if totals["ventas"] > 0 else 0
        return {
            "empresa": empresa,
            "articulo": str(args["articulo"]).strip(),
            "desde": self._date_arg(args.get("desde")) if args.get("desde") else None,
            "hasta": self._date_arg(args.get("hasta")) if args.get("hasta") else None,
            "modo_coste": mode_info,
            "totales": {
                "unidades": round(totals["unidades"], 4),
                "ventas": round(totals["ventas"], 2),
                "coste": round(totals["coste"], 2),
                "margen": round(totals["margen"], 2),
                "rentabilidad": round(profitability, 4),
            },
            "lineas_total": len(lines),
            "lineas_omitidas": skipped,
            "lineas": lines,
            "fuente_delphi": "ANAVEN_UR.ANALISIS_VENTAS + LIBESP_U.RENTABILIDAD_LINEA",
        }

    def rentabilidad_articulos_resumen(self, args: dict[str, Any]) -> dict[str, Any]:
        query_args = dict(args)
        query_args["limite"] = args.get("limite_lineas") or MAX_ROWS_LIMIT
        empresa, mode_info, rows = self._profit_sales_rows(query_args, False, MAX_ROWS_LIMIT)
        article_cache: dict[str, dict[str, Any]] = {}
        buckets: dict[str, dict[str, Any]] = {}
        skipped = 0
        for row in rows:
            values = self._profit_line_values(empresa, row, str(mode_info["modo"]), article_cache)
            if not values.get("incluida"):
                skipped += 1
                continue
            code = str(row.get("dmv_codart") or "")
            article = self._profit_article(empresa, code, article_cache)
            bucket = buckets.setdefault(code, {
                "articulo": code,
                "descripcion": row.get("dmv_descri") or article.get("art_descri"),
                "familia": row.get("art_codfam") if row.get("art_codfam") is not None else article.get("art_codfam"),
                "subfamilia": row.get("art_subfam") if row.get("art_subfam") is not None else article.get("art_subfam"),
                "proveedor": row.get("art_codpro") if row.get("art_codpro") is not None else article.get("art_codpro"),
                "lineas": 0,
                "unidades": 0.0,
                "ventas": 0.0,
                "coste": 0.0,
                "margen": 0.0,
            })
            bucket["lineas"] += 1
            bucket["unidades"] += self._to_float(row.get("dmv_cantid"), 0)
            bucket["ventas"] += self._to_float(values.get("ventas_raw"), 0)
            bucket["coste"] += self._to_float(values.get("coste_raw"), 0)
            bucket["margen"] += self._to_float(values.get("margen_raw"), 0)
        items: list[dict[str, Any]] = []
        totals = {"unidades": 0.0, "ventas": 0.0, "coste": 0.0, "margen": 0.0}
        for bucket in buckets.values():
            profitability = (bucket["margen"] * 100 / bucket["ventas"]) if bucket["ventas"] > 0 else 0
            item = {
                **bucket,
                "unidades": round(bucket["unidades"], 4),
                "ventas": round(bucket["ventas"], 2),
                "coste": round(bucket["coste"], 2),
                "margen": round(bucket["margen"], 2),
                "rentabilidad": round(profitability, 4),
            }
            items.append(item)
            totals["unidades"] += bucket["unidades"]
            totals["ventas"] += bucket["ventas"]
            totals["coste"] += bucket["coste"]
            totals["margen"] += bucket["margen"]
        order = str(args.get("orden") or "margen").strip().lower()
        order_key = {"margen": "margen", "ventas": "ventas", "rentabilidad": "rentabilidad", "unidades": "unidades"}.get(order, "margen")
        items.sort(key=lambda item: self._to_float(item.get(order_key), 0), reverse=True)
        limit = _positive_limit(args.get("limite"), 100)
        profitability_total = (totals["margen"] * 100 / totals["ventas"]) if totals["ventas"] > 0 else 0
        return {
            "empresa": empresa,
            "desde": self._date_arg(args.get("desde")) if args.get("desde") else None,
            "hasta": self._date_arg(args.get("hasta")) if args.get("hasta") else None,
            "modo_coste": mode_info,
            "totales": {
                "articulos": len(items),
                "unidades": round(totals["unidades"], 4),
                "ventas": round(totals["ventas"], 2),
                "coste": round(totals["coste"], 2),
                "margen": round(totals["margen"], 2),
                "rentabilidad": round(profitability_total, 4),
            },
            "lineas_leidas": len(rows),
            "lineas_omitidas": skipped,
            "orden": order_key,
            "items": items[:limit],
            "fuente_delphi": "ANAVEN_UR.ANALISIS_VENTAS + LIBESP_U.RENTABILIDAD_LINEA",
        }

    def ventas_documentos_detalle(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"), 250)
        where, params = self._anadoc_where(args, "C")
        base_expr = self._anadoc_base_expr("C")
        iva_expr = self._anadoc_tax_expr("C", "IVA")
        req_expr = self._anadoc_tax_expr("C", "REQ")
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   C.CBV_NUMEMP, C.CBV_CENTRO, C.CBV_TIPDOC, C.CBV_TIPAC, C.CBV_EJERCI,
                   C.CBV_SERIE, C.CBV_NUMDOC, C.CBV_CAJA, C.CBV_FECHA, C.CBV_CODCLI,
                   C.CBV_SUBCLI, C.CBV_NOMCLI, COALESCE(CL.CLI_NOMCLI, C.CBV_NOMCLI) AS CLI_NOMCOR,
                   C.CBV_CODMON, {base_expr} AS CBV_BASIMP, {iva_expr} AS CBV_IVA, {req_expr} AS CBV_RE,
                   C.CBV_TOTALD, C.CBV_IMPCOB, (COALESCE(C.CBV_TOTALD,0) - COALESCE(C.CBV_IMPCOB,0)) AS CBV_IMPPEN,
                   C.CBV_SITUAC, C.CBV_REFCLI, C.CBV_RETIRA, C.CBV_CODTAR, T.TAR_NOMBRE AS TARJETA,
                   C.CBV_CODPAG, FP.FPG_DESCRI AS FORMA_PAGO, C.CBV_FORCOB, C.CBV_USUMOD,
                   C.CBV_OBSERV, C.CBV_CIF, C.CBV_CODREP, R.REP_NOMBRE AS REPRESENTANTE,
                   COALESCE(C.CBV_POBLAC, CL.CLI_POBLAC) AS POBLACION,
                   EXTRACT(MONTH FROM C.CBV_FECHA) AS MES,
                   EXTRACT(WEEKDAY FROM C.CBV_FECHA) AS DIA_SEMANA,
                   EXTRACT(HOUR FROM C.CBV_FECMOD) AS HORA,
                   E.CBVE_FECVTO AS FECVTO
            FROM CABDOCV C
            LEFT JOIN CLIEN CL ON CL.CLI_NUMEMP = C.CBV_NUMEMP
                              AND CL.CLI_CODCLI = C.CBV_CODCLI
                              AND CL.CLI_SUBCLI = C.CBV_SUBCLI
            LEFT JOIN FORPAG FP ON FP.FPG_NUMEMP = C.CBV_NUMEMP AND FP.FPG_CODIGO = C.CBV_CODPAG
            LEFT JOIN REPRESE R ON R.REP_NUMEMP = C.CBV_NUMEMP AND R.REP_CODREP = C.CBV_CODREP
            LEFT JOIN CLITAR T ON T.TAR_NUMEMP = C.CBV_NUMEMP AND T.TAR_CODTAR = C.CBV_CODTAR
            LEFT JOIN (
                SELECT CBVE_NUMEMP, CBVE_CENTRO, CBVE_TIPDOC, CBVE_TIPAC,
                       CBVE_EJERCI, CBVE_SERIE, CBVE_NUMDOC, MIN(CBVE_FECVTO) AS CBVE_FECVTO
                FROM CABDOCVE
                GROUP BY CBVE_NUMEMP, CBVE_CENTRO, CBVE_TIPDOC, CBVE_TIPAC,
                         CBVE_EJERCI, CBVE_SERIE, CBVE_NUMDOC
            ) E ON E.CBVE_NUMEMP = C.CBV_NUMEMP
               AND E.CBVE_CENTRO = C.CBV_CENTRO
               AND E.CBVE_TIPDOC = C.CBV_TIPDOC
               AND E.CBVE_TIPAC = C.CBV_TIPAC
               AND E.CBVE_EJERCI = C.CBV_EJERCI
               AND E.CBVE_SERIE = C.CBV_SERIE
               AND E.CBVE_NUMDOC = C.CBV_NUMDOC
            WHERE {' AND '.join(where)}
            ORDER BY C.CBV_FECHA DESC, C.CBV_EJERCI DESC, C.CBV_SERIE, C.CBV_NUMDOC DESC
            """,
            tuple(params),
            limit,
        )
        items: list[dict[str, Any]] = []
        totals = {"base": 0.0, "iva": 0.0, "recargo": 0.0, "total": 0.0, "cobrado": 0.0, "pendiente": 0.0}
        for row in rows:
            tipdoc = str(row.get("cbv_tipdoc") or "").strip().upper()
            for key, source in (("base", "cbv_basimp"), ("iva", "cbv_iva"), ("recargo", "cbv_re"), ("total", "cbv_totald"), ("cobrado", "cbv_impcob"), ("pendiente", "cbv_imppen")):
                totals[key] += self._to_float(row.get(source), 0)
            forma_cobro = str(row.get("cbv_forcob") or "").strip().upper()
            items.append({
                "centro": row.get("cbv_centro"),
                "tipo_documento": tipdoc,
                "tipo_documento_descripcion": SALES_DOCUMENT_TYPES.get(tipdoc, tipdoc),
                "tipo_accion": row.get("cbv_tipac"),
                "ejercicio": row.get("cbv_ejerci"),
                "serie": row.get("cbv_serie"),
                "numero": row.get("cbv_numdoc"),
                "fecha": row.get("cbv_fecha"),
                "cliente": row.get("cbv_codcli"),
                "subcliente": row.get("cbv_subcli"),
                "nombre_cliente": row.get("cbv_nomcli"),
                "poblacion": row.get("poblacion"),
                "moneda": row.get("cbv_codmon"),
                "base": round(self._to_float(row.get("cbv_basimp"), 0), 2),
                "iva": round(self._to_float(row.get("cbv_iva"), 0), 2),
                "recargo": round(self._to_float(row.get("cbv_re"), 0), 2),
                "total": round(self._to_float(row.get("cbv_totald"), 0), 2),
                "cobrado": round(self._to_float(row.get("cbv_impcob"), 0), 2),
                "pendiente": round(self._to_float(row.get("cbv_imppen"), 0), 2),
                "situacion": row.get("cbv_situac"),
                "referencia_cliente": row.get("cbv_refcli"),
                "retira": row.get("cbv_retira"),
                "tarjeta": row.get("cbv_codtar"),
                "tarjeta_nombre": row.get("tarjeta"),
                "forma_pago": row.get("cbv_codpag"),
                "forma_pago_descripcion": row.get("forma_pago") or ("FACTURA TICKETS" if self._to_int(row.get("cbv_codpag"), 0) == -1 else None),
                "forma_cobro": forma_cobro,
                "forma_cobro_descripcion": SALES_COLLECTION_LABELS.get(forma_cobro, forma_cobro),
                "usuario": row.get("cbv_usumod"),
                "observaciones": row.get("cbv_observ"),
                "cif": row.get("cbv_cif"),
                "representante": row.get("cbv_codrep"),
                "representante_nombre": row.get("representante"),
                "mes": row.get("mes"),
                "dia_semana": row.get("dia_semana"),
                "hora": row.get("hora"),
                "vencimiento": row.get("fecvto"),
            })
        return {
            "empresa": empresa,
            "periodo": {"desde": self._date_arg(args.get("desde")) if args.get("desde") else None, "hasta": self._date_arg(args.get("hasta")) if args.get("hasta") else None},
            "documentos": len(items),
            "totales": {key: round(value, 2) for key, value in totals.items()},
            "items": items,
            "fuente_delphi": "ANADOC_UR.ANALISIS_DOCUMENTOS_VENTA",
        }

    def ventas_documentos_resumen(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"), 100)
        group_by = str(args.get("agrupar_por") or "tipo_documento").strip().lower()
        key_expr, name_expr, key_alias, name_alias = self._anadoc_group_expr(group_by)
        where, params = self._anadoc_where(args, "C")
        base_expr = self._anadoc_base_expr("C")
        iva_expr = self._anadoc_tax_expr("C", "IVA")
        req_expr = self._anadoc_tax_expr("C", "REQ")
        order = str(args.get("orden") or "total").strip().lower()
        order_expr = {"base": "BASE", "total": "TOTAL", "pendiente": "PENDIENTE", "documentos": "DOCUMENTOS"}.get(order, "TOTAL")
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   {key_expr} AS {key_alias}, MIN({name_expr}) AS {name_alias},
                   COUNT(*) AS DOCUMENTOS,
                   SUM({base_expr}) AS BASE,
                   SUM({iva_expr}) AS IVA,
                   SUM({req_expr}) AS RECARGO,
                   SUM(COALESCE(C.CBV_TOTALD,0)) AS TOTAL,
                   SUM(COALESCE(C.CBV_IMPCOB,0)) AS COBRADO,
                   SUM(COALESCE(C.CBV_TOTALD,0) - COALESCE(C.CBV_IMPCOB,0)) AS PENDIENTE
            FROM CABDOCV C
            LEFT JOIN CLIEN CL ON CL.CLI_NUMEMP = C.CBV_NUMEMP
                              AND CL.CLI_CODCLI = C.CBV_CODCLI
                              AND CL.CLI_SUBCLI = C.CBV_SUBCLI
            LEFT JOIN FORPAG FP ON FP.FPG_NUMEMP = C.CBV_NUMEMP AND FP.FPG_CODIGO = C.CBV_CODPAG
            LEFT JOIN REPRESE R ON R.REP_NUMEMP = C.CBV_NUMEMP AND R.REP_CODREP = C.CBV_CODREP
            LEFT JOIN CENTROS CE ON CE.CEN_NUMEMP = C.CBV_NUMEMP AND CE.CEN_CODCEN = C.CBV_CENTRO
            LEFT JOIN CLITAR T ON T.TAR_NUMEMP = C.CBV_NUMEMP AND T.TAR_CODTAR = C.CBV_CODTAR
            WHERE {' AND '.join(where)}
            GROUP BY {key_expr}
            ORDER BY {order_expr} DESC
            """,
            tuple(params),
            limit,
        )
        total_row = self.db.one(
            f"""
            SELECT COUNT(*) AS DOCUMENTOS,
                   SUM({base_expr}) AS BASE,
                   SUM({iva_expr}) AS IVA,
                   SUM({req_expr}) AS RECARGO,
                   SUM(COALESCE(C.CBV_TOTALD,0)) AS TOTAL,
                   SUM(COALESCE(C.CBV_IMPCOB,0)) AS COBRADO,
                   SUM(COALESCE(C.CBV_TOTALD,0) - COALESCE(C.CBV_IMPCOB,0)) AS PENDIENTE
            FROM CABDOCV C
            LEFT JOIN CLIEN CL ON CL.CLI_NUMEMP = C.CBV_NUMEMP
                              AND CL.CLI_CODCLI = C.CBV_CODCLI
                              AND CL.CLI_SUBCLI = C.CBV_SUBCLI
            LEFT JOIN FORPAG FP ON FP.FPG_NUMEMP = C.CBV_NUMEMP AND FP.FPG_CODIGO = C.CBV_CODPAG
            LEFT JOIN REPRESE R ON R.REP_NUMEMP = C.CBV_NUMEMP AND R.REP_CODREP = C.CBV_CODREP
            LEFT JOIN CENTROS CE ON CE.CEN_NUMEMP = C.CBV_NUMEMP AND CE.CEN_CODCEN = C.CBV_CENTRO
            LEFT JOIN CLITAR T ON T.TAR_NUMEMP = C.CBV_NUMEMP AND T.TAR_CODTAR = C.CBV_CODTAR
            WHERE {' AND '.join(where)}
            """,
            tuple(params),
        ) or {}
        items: list[dict[str, Any]] = []
        for row in rows:
            code = row.get(key_alias.lower())
            name = row.get(name_alias.lower())
            if group_by in {"tipo_documento", "documento"}:
                name = SALES_DOCUMENT_TYPES.get(str(code or "").strip().upper(), name)
            elif group_by == "forma_cobro":
                name = SALES_COLLECTION_LABELS.get(str(code or "").strip().upper(), name)
            elif group_by == "forma_pago" and self._to_int(code, 0) == -1:
                name = "FACTURA TICKETS"
            item = {
                "codigo": code,
                "nombre": name,
                "documentos": self._to_int(row.get("documentos"), 0),
                "base": round(self._to_float(row.get("base"), 0), 2),
                "iva": round(self._to_float(row.get("iva"), 0), 2),
                "recargo": round(self._to_float(row.get("recargo"), 0), 2),
                "total": round(self._to_float(row.get("total"), 0), 2),
                "cobrado": round(self._to_float(row.get("cobrado"), 0), 2),
                "pendiente": round(self._to_float(row.get("pendiente"), 0), 2),
            }
            items.append(item)
        totals = {
            "documentos": self._to_int(total_row.get("documentos"), 0),
            "base": round(self._to_float(total_row.get("base"), 0), 2),
            "iva": round(self._to_float(total_row.get("iva"), 0), 2),
            "recargo": round(self._to_float(total_row.get("recargo"), 0), 2),
            "total": round(self._to_float(total_row.get("total"), 0), 2),
            "cobrado": round(self._to_float(total_row.get("cobrado"), 0), 2),
            "pendiente": round(self._to_float(total_row.get("pendiente"), 0), 2),
        }
        return {
            "empresa": empresa,
            "periodo": {"desde": self._date_arg(args.get("desde")) if args.get("desde") else None, "hasta": self._date_arg(args.get("hasta")) if args.get("hasta") else None},
            "agrupar_por": group_by,
            "orden": order_expr.lower(),
            "totales": totals,
            "grupos_devueltos": len(items),
            "items": items,
            "fuente_delphi": "ANADOC_UR.ANALISIS_DOCUMENTOS_VENTA",
        }

    def venta_documento_alta_preparar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        tipdoc = self._sale_document_type(args.get("tipo_documento"), "P")
        tipac = str(args.get("tipo_accion") or "0").strip()[:1] or "0"
        cliente_code = int(args["cliente"])
        subcliente = self._to_int(args.get("subcliente"), 0)
        sale_date = self._date_arg(args.get("fecha"))
        ejercicio = self._to_int(args.get("ejercicio"), int(sale_date[:4]))
        serie = self._sale_series(empresa, centro, cliente_code, subcliente, tipdoc, args.get("serie"))
        if not serie:
            raise KofedasError("No se pudo resolver la serie de venta; informa serie")
        numero = self._next_sale_number(empresa, centro, tipdoc, tipac, ejercicio, serie, self._to_int(args.get("numero"), 0) or None)
        customer = self._sale_client(empresa, cliente_code, subcliente)
        now = datetime.now().replace(microsecond=0).isoformat(sep=" ")
        currency = str(args.get("moneda") or customer.get("cli_codmon") or "E").strip()[:1] or "E"
        header = {
            "CBV_NUMEMP": empresa,
            "CBV_CENTRO": centro,
            "CBV_TIPDOC": tipdoc,
            "CBV_TIPAC": tipac,
            "CBV_EJERCI": ejercicio,
            "CBV_SERIE": serie,
            "CBV_NUMDOC": numero,
            "CBV_CAJA": self._to_int(args.get("caja"), 1),
            "CBV_FECHA": sale_date,
            "CBV_FECHAE": self._date_arg(args.get("fecha_entrega"), sale_date),
            "CBV_CODCLI": cliente_code,
            "CBV_SUBCLI": subcliente,
            "CBV_CODREP": self._to_int(args.get("representante"), self._to_int(customer.get("cli_codrep"), 0)),
            "CBV_COMREP": self._to_float(args.get("comision"), 0),
            "CBV_NOMCLI": str(args.get("nombre_cliente") or customer.get("cli_razsoc") or customer.get("cli_nomcli") or "")[:40],
            "CBV_CIF": str(args.get("cif") or customer.get("cli_cif") or "")[:15],
            "CBV_DOMCLI": str(args.get("domicilio") or customer.get("cli_domici") or "")[:50],
            "CBV_CODPOS": self._to_int(args.get("codigo_postal"), self._to_int(customer.get("cli_codpos"), 0)),
            "CBV_POBLAC": str(args.get("poblacion") or customer.get("cli_poblac") or "")[:40],
            "CBV_CODPAG": self._to_int(args.get("forma_pago"), self._to_int(customer.get("cli_forpag"), 0)),
            "CBV_FORENV": self._to_int(args.get("forma_envio"), self._to_int(customer.get("cli_forenv"), 0)),
            "CBV_PORDTO": self._to_float(args.get("descuento"), self._to_float(customer.get("cli_dtoesp"), 0)),
            "CBV_IMPPOR": self._to_float(args.get("portes"), 0),
            "CBV_BASIMP1": 0, "CBV_PORIVA1": 0, "CBV_PORREQ1": 0,
            "CBV_BASIMP2": 0, "CBV_PORIVA2": 0, "CBV_PORREQ2": 0,
            "CBV_BASIMP3": 0, "CBV_PORIVA3": 0, "CBV_PORREQ3": 0,
            "CBV_BASIMP4": 0, "CBV_PORIVA4": 0, "CBV_PORREQ4": 0,
            "CBV_TOTALS": 0, "CBV_TOTALD": 0, "CBV_IMPCOB": 0,
            "CBV_CODMON": currency,
            "CBV_FORCOB": str(args.get("forma_cobro") or "")[:1],
            "CBV_NUMTAR": self._to_int(args.get("numero_tarjeta"), 0),
            "CBV_TIPVEN": self._sale_tipven(empresa, tipdoc, tipac),
            "CBV_SITUAC": str(args.get("situacion") or "P").strip().upper()[:1] or "P",
            "CBV_INDEDI": str(args.get("indice_edicion") or "N").strip().upper()[:1] or "N",
            "CBV_OBSERV": str(args.get("observaciones") or "")[:50],
            "CBV_EJERCID": self._to_int(args.get("documento_origen_ejercicio"), 0),
            "CBV_TIPDOCD": str(args.get("documento_origen_tipo") or "")[:1],
            "CBV_SERIED": str(args.get("documento_origen_serie") or "")[:2],
            "CBV_NUMDOCD": self._to_int(args.get("documento_origen_numero"), 0),
            "CBV_FECMOD": now,
            "CBV_USUMOD": str(args.get("usuario") or f"{centro} MCP")[:10],
            "CBV_REFCLI": str(args.get("referencia_cliente") or "")[:20],
            "CBV_RETIRA": str(args.get("retira") or "")[:40],
            "CBV_CODTAR": str(args.get("tarjeta") or "")[:15],
        }
        raw_lines = args.get("lineas") or []
        if not isinstance(raw_lines, list) or not raw_lines:
            raise KofedasError("lineas debe ser una lista no vacia")
        lines: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        for index, item in enumerate(raw_lines, start=1):
            if not isinstance(item, dict):
                errors.append({"linea": index, "error": "La linea debe ser un objeto"})
                continue
            try:
                line = self._sale_line_defaults(header, index * 10, item, customer)
                lines.append(line)
            except Exception as exc:
                errors.append({"linea": index, "error": str(exc)})
        header = self._sale_totals(header, lines)
        return {
            "numero": numero,
            "ejercicio": ejercicio,
            "serie": serie,
            "tipo_documento": tipdoc,
            "tipo_accion": tipac,
            "cabecera": header,
            "lineas": lines,
            "errores": errors,
            "fuente_delphi": {
                "cabecera": "CABDOCV_UDM.GRABAR_CABDOCV('G')",
                "detalle": "CABDOCV_UDM.GRABAR_DETMOV('G')",
                "precio": "LIBESP_U.PRECIO_VENTA / Precio_Cliente_Articulo adaptado a tablas Kofedas disponibles",
                "valoracion": "CABDOCV_UDM.VALORAR_DETMOV + FINALIZAR_CABDOCV",
            },
        }

    def venta_documento_alta(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        plan = self.venta_documento_alta_preparar(args)
        if plan["errores"]:
            if args.get("simular"):
                return {"simulado": True, **plan}
            raise KofedasError("El documento contiene errores; usa simular=true para revisar el detalle")
        header = plan["cabecera"]
        empresa = int(header["CBV_NUMEMP"])
        centro = int(header["CBV_CENTRO"])
        tipdoc = str(header["CBV_TIPDOC"])
        tipac = str(header["CBV_TIPAC"])
        ejercicio = int(header["CBV_EJERCI"])
        serie = str(header["CBV_SERIE"])
        numero = int(header["CBV_NUMDOC"])
        explicit_number = bool(args.get("numero"))
        if self._sale_document_exists(empresa, centro, tipdoc, tipac, ejercicio, serie, numero):
            if explicit_number:
                raise KofedasError("Ya existe el documento de venta indicado")
            while self._sale_document_exists(empresa, centro, tipdoc, tipac, ejercicio, serie, numero):
                numero += 1
            header["CBV_NUMDOC"] = numero
            plan["numero"] = numero
            for line in plan["lineas"]:
                line["DMV_NUMDOC"] = numero
        columns = self._table_columns("CABDOCV")
        line_columns = self._table_columns("DETMOV")
        statements: list[tuple[str, tuple[Any, ...]]] = [
            (
                "INSERT INTO CABDOCV (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
                tuple(header.get(column) for column in columns),
            )
        ]
        for line in plan["lineas"]:
            data = {column: line.get(column) for column in line_columns}
            statements.append((
                "INSERT INTO DETMOV (" + ", ".join(line_columns) + ") VALUES (" + ", ".join("?" for _ in line_columns) + ")",
                tuple(data.get(column) for column in line_columns),
            ))
        statements.append(self._sale_numera_statement(header))
        if args.get("simular"):
            return {
                "simulado": True,
                "ejercicio": ejercicio,
                "serie": serie,
                "numero": numero,
                "tipo_documento": tipdoc,
                "lineas": len(plan["lineas"]),
                "sentencias": len(statements),
                "plan": plan,
            }
        counts = self.db.execute_transaction(statements)
        return {
            "simulado": False,
            "ejercicio": ejercicio,
            "serie": serie,
            "numero": numero,
            "tipo_documento": tipdoc,
            "lineas": len(plan["lineas"]),
            "filas_afectadas": counts,
            "totales": {"base": header["CBV_TOTALS"], "total": header["CBV_TOTALD"]},
            "nota_stock": "Los pedidos/presupuestos no modifican existencias; para otros tipos esta funcion graba DETMOV pero no actualiza ARTICULE.",
        }

    def venta_pedido_alta(self, args: dict[str, Any]) -> dict[str, Any]:
        payload = dict(args)
        payload["tipo_documento"] = "P"
        payload["tipo_accion"] = payload.get("tipo_accion") or "0"
        return self.venta_documento_alta(payload)

    def pedido_crear(self, args: dict[str, Any]) -> dict[str, Any]:
        payload = dict(args)
        if payload.get("cliente") in (None, "") and payload.get("codcli") not in (None, ""):
            payload["cliente"] = payload["codcli"]
        if payload.get("subcliente") in (None, "") and payload.get("subcli") not in (None, ""):
            payload["subcliente"] = payload["subcli"]
        if not payload.get("lineas") and payload.get("texto"):
            try:
                parsed = json.loads(str(payload["texto"]))
                if isinstance(parsed, list):
                    payload["lineas"] = parsed
                elif isinstance(parsed, dict) and isinstance(parsed.get("lineas"), list):
                    payload.update({key: value for key, value in parsed.items() if key not in payload or payload.get(key) in (None, "")})
            except json.JSONDecodeError:
                pass
        if not payload.get("lineas"):
            raise KofedasError("pedido_crear necesita lineas estructuradas o texto JSON con lineas")
        return self.venta_pedido_alta(payload)

    def pedido_listar(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"), 100)
        where = ["C.CBV_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") not in (None, ""):
            where.append("C.CBV_CENTRO = ?")
            params.append(self._to_int(args.get("centro"), self.centro))
        elif args.get("cliente", args.get("codcli")) in (None, ""):
            where.append("C.CBV_CENTRO = ?")
            params.append(self.centro)
        raw_tipdoc = args.get("tipo_documento", args.get("tipdoc"))
        if raw_tipdoc not in (None, ""):
            tipdocs = [self._sale_document_type(part, "P") for part in re.split(r"[,; ]+", str(raw_tipdoc)) if part.strip()]
        else:
            tipdocs = ["P", "R"]
        where.append("C.CBV_TIPDOC IN (" + ", ".join("?" for _ in tipdocs) + ")")
        params.extend(tipdocs)
        if args.get("tipo_accion", args.get("tipac")) not in (None, ""):
            where.append("C.CBV_TIPAC = ?")
            params.append(str(args.get("tipo_accion", args.get("tipac"))).strip()[:1])
        for arg_name, column in (("ejercicio", "C.CBV_EJERCI"), ("numero", "C.CBV_NUMDOC")):
            if args.get(arg_name) not in (None, ""):
                where.append(f"{column} = ?")
                params.append(self._to_int(args.get(arg_name), 0))
        if args.get("ejerci") not in (None, "") and args.get("ejercicio") in (None, ""):
            where.append("C.CBV_EJERCI = ?")
            params.append(self._to_int(args.get("ejerci"), 0))
        if args.get("numdoc") not in (None, "") and args.get("numero") in (None, ""):
            where.append("C.CBV_NUMDOC = ?")
            params.append(self._to_int(args.get("numdoc"), 0))
        if args.get("serie") not in (None, ""):
            where.append("C.CBV_SERIE = ?")
            params.append(str(args.get("serie")).strip()[:2])
        cliente = args.get("cliente", args.get("codcli"))
        subcliente = args.get("subcliente", args.get("subcli"))
        if cliente not in (None, ""):
            where.append("C.CBV_CODCLI = ?")
            params.append(self._to_int(cliente, 0))
        if subcliente not in (None, ""):
            where.append("C.CBV_SUBCLI = ?")
            params.append(self._to_int(subcliente, 0))
        if args.get("estado") not in (None, ""):
            where.append("C.CBV_SITUAC = ?")
            params.append(str(args.get("estado")).strip().upper()[:1])
        if args.get("desde"):
            where.append("C.CBV_FECHA >= ?")
            params.append(self._date_arg(args.get("desde")))
        if args.get("hasta"):
            where.append("C.CBV_FECHA <= ?")
            params.append(self._date_arg(args.get("hasta")))
        if args.get("articulo"):
            where.append(
                """EXISTS (
                    SELECT 1 FROM DETMOV D
                    WHERE D.DMV_NUMEMP = C.CBV_NUMEMP AND D.DMV_CENTRO = C.CBV_CENTRO
                      AND D.DMV_TIPDOC = C.CBV_TIPDOC AND D.DMV_TIPAC = C.CBV_TIPAC
                      AND D.DMV_EJERCI = C.CBV_EJERCI AND D.DMV_SERIE = C.CBV_SERIE
                      AND D.DMV_NUMDOC = C.CBV_NUMDOC AND D.DMV_CODART = ?
                )"""
            )
            params.append(str(args.get("articulo")).strip())
        if args.get("texto"):
            pattern = _like(str(args.get("texto")))
            where.append("(UPPER(C.CBV_NOMCLI) LIKE ? OR UPPER(C.CBV_REFCLI) LIKE ? OR UPPER(C.CBV_OBSERV) LIKE ?)")
            params.extend([pattern, pattern, pattern])
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   C.CBV_NUMEMP, C.CBV_CENTRO, C.CBV_TIPDOC, C.CBV_TIPAC, C.CBV_EJERCI,
                   C.CBV_SERIE, C.CBV_NUMDOC, C.CBV_FECHA, C.CBV_FECHAE, C.CBV_CODCLI,
                   C.CBV_SUBCLI, C.CBV_NOMCLI, C.CBV_REFCLI, C.CBV_RETIRA, C.CBV_SITUAC,
                   C.CBV_TOTALS, C.CBV_TOTALD, C.CBV_IMPCOB, C.CBV_OBSERV,
                   COUNT(D.DMV_NUMLIN) AS LINEAS,
                   SUM(COALESCE(D.DMV_CANTID,0)) AS CANTIDAD
            FROM CABDOCV C
            LEFT JOIN DETMOV D ON D.DMV_NUMEMP = C.CBV_NUMEMP AND D.DMV_CENTRO = C.CBV_CENTRO
                              AND D.DMV_TIPDOC = C.CBV_TIPDOC AND D.DMV_TIPAC = C.CBV_TIPAC
                              AND D.DMV_EJERCI = C.CBV_EJERCI AND D.DMV_SERIE = C.CBV_SERIE
                              AND D.DMV_NUMDOC = C.CBV_NUMDOC
            WHERE {' AND '.join(where)}
            GROUP BY C.CBV_NUMEMP, C.CBV_CENTRO, C.CBV_TIPDOC, C.CBV_TIPAC, C.CBV_EJERCI,
                     C.CBV_SERIE, C.CBV_NUMDOC, C.CBV_FECHA, C.CBV_FECHAE, C.CBV_CODCLI,
                     C.CBV_SUBCLI, C.CBV_NOMCLI, C.CBV_REFCLI, C.CBV_RETIRA, C.CBV_SITUAC,
                     C.CBV_TOTALS, C.CBV_TOTALD, C.CBV_IMPCOB, C.CBV_OBSERV
            ORDER BY C.CBV_FECHA DESC, C.CBV_EJERCI DESC, C.CBV_SERIE, C.CBV_NUMDOC DESC
            """,
            tuple(params),
            limit,
        )
        return {"total": len(rows), "items": rows}

    def pedido_detalle(self, args: dict[str, Any]) -> dict[str, Any]:
        key = self._pedido_key(args)
        modo = str(args.get("modo") or "normal").strip().lower()
        if modo not in {"normal", "preparacion"}:
            raise KofedasError("modo debe ser normal o preparacion")
        return self._pedido_document_result(key, modo)

    def pedido_situacion_actualizar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        key = self._pedido_key(args)
        situation = str(args.get("situacion", args.get("situac")) or "").strip().upper()[:1]
        if not situation:
            raise KofedasError("Debe informar situacion")
        where, params = self._pedido_where_from_key(key, "")
        updates = {
            "CBV_SITUAC": situation,
            "CBV_FECMOD": datetime.now().replace(microsecond=0).isoformat(sep=" "),
            "CBV_USUMOD": str(args.get("usuario") or f"{key['centro']} MCP")[:10],
        }
        statement = self._pedido_update_statement("CABDOCV", updates, where, params)
        if not statement:
            raise KofedasError("CABDOCV no contiene campos actualizables de situacion")
        if args.get("simular"):
            return {"simulado": True, "pedido": key, "situacion": situation, "sentencias": [statement[0]]}
        count = self.db.execute(statement[0], statement[1])
        return {"simulado": False, "pedido": key, "situacion": situation, "filas_afectadas": count}

    def pedido_retirada_actualizar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        key = self._pedido_key(args)
        where, params = self._pedido_where_from_key(key, "")
        updates = {
            "CBV_RETIRA": str(args.get("retirado", args.get("retira")) or "")[:40],
            "CBV_REFCLI": str(args.get("referencia", args.get("referencia_cliente")) or "")[:20],
            "CBV_FECMOD": datetime.now().replace(microsecond=0).isoformat(sep=" "),
            "CBV_USUMOD": str(args.get("usuario") or f"{key['centro']} MCP")[:10],
        }
        statement = self._pedido_update_statement("CABDOCV", updates, where, params)
        if not statement:
            raise KofedasError("CABDOCV no contiene campos actualizables de retirada")
        if args.get("simular"):
            return {"simulado": True, "pedido": key, "retira": updates["CBV_RETIRA"], "referencia": updates["CBV_REFCLI"], "sentencias": [statement[0]]}
        count = self.db.execute(statement[0], statement[1])
        return {"simulado": False, "pedido": key, "retira": updates["CBV_RETIRA"], "referencia": updates["CBV_REFCLI"], "filas_afectadas": count}

    def pedido_cerrar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        key = self._pedido_key(args)
        if key["tipdoc"] not in {"P", "R"}:
            raise KofedasError("pedido_cerrar solo puede cerrar pedidos (P) o presupuestos (R)")
        self._pedido_header(key)
        closed_key = {**key, "tipdoc": "S"}
        if self._sale_document_exists(closed_key["empresa"], closed_key["centro"], closed_key["tipdoc"], closed_key["tipac"], closed_key["ejercicio"], closed_key["serie"], closed_key["numero"]):
            raise KofedasError("Ya existe un documento historico S con las mismas claves")
        now = datetime.now().replace(microsecond=0).isoformat(sep=" ")
        user = str(args.get("usuario") or f"{key['centro']} MCP")[:10]
        where_h, params_h = self._pedido_where_from_key(key, "")
        statements: list[tuple[str, tuple[Any, ...]]] = [
            ("UPDATE CABDOCV SET CBV_TIPDOC = 'S', CBV_SITUAC = 'C', CBV_FECMOD = ?, CBV_USUMOD = ? WHERE " + where_h, (now, user, *params_h)),
            (
                """
                UPDATE DETMOV SET DMV_TIPDOC = 'S'
                WHERE DMV_NUMEMP = ? AND DMV_CENTRO = ? AND DMV_TIPDOC = ? AND DMV_TIPAC = ?
                  AND DMV_EJERCI = ? AND DMV_SERIE = ? AND DMV_NUMDOC = ?
                """,
                (key["empresa"], key["centro"], key["tipdoc"], key["tipac"], key["ejercicio"], key["serie"], key["numero"]),
            ),
        ]
        if args.get("simular"):
            return {"simulado": True, "pedido": key, "pedido_historico": closed_key, "sentencias": [sql for sql, _ in statements]}
        counts = self.db.execute_transaction(statements)
        return {"simulado": False, "pedido": key, "pedido_historico": closed_key, "filas_afectadas": counts}

    def pedido_marcar_preparado(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        key = self._pedido_key(args)
        self._pedido_header(key)
        columns = set(self._table_columns("DETMOV"))
        prepared_column = next((column for column in ("DMV_CANPREA", "DMV_CANPRE", "DMV_CANSER", "DMV_CANPREPARADA") if column in columns), None)
        statements: list[tuple[str, tuple[Any, ...]]] = []
        raw_lines = args.get("lineas") or []
        if prepared_column and raw_lines:
            for item in raw_lines:
                if not isinstance(item, dict):
                    continue
                line_number = self._to_int(item.get("linea", item.get("numlin")), 0)
                if not line_number:
                    continue
                quantity = self._to_float(item.get("cantidad_preparada", item.get("cantidad", item.get("cantid"))), 0)
                statements.append((
                    f"""
                    UPDATE DETMOV SET {prepared_column} = ?
                    WHERE DMV_NUMEMP = ? AND DMV_CENTRO = ? AND DMV_TIPDOC = ? AND DMV_TIPAC = ?
                      AND DMV_EJERCI = ? AND DMV_SERIE = ? AND DMV_NUMDOC = ? AND DMV_NUMLIN = ?
                    """,
                    (quantity, key["empresa"], key["centro"], key["tipdoc"], key["tipac"], key["ejercicio"], key["serie"], key["numero"], line_number),
                ))
        situation = str(args.get("situacion") or "P").strip().upper()[:1]
        where_h, params_h = self._pedido_where_from_key(key, "")
        header_statement = self._pedido_update_statement(
            "CABDOCV",
            {
                "CBV_SITUAC": situation,
                "CBV_FECMOD": datetime.now().replace(microsecond=0).isoformat(sep=" "),
                "CBV_USUMOD": str(args.get("usuario") or f"{key['centro']} MCP")[:10],
            },
            where_h,
            params_h,
        )
        if header_statement:
            statements.append(header_statement)
        if args.get("simular"):
            return {"simulado": True, "pedido": key, "campo_preparacion": prepared_column, "sentencias": [sql for sql, _ in statements]}
        counts = self.db.execute_transaction(statements) if statements else []
        return {"simulado": False, "pedido": key, "campo_preparacion": prepared_column, "filas_afectadas": counts}

    def pedido_finalizar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        key = self._pedido_key(args)
        detail = self._pedido_document_result(key, "preparacion")
        situation = str(args.get("situacion") or "").strip().upper()[:1]
        if not situation:
            return {"simulado": bool(args.get("simular")), "pedido": key, "diagnostico": detail["totales_preparacion"], "detalle": detail}
        result = self.pedido_situacion_actualizar({**args, "situacion": situation})
        return {"pedido": key, "diagnostico": detail["totales_preparacion"], "actualizacion": result}

    def pedido_linea_mover(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        key = self._pedido_key(args)
        line_number = self._to_int(args.get("linea", args.get("numlin")), 0)
        if not line_number:
            raise KofedasError("Debe informar linea")
        columns = set(self._table_columns("DETMOV"))
        zone_column = next((column for column in ("DMV_ZONA", "DMV_ZONPRE", "DMV_ZONPREP") if column in columns), None)
        if not zone_column:
            raise KofedasError("DETMOV no tiene campos de zona de preparacion conocidos")
        params: tuple[Any, ...] = (
            self._to_int(args.get("zona_destino"), 0),
            key["empresa"], key["centro"], key["tipdoc"], key["tipac"], key["ejercicio"], key["serie"], key["numero"], line_number,
        )
        sql = (
            f"UPDATE DETMOV SET {zone_column} = ? "
            "WHERE DMV_NUMEMP = ? AND DMV_CENTRO = ? AND DMV_TIPDOC = ? AND DMV_TIPAC = ? "
            "AND DMV_EJERCI = ? AND DMV_SERIE = ? AND DMV_NUMDOC = ? AND DMV_NUMLIN = ?"
        )
        if args.get("zona_origen") not in (None, ""):
            sql += f" AND {zone_column} = ?"
            params = (*params, self._to_int(args.get("zona_origen"), 0))
        if args.get("simular"):
            return {"simulado": True, "pedido": key, "linea": line_number, "campo_zona": zone_column, "sentencias": [sql]}
        count = self.db.execute(sql, params)
        return {"simulado": False, "pedido": key, "linea": line_number, "campo_zona": zone_column, "filas_afectadas": count}

    def pedido_albaranar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        order_key = self._pedido_key(args)
        header = self._pedido_header(order_key)
        lines = self._pedido_lines(order_key)
        selected_quantities: dict[int, float] = {}
        for item in args.get("lineas") or []:
            if isinstance(item, dict):
                selected_quantities[self._to_int(item.get("linea", item.get("numlin")), 0)] = self._to_float(item.get("cantidad", item.get("cantid")), 0)
        delivery_lines: list[dict[str, Any]] = []
        for line in lines:
            if str(line.get("dmv_tiplin") or "").upper() not in {"D", "X"}:
                continue
            line_number = self._to_int(line.get("dmv_numlin"), 0)
            quantity = selected_quantities.get(line_number, self._to_float(line.get("dmv_canpen"), 0) or self._to_float(line.get("dmv_cantid"), 0))
            if quantity <= 0:
                continue
            delivery_lines.append({
                "articulo": line.get("dmv_codart"),
                "descripcion": line.get("dmv_descri"),
                "cantidad": quantity,
                "precio": line.get("dmv_preven"),
                "dto1": line.get("dmv_dto1"),
                "dto2": line.get("dmv_dto2"),
                "unidad": line.get("dmv_unimed"),
                "origen_ejercicio": order_key["ejercicio"],
                "origen_tipo_documento": order_key["tipdoc"],
                "origen_serie": order_key["serie"],
                "origen_numero": order_key["numero"],
                "origen_linea": line_number,
            })
        if not delivery_lines:
            raise KofedasError("No hay lineas con cantidad para albaranar")
        payload = {
            "empresa": order_key["empresa"],
            "centro": order_key["centro"],
            "tipo_documento": "A",
            "tipo_accion": order_key["tipac"],
            "serie": args.get("serie_albaran"),
            "cliente": header.get("cbv_codcli"),
            "subcliente": header.get("cbv_subcli"),
            "fecha": args.get("fecha"),
            "fecha_entrega": args.get("fecha"),
            "observaciones": "Albaran generado desde pedido "
            + f"{order_key['ejercicio']}/{order_key['serie']}/{order_key['numero']}",
            "referencia_cliente": header.get("cbv_refcli"),
            "retira": header.get("cbv_retira"),
            "lineas": delivery_lines,
            "simular": args.get("simular"),
        }
        result = self.venta_documento_alta(payload)
        if args.get("cerrar_pedido") and not args.get("simular"):
            result["cierre_pedido"] = self.pedido_cerrar({**args, "simular": False})
        return {"pedido": order_key, "albaran": result}

    def pedido_pdf_gestion(self, args: dict[str, Any]) -> dict[str, Any]:
        action = str(args.get("accion") or "").strip().lower()
        if action not in {"generar", "obtener"}:
            raise KofedasError("accion debe ser generar u obtener")
        key = self._pedido_key(args)
        html = self._pedido_html(key)
        encoded = base64.b64encode(html.encode("utf-8")).decode("ascii")
        return {
            "accion": action,
            "pedido": key,
            "formato": "html",
            "mime_type": "text/html; charset=utf-8",
            "nombre_fichero": f"pedido_{key['ejercicio']}_{key['serie']}_{key['numero']}.html",
            "content_base64": encoded,
            "nota": "Kofedas MCP genera HTML autonomo; no ejecuta el motor de informes Delphi.",
        }

    def pedido_enviar(self, args: dict[str, Any]) -> dict[str, Any]:
        key = self._pedido_key(args)
        header = self._pedido_header(key)
        attachment = self.pedido_pdf_gestion({**args, "accion": "generar"})
        email = str(args.get("email") or header.get("cli_email") or "").strip()
        subject = str(args.get("asunto") or f"Pedido {key['ejercicio']}/{key['serie']}/{key['numero']}").strip()
        return {
            "enviado": False,
            "pedido": key,
            "destinatario": email,
            "asunto": subject,
            "adjunto": attachment,
            "nota": "Funcion preparada para integracion: devuelve el contenido, pero no envia correo desde el MCP.",
        }

    def cartera_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(CARTERA_TABLES.items()):
            try:
                columns = self._table_columns(table)
                keys = self.db.primary_key(table)
                available = True
            except Exception as exc:
                columns = []
                keys = []
                available = False
                description = f"{description} (no disponible: {exc})"
            result.append({"tabla": table, "descripcion": description, "disponible": available, "claves": keys, "columnas": columns})
        return result

    def cartera_efectos_detalle(self, args: dict[str, Any]) -> dict[str, Any]:
        effects, meta = self._cartera_effect_rows(args)
        return {"filtros": args, **meta, "efectos": effects}

    def cartera_deuda_cliente(self, args: dict[str, Any]) -> dict[str, Any]:
        payload = {**args, "situacion": args.get("situacion") or "pendiente"}
        effects, meta = self._cartera_effect_rows(payload)
        customer = None
        if effects:
            customer = effects[0]["cliente"]
        else:
            row = self.db.one(
                "SELECT FIRST 1 CLI_CODCLI, CLI_SUBCLI, CLI_NOMCLI, CLI_RAZSOC, CLI_CIF FROM CLIEN WHERE CLI_NUMEMP = ? AND CLI_CODCLI = ?",
                (self._empresa(args), int(args["cliente"])),
            )
            if row:
                customer = {
                    "codigo": row.get("cli_codcli"),
                    "subcliente": row.get("cli_subcli"),
                    "nombre": row.get("cli_razsoc") or row.get("cli_nomcli"),
                    "cif": row.get("cli_cif"),
                }
        return {"cliente": customer or {"codigo": int(args["cliente"]), "subcliente": args.get("subcliente")}, **meta, "efectos": effects}

    def cartera_deuda_por_cliente(self, args: dict[str, Any]) -> dict[str, Any]:
        payload = {**args, "situacion": args.get("situacion") or "pendiente", "limite": args.get("limite") or 1000}
        effects, meta = self._cartera_effect_rows(payload)
        groups: dict[tuple[int, int], dict[str, Any]] = {}
        for effect in effects:
            cli = effect["cliente"]
            key = (int(cli["codigo"] or 0), int(cli["subcliente"] or 0))
            group = groups.setdefault(key, {
                "cliente": cli,
                "efectos": 0,
                "pendiente": 0.0,
                "vencido": 0.0,
                "remesado": 0.0,
                "no_remesado": 0.0,
                "proximo_vencimiento": None,
            })
            pending = self._to_float(effect.get("pendiente"), 0)
            group["efectos"] += 1
            group["pendiente"] += pending
            if effect["situacion"] == "vencido":
                group["vencido"] += pending
            if effect["remesado"]:
                group["remesado"] += pending
            else:
                group["no_remesado"] += pending
            due = effect.get("vencimiento")
            if due and (group["proximo_vencimiento"] is None or str(due) < str(group["proximo_vencimiento"])):
                group["proximo_vencimiento"] = due
        items = []
        for group in groups.values():
            for field in ("pendiente", "vencido", "remesado", "no_remesado"):
                group[field] = round(group[field], 2)
            items.append(group)
        items.sort(key=lambda item: self._to_float(item.get("pendiente"), 0), reverse=True)
        limit = _positive_limit(args.get("limite_clientes"), 100)
        return {**meta, "clientes": items[:limit]}

    def cartera_pendiente_remesar(self, args: dict[str, Any]) -> dict[str, Any]:
        effects, meta = self._cartera_effect_rows({**args, "situacion": "pendiente", "remesado": "N"})
        return {"criterio": "CBVE_FECCAN IS NULL AND CBVE_EJEREM=0 AND CBVE_CODREM=0", **meta, "efectos": effects}

    def cartera_deuda_por_tipo(self, args: dict[str, Any]) -> dict[str, Any]:
        payload = {**args, "situacion": args.get("situacion") or "pendiente", "limite": args.get("limite") or 1000}
        effects, meta = self._cartera_effect_rows(payload)
        groups: dict[tuple[str, str, str, str], dict[str, Any]] = {}
        for effect in effects:
            key = (
                str(effect["documento"]["tipo"] or ""),
                str(effect["tipo_efecto"] or ""),
                str(effect["situacion"] or ""),
                "S" if effect["remesado"] else "N",
            )
            group = groups.setdefault(key, {
                "tipo_documento": key[0],
                "tipo_documento_descripcion": SALES_DOCUMENT_TYPES.get(key[0], key[0]),
                "tipo_efecto": key[1],
                "tipo_efecto_nombre": self._cartera_effect_type_label(key[1]),
                "situacion": key[2],
                "remesado": key[3] == "S",
                "efectos": 0,
                "nominal": 0.0,
                "cobrado": 0.0,
                "pendiente": 0.0,
                "vencido": 0.0,
            })
            group["efectos"] += 1
            group["nominal"] += self._to_float(effect.get("nominal"), 0)
            group["cobrado"] += self._to_float(effect.get("cobrado"), 0)
            group["pendiente"] += self._to_float(effect.get("pendiente"), 0)
            if effect["situacion"] == "vencido":
                group["vencido"] += self._to_float(effect.get("pendiente"), 0)
        items = []
        for group in groups.values():
            for field in ("nominal", "cobrado", "pendiente", "vencido"):
                group[field] = round(group[field], 2)
            items.append(group)
        items.sort(key=lambda item: self._to_float(item.get("pendiente"), 0), reverse=True)
        return {**meta, "resumen": items}

    def dashboard_resumen(self, args: dict[str, Any]) -> dict[str, Any]:
        start, end = self._dashboard_period(args)
        sales = self._dashboard_sales_totals(args, start, end)
        purchases = self._dashboard_purchase_totals(args, start, end)
        cartera = self.cartera_deuda_por_cliente({"empresa": self._empresa(args), "centro": args.get("centro"), "limite": 1000, "limite_clientes": 10})
        pending_purchase_orders = self.orden_compra_listar({"empresa": self._empresa(args), "centro": args.get("centro"), "estado": "pendiente", "limite": 100})
        pending_entries = self.entrada_almacen_pendientes_facturar({"empresa": self._empresa(args), "centro": args.get("centro"), "limite": 100})
        margin = round(sales["base"] - purchases["base"], 2)
        return {
            "periodo": {"desde": start, "hasta": end},
            "ventas": sales,
            "compras": purchases,
            "margen_bruto_aproximado": margin,
            "margen_bruto_pct": round((margin * 100 / sales["base"]) if sales["base"] else 0, 2),
            "cartera": cartera.get("totales", {}),
            "pendientes": {
                "ordenes_compra": {
                    "documentos": len(pending_purchase_orders),
                    "importe_pendiente": round(sum(self._to_float(row.get("coc_imppen"), 0) for row in pending_purchase_orders), 2),
                },
                "entradas_facturar": {
                    "documentos": len(pending_entries.get("documentos", [])) if isinstance(pending_entries, dict) else 0,
                    "total": (pending_entries.get("totales", {}) or {}).get("total") if isinstance(pending_entries, dict) else 0,
                },
            },
        }

    def ventas_resumen(self, args: dict[str, Any]) -> dict[str, Any]:
        start, end = self._dashboard_period(args)
        group_by = str(args.get("agrupar_por") or "mes").strip().lower()
        key_expr, name_expr, key_alias, name_alias = self._dashboard_group_expr("ventas", group_by)
        where, params, types = self._dashboard_sales_where(args, start, end, "C")
        limit = _positive_limit(args.get("limite"), 100)
        joins = ""
        value_expr = "CASE WHEN C.CBV_TIPDOC = 'C' THEN -D.DMV_VALLINS ELSE D.DMV_VALLINS END"
        qty_expr = "CASE WHEN C.CBV_TIPDOC = 'C' THEN -D.DMV_CANTID ELSE D.DMV_CANTID END"
        if group_by in {"articulo", "familia"}:
            joins += """
            JOIN DETMOV D ON D.DMV_NUMEMP = C.CBV_NUMEMP AND D.DMV_CENTRO = C.CBV_CENTRO
                         AND D.DMV_TIPDOC = C.CBV_TIPDOC AND D.DMV_TIPAC = C.CBV_TIPAC
                         AND D.DMV_EJERCI = C.CBV_EJERCI AND D.DMV_SERIE = C.CBV_SERIE
                         AND D.DMV_NUMDOC = C.CBV_NUMDOC AND D.DMV_TIPLIN IN ('D','X')
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = D.DMV_NUMEMP AND A.ART_CODART = D.DMV_CODART
            LEFT JOIN FAMILI F ON F.FAM_NUMEMP = A.ART_NUMEMP AND F.FAM_CODIGO = A.ART_CODFAM
            """
            select_value = f"SUM({value_expr})"
            select_units = f"SUM({qty_expr})"
            documents_count = "COUNT(*)"
        else:
            select_value = "SUM(CASE WHEN C.CBV_TIPDOC = 'C' THEN -C.CBV_TOTALS ELSE C.CBV_TOTALS END)"
            select_units = "0"
            documents_count = "COUNT(*)"
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   {key_expr} AS {key_alias}, {name_expr} AS {name_alias},
                   {documents_count} AS DOCUMENTOS,
                   {select_value} AS BASE,
                   {select_units} AS UNIDADES
            FROM CABDOCV C
            {joins}
            WHERE {' AND '.join(where)}
            GROUP BY {key_expr}, {name_expr}
            ORDER BY BASE DESC
            """,
            tuple(params),
            limit,
        )
        items = [{
            "codigo": row.get(key_alias.lower()),
            "nombre": row.get(name_alias.lower()),
            "documentos": self._to_int(row.get("documentos"), 0),
            "base": round(self._to_float(row.get("base"), 0), 2),
            "unidades": round(self._to_float(row.get("unidades"), 0), 4),
        } for row in rows]
        return {"periodo": {"desde": start, "hasta": end}, "agrupar_por": group_by, "tipos_documento": types, "totales": self._dashboard_sales_totals(args, start, end), "items": items}

    def compras_resumen(self, args: dict[str, Any]) -> dict[str, Any]:
        start, end = self._dashboard_period(args)
        group_by = str(args.get("agrupar_por") or "mes").strip().lower()
        key_expr, name_expr, key_alias, name_alias = self._dashboard_group_expr("compras", group_by)
        where, params = self._dashboard_purchase_where(args, start, end, "C")
        limit = _positive_limit(args.get("limite"), 100)
        joins = ""
        if group_by in {"articulo", "familia"}:
            joins += """
            JOIN DETMOVM D ON D.DMM_NUMEMP = C.CBM_NUMEMP AND D.DMM_CENTRO = C.CBM_CENTRO
                          AND D.DMM_EJERCI = C.CBM_EJERCI AND D.DMM_SERIE = C.CBM_SERIE
                          AND D.DMM_NUMDOC = C.CBM_NUMDOC AND D.DMM_TIPLIN IN ('D','X')
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = D.DMM_NUMEMP AND A.ART_CODART = D.DMM_CODART
            LEFT JOIN FAMILI F ON F.FAM_NUMEMP = A.ART_NUMEMP AND F.FAM_CODIGO = A.ART_CODFAM
            """
            select_value = "SUM(D.DMM_VALLIN)"
            select_units = "SUM(D.DMM_CANTID)"
            documents_count = "COUNT(*)"
        else:
            select_value = "SUM(C.CBM_TOTALS)"
            select_units = "0"
            documents_count = "COUNT(*)"
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   {key_expr} AS {key_alias}, {name_expr} AS {name_alias},
                   {documents_count} AS DOCUMENTOS,
                   {select_value} AS BASE,
                   {select_units} AS UNIDADES
            FROM CABDOCM C
            {joins}
            WHERE {' AND '.join(where)}
            GROUP BY {key_expr}, {name_expr}
            ORDER BY BASE DESC
            """,
            tuple(params),
            limit,
        )
        items = [{
            "codigo": row.get(key_alias.lower()),
            "nombre": row.get(name_alias.lower()),
            "documentos": self._to_int(row.get("documentos"), 0),
            "base": round(self._to_float(row.get("base"), 0), 2),
            "unidades": round(self._to_float(row.get("unidades"), 0), 4),
        } for row in rows]
        return {"periodo": {"desde": start, "hasta": end}, "agrupar_por": group_by, "totales": self._dashboard_purchase_totals(args, start, end), "items": items}

    def dashboard_evolucion_anual(self, args: dict[str, Any]) -> dict[str, Any]:
        current_year = date.today().year
        year_from = self._to_int(args.get("anio_desde"), current_year - 4)
        year_to = self._to_int(args.get("anio_hasta"), current_year)
        if year_from > year_to:
            raise KofedasError("anio_desde no puede ser posterior a anio_hasta")
        items: list[dict[str, Any]] = []
        for year in range(year_from, year_to + 1):
            payload = {**args, "desde": f"{year}-01-01", "hasta": f"{year}-12-31"}
            sales = self._dashboard_sales_totals(payload, payload["desde"], payload["hasta"])
            purchases = self._dashboard_purchase_totals(payload, payload["desde"], payload["hasta"])
            margin = round(sales["base"] - purchases["base"], 2)
            items.append({
                "anio": year,
                "ventas_base": sales["base"],
                "ventas_total": sales["total"],
                "ventas_documentos": sales["documentos"],
                "compras_base": purchases["base"],
                "compras_total": purchases["total"],
                "compras_documentos": purchases["documentos"],
                "margen_bruto_aproximado": margin,
                "margen_bruto_pct": round((margin * 100 / sales["base"]) if sales["base"] else 0, 2),
            })
        return {"anio_desde": year_from, "anio_hasta": year_to, "items": items}

    def dashboard_series_temporales(self, args: dict[str, Any]) -> dict[str, Any]:
        group = str(args.get("agrupar_por") or "mes").strip().lower()
        if group not in {"mes", "anio"}:
            raise KofedasError("dashboard_series_temporales solo admite agrupar_por mes o anio")
        sales = self.ventas_resumen({**args, "agrupar_por": group, "limite": args.get("limite") or 500})
        purchases = self.compras_resumen({**args, "agrupar_por": group, "limite": args.get("limite") or 500})
        merged: dict[str, dict[str, Any]] = {}
        for item in sales["items"]:
            key = str(item.get("codigo"))
            merged.setdefault(key, {"periodo": item.get("codigo"), "ventas_base": 0.0, "compras_base": 0.0, "ventas_documentos": 0, "compras_documentos": 0})
            merged[key]["ventas_base"] = item.get("base", 0)
            merged[key]["ventas_documentos"] = item.get("documentos", 0)
        for item in purchases["items"]:
            key = str(item.get("codigo"))
            merged.setdefault(key, {"periodo": item.get("codigo"), "ventas_base": 0.0, "compras_base": 0.0, "ventas_documentos": 0, "compras_documentos": 0})
            merged[key]["compras_base"] = item.get("base", 0)
            merged[key]["compras_documentos"] = item.get("documentos", 0)
        items = sorted(merged.values(), key=lambda item: str(item["periodo"]))
        for item in items:
            item["margen_bruto_aproximado"] = round(self._to_float(item["ventas_base"], 0) - self._to_float(item["compras_base"], 0), 2)
        return {"periodo": sales["periodo"], "agrupar_por": group, "items": items}

    def dashboard_rankings(self, args: dict[str, Any]) -> dict[str, Any]:
        limit = _positive_limit(args.get("limite"), 10)
        sales_clients = self.ventas_resumen({**args, "agrupar_por": "cliente", "limite": limit})
        sales_articles = self.ventas_resumen({**args, "agrupar_por": "articulo", "limite": limit})
        purchase_providers = self.compras_resumen({**args, "agrupar_por": "proveedor", "limite": limit})
        purchase_articles = self.compras_resumen({**args, "agrupar_por": "articulo", "limite": limit})
        return {
            "periodo": sales_clients["periodo"],
            "clientes_venta": sales_clients["items"],
            "articulos_venta": sales_articles["items"],
            "proveedores_compra": purchase_providers["items"],
            "articulos_compra": purchase_articles["items"],
        }

    def regularizacion_tablas(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        del args
        result: list[dict[str, Any]] = []
        for table, description in sorted(REGULARIZATION_TABLES.items()):
            result.append({
                "tabla": table,
                "descripcion": description,
                "claves": self._table_pk(table),
                "columnas": self._table_columns(table),
            })
        return result

    def regularizacion_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        limit = _positive_limit(args.get("limite"))
        where = ["C.CBR_NUMEMP = ?"]
        params: list[Any] = [empresa]
        if args.get("centro") is not None:
            where.append("C.CBR_CENTRO = ?")
            params.append(int(args["centro"]))
        if str(args.get("tipo") or "").strip():
            where.append("C.CBR_TIPO = ?")
            params.append(str(args.get("tipo")).strip().upper()[:1])
        if args.get("desde"):
            where.append("C.CBR_FECHA >= ?")
            params.append(self._date_arg(args.get("desde")))
        if args.get("hasta"):
            where.append("C.CBR_FECHA <= ?")
            params.append(self._date_arg(args.get("hasta")))
        articulo = str(args.get("articulo") or "").strip()
        if articulo:
            where.append(
                """
                EXISTS (
                    SELECT 1 FROM DETMOVR D
                    WHERE D.DMR_NUMEMP = C.CBR_NUMEMP
                      AND D.DMR_CENTRO = C.CBR_CENTRO
                      AND D.DMR_EJERCI = C.CBR_EJERCI
                      AND D.DMR_SERIE = C.CBR_SERIE
                      AND D.DMR_NUMDOC = C.CBR_NUMDOC
                      AND D.DMR_CODART = ?
                )
                """
            )
            params.append(articulo)
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   C.CBR_NUMEMP, C.CBR_CENTRO, C.CBR_EJERCI, C.CBR_SERIE, C.CBR_NUMDOC,
                   C.CBR_FECHA, C.CBR_TIPO, C.CBR_CENREL, C.CBR_FECINF, C.CBR_FECSUP,
                   C.CBR_OBSERV, C.CBR_NUMDOCE, C.CBR_FECMOD, C.CBR_USUMOD,
                   COUNT(D.DMR_NUMLIN) AS LINEAS, SUM(D.DMR_CANTID) AS CANTIDAD
            FROM CABDOCR C
            LEFT JOIN DETMOVR D ON D.DMR_NUMEMP = C.CBR_NUMEMP
                              AND D.DMR_CENTRO = C.CBR_CENTRO
                              AND D.DMR_EJERCI = C.CBR_EJERCI
                              AND D.DMR_SERIE = C.CBR_SERIE
                              AND D.DMR_NUMDOC = C.CBR_NUMDOC
            WHERE {' AND '.join(where)}
            GROUP BY C.CBR_NUMEMP, C.CBR_CENTRO, C.CBR_EJERCI, C.CBR_SERIE, C.CBR_NUMDOC,
                     C.CBR_FECHA, C.CBR_TIPO, C.CBR_CENREL, C.CBR_FECINF, C.CBR_FECSUP,
                     C.CBR_OBSERV, C.CBR_NUMDOCE, C.CBR_FECMOD, C.CBR_USUMOD
            ORDER BY C.CBR_FECHA DESC, C.CBR_CENTRO, C.CBR_EJERCI DESC, C.CBR_SERIE, C.CBR_NUMDOC DESC
            """,
            tuple(params),
            limit,
        )
        for row in rows:
            tipo = str(row.get("cbr_tipo") or "").strip().upper()
            row["tipo_descripcion"] = REGULARIZATION_TYPE_LABELS.get(tipo, tipo)
        return rows

    def stock_por_almacen(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        articulo = str(args["articulo"]).strip()
        where = ["E.ARTE_NUMEMP = ?", "E.ARTE_CODART = ?"]
        params: list[Any] = [empresa, articulo]
        if args.get("centro") is not None:
            where.append("E.ARTE_CENTRO = ?")
            params.append(int(args["centro"]))
        return self.db.query(
            f"""
            SELECT E.ARTE_NUMEMP, E.ARTE_CODART, A.ART_DESCRI, E.ARTE_CENTRO,
                   C.CEN_NOMCEN, E.ARTE_EXIST, E.ARTE_MINIMO, E.ARTE_MAXIMO,
                   E.ARTE_FECCOM, E.ARTE_FECVEN, E.ARTE_FECMOV
            FROM ARTICULE E
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = E.ARTE_NUMEMP AND A.ART_CODART = E.ARTE_CODART
            LEFT JOIN CENTROS C ON C.CEN_NUMEMP = E.ARTE_NUMEMP AND C.CEN_CODCEN = E.ARTE_CENTRO
            WHERE {' AND '.join(where)}
            ORDER BY E.ARTE_CENTRO
            """,
            tuple(params),
            MAX_ROWS_LIMIT,
        )

    def stock_a_fecha(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        articulo = str(args["articulo"]).strip()
        centro = int(args.get("centro") if args.get("centro") is not None else self.centro)
        query_date = self._date_arg(args.get("fecha")) if args.get("fecha") else date.today().isoformat()
        current = self._stock_current(empresa, articulo, centro)
        article = self._article_basic(empresa, articulo)
        if str(article.get("art_indinv") or "").upper() == "N":
            return {"articulo": articulo, "centro": centro, "fecha": query_date, "stock_actual": current, "stock_fecha": 0, "inventariable": False}
        if query_date == date.today().isoformat():
            return {"articulo": articulo, "centro": centro, "fecha": query_date, "stock_actual": current, "stock_fecha": current, "inventariable": True}
        if centro == -1:
            entrada_filter = "DMM_NUMEMP = ? AND DMM_CODART = ? AND DMM_FECMOV >= ?"
            regular_filter = "DMR_NUMEMP = ? AND DMR_CODART = ? AND DMR_FECMOV >= ?"
            venta_filter = "DMV_NUMEMP = ? AND DMV_CODART = ? AND DMV_SIGNO = '1' AND DMV_FECMOV >= ?"
            movement_params = (empresa, articulo, query_date)
        else:
            entrada_filter = "DMM_NUMEMP = ? AND DMM_CENTRO = ? AND DMM_CODART = ? AND DMM_FECMOV >= ?"
            regular_filter = "DMR_NUMEMP = ? AND DMR_CENTRO = ? AND DMR_CODART = ? AND DMR_FECMOV >= ?"
            venta_filter = "DMV_NUMEMP = ? AND DMV_CENTRO = ? AND DMV_CODART = ? AND DMV_SIGNO = '1' AND DMV_FECMOV >= ?"
            movement_params = (empresa, centro, articulo, query_date)
        entradas = self._movement_sum(f"SELECT SUM(DMM_CANTID) AS EXISTENCIAS FROM DETMOVM WHERE {entrada_filter}", movement_params)
        regularizaciones = self._movement_sum(f"SELECT SUM(DMR_CANTID) AS EXISTENCIAS FROM DETMOVR WHERE {regular_filter}", movement_params)
        ventas = self._movement_sum(f"SELECT SUM(DMV_CANTID) AS EXISTENCIAS FROM DETMOV WHERE {venta_filter}", movement_params)
        stock_fecha = current - entradas - regularizaciones + ventas
        return {
            "articulo": articulo,
            "centro": centro,
            "fecha": query_date,
            "stock_actual": current,
            "movimientos_desde_fecha": {"entradas": entradas, "regularizaciones": regularizaciones, "ventas": ventas},
            "stock_fecha": round(stock_fecha, 4),
            "inventariable": True,
            "fuente_delphi": "LIBESP_U.UTL_STOCK",
        }

    def inventario_valorar_articulos(self, args: dict[str, Any]) -> dict[str, Any]:
        empresa = self._empresa(args)
        centro = int(args.get("centro") if args.get("centro") is not None else self.centro)
        query_date = self._date_arg(args.get("fecha")) if args.get("fecha") else date.today().isoformat()
        limit = _positive_limit(args.get("limite"), 100)
        mode_info = self._inventory_cost_mode(empresa, args.get("modo_coste"))
        where = ["A.ART_NUMEMP = ?"]
        join_params: list[Any] = []
        params: list[Any] = [empresa]
        join_extra = ""
        if centro != -1:
            join_extra = " AND E.ARTE_CENTRO = ?"
            join_params.append(centro)
        if args.get("articulo"):
            where.append("A.ART_CODART = ?")
            params.append(str(args["articulo"]).strip())
        if args.get("texto"):
            where.append("(UPPER(A.ART_CODART) LIKE ? OR UPPER(A.ART_DESCRI) LIKE ?)")
            like = _like(str(args["texto"]))
            params.extend([like, like])
        if args.get("familia"):
            where.append("A.ART_CODFAM = ?")
            params.append(str(args["familia"]).strip())
        if args.get("subfamilia"):
            where.append("A.ART_SUBFAM = ?")
            params.append(str(args["subfamilia"]).strip())
        if args.get("proveedor") is not None:
            where.append("A.ART_CODPRO = ?")
            params.append(int(args["proveedor"]))
        if not args.get("incluir_no_inventariables"):
            where.append("COALESCE(A.ART_INDINV, 'S') <> 'N'")
        having = "HAVING COALESCE(SUM(E.ARTE_EXIST), 0) <> 0" if args.get("solo_con_stock", True) else ""
        rows = self.db.query(
            f"""
            SELECT FIRST {limit}
                   A.ART_CODART, A.ART_DESCRI, A.ART_INDINV, A.ART_CODFAM, A.ART_SUBFAM,
                   A.ART_CODPRO, A.ART_PRECOS, A.ART_PREBAS, A.ART_CANPRE, A.ART_CODMON,
                   COALESCE(SUM(E.ARTE_EXIST), 0) AS STOCK_ACTUAL
            FROM ARTICUL A
            LEFT JOIN ARTICULE E ON E.ARTE_NUMEMP = A.ART_NUMEMP
                                AND E.ARTE_CODART = A.ART_CODART
                                {join_extra}
            WHERE {' AND '.join(where)}
            GROUP BY A.ART_CODART, A.ART_DESCRI, A.ART_INDINV, A.ART_CODFAM, A.ART_SUBFAM,
                     A.ART_CODPRO, A.ART_PRECOS, A.ART_PREBAS, A.ART_CANPRE, A.ART_CODMON
            {having}
            ORDER BY A.ART_DESCRI, A.ART_CODART
            """,
            tuple(join_params + params),
            limit,
        )
        items: list[dict[str, Any]] = []
        total_stock = 0.0
        total_value = 0.0
        for row in rows:
            articulo = str(row.get("art_codart") or "")
            stock = self._to_float(row.get("stock_actual"), 0)
            if query_date != date.today().isoformat():
                stock = self._stock_at_date(empresa, articulo, centro, query_date)
            if args.get("solo_con_stock", True) and stock == 0:
                continue
            cost = self._article_inventory_cost(empresa, row, query_date, str(mode_info["modo"]))
            value = stock * self._to_float(cost["coste_unitario"], 0)
            total_stock += stock
            total_value += value
            items.append({
                "articulo": articulo,
                "descripcion": row.get("art_descri"),
                "familia": row.get("art_codfam"),
                "subfamilia": row.get("art_subfam"),
                "proveedor": row.get("art_codpro"),
                "inventariable": str(row.get("art_indinv") or "S").upper() != "N",
                "stock": round(stock, 4),
                "coste_unitario": cost["coste_unitario"],
                "valor_coste": round(value, 2),
                "origen_coste": cost["origen_coste"],
            })
        return {
            "empresa": empresa,
            "centro": centro,
            "fecha": query_date,
            "modo_coste": mode_info,
            "total_articulos": len(items),
            "totales": {"stock": round(total_stock, 4), "valor_coste": round(total_value, 2)},
            "items": items,
            "fuente_delphi": "ARTICUL_UB.PRECIO_COSTE_ARTICUL + PARAMETROS.RENTAB",
        }

    def articulo_regularizar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        centro = self._centro(args)
        articulo = str(args["articulo"]).strip()
        target = self._to_float(args.get("cantidad"), 0)
        current = self._stock_current(empresa, articulo, centro)
        difference = target - current
        article = self._article_basic(empresa, articulo)
        if difference == 0:
            return {"simulado": bool(args.get("simular")), "accion": "sin_cambios", "articulo": articulo, "centro": centro, "stock_actual": current, "stock_objetivo": target, "diferencia": 0}
        fecha = self._date_arg(args.get("fecha"), (date.today() - timedelta(days=1)).isoformat())
        serie = self._regularization_series(empresa, centro, args.get("serie"))
        header = self._regularization_header(empresa, centro, fecha, "R", serie, args.get("observaciones") or "Regularizacion MCP")
        line = self._detmovr_row(
            empresa, centro, int(header["CBR_EJERCI"]), serie, int(header["CBR_NUMDOC"]), 1, fecha,
            articulo, str(article.get("art_descri") or ""), str(article.get("art_unimed") or ""), difference,
        )
        columns = self._table_columns("CABDOCR")
        dcols = self._table_columns("DETMOVR")
        statements: list[tuple[str, tuple[Any, ...]]] = [(
            "INSERT INTO CABDOCR (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")",
            tuple(header.get(column) for column in columns),
        )]
        if str(article.get("art_indinv") or "").upper() == "S":
            statements.extend(self._regularization_stock_statements(empresa, centro, articulo, difference, fecha))
        statements.append((
            "INSERT INTO DETMOVR (" + ", ".join(dcols) + ") VALUES (" + ", ".join("?" for _ in dcols) + ")",
            tuple(line.get(column) for column in dcols),
        ))
        if args.get("simular"):
            return {"simulado": True, "cabecera": header, "linea": line, "stock_actual": current, "stock_objetivo": target, "diferencia": difference, "sentencias": len(statements)}
        counts = self.db.execute_transaction(statements)
        return {"simulado": False, "cabecera": header, "linea": line, "stock_anterior": current, "stock_objetivo": target, "diferencia": difference, "filas_afectadas": counts}

    def trasvase_generar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        origin = int(args["centro_origen"])
        destination = int(args["centro_destino"])
        if origin == destination:
            raise KofedasError("centro_origen y centro_destino no pueden ser iguales")
        fecha = self._date_arg(args.get("fecha"))
        serie = self._regularization_series(empresa, origin, args.get("serie"))
        raw_lines = args.get("lineas") or []
        if not isinstance(raw_lines, list) or not raw_lines:
            raise KofedasError("lineas debe ser una lista no vacia")
        out_header = self._regularization_header(empresa, origin, fecha, "S", serie, args.get("observaciones") or f"Trasvase Centro {destination}", destination)
        in_header = self._regularization_header(empresa, destination, fecha, "E", serie, args.get("observaciones") or f"Trasvase Centro {origin}", origin, int(out_header["CBR_NUMDOC"]))
        out_header["CBR_NUMDOCE"] = int(in_header["CBR_NUMDOC"])
        lines_out: list[dict[str, Any]] = []
        lines_in: list[dict[str, Any]] = []
        statements: list[tuple[str, tuple[Any, ...]]] = []
        columns = self._table_columns("CABDOCR")
        dcols = self._table_columns("DETMOVR")
        statements.append(("INSERT INTO CABDOCR (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")", tuple(out_header.get(column) for column in columns)))
        statements.append(("INSERT INTO CABDOCR (" + ", ".join(columns) + ") VALUES (" + ", ".join("?" for _ in columns) + ")", tuple(in_header.get(column) for column in columns)))
        for index, raw in enumerate(raw_lines, start=1):
            if not isinstance(raw, dict):
                raise KofedasError("Cada linea de trasvase debe ser un objeto")
            articulo = str(raw.get("articulo") or raw.get("codigo") or "").strip()
            article = self._article_basic(empresa, articulo)
            qty = self._to_float(raw.get("cantidad"), 0)
            if qty <= 0:
                raise KofedasError("La cantidad debe ser mayor que cero para " + articulo)
            desc = str(raw.get("descripcion") or article.get("art_descri") or "")
            unit = str(raw.get("unidad") or article.get("art_unimed") or "")
            out_line = self._detmovr_row(empresa, origin, int(out_header["CBR_EJERCI"]), serie, int(out_header["CBR_NUMDOC"]), index, fecha, articulo, desc, unit, -qty)
            in_line = self._detmovr_row(empresa, destination, int(in_header["CBR_EJERCI"]), serie, int(in_header["CBR_NUMDOC"]), index, fecha, articulo, desc, unit, qty, {"ejercicio": out_header["CBR_EJERCI"], "serie": serie, "numero": out_header["CBR_NUMDOC"], "linea": index})
            lines_out.append(out_line)
            lines_in.append(in_line)
            if str(article.get("art_indinv") or "").upper() == "S":
                statements.extend(self._regularization_stock_statements(empresa, origin, articulo, -qty, fecha))
                statements.extend(self._regularization_stock_statements(empresa, destination, articulo, qty, fecha))
            statements.append(("INSERT INTO DETMOVR (" + ", ".join(dcols) + ") VALUES (" + ", ".join("?" for _ in dcols) + ")", tuple(out_line.get(column) for column in dcols)))
            statements.append(("INSERT INTO DETMOVR (" + ", ".join(dcols) + ") VALUES (" + ", ".join("?" for _ in dcols) + ")", tuple(in_line.get(column) for column in dcols)))
        if args.get("simular"):
            return {"simulado": True, "salida": out_header, "entrada": in_header, "lineas_salida": lines_out, "lineas_entrada": lines_in, "sentencias": len(statements)}
        counts = self.db.execute_transaction(statements)
        return {"simulado": False, "salida": out_header, "entrada": in_header, "lineas": len(lines_out), "filas_afectadas": counts}

    def recuento_listar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        empresa = self._empresa(args)
        centro = self._centro(args)
        limit = _positive_limit(args.get("limite"))
        where = ["R.REC_NUMEMP = ?", "R.REC_CENTRO = ?"]
        params: list[Any] = [empresa, centro]
        if str(args.get("articulo") or "").strip():
            where.append("R.REC_CODART = ?")
            params.append(str(args.get("articulo")).strip())
        return self.db.query(
            f"""
            SELECT FIRST {limit} R.*, A.ART_DESCRI AS ART_DESCRI_MAESTRA, E.ARTE_EXIST AS STOCK_ACTUAL
            FROM RECUENTO R
            LEFT JOIN ARTICUL A ON A.ART_NUMEMP = R.REC_NUMEMP AND A.ART_CODART = R.REC_CODART
            LEFT JOIN ARTICULE E ON E.ARTE_NUMEMP = R.REC_NUMEMP AND E.ARTE_CODART = R.REC_CODART AND E.ARTE_CENTRO = R.REC_CENTRO
            WHERE {' AND '.join(where)}
            ORDER BY R.REC_CODART
            """,
            tuple(params),
            limit,
        )

    def recuento_grabar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        centro = self._centro(args)
        articulo = str(args["articulo"]).strip()
        article = self._article_basic(empresa, articulo)
        amount = self._to_float(args.get("cantidad"), 0)
        fecha = self._date_arg(args.get("fecha"))
        existing = self.db.one(
            "SELECT REC_EXIST FROM RECUENTO WHERE REC_NUMEMP = ? AND REC_CENTRO = ? AND REC_CODART = ?",
            (empresa, centro, articulo),
        )
        aumentar = bool(args.get("aumentar", False))
        final_amount = self._to_float((existing or {}).get("rec_exist"), 0) + amount if existing and aumentar else amount
        statements: list[tuple[str, tuple[Any, ...]]]
        if existing:
            statements = [(
                """
                UPDATE RECUENTO
                SET REC_DESCRI = ?, REC_EXIST = ?, REC_FECHA = ?, REC_UNIMED = ?
                WHERE REC_NUMEMP = ? AND REC_CENTRO = ? AND REC_CODART = ?
                """,
                (str(article.get("art_descri") or "")[:50], final_amount, fecha, str(article.get("art_unimed") or "")[:4], empresa, centro, articulo),
            )]
        else:
            statements = [(
                "INSERT INTO RECUENTO (REC_NUMEMP, REC_CENTRO, REC_CODART, REC_DESCRI, REC_EXIST, REC_FECHA, REC_UNIMED) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (empresa, centro, articulo, str(article.get("art_descri") or "")[:50], final_amount, fecha, str(article.get("art_unimed") or "")[:4]),
            )]
        if args.get("simular"):
            return {"simulado": True, "articulo": articulo, "centro": centro, "existia": bool(existing), "cantidad_anterior": self._to_float((existing or {}).get("rec_exist"), 0), "cantidad_final": final_amount, "sentencias": len(statements)}
        counts = self.db.execute_transaction(statements)
        return {"simulado": False, "articulo": articulo, "centro": centro, "cantidad_final": final_amount, "filas_afectadas": counts}

    def recuento_borrar(self, args: dict[str, Any]) -> dict[str, Any]:
        self._require_write()
        empresa = self._empresa(args)
        centro = self._centro(args)
        articulo = str(args["articulo"]).strip()
        exists = self.db.one(
            "SELECT FIRST 1 1 AS EXISTE FROM RECUENTO WHERE REC_NUMEMP = ? AND REC_CENTRO = ? AND REC_CODART = ?",
            (empresa, centro, articulo),
        ) is not None
        statements = [("DELETE FROM RECUENTO WHERE REC_NUMEMP = ? AND REC_CENTRO = ? AND REC_CODART = ?", (empresa, centro, articulo))]
        if args.get("simular"):
            return {"simulado": True, "articulo": articulo, "centro": centro, "existia": exists, "sentencias": len(statements)}
        counts = self.db.execute_transaction(statements)
        return {"simulado": False, "articulo": articulo, "centro": centro, "existia": exists, "filas_afectadas": counts}
